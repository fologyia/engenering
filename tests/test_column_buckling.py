"""Flambagem de colunas: a barra inteira, todos os eixos e modos, três normas e o critério Anglo.

Os valores de aceite vêm do pedido de robustecimento do módulo: o exemplo do programa (barra
circular Ø50), o AISC Design Guide 29 (HSS8×8×1/2), o U idealizado do critério e os limites da
Anglo. Os demais testes travam o contrato: o módulo concorda com ``core.nbr8800`` em todo o
catálogo, bloqueia o que a norma não cobre em vez de estimar, e a interação N + M_x + M_y é uma
equação só.
"""

from __future__ import annotations

import math
from dataclasses import replace

import pytest

from core import column_buckling as cb
from core import column_design as cd
from core import nbr8800
from core import section_catalog as catalogo
from core import steel_sections as ss

KSI, INCH, KIP = 6.894757, 25.4, 4.448222
E, G = 200_000.0, 77_000.0


def perto(a: float, b: float, rel: float = 1e-3) -> bool:
    return abs(a - b) <= rel * abs(b)


def entrada(secao, **mudancas) -> cb.EntradaColuna:
    base = dict(secao=secao, fy_MPa=250.0, Lx_mm=2000.0, Ly_mm=2000.0, N_Sd_kN=10.0)
    base.update(mudancas)
    return cb.EntradaColuna(**base)


def linha(resultado: cb.ResultadoColuna, trecho: str):
    achadas = [v for v in resultado.verificacoes if trecho in v.nome]
    assert len(achadas) == 1, f"{trecho!r}: {[v.nome for v in resultado.verificacoes]}"
    return achadas[0]


def tem_linha(resultado: cb.ResultadoColuna, trecho: str) -> bool:
    return any(trecho in v.nome for v in resultado.verificacoes)


def secao_generica_com_mrd(**mudancas) -> cd.Secao:
    base = dict(
        A=5000.0, Ix=1e8, Iy=1e8, Wx=1e6, Wy=1e6, Zx=1.1e6, Zy=1.1e6, torcao_relevante=False
    )
    base.update(mudancas)
    return cd.secao_generica("genérica", **base)


# ------------------------------------------------------------------ ações de cálculo e tabelas de K
class TestAcoesEK:
    def test_majoracao_padrao(self):
        assert cb.forca_de_calculo(30_000.0, 20_000.0) == pytest.approx(70_000.0)
        assert cb.forca_de_calculo(30e3, 20e3, gamma_g=1.4, gamma_q=1.5) == pytest.approx(72_000.0)

    def test_rejeita_negativos(self):
        with pytest.raises(ValueError):
            cb.forca_de_calculo(-1.0)

    def test_k_recomendado_nunca_e_menor_que_o_teorico(self):
        assert set(cb.CONDICOES_APOIO) == set(cb.CONDICOES_APOIO_RECOMENDADAS)
        for nome, teorico in cb.CONDICOES_APOIO.items():
            assert cb.CONDICOES_APOIO_RECOMENDADAS[nome] >= teorico

    def test_tabelas_concordam_com_a_tabela_E1_do_modulo_de_normas(self):
        # Mesma tabela em dois lugares: se uma mudar sem a outra, o teste aponta.
        correspondencia = {
            "Biapoiada (pino-pino)": "rotula-rotula",
            "Engastada-livre (em balanço)": "engaste-livre",
            "Engastada-pino": "engaste-rotula",
            "Biengastada": "engaste-engaste",
            "Biengastada com translação (deslocável)": "engaste-engaste_translacao_livre",
            "Engastada-pino com translação (deslocável)": "rotula-engaste_translacao_livre",
        }
        assert set(correspondencia) == set(cb.CONDICOES_APOIO)
        for rotulo, chave in correspondencia.items():
            teorico, recomendado = cd.K_TABELA_E1[chave]
            assert cb.CONDICOES_APOIO[rotulo] == teorico
            assert cb.CONDICOES_APOIO_RECOMENDADAS[rotulo] == recomendado
            assert cd.K_recomendado(chave) == recomendado


class TestMaoFrancesa:
    BALANCO, BIAPOIADA, PROPPED = cb.VINCULOS_MAO_FRANCESA

    def test_decompoe_a_forca_no_angulo(self):
        mf = cb.esforcos_mao_francesa(20_000.0, 30.0, 800.0, 3_000.0, self.BALANCO)
        assert mf.componente_horizontal_N == pytest.approx(20_000.0 * math.sin(math.radians(30.0)))
        assert mf.componente_vertical_N == pytest.approx(20_000.0 * math.cos(math.radians(30.0)))

    def test_vinculos(self):
        balanco = cb.esforcos_mao_francesa(20_000.0, 45.0, 800.0, 3_000.0, self.BALANCO)
        biapoiada = cb.esforcos_mao_francesa(20_000.0, 45.0, 800.0, 3_000.0, self.BIAPOIADA)
        apoiada = cb.esforcos_mao_francesa(20_000.0, 45.0, 800.0, 3_000.0, self.PROPPED)
        h = balanco.componente_horizontal_N
        assert balanco.momento_Nmm == pytest.approx(h * 800.0)
        assert biapoiada.momento_Nmm == pytest.approx(h * 800.0 * 2_200.0 / 3_000.0)
        assert apoiada.momento_Nmm < balanco.momento_Nmm

    def test_excentricidade_soma_v_vezes_e(self):
        mf = cb.esforcos_mao_francesa(20_000.0, 45.0, 800.0, 3_000.0, self.BALANCO, 50.0)
        assert mf.momento_excentricidade_Nmm == pytest.approx(mf.componente_vertical_N * 50.0)
        assert mf.momento_Nmm == pytest.approx(
            mf.momento_horizontal_Nmm + mf.momento_excentricidade_Nmm
        )

    def test_entradas_invalidas(self):
        for argumentos in (
            (1.0, 0.0, 800.0, 3_000.0),
            (1.0, 90.0, 800.0, 3_000.0),
            (1.0, 45.0, 3_500.0, 3_000.0),
            (1.0, 45.0, 800.0, 3_000.0, "apoio qualquer"),
        ):
            with pytest.raises(ValueError):
                cb.esforcos_mao_francesa(*argumentos)

    def test_o_momento_entra_na_interacao_no_eixo_certo_e_avisa_do_cortante(self):
        sec = cd.secao_I(300, 150, 12.5, 8.0)
        mf = cb.esforcos_mao_francesa(20_000.0, 45.0, 800.0, 3_000.0, self.BALANCO)
        sem = cb.verificar_coluna(entrada(sec, N_Sd_kN=100.0))
        com = cb.verificar_coluna(
            entrada(
                sec, N_Sd_kN=100.0, Mx_mao_francesa_kNm=mf.momento_Nmm / 1e6, com_mao_francesa=True
            )
        )
        assert sem.indice_interacao is None and com.indice_interacao is not None
        assert com.momento_x_primeira_ordem_kNm == pytest.approx(mf.momento_Nmm / 1e6)
        assert com.momento_y_primeira_ordem_kNm == 0.0
        assert linha(com, "Cisalhamento").status == "ALERTA"
        assert not tem_linha(sem, "Cisalhamento")


# ------------------------------------------------------------------ critérios de aceite
class TestCriteriosDeAceite:
    def test_barra_circular_50(self):
        r = cb.verificar_coluna(entrada(cd.secao_circular_macica(50), N_Sd_kN=72.0))
        comp = r.compressao
        assert perto(comp["Nc_Rd"], 120.71, 1e-3) and perto(comp["Nex"], 151.4, 1e-3)
        assert perto(comp["lambda0"], 1.801, 1e-3) and perto(comp["chi"], 0.270, 2e-3)
        assert perto(r.aproveitamento_max, 0.596, 2e-3) and r.status_geral == "OK"
        esbeltez = linha(r, "Esbeltez limite")
        assert esbeltez.solicitante == pytest.approx(160.0) and esbeltez.status == "OK"
        assert linha(r, "N_c,Rd").status == "OK"

    def test_nsd_pelo_programa_1_40_g_mais_1_50_q(self):
        assert cd.combinacao_ultima([(30, 1.40)], (20, 1.50)) == pytest.approx(72.0)

    def test_hss8x8x1_2_do_design_guide_29(self):
        t = 0.465 * INCH
        area, raio = 13.5 * INCH**2, 3.04 * INCH
        sec = cd.secao_generica(
            "HSS8x8x1/2",
            area,
            area * raio**2,
            area * raio**2,
            torcao_relevante=False,
            elementos=[dict(tipo="AA", grupo=1, b=14.2 * t, t=t, n=4, tubo_ret=True)],
        )
        comprimento = 24 * 12 * INCH
        r = cb.verificar_coluna(
            cb.EntradaColuna(
                secao=sec,
                fy_MPa=46 * KSI,
                norma="AISC360_16_LRFD",
                E_MPa=29_000 * KSI,
                Lx_mm=comprimento,
                Ly_mm=comprimento,
                N_Sd_kN=1000.0,
                secao_compacta_confirmada=True,
            )
        )
        assert perto(r.compressao["Nc_Rd"] / KIP, 306.0, 5e-3)
        assert perto(r.compressao["Nc_Rd"], 1361.0, 5e-3)
        assert r.compressao["Q"] == 1.0  # b/t = 14,2 < 1,40√(E/F_y)

    def test_u_idealizado_a36_com_todos_os_modos(self):
        u = cd.secao_U(152.4, 48.8, 8.71, 5.08)
        r = cb.verificar_coluna(entrada(u, Lx_mm=3000.0, Ly_mm=1500.0, Lz_mm=3000.0, kz=1.0))
        comp = r.compressao
        assert perto(comp["Ney"], 308.4, 1e-3) and comp["modo"] == "flexão y"
        assert perto(comp["modos"]["flexo-torção xz"], 457.4, 1e-3)
        assert perto(comp["Nc_Rd"], 207.3, 1e-3)
        # todos os N_e aparecem na tabela
        for trecho in ("N_ex", "N_ey", "N_ez", "N_e flexo-torção xz", "N_e adotado"):
            assert tem_linha(r, trecho), trecho

    def test_interacao_biaxial_reprova_o_que_eixo_a_eixo_aprovaria(self):
        sec = secao_generica_com_mrd()
        base = entrada(sec, Lx_mm=1000.0, Ly_mm=1000.0, secao_compacta_confirmada=True)
        n_rd = cb.verificar_coluna(replace(base, N_Sd_kN=1.0)).compressao["Nc_Rd"]
        comum = dict(N_Sd_kN=0.3 * n_rd, MRd_x_informado_kNm=100.0, MRd_y_informado_kNm=100.0)
        so_x = cb.verificar_coluna(replace(base, Mx_kNm=50.0, **comum))
        so_y = cb.verificar_coluna(replace(base, My_kNm=40.0, **comum))
        juntos = cb.verificar_coluna(replace(base, Mx_kNm=50.0, My_kNm=40.0, **comum))
        # B₁ ≈ 1,002 desloca os valores de 0,744 e 0,656 só na terceira casa
        assert so_x.indice_interacao < 0.75 and so_y.indice_interacao < 0.66
        assert perto(juntos.indice_interacao, 1.10, 3e-3) and juntos.status_geral == "NÃO OK"
        assert linha(juntos, "Interação").status == "NÃO OK"
        assert "Interação" in juntos.governante

    @pytest.mark.parametrize(
        ("esbeltez", "status"), [(199.0, "OK"), (200.0, "OK"), (201.0, "NÃO OK")]
    )
    def test_limite_de_esbeltez_da_anglo(self, esbeltez, status):
        sec = cd.secao_circular_macica(20)  # r = 5 mm
        r = cb.verificar_coluna(
            entrada(sec, Lx_mm=esbeltez * 5.0, Ly_mm=esbeltez * 5.0, N_Sd_kN=0.01)
        )
        assert linha(r, "Esbeltez limite").status == status
        assert r.status_geral == status
        if status == "NÃO OK":
            # a resistência sozinha não reprova: o limite é que reprova — e o status geral segue
            assert r.aproveitamento_max < 1.0 and r.reprovadas[0].nome.startswith("Esbeltez limite")

    def test_espessura_minima_do_u_laminado(self):
        u = cd.secao_U(152.4, 48.8, 8.71, 4.5)  # t_w = 4,5 mm < 4,80
        categoria, t = cb.categoria_espessura_anglo(u)
        assert categoria == "perfil_laminado_L_U" and t == pytest.approx(4.5)
        r = cb.verificar_coluna(entrada(u, espessuras_anglo={categoria: t}))
        v = linha(r, "Espessura mínima")
        assert v.status == "NÃO OK" and "laminado L/U" in v.nome and r.status_geral == "NÃO OK"
        ok = cb.verificar_coluna(entrada(u, espessuras_anglo={categoria: 4.8}))
        assert linha(ok, "Espessura mínima").status == "OK"


# ------------------------------------------------------------------ seções
class TestSecoes:
    def test_perfil_do_catalogo_usa_as_propriedades_dele(self):
        perfil = catalogo.obter_perfil("W 310 x 44,5")
        sec = cb.secao_de_perfil(perfil)
        assert sec.tipo == "I" and sec.nome == perfil.nome
        assert (sec.A, sec.Ix, sec.Iy, sec.J, sec.Cw) == (
            perfil.area_mm2,
            perfil.ix_mm4,
            perfil.iy_mm4,
            perfil.j_mm4,
            perfil.cw_mm6,
        )
        assert (sec.Zx, sec.Zy) == (perfil.zx_mm3, perfil.zy_mm3)
        assert [e["grupo"] for e in sec.elementos] == [4, 2]
        assert cb.secao_de_perfil(perfil, soldado=True).elementos[0]["grupo"] == 5

    def test_todas_as_familias_do_catalogo_viram_secao(self):
        tipos = set()
        for nome, perfil in catalogo.listar_perfis().items():
            sec = cb.secao_de_perfil(perfil)
            tipos.add(sec.tipo)
            assert sec.A > 0 and sec.Ix > 0 and sec.Iy > 0, nome
            if sec.tipo == "generica" and not sec.elementos:
                assert sec.dims.get("aviso"), nome  # nunca assume Q = 1 em silêncio
        assert {"I", "U", "tubo_circ", "tubo_ret", "circ_macica", "ret_macica", "generica"} <= tipos

    def test_u_do_catalogo_tem_centro_de_cisalhamento_e_wy_da_ponta_da_mesa(self):
        perfil = catalogo.obter_perfil('U 6" x 12,20')
        sec = cb.secao_de_perfil(perfil)
        assert sec.tipo == "U" and sec.x0 > 0 and sec.Cw > 0
        xb = ss.centroide_do_perfil(perfil)[0]
        assert sec.Wy == pytest.approx(perfil.iy_mm4 / max(xb, perfil.largura_mm - xb))
        assert sec.Wy < perfil.iy_mm4 / (perfil.largura_mm / 2)  # menor que o do meio da largura

    def test_t_e_generico_com_talao_e_centro_de_cisalhamento(self):
        perfil = catalogo.obter_perfil('T 3/4x1/8"')
        sec = cb.secao_de_perfil(perfil)
        assert sec.tipo == "generica" and sec.y0 > 0 and sec.x0 == 0.0
        assert [e["grupo"] for e in sec.elementos] == [4, 6]

    def test_formado_a_frio_e_marcado_como_fora_do_escopo(self):
        perfil = ss.CATALOGO_PERFIS["C ideal 100×50×17×2.65"]
        sec = cb.secao_de_perfil(perfil)
        assert sec.tipo == "generica" and "NBR 14762" in sec.dims["aviso"]
        r = cb.verificar_coluna(entrada(sec))
        assert linha(r, "Seção fora do escopo").status == "ALERTA"
        assert r.status_geral == "ALERTA"

    def test_secao_direta(self):
        sec = cb.secao_direta(1_963.5, 12.5, distancia_fibra_x_mm=25.0)
        assert sec.rx == pytest.approx(12.5) and sec.ry == pytest.approx(12.5)
        assert sec.Wx == pytest.approx(1_963.5 * 12.5**2 / 25.0) and sec.Wy is None
        assert sec.torcao_relevante is False and sec.elementos == []

    def test_elementos_manuais_validam_tipo_e_grupo(self):
        ok = cb.elementos_manuais(
            [
                dict(tipo="al", grupo=5, b=75.0, t=9.5, n=4, kc=0.5),
                dict(tipo="AA", grupo=2, b=280.0, t=6.3, n=1),
                dict(tipo="AA", grupo=1, b=100.0, t=4.0, n=2),
            ]
        )
        assert ok[0]["kc"] == 0.5 and ok[2]["tubo_ret"] is True and ok[0]["tipo"] == "AL"
        for ruim in (
            dict(tipo="AL", grupo=2, b=1, t=1),  # grupo 2 é AA
            dict(tipo="AA", grupo=4, b=1, t=1),  # grupo 4 é AL
            dict(tipo="AL", grupo=5, b=1, t=1),  # sem k_c
            dict(tipo="AL", grupo=4, b=0, t=1),
            dict(tipo="XX", grupo=4, b=1, t=1),
        ):
            with pytest.raises(ValueError):
                cb.elementos_manuais([ruim])

    def test_categoria_de_espessura_por_tipo(self):
        assert (
            cb.categoria_espessura_anglo(cd.secao_I(300, 150, 9.5, 6.3))[0] == "perfil_laminado_H_W"
        )
        assert (
            cb.categoria_espessura_anglo(cd.secao_I(300, 150, 9.5, 6.3), soldado=True)[0]
            == "perfil_soldado"
        )
        assert cb.categoria_espessura_anglo(cd.secao_circular_macica(50)) == (None, None)
        assert cb.categoria_espessura_anglo(cd.secao_tubo_retangular(100, 100, 4)) == (None, None)


class TestSecoesPorDimensoesEInformadas:
    def test_cada_tipo_vira_a_secao_do_nucleo(self):
        casos = {
            "circular_macica": (dict(d=50.0), cd.secao_circular_macica(50.0)),
            "retangular_macica": (dict(b=50.0, h=100.0), cd.secao_retangular_macica(50.0, 100.0)),
            "tubo_circular": (dict(D=60.0, t=4.0), cd.secao_tubo_circular(60.0, 4.0)),
            "tubo_retangular": (
                dict(B=100.0, H=150.0, t=5.0),
                cd.secao_tubo_retangular(100.0, 150.0, 5.0),
            ),
            "I": (dict(d=300.0, bf=150.0, tf=12.5, tw=8.0), cd.secao_I(300, 150, 12.5, 8.0)),
            "U": (dict(d=152.4, bf=48.8, tf=8.71, tw=5.08), cd.secao_U(152.4, 48.8, 8.71, 5.08)),
        }
        assert set(casos) == set(cb.TIPOS_SECAO_POR_DIMENSOES)
        for tipo, (dimensoes, esperada) in casos.items():
            secao = cb.secao_por_dimensoes(tipo, **dimensoes)
            assert secao.nome == esperada.nome and secao.tipo == esperada.tipo, tipo
            assert (secao.A, secao.Ix, secao.Iy, secao.J) == (
                esperada.A,
                esperada.Ix,
                esperada.Iy,
                esperada.J,
            ), tipo

    def test_soldado_muda_o_grupo_da_mesa(self):
        dimensoes = dict(d=300.0, bf=150.0, tf=12.5, tw=8.0)
        laminado = cb.secao_por_dimensoes("I", **dimensoes)
        soldado = cb.secao_por_dimensoes("I", soldado=True, **dimensoes)
        assert [e["grupo"] for e in laminado.elementos] == [4, 2]
        assert [e["grupo"] for e in soldado.elementos] == [5, 2]
        assert cb.categoria_espessura_anglo(soldado)[0] == "perfil_soldado"

    @pytest.mark.parametrize(
        ("tipo", "dimensoes", "trecho"),
        [
            ("tubo_circular", dict(D=20.0, t=10.0), "metade do diâmetro"),
            ("tubo_circular", dict(D=20.0, t=12.0), "metade do diâmetro"),
            ("tubo_retangular", dict(B=100.0, H=60.0, t=30.0), "metade do menor lado"),
            ("I", dict(d=20.0, bf=100.0, tf=10.0, tw=5.0), "mesas"),
            ("I", dict(d=300.0, bf=100.0, tf=10.0, tw=100.0), "alma"),
            ("U", dict(d=20.0, bf=50.0, tf=12.0, tw=5.0), "mesas"),
            ("circular_macica", dict(d=0.0), "maior que zero"),
            ("circular_macica", dict(d=float("nan")), "finito"),
            ("retangular_macica", dict(b=50.0), "Faltam"),
            ("hexagonal", dict(d=10.0), "desconhecido"),
        ],
    )
    def test_geometria_impossivel_ou_incompleta_e_recusada(self, tipo, dimensoes, trecho):
        with pytest.raises(ValueError, match=trecho):
            cb.secao_por_dimensoes(tipo, **dimensoes)

    def test_informada_zero_em_w_e_z_quer_dizer_nao_informado(self):
        secao = cb.secao_informada("Minha seção", 5_000.0, 1e8, 1e7, J_mm4=2e5, Wx_mm3=6e5)
        assert secao.nome == "Minha seção" and secao.tipo == "generica"
        assert (secao.Wx, secao.Wy, secao.Zx, secao.Zy) == (6e5, None, None, None)
        assert secao.J == 2e5 and secao.elementos == [] and secao.torcao_relevante is True
        assert cb.secao_informada("  ", 5_000.0, 1e8, 1e7).nome == "Seção genérica"

    def test_informada_valida_propriedades_e_paredes(self):
        for argumentos in (
            (0.0, 1e8, 1e7),
            (5_000.0, -1.0, 1e7),
            (5_000.0, 1e8, float("inf")),
        ):
            with pytest.raises(ValueError):
                cb.secao_informada("x", *argumentos)
        with pytest.raises(ValueError, match="J"):
            cb.secao_informada("x", 5_000.0, 1e8, 1e7, J_mm4=-1.0)
        with pytest.raises(ValueError, match="centro de cisalhamento"):
            cb.secao_informada("x", 5_000.0, 1e8, 1e7, x0_mm=float("nan"))
        with pytest.raises(ValueError, match="grupo 4"):
            cb.secao_informada(
                "x", 5_000.0, 1e8, 1e7, elementos=[dict(tipo="AA", grupo=4, b=10.0, t=2.0)]
            )
        secao = cb.secao_informada(
            "x",
            5_000.0,
            1e8,
            1e7,
            elementos=[dict(tipo="AL", grupo=4, b=50.0, t=8.0, n=4)],
            torcao_relevante=False,
        )
        assert secao.elementos[0]["grupo"] == 4 and secao.torcao_relevante is False

    def test_informada_sem_paredes_nunca_fica_ok_sem_confirmacao(self):
        secao = cb.secao_informada("x", 5_000.0, 1e8, 1e7, torcao_relevante=False)
        assert cb.verificar_coluna(entrada(secao)).status_geral == "ALERTA"
        assert (
            cb.verificar_coluna(entrada(secao, secao_compacta_confirmada=True)).status_geral == "OK"
        )


# ------------------------------------------------------------------ concordância com core.nbr8800
FAMILIAS_IDENTICAS = {"W", "HP", "I", "U", "T", "Barra", "Tubo"}


class TestConcordaComNbr8800:
    """O núcleo novo e o do módulo de barras (página Estruturas de aço) calculam o mesmo."""

    @staticmethod
    def _comparar(perfil, fy, lx, ly, lz, soldado):
        sec = cb.secao_de_perfil(perfil, soldado=soldado)
        r = cb.verificar_coluna(
            entrada(sec, fy_MPa=fy, Lx_mm=lx, Ly_mm=ly, Lz_mm=lz, kz=1.0 if lz else None)
        )
        velho = nbr8800.verificar_compressao(
            perfil, fy, E, G, 1.0, lx, ly, 1.0, kz=lz if lz else None, soldado=soldado
        )
        return r.compressao, velho

    @pytest.mark.parametrize("fy", [250.0, 345.0])
    @pytest.mark.parametrize(
        ("lx", "ly", "lz"),
        [(2000.0, 2000.0, None), (4500.0, 3000.0, 4500.0), (1500.0, 6000.0, None)],
    )
    def test_compressao_identica_em_todo_o_catalogo(self, fy, lx, ly, lz):
        n = 0
        for nome, perfil in catalogo.listar_perfis().items():
            familia = perfil.familia.split()[0]
            retangular = ss.e_tubo_retangular(perfil)
            if familia not in FAMILIAS_IDENTICAS or retangular:
                continue
            for soldado in (False, True):
                novo, velho = self._comparar(perfil, fy, lx, ly, lz, soldado)
                assert novo["Ne"] == pytest.approx(velho.ne_N / 1e3, rel=1e-9), nome
                assert novo["Q"] == pytest.approx(velho.fator_q, rel=1e-9), nome
                assert novo["chi"] == pytest.approx(velho.chi, rel=1e-9), nome
                assert novo["Nc_Rd"] == pytest.approx(velho.resistencia_N / 1e3, rel=1e-9), nome
                n += 1
        assert n > 200

    def test_tubo_retangular_difere_so_pela_largura_plana(self):
        # nbr8800 usa lado − 2t; o módulo de normas usa lado − 3t (AISC B4.1b), 1% mais favorável.
        perfil = ss.CATALOGO_PERFIS["TR ideal 250×150×6"]
        novo, velho = self._comparar(perfil, 345.0, 2000.0, 2000.0, None, False)
        assert novo["Ne"] == pytest.approx(velho.ne_N / 1e3, rel=1e-9)
        assert velho.fator_q < novo["Q"] < velho.fator_q * 1.02
        # com a mesma largura plana os dois coincidem
        t, b, h = 6.0, 150.0, 250.0
        igual = cd.secao_tubo_retangular(b, h, t, b_plano=b - 2 * t, h_plano=h - 2 * t)
        r = cb.verificar_coluna(entrada(igual, fy_MPa=345.0))
        assert r.compressao["Q"] == pytest.approx(velho.fator_q, rel=1e-9)

    def test_coluna_muito_esbelta_nao_perde_a_alma_inteira(self):
        # Com σ = χ·f_y minúsculo a alma levemente esbelta não flamba localmente; o cálculo antigo
        # zerava a largura efetiva e derrubava Q para 0,5.
        perfil = catalogo.obter_perfil('U 10" x 22,77')
        novo, velho = self._comparar(perfil, 345.0, 1500.0, 6000.0, None, False)
        assert novo["Q"] == 1.0 and velho.fator_q == 1.0

    @pytest.mark.parametrize("lb", [1000.0, 3000.0, 8000.0])
    @pytest.mark.parametrize("cb_valor", [1.0, 1.3])
    def test_flexao_em_x_identica_para_i_w_e_hp(self, lb, cb_valor):
        n = 0
        for nome, perfil in catalogo.listar_perfis().items():
            if perfil.familia.split()[0] not in {"W", "HP", "I"}:
                continue
            sec = cb.secao_de_perfil(perfil)
            novo = cd.momento_resistente(sec, 345.0, "x", lb, cb_valor, "NBR8800_2008")["MRd"]
            velho = nbr8800.verificar_flexao(
                perfil, 345.0, E, G, 0.0, eixo="x", comprimento_destravado_mm=lb, cb=cb_valor
            )
            assert novo == pytest.approx(velho.resistencia_Nmm / 1e6, rel=2e-4), nome
            n += 1
        assert n > 100

    def test_flexao_em_y_da_mesa_difere_so_pelo_teto_do_m_pl(self):
        # O módulo de normas usa M_pl = Z·f_y e limita só o resultado final a 1,5·W·f_y (5.4.2.2);
        # o antigo limitava M_pl antes de interpolar. Nunca desfavorável, e < 4%.
        for nome in ("W 150 x 13,0", "W 310 x 44,5", "W 460 x 68,0"):
            perfil = catalogo.obter_perfil(nome)
            sec = cb.secao_de_perfil(perfil)
            novo = cd.momento_resistente(sec, 345.0, "y")["MRd"]
            velho = (
                nbr8800.verificar_flexao(perfil, 345.0, E, G, 0.0, eixo="y").resistencia_Nmm / 1e6
            )
            assert velho * (1 - 1e-9) <= novo <= velho * 1.04, nome

    def test_u_em_y_usa_o_modulo_elastico_da_ponta_da_mesa(self):
        perfil = catalogo.obter_perfil('U 12" x 29,76')
        novo = cd.momento_resistente(cb.secao_de_perfil(perfil), 250.0, "y")["MRd"]
        velho = nbr8800.verificar_flexao(perfil, 250.0, E, G, 0.0, eixo="y").resistencia_Nmm / 1e6
        assert novo < velho  # o antigo usava W = I/(b_f/2), que superestima o do U


# ------------------------------------------------------------------ o que a norma não cobre vira linha NÃO OK
class TestBloqueios:
    def test_tubo_circular_com_d_t_acima_de_045_e_fy(self):
        r = cb.verificar_coluna(entrada(cd.secao_tubo_circular(200, 0.5)))  # D/t = 400
        assert r.status_geral == "NÃO OK" and r.compressao is None and r.indice_interacao is None
        assert any("D/t > 0,45" in b for b in r.bloqueios)
        assert r.reprovadas and not tem_linha(r, "N_c,Rd")

    def test_alma_esbelta_na_flexao(self):
        r = cb.verificar_coluna(
            entrada(cd.secao_I(900, 200, 8, 4, soldado=True), fy_MPa=345.0, Mx_kNm=5.0)
        )
        assert linha(r, "Flexão em x não calculada").status == "NÃO OK"
        assert "Alma esbelta" in r.bloqueios[0] and r.indice_interacao is None
        assert r.status_geral == "NÃO OK"

    def test_tubo_nao_compacto_na_flexao_nao_e_estimado(self):
        sec = cd.secao_tubo_circular(200, 2.5)  # D/t = 80 > 0,07·E/f_y = 56
        r = cb.verificar_coluna(entrada(sec, Mx_kNm=2.0))
        assert linha(r, "Flexão em x não calculada").status == "NÃO OK"
        assert "não compacto" in r.bloqueios[0]
        # M_Rd informado libera o cálculo — e a linha diz que veio do usuário
        liberado = cb.verificar_coluna(entrada(sec, Mx_kNm=2.0, MRd_x_informado_kNm=30.0))
        assert linha(liberado, "informado pelo usuário").resistente == 30.0
        assert liberado.indice_interacao is not None and liberado.status_geral == "OK"

    def test_tubo_retangular_nao_compacto(self):
        r = cb.verificar_coluna(entrada(cd.secao_tubo_retangular(200, 200, 3), Mx_kNm=2.0))
        assert r.status_geral == "NÃO OK" and "não compacto" in r.bloqueios[0]

    def test_nsd_acima_de_ne_no_b1(self):
        sec = cd.secao_circular_macica(50)
        r = cb.verificar_coluna(entrada(sec, N_Sd_kN=200.0, Mx_kNm=1.0))
        assert linha(r, "Amplificação B₁ em x").status == "NÃO OK"
        assert "N_Sd1 ≥ N_e" in r.bloqueios[0] and r.indice_interacao is None
        assert r.status_geral == "NÃO OK"

    def test_secao_sem_z(self):
        sec = cd.secao_generica("sem Z", 5000.0, 1e8, 1e8, Wx=1e6, Wy=1e6, torcao_relevante=False)
        r = cb.verificar_coluna(entrada(sec, Mx_kNm=2.0, secao_compacta_confirmada=True))
        assert "Z" in r.bloqueios[0] and r.status_geral == "NÃO OK"

    def test_secao_sem_rotina_de_mrd(self):
        r = cb.verificar_coluna(entrada(secao_generica_com_mrd(), Mx_kNm=2.0))
        assert "sem rotina de M_Rd" in r.bloqueios[0]
        informado = cb.verificar_coluna(
            entrada(
                secao_generica_com_mrd(),
                Mx_kNm=2.0,
                MRd_x_informado_kNm=120.0,
                secao_compacta_confirmada=True,
            )
        )
        assert informado.status_geral == "OK"

    def test_o_bloqueio_de_y_nao_apaga_o_calculo_de_x(self):
        sec = cd.secao_I(300, 150, 12.5, 8.0)
        r = cb.verificar_coluna(entrada(sec, Mx_kNm=20.0, My_kNm=5.0, N_Sd_kN=50.0))
        assert r.indice_interacao is not None  # I tem rotina para x e para y

    def test_entradas_invalidas(self):
        sec = cd.secao_circular_macica(50)
        for mudanca in (
            {"norma": "XYZ"},
            {"fy_MPa": 0.0},
            {"Lx_mm": -1.0},
            {"kx": 0.0},
            {"Cb": 0.9},
            {"N_Sd_kN": -1.0},
            {"Mx_kNm": -1.0},
            {"razao_m1_m2_x": 1.5},
            {"espessuras_anglo": {"inexistente": 5.0}},
            {"momentos_cb_kNm": (0.0, 1.0, 1.0, 1.0)},
        ):
            with pytest.raises(ValueError):
                cb.verificar_coluna(entrada(sec, **mudanca))


class TestSecaoGenericaNuncaOK:
    def test_area_e_raio_diretos_sem_confirmacao_e_alerta(self):
        sec = cb.secao_direta(1_963.5, 12.5)
        r = cb.verificar_coluna(entrada(sec, N_Sd_kN=50.0))
        assert linha(r, "Flambagem local não verificada").status == "ALERTA"
        assert linha(r, "Torção e flexo-torção não verificadas").status == "ALERTA"
        assert r.status_geral == "ALERTA" and r.atende

    def test_confirmada_pode_ser_ok(self):
        sec = cb.secao_direta(1_963.5, 12.5)
        r = cb.verificar_coluna(entrada(sec, N_Sd_kN=50.0, secao_compacta_confirmada=True))
        assert r.status_geral == "OK" and linha(r, "confirmado").status == "INFO"

    def test_com_b_t_e_sem_torcao_so_a_torcao_alerta(self):
        sec = cd.secao_generica(
            "com paredes", 5000.0, 1e8, 1e8, elementos=[dict(tipo="AA", grupo=2, b=200.0, t=10.0, n=1)],
            torcao_relevante=False,
        )  # fmt: skip
        r = cb.verificar_coluna(entrada(sec))
        assert not tem_linha(r, "Flambagem local não verificada")
        assert linha(r, "Torção e flexo-torção não verificadas").status == "ALERTA"

    def test_perfis_do_catalogo_conhecidos_nao_geram_alerta_de_secao(self):
        r = cb.verificar_coluna(entrada(cb.secao_de_perfil(catalogo.obter_perfil("W 310 x 44,5"))))
        assert r.status_geral == "OK" and not tem_linha(r, "não verificada")


# ------------------------------------------------------------------ três normas
class TestNormas:
    def test_secao_compacta_2008_igual_ao_projeto_2024_e_aisc_99_por_cento(self):
        sec = cd.secao_I(300, 150, 12.5, 8.0)
        r08, r24, raisc = (
            cb.verificar_coluna(entrada(sec, Lx_mm=4000.0, Ly_mm=4000.0, norma=n))
            for n in cb.NORMAS
        )
        assert r08.compressao["Q"] == 1.0
        assert r08.compressao["Nc_Rd"] == pytest.approx(r24.compressao["Nc_Rd"])
        assert raisc.compressao["Nc_Rd"] / r08.compressao["Nc_Rd"] == pytest.approx(0.99, abs=1e-6)

    def test_secao_esbelta_reduz_a_resistencia_nas_duas_formas(self):
        sec = cd.secao_I(600, 200, 6.3, 5.0, soldado=True)
        for norma in cb.NORMAS:
            comp = cb.verificar_coluna(
                entrada(sec, fy_MPa=345.0, Lx_mm=3000.0, Ly_mm=3000.0, norma=norma)
            ).compressao
            assert comp["A_ef"] < sec.A
            assert comp["Nc_Rd"] < comp["chi"] * sec.A * 345.0 / 1.10 / 1000
        r = cb.verificar_coluna(entrada(sec, fy_MPa=345.0, Lx_mm=3000.0, Ly_mm=3000.0))
        assert any("Flambagem local reduz" in a for a in r.avisos)

    def test_linha_de_flambagem_local_muda_com_a_norma(self):
        sec = cd.secao_I(300, 150, 12.5, 8.0)
        assert tem_linha(cb.verificar_coluna(entrada(sec)), "fator Q = Q_s·Q_a")
        assert tem_linha(
            cb.verificar_coluna(entrada(sec, norma="NBR8800_2024")), "área efetiva A_ef"
        )

    def test_aviso_de_k_so_nas_normas_novas(self):
        sec = cd.secao_I(300, 150, 12.5, 8.0)
        com_k = dict(kx=0.8, ky=0.8)
        assert not tem_linha(cb.verificar_coluna(entrada(sec, **com_k)), "K diferente de 1,0")
        assert not cb.verificar_coluna(entrada(sec, **com_k)).avisos
        for norma in ("NBR8800_2024", "AISC360_16_LRFD"):
            resultado = cb.verificar_coluna(entrada(sec, norma=norma, **com_k))
            v = linha(resultado, "K diferente de 1,0")
            assert v.status == "INFO" and "análise direta" in v.formula
            assert any(a.startswith("K ≠ 1,0") and "análise direta" in a for a in resultado.avisos)
        assert not tem_linha(
            cb.verificar_coluna(entrada(sec, norma="AISC360_16_LRFD")), "K diferente"
        )

    def test_comparar_normas_traz_as_tres_e_nao_esconde_bloqueio(self):
        sec = cd.secao_I(300, 150, 12.5, 8.0)
        tabela = cb.comparar_normas(entrada(sec, Mx_kNm=20.0, N_Sd_kN=100.0))
        assert [x["norma"] for x in tabela] == list(cb.NORMAS)
        assert all(x["Nc_Rd_kN"] > 0 and x["MRd_x_kNm"] > 0 for x in tabela)
        assert tabela[2]["Nc_Rd_kN"] < tabela[0]["Nc_Rd_kN"]  # φ = 0,90 < 1/1,10
        bloqueado = cb.comparar_normas(entrada(cd.secao_tubo_circular(200, 0.5)))
        assert all(x["status"] == "NÃO OK" and "D/t > 0,45" in x["erro"] for x in bloqueado)

    def test_o_aisc_usa_phi_090(self):
        sec = cd.secao_circular_macica(50)
        r = cb.verificar_coluna(entrada(sec, norma="AISC360_16_LRFD", N_Sd_kN=50.0))
        assert "φ" in linha(r, "N_c,Rd").formula


class TestRevisao:
    """Defeitos achados na revisão do código: cada um com o caso que o mostrava."""

    def test_b1_bloqueado_nao_passa_por_resultado(self):
        sec = cd.secao_circular_macica(50)
        e = entrada(sec, N_Sd_kN=200.0, Mx_kNm=1.0)  # N_ex = 151,4 kN
        r = cb.verificar_coluna(e)
        assert math.isinf(r.b1_x) and math.isinf(r.momento_x_amplificado_kNm)
        assert r.indice_interacao is None and r.mrd_x_kNm is None
        resultados = cb.registro_coluna(e, r)["resultados"]
        assert resultados["b1_x"] == "∞" and resultados["momento_x_amplificado_kNm"] == "∞"
        assert "b1_y" not in resultados  # sem momento em y

    def test_mrd_informado_vai_ao_registro_e_ao_resultado(self):
        sec = cd.secao_tubo_circular(200, 2.5)  # não compacto: a norma não estima M_Rd
        e = entrada(sec, N_Sd_kN=10.0, Mx_kNm=2.0, MRd_x_informado_kNm=30.0)
        r = cb.verificar_coluna(e)
        assert r.flexao_x is None and r.mrd_x_kNm == 30.0 and r.mrd_y_kNm is None
        registro = cb.registro_coluna(e, r)
        assert registro["entradas"]["MRd_x_informado_kNm"] == 30.0
        assert registro["resultados"]["momento_resistente_x_kNm"] == 30.0
        da_norma = cb.verificar_coluna(
            entrada(cd.secao_I(300, 150, 12.5, 8.0), N_Sd_kN=10.0, Mx_kNm=20.0)
        )
        assert da_norma.flexao_x is not None
        assert da_norma.mrd_x_kNm == pytest.approx(da_norma.flexao_x["MRd"])

    def test_retangular_deitada_tem_flt_no_eixo_forte_y_com_lx(self):
        em_pe = cd.secao_retangular_macica(50, 100)
        deitada = cd.secao_retangular_macica(100, 50)
        # a flexão em y da deitada trava onde se impede a flexão em x (L_x); a da peça em pé, em L_y
        a = cb.verificar_coluna(entrada(em_pe, Lx_mm=1500.0, Ly_mm=4000.0, Mx_kNm=5.0))
        b = cb.verificar_coluna(entrada(deitada, Lx_mm=4000.0, Ly_mm=1500.0, My_kNm=5.0))
        assert a.flexao_x is not None and b.flexao_y is not None
        assert "FLT" in a.flexao_x["estados"] and "FLT" in b.flexao_y["estados"]
        assert b.mrd_y_kNm == pytest.approx(a.mrd_x_kNm, rel=1e-9)
        assert b.indice_interacao == pytest.approx(a.indice_interacao, rel=0.02)


# ------------------------------------------------------------------ curva, interação e avisos de entrada
class TestComportamentoDaCurvaEAvisos:
    def test_coluna_mais_longa_resiste_menos_e_o_balanco_equivale_ao_dobro_do_comprimento(self):
        sec = cd.secao_I(300, 150, 12.5, 8.0)

        def n_rd(**mudancas) -> float:
            resultado = cb.verificar_coluna(entrada(sec, **mudancas)).compressao
            assert resultado is not None
            return resultado["Nc_Rd"]

        curta = n_rd(Lx_mm=2000.0, Ly_mm=2000.0)
        longa = n_rd(Lx_mm=4000.0, Ly_mm=4000.0)
        balanco = n_rd(Lx_mm=2000.0, Ly_mm=2000.0, kx=2.0, ky=2.0)
        assert longa < curta
        assert balanco < curta
        assert balanco == pytest.approx(longa, rel=1e-12)  # K·L = 4000 mm nos dois casos

    def test_ramos_da_curva_chi_e_perda_em_relacao_a_euler(self):
        sec = cd.secao_circular_macica(50)
        curta = cb.verificar_coluna(entrada(sec, Lx_mm=500.0, Ly_mm=500.0)).compressao
        longa = cb.verificar_coluna(entrada(sec, Lx_mm=2000.0, Ly_mm=2000.0)).compressao
        assert curta is not None and longa is not None
        assert curta["lambda0"] <= 1.5 < longa["lambda0"]
        assert curta["chi"] == pytest.approx(0.658 ** (curta["lambda0"] ** 2))
        assert longa["chi"] == pytest.approx(0.877 / longa["lambda0"] ** 2)
        # Mesmo na coluna longa a norma dá 0,877·N_e/γ_a1 (80% de Euler): imperfeições e tensões
        # residuais.
        assert longa["Nc_Rd"] == pytest.approx(0.877 * longa["Ne"] / 1.10, rel=1e-12)

    def test_eixo_governante_e_o_mais_esbelto(self):
        sec = cd.secao_I(300, 150, 12.5, 8.0)  # r_y < r_x
        igual = entrada(sec, Lx_mm=3000.0, Ly_mm=3000.0)
        travado_em_y = entrada(sec, Lx_mm=3000.0, Ly_mm=800.0)
        eixos = {}
        for nome, e in (("igual", igual), ("travado_em_y", travado_em_y)):
            registro = cb.registro_coluna(e, cb.verificar_coluna(e))
            eixos[nome] = registro["resultados"]["eixo_governante"]
        assert eixos == {"igual": "y", "travado_em_y": "x"}

    def test_ramos_da_equacao_de_interacao_no_nivel_da_pagina(self):
        sec = secao_generica_com_mrd()
        base = entrada(
            sec,
            Lx_mm=1000.0,
            Ly_mm=1000.0,
            secao_compacta_confirmada=True,
            MRd_x_informado_kNm=100.0,
            Mx_kNm=50.0,
        )
        n_rd = cb.verificar_coluna(replace(base, N_Sd_kN=1.0, Mx_kNm=0.0)).compressao["Nc_Rd"]
        pequeno = cb.verificar_coluna(replace(base, N_Sd_kN=0.1 * n_rd))
        medio = cb.verificar_coluna(replace(base, N_Sd_kN=0.5 * n_rd))
        # N/N_Rd < 0,2:  N/(2 N_Rd) + M/M_Rd;  acima:  N/N_Rd + 8/9·M/M_Rd.
        assert pequeno.indice_interacao == pytest.approx(
            0.05 + pequeno.momento_x_amplificado_kNm / 100.0, rel=1e-9
        )
        assert medio.indice_interacao == pytest.approx(
            0.5 + 8.0 / 9.0 * medio.momento_x_amplificado_kNm / 100.0, rel=1e-9
        )

    def test_avisos_de_k_fora_da_faixa_e_de_unidades_implausiveis(self):
        sec = cd.secao_circular_macica(50)
        assert not cb.verificar_coluna(entrada(sec)).avisos
        assert any(
            "fora da faixa física" in a for a in cb.verificar_coluna(entrada(sec, kx=5.0)).avisos
        )
        assert any("GPa" in a for a in cb.verificar_coluna(entrada(sec, E_MPa=200.0)).avisos)
        assert any("f_y/E" in a for a in cb.verificar_coluna(entrada(sec, fy_MPa=30_000.0)).avisos)


# ------------------------------------------------------------------ B₁, C_m, C_b
class TestAmplificacaoECb:
    def test_b1_usa_o_comprimento_real_e_k_igual_a_1(self):
        sec = cd.secao_I(300, 150, 12.5, 8.0)
        a = cb.verificar_coluna(
            entrada(sec, kx=1.0, ky=1.0, Lx_mm=4000.0, N_Sd_kN=200.0, Mx_kNm=20.0)
        )
        b = cb.verificar_coluna(
            entrada(sec, kx=2.0, ky=1.0, Lx_mm=4000.0, N_Sd_kN=200.0, Mx_kNm=20.0)
        )
        ne_real = math.pi**2 * E * sec.Ix / 4000.0**2 / 1000.0
        assert a.b1_x == pytest.approx(1.0 / (1.0 - 200.0 / ne_real))
        assert b.b1_x == pytest.approx(a.b1_x)  # K não entra no B₁
        assert b.compressao["Nc_Rd"] < a.compressao["Nc_Rd"]  # mas entra em N_c,Rd
        assert a.momento_x_amplificado_kNm == pytest.approx(20.0 * a.b1_x)

    def test_cm_pela_razao_m1_m2_e_forcas_transversais(self):
        sec = cd.secao_I(300, 150, 12.5, 8.0)
        base = entrada(sec, N_Sd_kN=100.0, Mx_kNm=20.0, Lx_mm=4000.0)
        assert cb.verificar_coluna(base).cm_x == 1.0  # sem M₁/M₂: conservador
        reversa = cb.verificar_coluna(replace(base, razao_m1_m2_x=0.5))
        assert reversa.cm_x == pytest.approx(0.60 - 0.40 * 0.5)
        curvatura_simples = cb.verificar_coluna(replace(base, razao_m1_m2_x=-1.0))
        assert curvatura_simples.cm_x == pytest.approx(1.0)
        transversal = cb.verificar_coluna(
            replace(base, razao_m1_m2_x=0.5, forcas_transversais_x=True)
        )
        assert transversal.cm_x == 1.0
        assert reversa.b1_x == 1.0  # B₁ ≥ 1 mesmo com C_m pequeno
        assert linha(transversal, "Cisalhamento").status == "ALERTA"

    def test_b1_so_aparece_para_o_eixo_com_momento(self):
        sec = cd.secao_I(300, 150, 12.5, 8.0)
        r = cb.verificar_coluna(entrada(sec, N_Sd_kN=100.0, My_kNm=5.0))
        assert tem_linha(r, "B₁y") and not tem_linha(r, "B₁x")

    def test_excentricidade_vira_momento_amplificado(self):
        sec = cd.secao_circular_macica(50)
        r = cb.verificar_coluna(entrada(sec, N_Sd_kN=50.0, ex_mm=5.0))
        assert r.momento_x_primeira_ordem_kNm == pytest.approx(0.25)
        ne_real = math.pi**2 * E * sec.Ix / 2000.0**2 / 1000.0
        assert r.b1_x == pytest.approx(1.0 / (1.0 - 50.0 / ne_real))
        assert r.indice_interacao is not None and "Interação" in r.governante

    def test_cb_pelo_diagrama_de_momentos_e_teto_de_3(self):
        sec = cd.secao_I(300, 150, 12.5, 8.0)
        base = entrada(sec, N_Sd_kN=50.0, Mx_kNm=20.0, Lb_mm=3000.0)
        com = cb.verificar_coluna(replace(base, momentos_cb_kNm=(20.0, 10.0, 15.0, 10.0)))
        esperado = min(3.0, cd.fator_Cb(20.0, 10.0, 15.0, 10.0))
        assert com.cb_usado == pytest.approx(esperado) and com.lb_usado_mm == 3000.0
        assert (
            cb.verificar_coluna(replace(base, momentos_cb_kNm=(20.0, 0.0, 0.0, 0.0))).cb_usado
            == 3.0
        )

    def test_lb_padrao_e_ly(self):
        sec = cd.secao_I(300, 150, 12.5, 8.0)
        r = cb.verificar_coluna(entrada(sec, Ly_mm=2500.0, Mx_kNm=10.0))
        assert r.lb_usado_mm == 2500.0

    def test_u_na_nbr_2008_toma_cb_igual_a_1(self):
        u = cd.secao_U(152.4, 48.8, 8.71, 5.08)
        r = cb.verificar_coluna(entrada(u, Mx_kNm=3.0, Cb=1.3, Lb_mm=3000.0))
        assert r.cb_usado == 1.0 and any("Seção U: C_b" in a for a in r.avisos)
        livre = cb.verificar_coluna(
            entrada(u, Mx_kNm=3.0, Cb=1.3, Lb_mm=3000.0, norma="NBR8800_2024")
        )
        assert livre.cb_usado == 1.3

    def test_flexao_mostra_flt_flm_fla_e_o_limite(self):
        sec = cd.secao_I(300, 150, 12.5, 8.0)
        r = cb.verificar_coluna(entrada(sec, Mx_kNm=40.0, N_Sd_kN=50.0, Lb_mm=6000.0))
        for estado in ("FLT", "FLM", "FLA"):
            assert tem_linha(r, f"Mx,Rd — {estado}")
        assert tem_linha(r, "limite 1,5·W·f_y")
        assert "FLT" in linha(r, "Flexão em torno de x").formula


# ------------------------------------------------------------------ tabela de verificações, elementos e registro
class TestTabelaERegistro:
    def test_toda_linha_tem_referencia_e_status_valido(self):
        sec = cd.secao_I(300, 150, 12.5, 8.0)
        r = cb.verificar_coluna(
            entrada(
                sec,
                N_Sd_kN=100.0,
                Mx_kNm=20.0,
                My_kNm=3.0,
                espessuras_anglo={"perfil_laminado_H_W": 8.0},
            )
        )
        assert len(r.verificacoes) >= 20
        for v in r.verificacoes:
            assert v.status in {"OK", "NÃO OK", "ALERTA", "INFO", "N/A"} and v.referencia and v.nome
            assert v.tipo in {"resistencia", "limite", "informativo"}
            if v.status == "INFO":
                assert v.aproveitamento is None  # informação nunca entra no aproveitamento

    def test_aproveitamento_maximo_so_conta_resistencia(self):
        sec = cb.secao_direta(1_963.5, 12.5)
        r = cb.verificar_coluna(entrada(sec, N_Sd_kN=50.0, Lx_mm=2600.0, Ly_mm=2600.0))  # λ = 208
        assert r.status_geral == "NÃO OK" and r.aproveitamento_max < 1.0

    def test_tabela_dos_elementos_da_parede(self):
        sec = cd.secao_I(600, 200, 6.3, 5.0, soldado=True)
        comp = cb.verificar_coluna(
            entrada(sec, fy_MPa=345.0, Lx_mm=3000.0, Ly_mm=3000.0)
        ).compressao
        linhas = cb.tabela_elementos(sec, 345.0, "NBR8800_2008", comp)
        assert [x["Tipo"] for x in linhas] == ["AL", "AA"]
        assert linhas[0]["Elemento"] == "mesa de perfil soldado"
        alma = linhas[1]
        assert alma["Esbelto"] == "sim" and alma["b_ef (mm)"] < alma["b (mm)"]
        assert linhas[0]["Q_s"] is not None and alma["Q_s"] is None

    def test_registro_mantem_as_chaves_dos_graficos_e_leva_a_tabela(self):
        sec = cd.secao_circular_macica(50)
        e = entrada(sec, N_Sd_kN=72.0)
        r = cb.verificar_coluna(e)
        registro = cb.registro_coluna(
            e, r, contexto={"condicao_apoio": "Biapoiada"}, materiais_ids=["m1"]
        )
        resultados, entradas = registro["resultados"], registro["entradas"]
        for chave in (
            "esbeltez_governante",
            "chi",
            "lambda_0",
            "fator_q",
            "resistencia_kN",
            "utilizacao",
            "modo_governante",
        ):
            assert chave in resultados, chave
        for chave in (
            "modulo_elasticidade_MPa",
            "escoamento_MPa",
            "area_mm2",
            "forca_solicitante_kN",
            "gamma_a1",
        ):
            assert chave in entradas, chave
        assert entradas["condicao_apoio"] == "Biapoiada" and registro["materiais_ids"] == ["m1"]
        assert perto(resultados["resistencia_kN"], 120.71, 1e-3) and perto(
            resultados["utilizacao_maxima"], 0.596, 3e-3
        )
        tabela = resultados["verificações"]
        assert len(tabela) == len(r.verificacoes) and set(tabela[0]) == {
            "verificação", "solicitante", "resistente", "unidade",
            "aproveitamento_pct", "status", "fórmula", "referência",
        }  # fmt: skip
        assert registro["status"] == "Atende" and registro["modulo_id"] == "flambagem_colunas"

    def test_registro_sobrevive_ao_hash_com_infinito_e_bloqueio(self):
        e = entrada(cd.secao_tubo_circular(200, 0.5))
        registro = cb.registro_coluna(e, cb.verificar_coluna(e))
        assert registro["status"] == "Não atende" and registro["hash_calculo"]
        assert (
            registro["resultados"]["bloqueios"] and "resistencia_kN" not in registro["resultados"]
        )
        reprovada = entrada(
            cd.secao_circular_macica(20), Lx_mm=1500.0, Ly_mm=1500.0, N_Sd_kN=0.1
        )  # λ = 300
        registro = cb.registro_coluna(reprovada, cb.verificar_coluna(reprovada))
        assert registro["status"] == "Não atende" and registro["hash_calculo"]

    def test_o_grafico_da_curva_de_flambagem_le_o_registro_novo(self):
        from core.record_charts import imagens_flambagem

        for norma in cb.NORMAS:
            e = entrada(cd.secao_I(300, 150, 12.5, 8.0), norma=norma, N_Sd_kN=100.0)
            registro = cb.registro_coluna(e, cb.verificar_coluna(e))
            imagens = imagens_flambagem(registro["entradas"], registro["resultados"])
            assert len(imagens) == 1 and imagens[0].png.startswith(b"\x89PNG"), norma

    def test_memorial_word_e_pdf_imprimem_a_tabela(self):
        from io import BytesIO

        from docx import Document
        from pypdf import PdfReader

        from core.project_report import (
            gerar_relatorio_industrial_pdf,
            gerar_relatorio_industrial_word,
        )
        from core.technical_records import normalizar_registro_tecnico
        from tests.test_project_validation import _projeto_documentado

        e = entrada(
            cd.secao_I(300, 150, 12.5, 8.0),
            N_Sd_kN=100.0,
            Mx_kNm=20.0,
            espessuras_anglo={"perfil_laminado_H_W": 4.0},
        )
        r = cb.verificar_coluna(e)
        projeto = _projeto_documentado()
        projeto["registros_tecnicos"] = [normalizar_registro_tecnico(cb.registro_coluna(e, r))]
        pdf = gerar_relatorio_industrial_pdf(projeto, secoes_incluidas=["registros"])
        texto = "\n".join(p.extract_text() for p in PdfReader(BytesIO(pdf)).pages)
        for trecho in (
            "Compress",
            "Interação",
            "Anglo 8.8",
            "aproveitamento",
            "Espessura mínima",
            "O que passou",
            "Não passou",
        ):
            assert trecho in texto, trecho
        word = Document(
            BytesIO(gerar_relatorio_industrial_word(projeto, secoes_incluidas=["registros"]))
        )
        celulas = " ".join(c.text for t in word.tables for linha_ in t.rows for c in linha_.cells)
        paragrafos = "\n".join(p.text for p in word.paragraphs)
        assert "Compressão axial N_c,Rd" in celulas
        # A espessura de 4,0 mm reprova o critério Anglo: o capítulo diz, no fim, o que não passou.
        assert "Resultado: NÃO ATENDE." in paragrafos
        assert paragrafos.index("O que passou") < paragrafos.index("Não passou (")

    def test_sensibilidade_le_o_registro_novo(self):
        from core.sensitivity import sugerir_de_registro

        e = entrada(
            cd.secao_I(300, 150, 12.5, 8.0), Lx_mm=4000.0, Ly_mm=2500.0, ky=0.8, N_Sd_kN=100.0
        )
        registro = cb.registro_coluna(e, cb.verificar_coluna(e))
        modelo = sugerir_de_registro(registro)
        assert modelo is not None and modelo["modelo_id"] == "carga_critica_euler"
        entradas = modelo["entradas"]
        assert entradas["L_mm"] == 2500.0 and entradas["K"] == 0.8  # y governa (λ maior)
        assert entradas["I_mm4"] == pytest.approx(e.secao.Iy) and entradas["E_GPa"] == 200.0
        # Registro no formato antigo (um comprimento só, raio de giração) continua sendo lido.
        antigo = {
            "modulo_id": "flambagem_colunas",
            "entradas": {
                "area_mm2": 1_000.0,
                "raio_giracao_x_mm": 20.0,
                "modulo_elasticidade_MPa": 200_000.0,
                "kx": 1.0,
                "comprimento_mm": 3_000.0,
            },
            "resultados": {"eixo_governante": "x"},
        }
        lido = sugerir_de_registro(antigo)
        assert lido is not None
        assert (
            lido["entradas"]["I_mm4"] == pytest.approx(4e5) and lido["entradas"]["L_mm"] == 3_000.0
        )
