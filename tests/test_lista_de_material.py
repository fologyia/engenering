"""Lista de material: massa, pintura, barras, lista de corte do SolidWorks e o PP do modelo."""

from __future__ import annotations

import io
import json
import math

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from core import lista_de_material as lm
from core import placa_base_pilares as pb
from core import plano_de_cargas as pc
from core import section_catalog as sc
from tests.test_esforcos_modelo import PILAR, importar
from tests.test_quadro_fundacoes import PILAR_2, portico

ACO = 7850e-6  # kg/m por mm² de seção


@pytest.mark.parametrize(
    ("descricao", "nome"),
    [
        ("W 200 x 35,9 (H)", "W 200 x 35,9 (H)"),
        ("W 200 X 35.9", "W 200 x 35,9 (H)"),
        ("W8X31", "W 200 x 46,1 (H)"),
        ("C8X13.75", 'U 8" x 20,50'),
        ("L2X2X1/4", 'L 2" × 1/4"'),
        ("L 2 1/2 x 2 1/2 x 1/4", 'L 2 1/2" × 1/4"'),
        ("Ângulo l L2.5X2.5X0.25", 'L 2 1/2" × 1/4"'),
        # Nomes dos itens da lista de corte na árvore do SolidWorks.
        ("C8x13.75<1>", 'U 8" x 20,50'),
        ("C6x8.2<1>(4)", 'U 6" x 12,20'),
        ("L2.5x2.5x0.25<3>", 'L 2 1/2" × 1/4"'),
        ("W8x31<1>(2)", "W 200 x 46,1 (H)"),
    ],
)
def test_perfil_pelo_nome_do_catalogo_ou_do_solidworks(descricao, nome):
    assert lm.perfil_conhecido(descricao).nome == nome


@pytest.mark.parametrize(
    ("descricao", "area_mm2", "perimetro_mm"),
    [
        ("TUBO QUADRADO 50 X 50 X 3", 2 * 3 * (50 + 50 - 6), 200.0),
        ("Tubo retangular 100x50x3,0", 2 * 3 * (100 + 50 - 6), 300.0),
        ("TUBO REDONDO 60,3 X 3,6", math.pi * 3.6 * (60.3 - 3.6), math.pi * 60.3),
        ("BARRA REDONDA 16", math.pi * 16**2 / 4, math.pi * 16),
        ("CANTONEIRA 50 X 50 X 5", 5 * (50 + 50 - 5), 200.0),
    ],
)
def test_tubos_barras_e_cantoneiras_pelas_medidas(descricao, area_mm2, perimetro_mm):
    p = lm.perfil_conhecido(descricao)
    assert p.massa_kg_m == pytest.approx(area_mm2 * ACO)
    assert p.perimetro_mm == pytest.approx(perimetro_mm)


def test_descricao_desconhecida_nao_vira_perfil():
    assert lm.perfil_conhecido("PERFIL ESQUISITO") is None
    assert lm.perfil_conhecido("") is None


def test_perimetro_de_pintura_por_familia():
    w = lm.perfil_conhecido("W 200 x 35,9 (H)")
    assert w.perimetro_mm == pytest.approx(2 * 201 + 4 * 165 - 2 * 6.2)  # faces das mesas e da alma
    perfis = sc.listar_perfis()
    tubo = next(n for n, p in perfis.items() if p.familia == "Tubo circular")
    assert lm.perfil_conhecido(tubo).perimetro_mm == pytest.approx(math.pi * perfis[tubo].altura_mm)
    te = next(n for n, p in perfis.items() if p.familia.startswith("T ("))
    assert lm.perfil_conhecido(te).perimetro_mm == pytest.approx(
        2 * perfis[te].altura_mm + 2 * perfis[te].largura_mm
    )


def test_massa_e_pintura_de_cada_tipo_de_item():
    perfil = lm.calcular_linha(lm.ItemDaLista("P1", lm.TIPO_PERFIL, "W 200 x 35,9 (H)", 2, 4.0))
    assert perfil.massa_kg == pytest.approx(2 * 4.0 * 35.9)
    assert perfil.area_pintura_m2 == pytest.approx(2 * 4.0 * 1.0496)
    chapa = lm.calcular_linha(lm.ItemDaLista("PB", lm.TIPO_CHAPA, "Placa", 2, 0.35, 300, 19))
    assert chapa.massa_kg == pytest.approx(2 * 0.35 * 0.3 * 0.019 * 7850)
    assert chapa.area_pintura_m2 == pytest.approx(2 * 2 * 0.35 * 0.3)
    grade = lm.calcular_linha(
        lm.ItemDaLista("G1", lm.TIPO_GRADE, "Grade 30x3", 3, 2.0, 1000, massa_unitaria=24.0)
    )
    assert grade.massa_kg == pytest.approx(3 * 2.0 * 1.0 * 24.0) and grade.area_pintura_m2 == 0
    outro = lm.calcular_linha(
        lm.ItemDaLista("CH", lm.TIPO_OUTRO, "Chumbador 3/4", 8, massa_unitaria=1.6)
    )
    assert outro.massa_kg == pytest.approx(12.8)
    fora = lm.calcular_linha(
        lm.ItemDaLista("X", lm.TIPO_PERFIL, "PERFIL ESPECIAL", 1, 3.0, massa_unitaria=10.0)
    )
    assert fora.massa_kg == pytest.approx(30.0) and "pintura" in fora.pendencia


@pytest.mark.parametrize(
    ("item", "trecho"),
    [
        (lm.ItemDaLista("1", lm.TIPO_PERFIL, "W 200 x 35,9 (H)", 1, 0.0), "comprimento"),
        (lm.ItemDaLista("2", lm.TIPO_PERFIL, "PERFIL ESQUISITO", 1, 3.0), "kg/m"),
        (lm.ItemDaLista("3", lm.TIPO_CHAPA, "Chapa 12,5", 1, 1.0, 0, 12.5), "largura"),
        (lm.ItemDaLista("4", lm.TIPO_GRADE, "Grade", 1, 1.0, 1000), "kg/m²"),
        (lm.ItemDaLista("5", lm.TIPO_OUTRO, "Parafuso", 10), "massa por unidade"),
    ],
)
def test_item_sem_dado_fica_pendente(item, trecho):
    linha = lm.calcular_linha(item)
    assert linha.massa_kg is None and trecho in linha.pendencia


def test_resumo_com_acrescimo_barras_e_chapas():
    lista = lm.ListaDeMaterial(
        itens=(
            lm.ItemDaLista("P1", lm.TIPO_PERFIL, "W 200 x 35,9 (H)", 2, 4.0),
            lm.ItemDaLista("P2", lm.TIPO_PERFIL, "W 200 x 35,9 (H)", 2, 4.5),
            lm.ItemDaLista("V1", lm.TIPO_PERFIL, "W 310 x 32,7", 1, 6.0),
            lm.ItemDaLista("PB", lm.TIPO_CHAPA, "Placa", 2, 0.35, 300, 19),
            lm.ItemDaLista("CH", lm.TIPO_CHAPA, "Chapa de ligação", 4, 0.2, 150, 9.5),
            lm.ItemDaLista("?", lm.TIPO_OUTRO, "sem massa", 1),
        ),
        acrescimo_pct=10.0,
    )
    r = lm.resumir(lista)
    itens = 2 * 4 * 35.9 + 2 * 4.5 * 35.9 + 6 * 32.7 + 2 * 0.35 * 0.3 * 0.019 * 7850
    itens += 4 * 0.2 * 0.15 * 0.0095 * 7850
    assert r.massa_itens_kg == pytest.approx(itens)
    assert r.massa_total_kg == pytest.approx(1.1 * itens)
    assert r.peso_total_kN == pytest.approx(1.1 * itens * 9.80665 / 1e3)
    w200 = next(p for p in r.perfis if p.perfil == "W 200 x 35,9 (H)")
    assert w200.pecas == 4 and w200.comprimento_m == pytest.approx(17.0) and w200.barras == 2
    assert [c.espessura_mm for c in r.chapas] == [9.5, 19.0]
    assert len(r.pendentes) == 1
    assert r.por_tipo[lm.TIPO_CHAPA] == pytest.approx(sum(c.massa_kg for c in r.chapas))


@pytest.mark.parametrize(
    ("texto", "unidade", "esperado"),
    [
        ("1234,5", "mm", (1234.5, "mm")),
        ("1.234,5", "mm", (1234.5, "mm")),
        ("1234.5", "mm", (1234.5, "mm")),
        ("6.000", "mm", (6000.0, "mm")),
        ("1.250", "m", (1.25, "m")),
        ("1 234,5", "mm", (1234.5, "mm")),
        ("2,5 m", "mm", (2.5, "m")),
        ("157.48 in", "mm", (157.48, "in")),
        ("abc", "mm", (None, "mm")),
    ],
)
def test_numeros_da_lista_de_corte(texto, unidade, esperado):
    assert lm.numero_da_celula(texto, unidade) == esperado


CSV_PT = (
    "Nº DO ITEM;QTD.;DESCRIÇÃO;COMPRIMENTO\r\n"
    "1;2;W 200 X 35.9;4000,00\r\n"
    "2;4;C8X13.75;6.000,0\r\n"
    "3;8;TUBO QUADRADO 50 X 50 X 3;1250\r\n"
    "4;4;PLACA 19;\r\n"
    "5;1;PERFIL ESQUISITO;3000\r\n"
    ";;;\r\n"
).encode("cp1252")


def test_lista_de_corte_do_solidworks_em_portugues():
    itens, avisos = lm.ler_lista_de_corte(CSV_PT, "lista.csv")
    assert [i.descricao for i in itens] == [
        "W 200 x 35,9 (H)",
        'U 8" x 20,50',
        "TUBO QUADRADO 50 X 50 X 3",
        "PLACA 19",
        "PERFIL ESQUISITO",
    ]
    assert [i.comprimento_m for i in itens] == [4.0, 6.0, 1.25, 0.0, 3.0]
    assert [i.quantidade for i in itens] == [2, 4, 8, 4, 1]
    assert itens[3].tipo == lm.TIPO_CHAPA and itens[3].espessura_mm == 19.0
    assert itens[0].marca == "1"
    assert any("PERFIL ESQUISITO" in a for a in avisos)


def test_lista_de_corte_em_ingles_com_polegadas_e_excel():
    csv_en = b"ITEM NO.,QTY.,DESCRIPTION,LENGTH\n1,2,W8X31,157.48 in\n2,3,L2X2X1/4,1500\n"
    itens, avisos = lm.ler_lista_de_corte(csv_en, "cut.csv")
    assert not avisos
    assert itens[0].comprimento_m == pytest.approx(4.0, abs=1e-4)
    assert itens[1].descricao == 'L 2" × 1/4"' and itens[1].comprimento_m == 1.5
    livro = Workbook()
    folha = livro.active
    folha.append(["Lista de corte"])
    folha.append(["Item", "Quantidade", "Perfil", "Comprimento (m)"])
    folha.append([1, 3, "W 310 x 32,7", 6])
    saida = io.BytesIO()
    livro.save(saida)
    itens, _ = lm.ler_lista_de_corte(saida.getvalue(), "lista.xlsx")
    assert itens == [lm.ItemDaLista("1", lm.TIPO_PERFIL, "W 310 x 32,7", 3, 6.0)]


#: O que a macro ``macros_solidworks/exportar_lista_de_corte.bas`` grava (ANSI, ponto e vírgula,
#: todos os campos entre aspas, comprimento com a unidade do documento e as propriedades no fim).
CSV_DA_MACRO = (
    'ITEM;QTD.;DESCRICAO;COMPRIMENTO;NOME NA LISTA DE CORTE;"DESCRIÇÃO";"COMPRIMENTO";'
    '"QUANTIDADE";"MATERIAL"\r\n'
    '1;6;"C8x13.75";"1828,80 mm";"C8x13.75<1>";"C8X13.75";"1828,80";"6";"ASTM A36 Aço"\r\n'
    '2;4;"C6x8.2";"914,40 mm";"C6x8.2<1>";"C6X8.2";"914,40";"4";"ASTM A36 Aço"\r\n'
    '3;8;"L2.5x2.5x0.25";"2400 mm";"L2.5x2.5x0.25<1>";"L2.5X2.5X0.25";"2400";"8";""\r\n'
    '4;2;"W8x31";"72.00 in";"W8x31<1>";"W8X31";"72.00";"2";""\r\n'
    '5;1;"Placa base";"";"Placa base<1>";"";"";"1";""\r\n'
).encode("cp1252")


def test_csv_da_macro_do_solidworks():
    itens, avisos = lm.ler_lista_de_corte(CSV_DA_MACRO, "estrutura_lista_de_corte.csv")
    assert [(i.descricao, i.quantidade, round(i.comprimento_m, 4)) for i in itens] == [
        ('U 8" x 20,50', 6, 1.8288),
        ('U 6" x 12,20', 4, 0.9144),
        ('L 2 1/2" × 1/4"', 8, 2.4),
        ("W 200 x 46,1 (H)", 2, 1.8288),
        ("Placa base", 1, 0.0),
    ]
    assert itens[4].tipo == lm.TIPO_CHAPA and [i.marca for i in itens] == ["1", "2", "3", "4", "5"]
    assert not avisos


#: A macro com as colunas da geometria; itens vazios (0 corpos, LENGTH sem resolver), itens com
#: nome automático e um perfil de alumínio da biblioteca do SolidWorks, como no modelo real.
CSV_DA_MACRO_COM_GEOMETRIA = (
    "ITEM;QTD.;DESCRICAO;COMPRIMENTO;NOME NA LISTA DE CORTE;VOLUME POR PECA (cm3);"
    'MASSA DO ACO POR PECA (kg);"COMPRIMENTO";"ÂNGULO1";"Descrição";"TOTAL LENGTH"\r\n'
    '1;2;"C8x13.75";"1800 mm";"C8x13.75<1>";"4650,0";"36,503";"1800";"0°";"C8x13.75";"3600"\r\n'
    '2;0;"C8x13.75";"""LENGTH@@@C8x13.75<2>@modelo.SLDPRT""";"C8x13.75<2>";"0,0";"0,000";'
    '"""LENGTH@@@C8x13.75<2>@modelo.SLDPRT""";"-";"C8x13.75";"48466.85"\r\n'
    '3;1;"AL T SECTION 2.50 x 2.50 x 1.77";"4000 mm";"AL T SECTION 2.50 x 2.50 x 1.77<1>";'
    '"3884,0";"30,489";"4000";"64.3°";"AL T SECTION 2.50 x 2.50 x 1.77";"4000"\r\n'
    '4;1;"Item da lista de corte48";"";"Item da lista de corte48";"1500,0";"11,775";"";"";"";""\r\n'
).encode("cp1252")


def test_macro_com_a_massa_da_geometria_completa_o_que_o_catalogo_nao_tem():
    itens, avisos = lm.ler_lista_de_corte(CSV_DA_MACRO_COM_GEOMETRIA, "modelo_lista_de_corte.csv")
    assert [(i.tipo, i.descricao, i.quantidade, i.massa_geometria_kg) for i in itens] == [
        (lm.TIPO_PERFIL, 'U 8" x 20,50', 2, 36.503),
        (lm.TIPO_PERFIL, "AL T SECTION 2.50 x 2.50 x 1.77", 1, 30.489),
        (lm.TIPO_OUTRO, "Item da lista de corte48", 1, 11.775),
    ]
    texto = " ".join(avisos)
    assert "1 item(ns) da lista de corte sem nenhum corpo (quantidade 0)" in texto
    assert "massa da geometria" in texto and "AL T SECTION" in texto
    assert "informe a massa" not in texto
    r = lm.resumir(lm.ListaDeMaterial(itens=tuple(itens)))
    assert not r.pendentes
    # O perfil do catálogo usa o kg/m do catálogo; os outros, a geometria.
    assert r.massa_itens_kg == pytest.approx(2 * 1.8 * 20.5 + 30.489 + 11.775)
    assert r.massa_geometria_kg == pytest.approx(2 * 36.503 + 30.489 + 11.775)
    assert ("Massa pela geometria do modelo (aço)", "115 kg") in lm.totais(r)
    # A geometria passa pela tabela editável da página e pelo projeto.
    lista = lm.ListaDeMaterial(itens=tuple(itens))
    tabela = lm.tabela_editavel(lista)
    assert lm.itens_da_tabela(tabela) == list(itens)
    assert lm.de_dicionario(json.loads(json.dumps(lm.para_dicionario(lista)))) == lista


def test_conferencia_usa_a_geometria_e_aponta_a_razao_de_10():
    dados = importar("PP")  # 9,0 kN de reação vertical
    item = lm.ItemDaLista("V", lm.TIPO_PERFIL, "W 310 x 32,7", 1, 6.0, massa_geometria_kg=917.7)
    r = lm.resumir(lm.ListaDeMaterial(itens=(item,)))  # catálogo 196 kg; geometria 918 kg
    texto = lm.conferir_com_o_modelo(r, dados)[0].texto
    assert "pela geometria do modelo" in texto and "bate com a lista" in texto
    dez_vezes = lm.resumir(
        lm.ListaDeMaterial(
            itens=(lm.ItemDaLista("V", lm.TIPO_OUTRO, "viga", 1, massa_unitaria=91.8),)
        )
    )
    texto = lm.conferir_com_o_modelo(dez_vezes, dados)[0].texto
    assert "quase 10 vezes" in texto and "gravidade" in texto


def test_macro_e_ascii_e_grava_as_colunas_que_a_leitura_procura():
    macro = lm.macro_da_lista_de_corte().decode("ascii")
    assert "\r\n" in macro and "\n" not in macro.replace("\r\n", "")
    assert '"ITEM;QTD.;DESCRICAO;COMPRIMENTO;NOME NA LISTA DE CORTE;"' in macro
    assert '"VOLUME POR PECA (cm3);MASSA DO ACO POR PECA (kg)"' in macro
    for chamada in (
        "UpdateCutList",
        "GetBodyCount",
        "CustomPropertyManager",
        "swUnitsLinear",
        "GetBodies",
        "GetMassProperties",
    ):
        assert chamada in macro, chamada


def test_lista_de_corte_sem_cabecalho_ou_arquivo_ruim():
    itens, avisos = lm.ler_lista_de_corte(b"a;b;c\n1;2;3\n", "x.csv")
    assert not itens and "cabeçalho" in avisos[0]
    with pytest.raises(ValueError, match="Excel"):
        lm.ler_lista_de_corte(b"isto nao e um excel", "x.xlsx")
    with pytest.raises(ValueError, match="CSV"):
        lm.ler_lista_de_corte(b"\xd0\xcf\x11\xe0", "lista.xls")
    itens, avisos = lm.ler_lista_de_corte(b"QTD;DESCRICAO\n2;W8X31\n", "x.csv")
    assert itens[0].comprimento_m == 0 and any("comprimento" in a for a in avisos)


def test_conferencia_com_o_caso_pp_do_modelo():
    dados = importar("PP")  # 9,0 kN de reação vertical
    certa = lm.resumir(
        lm.ListaDeMaterial(
            itens=(lm.ItemDaLista("V", lm.TIPO_OUTRO, "viga", 1, massa_unitaria=920.0),)
        )
    )
    itens = lm.conferir_com_o_modelo(certa, dados)
    assert itens[0].nivel == pc.NIVEL_OK and "bate com a lista" in itens[0].texto
    leve = lm.resumir(
        lm.ListaDeMaterial(
            itens=(lm.ItemDaLista("V", lm.TIPO_OUTRO, "viga", 1, massa_unitaria=600.0),)
        )
    )
    itens = lm.conferir_com_o_modelo(leve, dados)
    assert itens[0].nivel == pc.NIVEL_ATENCAO and "pesa mais que a lista" in itens[0].texto
    pesada = lm.resumir(
        lm.ListaDeMaterial(
            itens=(lm.ItemDaLista("V", lm.TIPO_OUTRO, "viga", 1, massa_unitaria=1500.0),)
        )
    )
    assert "pesa menos que a lista" in lm.conferir_com_o_modelo(pesada, dados)[0].texto
    sem_pp = lm.conferir_com_o_modelo(certa, importar("SC"))
    assert "importe em Esforços do modelo o caso PP" in sem_pp[0].texto


def test_placas_de_base_dos_pilares_entram_sem_duplicar():
    dados = portico(PILAR, PILAR_2)
    placa = pb.ParametrosDaPlaca(espessura_mm=22.4)
    from core import esforcos_modelo as em

    dados = em.com_placa(dados, pb.para_dicionario(placa))
    lista = lm.com_placas_de_base(lm.ListaDeMaterial(), dados)
    (item,) = lista.itens
    assert item.tipo == lm.TIPO_CHAPA and item.quantidade == 2 and item.espessura_mm == 22.4
    assert lm.com_placas_de_base(lista, dados).itens == lista.itens
    assert lm.itens_das_placas_de_base(importar("PP")) == []


def test_lista_vai_e_volta_pelo_projeto_e_pela_tabela_da_pagina():
    lista = lm.ListaDeMaterial(
        itens=(
            lm.ItemDaLista("P1", lm.TIPO_PERFIL, "W 200 x 35,9 (H)", 2, 4.0),
            lm.ItemDaLista("G1", lm.TIPO_GRADE, "Grade", 3, 2.0, 1000, massa_unitaria=24.0),
        ),
        acrescimo_pct=7.5,
        comprimento_barra_m=6.0,
    )
    volta = lm.de_dicionario(json.loads(json.dumps(lm.para_dicionario(lista))))
    assert volta == lista
    assert lm.lista_do_projeto({"lista_de_material": lm.para_dicionario(lista)}) == lista
    assert lm.lista_do_projeto(None) == lm.ListaDeMaterial()
    estranho = lm.de_dicionario(
        {
            "itens": [{"tipo": "??", "quantidade": "abc", "comprimento_m": "x"}, 5],
            "acrescimo_pct": -3,
        }
    )
    assert estranho.itens[0].tipo == lm.TIPO_OUTRO and estranho.itens[0].quantidade == 1
    assert estranho.acrescimo_pct == 0.0
    tabela = pd.DataFrame(lm.tabela_editavel(lista), columns=list(lm.COLUNAS_EDITAVEIS))
    tabela.loc[len(tabela)] = [None] * len(lm.COLUNAS_EDITAVEIS)  # linha nova, vazia
    tabela = tabela.astype(object).where(pd.notna(tabela), pd.NA)
    assert lm.itens_da_tabela(tabela.to_dict("records")) == list(lista.itens)


def test_excel_e_csv_da_lista():
    r = lm.resumir(
        lm.ListaDeMaterial(
            itens=(
                lm.ItemDaLista("P1", lm.TIPO_PERFIL, "W 200 x 35,9 (H)", 2, 4.0),
                lm.ItemDaLista("PB", lm.TIPO_CHAPA, "Placa", 2, 0.35, 300, 19),
            )
        )
    )
    livro = load_workbook(io.BytesIO(lm.xlsx_da_lista(r, titulo="Teste")))
    assert livro.sheetnames == ["Resumo", "Itens", "Por perfil", "Chapas"]
    assert len(list(livro["Itens"].values)) == 3
    texto = lm.csv_da_lista(r).decode("utf-8-sig")
    assert texto.splitlines()[0].split(";") == list(lm.COLUNAS)
