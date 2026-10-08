"""A ligação de contraventamento no memorial (Word e PDF): resultado, o que passou e o que não passou."""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfReader

from core import contraventamento_chapa as ch
from core import contraventamento_ligacao as lig
from core import contraventamento_registro as reg
from core.project_report import (
    gerar_relatorio_industrial_pdf,
    gerar_relatorio_industrial_word,
    montar_modelo_relatorio,
)
from core.project_store import novo_projeto_documento
from core.technical_records import avaliar_contrato_registro, identificar_peca_registro
from tests.test_contraventamento_ligacao import exemplo_5_1

PALAVRAS_PROIBIDAS = ("Não informado", "bloqueio", "Snapshot", "Apêndice")


def registro(**troca) -> dict:
    r = lig.calcular_ligacao(exemplo_5_1(**troca))
    return reg.registro_ligacao(
        r, contexto={"obra": "Obra X", "tag": "LC-1", "data_do_calculo": "2026-10-08"}
    )


def projeto_com(*registros: dict) -> dict:
    projeto = novo_projeto_documento("Contraventamento LC-1", codigo="MC-LC-01")
    projeto.update(
        {
            "objetivo": "Verificar a ligação do contraventamento.",
            "responsavel": "Eng. A",
            "verificador": "Eng. B",
            "aprovador": "Eng. C",
        }
    )
    projeto["registros_tecnicos"] = [
        identificar_peca_registro(r, peca=f"Nó N{i}") for i, r in enumerate(registros, start=1)
    ]
    return projeto


def texto_word(conteudo: bytes) -> str:
    documento = Document(BytesIO(conteudo))
    partes = [p.text for p in documento.paragraphs]
    partes += [c.text for t in documento.tables for linha in t.rows for c in linha.cells]
    return "\n".join(partes)


def texto_pdf(conteudo: bytes) -> str:
    return "\n".join(p.extract_text() or "" for p in PdfReader(BytesIO(conteudo)).pages)


def test_registro_cumpre_o_contrato_do_programa():
    r = registro()
    assert avaliar_contrato_registro(r)["valido"]
    assert r["modulo"] == reg.MODULO_TITULO
    assert "LC-1" in r["titulo"]


@pytest.fixture()
def modelo():
    return montar_modelo_relatorio(projeto_com(registro(), registro(t_chapa_mm=6.0)))


def capitulos(modelo: dict) -> list[dict]:
    return [s for s in modelo["secoes"] if "blocos" in s]


def test_cada_ligacao_registrada_vira_um_capitulo_com_resultado(modelo):
    caps = capitulos(modelo)
    assert len(caps) == 2
    assert caps[0]["blocos"][0]["tipo"] == "destaque"
    assert caps[0]["blocos"][0]["rotulo"] == "Resultado: ATENÇÃO."
    assert caps[1]["blocos"][0]["rotulo"] == "Resultado: NÃO ATENDE."


def test_capitulo_mostra_entradas_tabelas_e_fecha_no_que_nao_passou(modelo):
    blocos = capitulos(modelo)[1]["blocos"]
    textos = [str(b.get("texto") or b.get("rotulo")) for b in blocos]
    passou = next(i for i, x in enumerate(textos) if x.startswith("O que passou ("))
    nao_passou = next(i for i, x in enumerate(textos) if x.startswith("Não passou ("))
    assert nao_passou > passou
    todas = str(blocos)
    assert "Sistema de normas" in todas and "Espessura da chapa [mm]" in todas
    legendas = " ".join(b["legenda"] for b in blocos if b.get("tipo") == "tabela")
    assert "Forças de cálculo em cada interface" in legendas
    assert "Geometria do nó" in legendas


def test_forcas_do_ufm_nao_viram_outros_resultados_repetidos(modelo):
    blocos = capitulos(modelo)[0]["blocos"]
    outros = [
        b for b in blocos if b.get("tipo") == "tabela" and "Outros resultados" in b["legenda"]
    ]
    assert outros == []


def test_word_e_pdf_do_memorial_trazem_a_ligacao_sem_palavras_proibidas():
    projeto = projeto_com(registro())
    word = texto_word(gerar_relatorio_industrial_word(projeto))
    pdf = texto_pdf(gerar_relatorio_industrial_pdf(projeto))
    for texto in (word, pdf):
        assert "Ligação de contraventamento" in texto
        assert "O que passou" in texto
        for proibido in PALAVRAS_PROIBIDAS:
            assert proibido not in texto, proibido
    assert "AISC" in word


def test_unidade_do_metodo_aparece_no_registro():
    assert registro()["entradas"]["metodo"] == ch.METODOS[ch.METODO_LRFD]
