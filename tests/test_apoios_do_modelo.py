"""Reações em cada apoio pela macro (v12): leitura dos arquivos, gravação no projeto, quadro por
apoio (ação na fundação, comportamento, conferências) e a página."""

from __future__ import annotations

import io

import pytest

from core import apoios_do_modelo as am
from core import esforcos_modelo as em
from core import leitura_do_modelo as ld
from core.project_store import obter_projeto_ativo, salvar_projeto

#: Os pés dos 4 pilares do teste01 do usuário (m, Y para cima).
PES = {101: (0.0, -2.0, -1.0), 102: (0.0, -2.0, 0.0), 103: (1.5, -2.0, -1.0), 104: (1.5, -2.0, 0.0)}


def _n(v: float) -> str:
    return f"{v:.6f}".replace(".", ",")


def arquivo_de_apoios(
    caso: str,
    reacoes: dict[int, tuple[float, ...]],
    *,
    escala: float = 1.0,
    por_fixacao: tuple[float, ...] | None = None,
) -> bytes:
    """O -apoios.csv como a macro grava: reação (N, N·m) em cada nó e a soma por fixação."""
    linhas = [
        "sep=;",
        "19:28; 24/10/2026",
        "Nome do estudo:Análise estática 1",
        f"Caso de carga: {caso}",
        "Macro: Mecanica Toolkit - exportar_casos_de_carga v12",
        "",
        "Reacoes nos nos (N; N.m)",
        "No;X (m);Y (m);Z (m);Fx (N);Fy (N);Fz (N);Mx (N.m);My (N.m);Mz (N.m)",
    ]
    for no, valores in reacoes.items():
        x, y, z = (c * escala for c in PES[no])
        linhas.append(";".join([str(no), _n(x), _n(y), _n(z), *map(_n, valores)]))
    linhas += [
        "",
        "Reacoes por fixacao (N; N.m)",
        "Fixacao;Fx (N);Fy (N);Fz (N);Mx (N.m);My (N.m);Mz (N.m)",
    ]
    soma = por_fixacao or tuple(sum(v[k] for v in reacoes.values()) for k in range(6))
    linhas.append("Fixo-1;" + ";".join(map(_n, soma)))
    return ("\r\n".join(linhas) + "\r\n").encode("cp1252")


def arquivo_de_fixacoes() -> bytes:
    linhas = [
        "sep=;",
        "19:28; 24/10/2026",
        "Nome do estudo:Análise estática 1",
        "Macro: Mecanica Toolkit - exportar_casos_de_carga v12",
        "",
        "Fixacoes do estudo",
        "Fixacao;Tipo (codigo);Entidades;Entidade;Tipo da entidade;X (m);Y (m);Z (m)",
    ]
    for k, p in enumerate(PES.values(), start=1):
        linhas.append(f"Fixo-1;0;4;{k};IVertex;" + ";".join(map(_n, p)))
    return ("\r\n".join(linhas) + "\r\n").encode("cp1252")


#: PE do teste01: 360 N em cada pé (para cima), sem momento — e um caso de vento com momento.
PE = {no: (0.0, 360.0, 0.0, 0.0, 0.0, 0.0) for no in PES}
W0 = {
    101: (-455.0, 600.0, 0.0, 0.0, 0.0, -410.0),
    102: (-455.0, 600.0, 0.0, 0.0, 0.0, -410.0),
    103: (-455.0, -600.0, 0.0, 0.0, 0.0, -410.0),
    104: (-455.0, -600.0, 0.0, 0.0, 0.0, -410.0),
}


def geometria() -> ld.GeometriaDoModelo:
    corpos = tuple(
        ld.CorpoDoModelo(
            f"Pilar[{k}]", "W6x9<1>", "W6x9", p, (p[0], -0.2, p[2]), 1.8, forma=ld.FORMA_BARRA
        )
        for k, p in enumerate(PES.values(), start=1)
    )
    return ld.GeometriaDoModelo(arquivo="teste01_esqueleto.csv", corpos=corpos)


def dados() -> am.ApoiosDoModelo:
    d = am.ApoiosDoModelo()
    for caso, reacoes in (("PE", PE), ("W0", W0)):
        d = am.com_caso(d, am.ler_apoios(arquivo_de_apoios(caso, reacoes), f"{caso}-apoios.csv"))
    return am.com_fixacoes(d, am.ler_fixacoes(arquivo_de_fixacoes(), "fixacoes.csv"))


# ------------------------------------------------------------------ leitura
def test_reconhece_e_le_os_dois_arquivos():
    apoios = arquivo_de_apoios("PE", PE)
    assert am.tipo_do_arquivo(apoios) == am.TIPO_APOIOS
    assert am.tipo_do_arquivo(arquivo_de_fixacoes()) == am.TIPO_FIXACOES
    assert am.tipo_do_arquivo(b"sep=;\r\nNome da viga;Elemento\r\n") is None
    lido = am.ler_apoios(apoios, "x.csv")
    assert lido.caso == "PE" and lido.estudo == "Análise estática 1" and len(lido.nos) == 4
    assert lido.nos[0].forca_N == (0.0, 360.0, 0.0) and lido.nos[0].ponto == PES[101]
    assert lido.por_fixacao == (("Fixo-1", (0.0, 1440.0, 0.0, 0.0, 0.0, 0.0)),)
    fix = am.ler_fixacoes(arquivo_de_fixacoes())
    assert fix.nomes == ("Fixo-1",) and fix.entidades[2].ponto == PES[103]


def test_arquivo_sem_caso_ou_sem_nos_vira_erro_claro():
    sem_caso = arquivo_de_apoios("PE", PE).replace(b"Caso de carga: PE", b"")
    with pytest.raises(am.ArquivoDeApoiosInvalido, match="caso de carga"):
        am.ler_apoios(sem_caso)
    sem_nos = b"sep=;\r\nCaso de carga: PE\r\nReacoes nos nos (N; N.m)\r\nNo;X\r\n"
    with pytest.raises(am.ArquivoDeApoiosInvalido, match="nenhum nó"):
        am.ler_apoios(sem_nos)


def test_vai_e_volta_do_projeto_e_dado_estragado_nao_quebra():
    d = dados()
    assert am.de_dicionario(am.para_dicionario(d)) == d
    estragado = {"casos": {"PP": {"nos": [["x", None], "lixo", [1, 0, 0, 0, "a"]]}}, "fixacoes": 3}
    assert am.de_dicionario(estragado) == am.ApoiosDoModelo()
    assert am.de_dicionario(None) == am.ApoiosDoModelo()


# ------------------------------------------------------------------ quadro por apoio
def test_quadro_por_apoio_acao_na_fundacao_e_comportamento():
    q = am.montar_quadro(dados(), geometria=geometria())
    assert [a.rotulo for a in q.apoios] == ["A1", "A2", "A3", "A4"]
    a1 = q.apoios[0]
    assert a1.barra == "Pilar[1]" and a1.fixacao == "Fixo-1" and a1.tipo_da_fixacao == 0
    pe = q.carga("A1", "PE")
    assert pe is not None and q.n_kN(pe) == pytest.approx(0.36)  # compressão positiva
    assert pe.forca_kN == pytest.approx((0.0, -0.36, 0.0))  # a ação na fundação é para baixo
    w0 = q.carga("A3", "W0")
    assert w0 is not None and q.n_kN(w0) == pytest.approx(-0.60)  # arrancamento
    assert w0.forca_kN[0] == pytest.approx(0.455)  # o vento empurra a fundação para +X
    assert w0.momento_kNm[2] == pytest.approx(0.41)
    assert all(a.comportamento == "engastado" for a in q.apoios)  # momento no W0
    so_pe = am.montar_quadro(
        am.com_caso(am.ApoiosDoModelo(), am.ler_apoios(arquivo_de_apoios("PE", PE)))
    )
    assert all(a.comportamento == "articulado" for a in so_pe.apoios)


def test_coordenadas_em_mm_sao_reconhecidas_pelo_esqueleto():
    d = am.com_caso(am.ApoiosDoModelo(), am.ler_apoios(arquivo_de_apoios("PE", PE, escala=1000.0)))
    q = am.montar_quadro(d, geometria=geometria())
    assert q.escala == pytest.approx(0.001)
    assert q.apoios[0].ponto == pytest.approx(PES[101]) and q.apoios[0].barra == "Pilar[1]"


def test_nos_de_uma_mesma_face_viram_um_apoio():
    reacoes = dict(PE)
    PES[105] = (0.05, -2.0, -1.0)  # outro nó da mesma placa do pé A1
    try:
        reacoes[105] = (0.0, 100.0, 0.0, 0.0, 0.0, 0.0)
        d = am.com_caso(am.ApoiosDoModelo(), am.ler_apoios(arquivo_de_apoios("PE", reacoes)))
        q = am.montar_quadro(d)
    finally:
        del PES[105]
    assert len(q.apoios) == 4
    a1 = next(a for a in q.apoios if 101 in a.nos)
    assert set(a1.nos) == {101, 105}
    carga = q.carga(a1.rotulo, "PE")
    assert carga is not None and q.n_kN(carga) == pytest.approx(0.46)


def test_conferencias_contra_a_reacao_total_e_a_soma_por_fixacao():
    esforcos = em.EsforcosDoModelo(
        casos={
            "PE": em.CasoImportado(
                "PE", "", "", "", "", {}, reacoes=em.Reacoes(modelo=(0.0, 1.44, 0.0))
            ),
            "W0": em.CasoImportado(
                "W0", "", "", "", "", {}, reacoes=em.Reacoes(modelo=(-2.0, 0.0, 0.0))
            ),
            "SC": em.CasoImportado("SC", "", "", "", "", {}),
        }
    )
    q = am.montar_quadro(dados(), esforcos=esforcos)
    textos = {i.nivel: [] for i in q.conferencia}
    for i in q.conferencia:
        textos[i.nivel].append(i.texto)
    assert any("PE: a soma dos apoios bate" in t for t in textos["ok"])
    assert any(t.startswith("W0: a soma dos apoios") and "não bate" in t for t in textos["atencao"])
    assert any("Casos sem o arquivo de apoios: SC" in t for t in textos["atencao"])
    fora = am.com_caso(
        am.ApoiosDoModelo(),
        am.ler_apoios(arquivo_de_apoios("PE", PE, por_fixacao=(0.0, 1000.0, 0, 0, 0, 0))),
    )
    assert any("soma por fixação" in i.texto for i in am.montar_quadro(fora).conferencia)


def test_tabelas_csv_e_excel():
    q = am.montar_quadro(dados(), geometria=geometria())
    linhas = am.linhas_do_quadro(q)
    assert len(linhas) == 8 and all(len(lin) == len(am.colunas()) for lin in linhas)
    texto = am.csv_do_quadro(q).decode("utf-8-sig")
    assert texto.startswith("Apoio;Barra;X (m)") and "Anglo" in texto
    from openpyxl import load_workbook

    livro = load_workbook(io.BytesIO(am.xlsx_do_quadro(q, "Plataforma")))
    assert livro.sheetnames == ["Apoios", "Comportamento"]
    assert livro["Comportamento"].max_row == 5


# ------------------------------------------------------------------ página
def test_quadro_por_apoio_na_pagina(banco_com_projeto):
    from tests.test_esforcos_modelo import PILAR
    from tests.test_esforcos_modelo_pagina import abrir, gravar_no_projeto, sem_ajuda
    from tests.test_quadro_fundacoes import PILAR_2, portico

    gravar_no_projeto(portico(PILAR, PILAR_2))
    projeto = obter_projeto_ativo()
    projeto["apoios_do_modelo"] = am.para_dicionario(dados())
    salvar_projeto(projeto, motivo="teste")
    t = abrir()
    assert not sem_ajuda(t), sorted(sem_ajuda(t))
    assert "Reação em cada apoio (pela macro)" in [s.value for s in t.subheader]
    tabela = next(d.value for d in t.dataframe if "N (kN, compressão +)" in list(d.value.columns))
    assert set(tabela["Apoio"]) == {"A1", "A2", "A3", "A4"}
    assert any(d.value.columns[-1] == "Trabalha como" for d in t.dataframe)


def test_nos_todos_no_mesmo_ponto_nao_viram_um_apoio_so():
    """O que a macro v12 gravou no SOLIDWORKS 2026 (24/10): 0, 0, 0 em todos os nós. Os quatro pés
    não podem virar um apoio; com outro caso de posições boas, valem as dele."""
    zerado = arquivo_de_apoios("PE", PE, escala=0.0)
    so_zerado = am.montar_quadro(am.com_caso(am.ApoiosDoModelo(), am.ler_apoios(zerado)))
    assert len(so_zerado.apoios) == 4
    assert any(
        "mesmo ponto" in i.texto and "cada nó é um apoio" in i.texto for i in so_zerado.conferencia
    )
    d = am.com_caso(am.ApoiosDoModelo(), am.ler_apoios(zerado))
    d = am.com_caso(d, am.ler_apoios(arquivo_de_apoios("W0", W0, escala=0.001)))  # m ÷ 1000
    q = am.montar_quadro(d, geometria=geometria())
    assert len(q.apoios) == 4 and q.escala == pytest.approx(1000.0)
    assert {a.barra for a in q.apoios} == {f"Pilar[{k}]" for k in range(1, 5)}
