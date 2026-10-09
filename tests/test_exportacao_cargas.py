"""Exportação do plano de cargas: unidades, eixos, sinais, combinações, CSV, Excel e conferência."""

from __future__ import annotations

import csv
import io

import pytest
from openpyxl import load_workbook

from core import exportacao_cargas as ex
from core import load_combinations as comb
from core import plano_de_cargas as pc


def plano_completo() -> pc.PlanoDeCargas:
    plano = pc.PlanoDeCargas()
    plano = pc.com_acao(plano, pc.nova_acao("PP"))
    plano = pc.com_acao(
        plano,
        pc.nova_acao("SC", cargas=[pc.CargaDoModelo("Piso", 5.0, "kN/m²", "Z (vertical)")]),
    )
    vento = {
        "X": [pc.CargaDoModelo("P1, nível 1", 2.0, "kN", "X")],
        "Y": [pc.CargaDoModelo("P1, nível 1", 3.0, "kN", "Y")],
    }
    for acao in pc.acoes_de_vento("teste", vento):
        plano = pc.com_acao(plano, acao)
    return plano


def linha(plano, caso, sistema=ex.UNIDADES_KN_M, eixos=ex.EIXO_Z_PARA_CIMA):
    return next(item for item in ex.linhas_de_carga(plano, sistema, eixos) if item.caso == caso)


# ------------------------------------------------------------------ unidades, eixos e sinais
@pytest.mark.parametrize(
    ("sistema", "unidade", "esperado", "nova"),
    [
        (ex.UNIDADES_N_MM, "kN", 1000.0, "N"),
        (ex.UNIDADES_N_MM, "kN/m", 1.0, "N/mm"),
        (ex.UNIDADES_N_MM, "kN/m²", 0.001, "N/mm² (MPa)"),
        (ex.UNIDADES_N_MM, "kN·m", 1e6, "N·mm"),
        (ex.UNIDADES_N_M, "kN/m", 1000.0, "N/m"),
        (ex.UNIDADES_N_M, "kN/m²", 1000.0, "N/m² (Pa)"),
        (ex.UNIDADES_KN_M, "kN·m", 1.0, "kN·m"),
        (ex.UNIDADES_N_MM, "°C", 1.0, "°C"),
    ],
)
def test_conversao_de_unidades(sistema, unidade, esperado, nova):
    assert ex.converter(1.0, unidade, sistema) == pytest.approx((esperado, nova))


@pytest.mark.parametrize(
    ("direcao", "eixo"),
    [("X", "X"), ("+X", "X"), ("−Y", "Y"), ("-y", "Y"), ("Z (vertical)", "Z"), ("—", None)],
)
def test_eixo_da_direcao(direcao, eixo):
    assert ex.eixo_da_direcao(direcao) == eixo


def test_vertical_positivo_e_para_baixo_e_o_sinal_horizontal_vem_do_valor():
    assert ex.componentes_no_programa(pc.CargaDoModelo("piso", 5.0, "kN/m²", "Z (vertical)")) == (
        0.0,
        0.0,
        -5.0,
    )
    assert ex.componentes_no_programa(pc.CargaDoModelo("nó", -2.0, "kN", "−X")) == (-2.0, 0.0, 0.0)
    assert ex.componentes_no_programa(pc.CargaDoModelo("todo", 10.0, "°C", "—")) is None


def test_solidworks_com_y_para_cima_troca_os_eixos_sem_menos_zero():
    assert ex.para_os_eixos((1.0, 2.0, 3.0), ex.EIXO_Y_PARA_CIMA) == (1.0, 3.0, -2.0)
    assert ex.para_os_eixos((1.0, 2.0, 3.0), ex.EIXO_Z_PARA_CIMA) == (1.0, 2.0, 3.0)
    gravidade = ex.para_os_eixos((0.0, 0.0, -9.81), ex.EIXO_Y_PARA_CIMA)
    assert gravidade == (0.0, -9.81, 0.0)
    assert str(gravidade[2]) == "0.0"  # nada de -0.0 na planilha


def test_sentido_por_extenso():
    assert ex.sentido((0.0, -5.0, 0.0), ex.EIXO_Y_PARA_CIMA) == "−Y (para baixo)"
    assert ex.sentido((0.0, 0.0, -5.0), ex.EIXO_Z_PARA_CIMA) == "−Z (para baixo)"
    assert ex.sentido((2.0, 0.0, 0.0), ex.EIXO_Z_PARA_CIMA) == "+X"
    assert ex.sentido(None, ex.EIXO_Z_PARA_CIMA) == "—"


# ------------------------------------------------------------------ cargas por caso
def test_pp_sem_cargas_vira_a_gravidade_do_modelo():
    plano = plano_completo()
    pp = linha(plano, "PP", eixos=ex.EIXO_Y_PARA_CIMA)
    assert pp.tipo == ex.TIPO_GRAVIDADE and pp.valor == 9.81 and pp.unidade == "m/s²"
    assert (pp.fx, pp.fy, pp.fz) == (0.0, -9.81, 0.0)
    assert "Gravidade" in pp.solidworks


def test_sobrecarga_em_n_por_mm2_com_y_para_cima():
    sc = linha(plano_completo(), "SC", ex.UNIDADES_N_MM, ex.EIXO_Y_PARA_CIMA)
    assert sc.tipo == ex.TIPO_AREA and sc.unidade == "N/mm² (MPa)"
    assert sc.valor == pytest.approx(0.005)
    assert (sc.fx, sc.fy, sc.fz) == pytest.approx((0.0, -0.005, 0.0))
    assert sc.sentido == "−Y (para baixo)" and "Pressão" in sc.solidworks


@pytest.mark.parametrize(
    ("caso", "z_para_cima", "y_para_cima"),
    [
        ("W0", (2000.0, 0.0, 0.0), (2000.0, 0.0, 0.0)),
        ("W90", (0.0, 3000.0, 0.0), (0.0, 0.0, -3000.0)),
        ("W180", (-2000.0, 0.0, 0.0), (-2000.0, 0.0, 0.0)),
        ("W270", (0.0, -3000.0, 0.0), (0.0, 0.0, 3000.0)),
    ],
)
def test_vento_nas_quatro_direcoes_nos_dois_sistemas_de_eixos(caso, z_para_cima, y_para_cima):
    plano = plano_completo()
    for eixos, esperado in ((ex.EIXO_Z_PARA_CIMA, z_para_cima), (ex.EIXO_Y_PARA_CIMA, y_para_cima)):
        w = linha(plano, caso, ex.UNIDADES_N_M, eixos)
        assert w.tipo == ex.TIPO_NODAL and w.unidade == "N"
        assert (w.fx, w.fy, w.fz) == pytest.approx(esperado)


def test_temperatura_e_momento():
    plano = pc.com_acao(
        pc.PlanoDeCargas(),
        pc.nova_acao("T+", cargas=[pc.CargaDoModelo("Toda a estrutura", 10.0, "°C", "—")]),
    )
    plano = pc.com_acao(
        plano, pc.nova_acao("IM", cargas=[pc.CargaDoModelo("Base do motor", 4.0, "kN·m", "Y")])
    )
    t = linha(plano, "T+", ex.UNIDADES_N_MM)
    assert t.tipo == ex.TIPO_TEMPERATURA and t.valor == 10.0 and t.unidade == "°C"
    assert (t.fx, t.fy, t.fz) == (None, None, None) and t.sentido == "ΔT"
    m = linha(plano, "IM", ex.UNIDADES_N_MM)
    assert m.tipo == ex.TIPO_MOMENTO and m.valor == pytest.approx(4e6) and m.unidade == "N·mm"
    assert m.observacao.startswith("vetor do momento")


# ------------------------------------------------------------------ combinações
def test_matriz_e_lista_das_combinacoes_batem():
    plano = plano_completo()
    estados = [comb.ELU_NORMAL, comb.ELS_RARA]
    cabecalho, matriz = ex.matriz_de_combinacoes(plano, estados)
    assert cabecalho[:2] == ["Combinação", "Estado-limite"] and cabecalho[-1] == "Expressão"
    assert cabecalho[2:-1] == list(plano.codigos)
    lista_cab, lista = ex.lista_de_combinacoes(plano, estados)
    assert lista_cab == ["Combinação", "Estado-limite", "Tipo", "Caso", "Fator"]
    combinacoes = pc.combinacoes(plano, estados)
    assert len(matriz) == len(combinacoes)
    for linha_matriz, c in zip(matriz, combinacoes, strict=True):
        assert linha_matriz[0] == ex.nome_da_combinacao(c, len(combinacoes))
        da_lista = {r[3]: r[4] for r in lista if r[0] == linha_matriz[0]}
        nao_nulos = {k: v for k, v in zip(plano.codigos, linha_matriz[2:-1], strict=True) if v}
        assert da_lista == pytest.approx(nao_nulos)
    assert {r[2] for r in lista} == {"ELU", "ELS"}
    assert matriz[0][0] == "C01-ELU"


def test_nome_da_combinacao_alarga_com_mais_de_99():
    c = pc.LinhaDeCombinacao(7, "x", comb.ELS_FREQUENTE, "")
    assert ex.nome_da_combinacao(c, 12) == "C07-ELSF"
    assert ex.nome_da_combinacao(c, 120) == "C007-ELSF"


def test_csv_em_portugues():
    plano = plano_completo()
    for conteudo in (
        ex.csv_de_cargas(plano, ex.UNIDADES_N_MM, ex.EIXO_Y_PARA_CIMA),
        ex.csv_matriz(plano, comb.ESTADOS_PADRAO),
        ex.csv_lista(plano, comb.ESTADOS_PADRAO),
    ):
        assert conteudo.startswith("﻿".encode())
        linhas = list(csv.reader(io.StringIO(conteudo.decode("utf-8-sig")), delimiter=";"))
        assert len(linhas) > 2
    cargas = list(
        csv.reader(
            io.StringIO(
                ex.csv_de_cargas(plano, ex.UNIDADES_N_MM, ex.EIXO_Y_PARA_CIMA).decode("utf-8-sig")
            ),
            delimiter=";",
        )
    )
    assert cargas[0] == list(ex.COLUNAS_CARGAS)
    sc = next(r for r in cargas if r[0] == "SC")
    assert sc[4] == "0,0050" and sc[7] == "-0,0050"


def test_planilha_excel_tem_as_cinco_abas():
    plano = plano_completo()
    conteudo = ex.xlsx_do_plano(
        plano,
        sistema=ex.UNIDADES_N_MM,
        eixos=ex.EIXO_Y_PARA_CIMA,
        estados=comb.ESTADOS_PADRAO,
        titulo="Plataforma P1",
        anglo=True,
    )
    livro = load_workbook(io.BytesIO(conteudo))
    assert livro.sheetnames == [
        "Leia-me",
        "Ações",
        "Cargas",
        "Combinações",
        "Combinações (lista)",
    ]
    leia = " ".join(str(c.value) for c in livro["Leia-me"]["A"] if c.value)
    assert "NO SOLIDWORKS SIMULATION" in leia and "Plataforma P1" in leia and "5.9" in leia
    cargas = list(livro["Cargas"].values)
    assert cargas[0] == ex.COLUNAS_CARGAS
    assert len(cargas) == 1 + len(ex.linhas_de_carga(plano, ex.UNIDADES_N_MM, ex.EIXO_Y_PARA_CIMA))
    matriz = list(livro["Combinações"].values)
    assert matriz[0][2:-1] == plano.codigos
    assert all(v is None or v > 0 for linha_m in matriz[1:] for v in linha_m[2:-1])
    assert livro["Cargas"].freeze_panes == "E2"
    acoes = list(livro["Ações"].values)
    assert [a[0] for a in acoes[1:]] == list(plano.codigos)
    assert acoes[1][1] == "↓ g"


def test_leia_me_explica_o_eixo_do_solidworks_so_quando_y_e_vertical():
    com_y = " ".join(ex.texto_leia_me(ex.UNIDADES_N_MM, ex.EIXO_Y_PARA_CIMA))
    com_z = " ".join(ex.texto_leia_me(ex.UNIDADES_KN_M, ex.EIXO_Z_PARA_CIMA))
    assert "−Z" in com_y and "−Z" not in com_z
    assert "CRITÉRIO ANGLO" not in com_z


# ------------------------------------------------------------------ conferência e símbolos
def niveis(plano, **kw) -> dict[str, list[str]]:
    saida: dict[str, list[str]] = {}
    for item in pc.conferir_plano(plano, **kw):
        saida.setdefault(item.nivel, []).append(item.texto)
    return saida


def test_conferencia_do_plano_completo_e_ok():
    assert niveis(plano_completo()) == {
        pc.NIVEL_OK: ["Plano completo: peso próprio, sobrecarga e vento nas quatro direções."]
    }


def test_conferencia_do_plano_vazio_e_erro():
    assert pc.NIVEL_ERRO in niveis(pc.PlanoDeCargas())


def test_conferencia_aponta_o_que_falta_e_o_que_nao_fecha():
    plano = pc.com_acao(pc.PlanoDeCargas(), pc.nova_acao("W0"))
    plano = pc.com_acao(
        plano, pc.nova_acao("T+", cargas=[pc.CargaDoModelo("Toda", 10.0, "kN", "X")])
    )
    plano = pc.com_acao(
        plano,
        pc.nova_acao(
            "EQ",
            cargas=[
                pc.CargaDoModelo("Vaso", 0.0, "kN", "Z (vertical)"),
                pc.CargaDoModelo("Bomba", 5.0, "kN", "—"),
                pc.CargaDoModelo("Bomba", 5.0, "kN", "—"),
            ],
        ),
    )
    r = niveis(plano, anglo=True)
    atencao = " ".join(r[pc.NIVEL_ATENCAO])
    erro = " ".join(r[pc.NIVEL_ERRO])
    for trecho in (
        "Sem PP",
        "Sem sobrecarga",
        "faltam W180, W270, W90",
        "Temperatura num sentido só",
        "W0: sem cargas",
        "carga nula",
        "duas vezes",
        "Critério Anglo 5.9",
    ):
        assert trecho in atencao, trecho
    assert "T+: ação de temperatura com carga em kN" in erro
    assert "sem direção" in erro


def test_vento_com_carga_vertical_pede_conferencia():
    plano = pc.com_acao(
        pc.PlanoDeCargas(),
        pc.nova_acao("W0", cargas=[pc.CargaDoModelo("Cobertura", 0.4, "kN/m²", "Z (vertical)")]),
    )
    assert any("carga vertical" in t for t in niveis(plano)[pc.NIVEL_ATENCAO])


@pytest.mark.parametrize(
    ("codigo", "unidade", "direcao", "esperado"),
    [
        ("W90", "kN", "Y", "↑ +Y"),
        ("PP", "kN", "Z (vertical)", "↓ g"),
        ("T−", "°C", "—", "ΔT −"),
        ("ZZ", "°C", "—", "ΔT"),
        ("ZZ", "kN", "Z (vertical)", "↓"),
        ("ZZ", "kN", "X", "⇄"),
    ],
)
def test_simbolos(codigo, unidade, direcao, esperado):
    acao = pc.nova_acao(codigo, cargas=[pc.CargaDoModelo("e", 1.0, unidade, direcao)])
    assert pc.simbolo(acao) == esperado
