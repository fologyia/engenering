"""O Word do fim da página Degrau de escada em grade: tudo o que a página mostra, editável."""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from core import degrau_escada as de
from core import degrau_registro as reg
from core import escada_lances as el
from core import longarina_escada as le
from core.degrau_word import gerar_word_degrau

IDENT = {"obra": "Obra X", "tag": "ESC-01", "responsavel": "Eng. Y", "data": "07/10/2026"}
PALAVRAS_PROIBIDAS = ("Não informado", "bloqueio", "Snapshot", "Apêndice")


def texto(conteudo: bytes) -> str:
    documento = Document(BytesIO(conteudo))
    partes = [p.text for p in documento.paragraphs]
    partes += [c.text for t in documento.tables for linha in t.rows for c in linha.cells]
    return "\n".join(partes)


@pytest.fixture(scope="module")
def completo() -> tuple[de.ResultadoEscada, bytes]:
    r = de.calcular_escada(de.EntradaDegrau(H_mm=5400))
    esc = el.montar_escada(
        r, [el.PatamarInformado(1500, 1200, 1000), el.PatamarInformado(3700, 800, None, True)]
    )
    r = el.aplicar_na_escada(r, esc)
    lg = le.calcular_longarina(r, le.EntradaLongarina(trecho_topo_m=0.3), esc)
    return r, gerar_word_degrau(r, IDENT, escada=esc, longarina=lg)


def test_word_abre_e_traz_as_figuras(completo):
    _, conteudo = completo
    documento = Document(BytesIO(conteudo))
    # esquemático + 7 diagramas por longarina (3 longarinas)
    assert len(documento.inline_shapes) == 1 + 3 * 7
    assert len(documento.tables) >= 20


def test_word_traz_todas_as_secoes_da_pagina(completo):
    cheio = texto(completo[1])
    for trecho in (
        "Obra X",
        "ESC-01",
        "Eng. Y",
        "07/10/2026",
        "O que esta página não faz",
        "Dados de entrada",
        "Quadro-resumo",
        "Deslocamento do degrau",
        "Verificações (",
        "2. Geometria da escada",
        "2.1 Lances, patamares e esquemático",
        "3. Memorial do modelo adotado",
        "3.1 Longarina",
        "Cargas por longarina",
        "Trechos",
        "Reações",
        "Diagramas",
        "Os 64 modelos do catálogo",
        "Texto para a requisição de compra",
        "Conflitos entre normas",
        "Referências",
    ):
        assert trecho in cheio, trecho
    for proibido in PALAVRAS_PROIBIDAS:
        assert proibido not in cheio, proibido


def test_word_traz_cada_dado_de_entrada_cada_verificacao_e_cada_modelo(completo):
    r, conteudo = completo
    cheio = texto(conteudo)
    for rotulo in (
        "Desnível total H [mm]",
        "Sobrecarga q [kN/m²]",
        "Coeficiente γg",
        "Longarina: perfil",
        "Longarina: trecho horizontal no topo [m]",
    ):
        assert rotulo in cheio, rotulo
    for linha in reg.tabela_de_verificacoes(r):
        assert linha["Verificação"] in cheio
    assert all(m.nome in cheio for m in de.catalogo())
    assert r.especificacao in cheio


def test_word_sem_escada_nem_identificacao():
    conteudo = gerar_word_degrau(de.calcular_escada(de.EntradaDegrau()), None)
    cheio = texto(conteudo)
    assert "Quadro-resumo" in cheio and "3.1 Longarina" not in cheio
    assert "Não informado" not in cheio


def test_word_com_a_escada_em_3d():
    from core import cena_3d as c3
    from core import cena_escada as cesc

    r = de.calcular_escada(de.EntradaDegrau(H_mm=5400))
    esc = el.montar_escada(r, [el.PatamarInformado(1500, 1200, 1000)])
    r = el.aplicar_na_escada(r, esc)
    dados = cesc.para_o_registro(
        esc, modo=cesc.MODO_DETALHADO, secao_da_longarina=None, altura_do_guarda_corpo_m=1.2
    )
    dados["azimute"], dados["elevacao"] = c3.VISTAS[c3.VISTA_ISOMETRICA]
    conteudo = gerar_word_degrau(r, IDENT, escada=esc, escada_3d=dados)
    assert len(Document(BytesIO(conteudo)).inline_shapes) == 2  # esquemático + 3D
    assert "Escada em 3D" in texto(conteudo)
