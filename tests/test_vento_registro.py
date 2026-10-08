"""O registro técnico do vento e o seu capítulo no memorial (Word e PDF).

O registro leva a tabela de verificações (aplicabilidade e valores de apoio), as tabelas de pressões
por zona, das vedações e das cargas do pórtico e o texto de destaque; o memorial as mostra como
"o que passou", "valores de apoio" e "não passou", com as tabelas do módulo entre uma coisa e outra.
"""

from __future__ import annotations

import re
import xml.etree.ElementTree as ET
from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfReader

from components.wind_figures import cor_do_ce, svg_planta
from core import memorial_verificacoes as mv
from core import vento_coeficientes as coef
from core import vento_edificio as ve
from core import vento_nbr6123 as v6123
from core import vento_portico as vp
from core import vento_registro as vr
from core.project_report import (
    gerar_relatorio_industrial_pdf,
    gerar_relatorio_industrial_word,
    montar_modelo_relatorio,
)
from core.project_store import novo_projeto_documento
from core.technical_records import identificar_peca_registro


def galpao(**mudancas) -> ve.EntradaEdificio:
    base = dict(
        v0_m_s=40.0,
        comprimento_a_m=60.0,
        largura_b_m=30.0,
        altura_h_m=8.0,
        cobertura=ve.COBERTURA_DUAS_AGUAS,
        theta_graus=10.0,
        categoria="III",
        grupo_s3=3,
        espacamento_porticos_m=6.0,
    )
    base.update(mudancas)
    return ve.EntradaEdificio(**base)


def registro(**mudancas) -> dict:
    resultado = ve.calcular_edificacao(galpao(**mudancas))
    casos = vp.casos_do_portico(resultado)
    return vr.registro_edificacao(resultado, casos_do_portico=casos)


# ------------------------------------------------------------------------------- o registro
def test_registro_tem_identidade_entradas_e_tabelas():
    r = registro()
    assert r["modulo_id"] == "vento_nbr6123" and r["modulo"] == "Vento nas estruturas"
    assert r["status"] == "Calculado"
    assert r["entradas"]["v0_m_s"] == 40.0
    assert r["entradas"]["edicao_da_norma"] == "ABNT NBR 6123:2023"
    assert r["entradas"]["permeabilidade"] == coef.CENARIOS_PERMEABILIDADE["quatro_faces"]
    assert r["resultados"]["status_geral"] == "OK"
    assert len(r["resultados"]["verificações"]) >= 15
    assert "q = " in r["resultados"]["destaque_memorial"]
    assert any("V_k" in str(equacao) for equacao in r["equacoes"])
    assert any("0,613" in str(equacao) for equacao in r["equacoes"])
    assert r["referencias"] == ["ABNT NBR 6123:2023 — Forças devidas ao vento em edificações."]
    assert r["hash_calculo"]  # o registro v2 é assinado


def test_registro_sem_nan_nem_infinito():
    import json

    r = registro(altura_h_m=2.0, largura_b_m=40.0)
    json.dumps(r, allow_nan=False)


def test_tabelas_do_memorial_somam_a_largura_util_e_tem_uma_linha_por_zona():
    resultado = ve.calcular_edificacao(galpao())
    tabelas = vr.tabelas_de_pressoes(resultado)
    assert len(tabelas) == 2  # vento a 0° e a 90°
    for tabela in tabelas:
        assert sum(tabela["larguras"]) == vr.LARGURA_UTIL_DXA
        assert len(tabela["larguras"]) == len(tabela["cabecalhos"])
        assert all(len(linha) == len(tabela["cabecalhos"]) for linha in tabela["linhas"])
    assert len(tabelas[0]["linhas"]) == len(resultado.casos_do_angulo(0)[0].pressoes)
    assert tabelas[0]["cabecalhos"][-2:] == ["Δp, c_pi = -0,30", "Δp, c_pi = 0,00"]
    ved = vr.tabela_de_vedacoes(resultado)
    assert sum(ved["larguras"]) == vr.LARGURA_UTIL_DXA and len(ved["linhas"]) == len(
        resultado.vedacoes
    )
    casos = vp.casos_do_portico(resultado)
    porticos = vr.tabela_do_portico_para_memorial(resultado, casos)
    assert (
        len(porticos["linhas"]) == len(casos) and sum(porticos["larguras"]) == vr.LARGURA_UTIL_DXA
    )


def test_uma_agua_registra_tres_direcoes():
    resultado = ve.calcular_edificacao(
        ve.EntradaEdificio(
            v0_m_s=40.0,
            comprimento_a_m=40.0,
            largura_b_m=20.0,
            altura_h_m=6.0,
            cobertura=ve.COBERTURA_UMA_AGUA,
            theta_graus=10.0,
        )
    )
    assert len(vr.tabelas_de_pressoes(resultado)) == 3
    r = vr.registro_edificacao(resultado)
    assert {c["angulo_graus"] for c in r["resultados"]["casos_de_vento"]} == {0, 90, -90}


def test_alerta_vira_situacao_atencao():
    r = registro(comprimento_a_m=150.0)  # a/b = 5 > 4
    assert r["resultados"]["status_geral"] == "ALERTA" and r["status"] == "Atenção"
    assert any("a/b" in str(alerta) for alerta in r["alertas"])


def test_solucao_do_portico_entra_no_registro():
    resultado = ve.calcular_edificacao(galpao())
    casos = vp.casos_do_portico(resultado)
    secao = vp.SecaoDoElemento(20_000.0, 5.0e8)
    solucao = vp.resolver_portico(
        vp.geometria_do_portico(resultado.geometria), casos[0], pilar=secao, rafter=secao
    )
    r = vr.registro_edificacao(resultado, casos_do_portico=casos, solucao=solucao)
    assert r["resultados"]["solucao_do_portico"]["caso"] == casos[0].nome
    assert len(r["resultados"]["solucao_do_portico"]["reacoes"]) == 2


# ------------------------------------------------------------------------------- o memorial
def projeto_com_vento(**mudancas) -> dict:
    projeto = novo_projeto_documento("Galpão G1", codigo="MC-G1-001")
    projeto.update({"objetivo": "Ação do vento.", "responsavel": "Eng. A", "verificador": "Eng. B"})
    projeto["registros_tecnicos"] = [
        identificar_peca_registro(registro(**mudancas), peca="Galpão G1")
    ]
    return projeto


def capitulo_do_vento(modelo) -> dict:
    capitulos = [s for s in modelo["secoes"] if "blocos" in s]
    assert len(capitulos) == 1
    return capitulos[0]


def test_capitulo_comeca_pelo_resultado_com_as_pressoes_dinamicas():
    capitulo = capitulo_do_vento(montar_modelo_relatorio(projeto_com_vento()))
    primeiro = capitulo["blocos"][0]
    assert primeiro["rotulo"] == "Resultado: ATENDE."
    assert "Pressão dinâmica sobre a estrutura" in primeiro["texto"]
    assert "q = 0,843 kN/m²" in primeiro["texto"]  # decimal em português


def test_capitulo_traz_dados_de_entrada_com_rotulos_legiveis():
    capitulo = capitulo_do_vento(montar_modelo_relatorio(projeto_com_vento()))
    entradas = next(b for b in capitulo["blocos"] if b.get("legenda") == "Dados de entrada.")
    rotulos = {celula for linha in entradas["linhas"] for celula in (linha[0], linha[2])}
    assert {"Velocidade básica V₀ [m/s]", "Comprimento a [m]", "Largura b [m]"} <= rotulos
    assert "Categoria de rugosidade" in rotulos and "Cobertura" in rotulos
    assert not any(re.search(r"[a-z]+_[a-z]+", str(rotulo)) for rotulo in rotulos)  # sem códigos


def test_capitulo_mostra_as_tabelas_do_modulo_depois_dos_valores_de_apoio():
    capitulo = capitulo_do_vento(montar_modelo_relatorio(projeto_com_vento()))
    legendas = [b.get("legenda", "") for b in capitulo["blocos"] if b["tipo"] == "tabela"]
    assert any(legenda.startswith("Valores de apoio") for legenda in legendas)
    assert any(legenda.startswith("Vento a 0°") for legenda in legendas)
    assert any(legenda.startswith("Vento a 90°") for legenda in legendas)
    assert any(legenda.startswith("Pressões de projeto das vedações") for legenda in legendas)
    assert any(legenda.startswith("Cargas do vento no pórtico") for legenda in legendas)
    apoio = next(i for i, legenda in enumerate(legendas) if legenda.startswith("Valores de apoio"))
    primeira = next(i for i, legenda in enumerate(legendas) if legenda.startswith("Vento a 0°"))
    assert apoio < primeira
    ultimo = capitulo["blocos"][-1]  # o que não passou sempre fecha o capítulo
    assert ultimo["tipo"] in ("destaque", "bullets", "tabela")


def test_capitulo_termina_no_que_nao_passou():
    ok = capitulo_do_vento(montar_modelo_relatorio(projeto_com_vento()))
    rotulos = [b.get("rotulo") for b in ok["blocos"]]
    assert "Não passou: nenhuma." in rotulos
    depois = ok["blocos"][rotulos.index("Não passou: nenhuma.") + 1 :]
    # Depois da declaração só vem o que o módulo manda conferir antes de emitir.
    assert [b["tipo"] for b in depois] == ["subtitulo", "bullets"]
    assert depois[0]["texto"] == "Conferir antes de emitir"
    assert any("±0,03" in item for item in depois[1]["itens"])
    alerta = capitulo_do_vento(montar_modelo_relatorio(projeto_com_vento(comprimento_a_m=150.0)))
    textos = str(alerta["blocos"])
    assert "Atenção — passou com ressalva" in textos and "a/b" in textos
    assert alerta["blocos"][0]["rotulo"] == "Resultado: ATENÇÃO."


@pytest.fixture(scope="module")
def documentos():
    projeto = projeto_com_vento()
    return gerar_relatorio_industrial_word(projeto), gerar_relatorio_industrial_pdf(projeto)


def test_word_e_pdf_levam_as_tabelas_do_vento(documentos):
    word, pdf = documentos
    documento = Document(BytesIO(word))
    texto_docx = "\n".join(p.text for p in documento.paragraphs)
    texto_docx += "\n".join(
        c.text for t in documento.tables for linha in t.rows for c in linha.cells
    )
    texto_pdf = "\n".join(p.extract_text() or "" for p in PdfReader(BytesIO(pdf)).pages)
    for texto in (texto_docx, texto_pdf):
        assert "Vento nas estruturas" in texto
        assert "Resultado: ATENDE" in texto
        assert "Pressão dinâmica sobre a estrutura" in texto
    assert "Δp" in texto_docx and "A1 e B1" in texto_docx
    for proibida in ("bloqueio", "Snapshot", "Apêndice", "Não informado"):
        assert proibida not in texto_docx


def test_decimais_do_registro_saem_em_portugues_no_memorial():
    capitulo = capitulo_do_vento(montar_modelo_relatorio(projeto_com_vento()))
    apoio = next(
        b for b in capitulo["blocos"] if b.get("legenda", "").startswith("Valores de apoio")
    )
    celulas = [celula for linha in apoio["linhas"] for celula in linha]
    assert not any(isinstance(c, str) and mv._DECIMAL_COM_PONTO.search(c) for c in celulas)


# ------------------------------------------------------------------------------- o desenho
@pytest.mark.parametrize(
    "cobertura", [ve.COBERTURA_PLANA, ve.COBERTURA_DUAS_AGUAS, ve.COBERTURA_UMA_AGUA]
)
def test_desenho_da_planta_e_um_svg_valido_para_cada_direcao(cobertura):
    theta = 0.0 if cobertura == ve.COBERTURA_PLANA else 10.0
    resultado = ve.calcular_edificacao(galpao(cobertura=cobertura, theta_graus=theta))
    for alpha in resultado.alphas:
        svg = svg_planta(resultado, alpha)
        raiz = ET.fromstring(svg)
        assert raiz.tag.endswith("svg")
        assert f"{alpha} graus" in svg
        rotulos_das_paredes = ("A1/B1", "A3/B3") if alpha == 0 else ("C1/D1", "C2/D2")
        for rotulo in rotulos_das_paredes:
            assert rotulo in svg, (alpha, rotulo)
    with pytest.raises(ValueError):
        svg_planta(resultado, 45)


def test_cores_vao_do_azul_da_sucao_ao_vermelho_da_pressao():
    assert cor_do_ce(-2.0) != cor_do_ce(-0.2) != cor_do_ce(0.0) != cor_do_ce(0.7)
    assert cor_do_ce(0.0) == "rgb(255,255,255)"
    azul = cor_do_ce(-2.0)
    vermelho = cor_do_ce(1.0)
    assert azul.endswith(",255)") and vermelho.startswith("rgb(255,")


# ------------------------------------------------------------------------------- auxiliares
@pytest.mark.parametrize(
    ("texto", "esperado"),
    [
        ("-0,3; 0,2", (-0.3, 0.2)),
        ("-0.3, 0.2", (-0.3, 0.2)),
        ("−0,3", (-0.3,)),
        ("0,2", (0.2,)),
        ("-0,3;0,2;-0,3", (-0.3, 0.2)),
        ("-0,3, 0,2, 0", (-0.3, 0.2, 0.0)),
        ("0,2 0,5", (0.2, 0.5)),
    ],
)
def test_interpretar_cpis(texto, esperado):
    assert coef.interpretar_cpis(texto) == pytest.approx(esperado)


@pytest.mark.parametrize("texto", ["", "abc", "3", "0,2; x"])
def test_interpretar_cpis_rejeita_o_invalido(texto):
    with pytest.raises(ValueError):
        coef.interpretar_cpis(texto)


def test_s3_do_projeto_nunca_fica_abaixo_do_grupo():
    assert v6123.fator_s3_do_projeto(3) == (1.0, None)
    valor, aviso = v6123.fator_s3_do_projeto(3, 0.63, 50.0)  # 0,999: arredondamento da norma
    assert valor == 1.0 and aviso is None
    valor, aviso = v6123.fator_s3_do_projeto(3, 0.5, 25.0)
    assert valor == 1.0 and aviso and "abaixo do mínimo do grupo 3" in aviso
    valor, aviso = v6123.fator_s3_do_projeto(1, 0.1, 100.0)
    assert valor == pytest.approx(v6123.fator_s3_estatistico(0.1, 100.0)) and aviso is None
    valor, _ = v6123.fator_s3_do_projeto(4, 0.63, 50.0)  # 0,999 > 0,95 do grupo 4
    assert valor == pytest.approx(0.9989, abs=1e-3)


def test_tabela_compacta_do_portico():
    resultado = ve.calcular_edificacao(galpao())
    casos = vp.casos_do_portico(resultado)
    linhas = vp.tabela_compacta_do_portico(casos)
    assert len(linhas) == len(casos)
    assert linhas[-1]["Pilar esq. (kN/m)"] == casos[-1].carga(vp.PILAR_ESQUERDO).carga_kN_m
    assert {"Caso", "Faixa (da empena de barlavento)", "Água dir. (kN/m)"} <= set(linhas[0])


def test_vento_tem_o_seu_fluxo_sugerido_sem_mudar_o_dos_outros_modulos():
    from core import module_sequencing as seq

    assert seq.fluxos_do_modulo("vento_nbr6123") == ["vento", "contraventamento"]
    etapas = seq.montar_sequencia(None, "vento_nbr6123")
    assert [e.modulo_id for e in etapas] == [
        "casos_carga",
        "vento_nbr6123",
        "estruturas_aco",
        "flambagem_colunas",
    ]
    assert {e.modulo_id: e.situacao for e in etapas}["vento_nbr6123"] == "atual"
    # Quem já tinha fluxo continua com o primeiro deles quando nada foi concluído.
    assert [e.modulo_id for e in seq.montar_sequencia(None, "flambagem_colunas")] == [
        "casos_carga",
        "assistente_cargas",
        "flambagem_colunas",
        "analise_sensibilidade",
    ]
