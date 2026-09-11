"""Identificação da peça nos registros e agrupamento por peça no memorial."""

from io import BytesIO

from docx import Document
from pypdf import PdfReader

from core.project_report import (
    gerar_relatorio_industrial_pdf,
    gerar_relatorio_industrial_word,
    montar_modelo_relatorio,
)
from core.project_store import criar_item
from core.technical_records import (
    agrupar_registros_por_componente,
    criar_registro_tecnico,
    identificar_peca_registro,
    normalizar_registro_tecnico,
    rotulo_componente,
)
from tests.test_project_validation import _projeto_documentado


def _registro(titulo: str, **extras):
    return criar_registro_tecnico(
        modulo="Flambagem de colunas",
        titulo=titulo,
        status="Atende",
        resumo="Verificação de flambagem.",
        entradas={"L_mm": 3000},
        resultados={"fator_seguranca": 2.1},
        conclusao="Atende.",
        **extras,
    )


def test_identificar_peca_prefixa_titulo_e_vincula_componente():
    registro = _registro("Flambagem de coluna — W 200 x 26,6")
    identificado = identificar_peca_registro(registro, peca="Coluna P1", componentes_ids=["c1"])
    assert identificado["titulo"] == "Coluna P1 — Flambagem de coluna — W 200 x 26,6"
    assert identificado["peca"] == "Coluna P1"
    assert identificado["componentes_ids"] == ["c1"]

    # Aplicar duas vezes não duplica o prefixo nem o vínculo.
    repetido = identificar_peca_registro(identificado, peca="Coluna P1", componentes_ids=["c1"])
    assert repetido["titulo"] == identificado["titulo"]
    assert repetido["componentes_ids"] == ["c1"]

    # Sem nome de peça o título fica como está; só o vínculo é adicionado.
    so_vinculo = identificar_peca_registro(registro, componentes_ids=["c2"])
    assert so_vinculo["titulo"] == registro["titulo"]
    assert so_vinculo["componentes_ids"] == ["c2"]
    assert normalizar_registro_tecnico(so_vinculo)["peca"] == ""


def test_duas_colunas_iguais_geram_registros_distintos():
    base = _registro("Flambagem de coluna — W 200 x 26,6")
    p1 = normalizar_registro_tecnico(identificar_peca_registro(base, peca="Coluna P1"))
    p2 = normalizar_registro_tecnico(identificar_peca_registro(base, peca="Coluna P2"))
    assert p1["hash_calculo"] != p2["hash_calculo"]


def test_registros_legados_mantem_assinatura():
    legado = normalizar_registro_tecnico(_registro("Registro antigo"))
    assinatura = legado["hash_calculo"]
    # Um registro salvo antes do campo ``peca`` existir precisa continuar
    # validando a assinatura ao ser relido.
    sem_peca = {chave: valor for chave, valor in legado.items() if chave != "peca"}
    assert normalizar_registro_tecnico(sem_peca)["hash_calculo"] == assinatura


def test_agrupamento_segue_ordem_do_escopo_e_deixa_soltos_ao_final():
    coluna = criar_item(tag="P1", descricao="Coluna principal")
    viga = criar_item(tag="V1", descricao="Viga contínua")
    componentes = [coluna, viga]
    registros = [
        _registro("Viga", componentes_ids=[viga["id"]]),
        _registro("Solto"),
        _registro("Coluna", componentes_ids=[coluna["id"]]),
        _registro("Compartilhado", componentes_ids=[coluna["id"], viga["id"]]),
        _registro("Vínculo órfão", componentes_ids=["nao-existe"]),
    ]
    grupos = agrupar_registros_por_componente(registros, componentes)
    assert [grupo["rotulo"] for grupo in grupos] == [
        "P1 · Coluna principal",
        "V1 · Viga contínua",
        "Registros sem peça vinculada",
    ]
    assert [item["titulo"] for item in grupos[0]["registros"]] == ["Coluna", "Compartilhado"]
    assert [item["titulo"] for item in grupos[1]["registros"]] == ["Viga", "Compartilhado"]
    assert [item["titulo"] for item in grupos[2]["registros"]] == ["Solto", "Vínculo órfão"]
    assert rotulo_componente({"tag": "P1"}) == "P1"
    assert rotulo_componente({"descricao": "Só descrição"}) == "Só descrição"
    assert rotulo_componente({}) == "Componente sem identificação"


def _projeto_mezanino():
    projeto = _projeto_documentado()
    coluna_p1 = criar_item(tag="P1", descricao="Coluna principal esquerda", servico="Mezanino", material="ASTM A572", desenho="DES-01", criticidade="Alta")
    coluna_p2 = criar_item(tag="P2", descricao="Coluna principal direita", servico="Mezanino", material="ASTM A572", desenho="DES-01", criticidade="Alta")
    viga = criar_item(tag="V1", descricao="Viga principal com balanço", servico="Mezanino", material="ASTM A572", desenho="DES-02", criticidade="Alta")
    projeto["componentes"] = [coluna_p1, coluna_p2, viga]
    base = _registro("Flambagem de coluna — W 200 x 26,6")
    projeto["registros_tecnicos"] = [
        normalizar_registro_tecnico(identificar_peca_registro(base, peca="Coluna P2", componentes_ids=[coluna_p2["id"]])),
        normalizar_registro_tecnico(identificar_peca_registro(base, peca="Coluna P1", componentes_ids=[coluna_p1["id"]])),
        normalizar_registro_tecnico(identificar_peca_registro(_registro("Viga contínua — 3 vãos + balanço"), peca="Viga V1", componentes_ids=[viga["id"]])),
        normalizar_registro_tecnico(identificar_peca_registro(_registro("Mão francesa — L 50x50x5"), peca="Mão francesa MF-1", componentes_ids=[coluna_p1["id"]])),
        normalizar_registro_tecnico(_registro("Combinações de ações ELU e ELS")),
    ]
    return projeto


def test_memorial_agrupa_registros_por_peca():
    projeto = _projeto_mezanino()
    modelo = montar_modelo_relatorio(projeto, secoes_incluidas=["plano_calculo", "registros"])
    secoes = modelo["secoes"]
    titulos = [item["titulo"] for item in secoes]

    numero_registros = next(item["titulo"] for item in secoes if item.get("nivel", 1) == 1 and "Registros" in item["titulo"]).split(".")[0]
    pecas = [item["titulo"] for item in secoes if item.get("nivel") == 2]
    assert pecas == [
        f"{numero_registros}.1 P1 · Coluna principal esquerda",
        f"{numero_registros}.2 P2 · Coluna principal direita",
        f"{numero_registros}.3 V1 · Viga principal com balanço",
        f"{numero_registros}.4 Registros sem peça vinculada",
    ]
    capitulos = [item["titulo"] for item in secoes if item.get("nivel") == 3]
    assert capitulos == [
        f"{numero_registros}.1.1 Coluna P1 — Flambagem de coluna — W 200 x 26,6",
        f"{numero_registros}.1.2 Mão francesa MF-1 — Mão francesa — L 50x50x5",
        f"{numero_registros}.2.1 Coluna P2 — Flambagem de coluna — W 200 x 26,6",
        f"{numero_registros}.3.1 Viga V1 — Viga contínua — 3 vãos + balanço",
        f"{numero_registros}.4.1 Combinações de ações ELU e ELS",
    ]
    assert any("Peça: Coluna P1." in paragrafo for item in secoes for paragrafo in item.get("paragrafos", []))
    # Os capítulos de cada peça vêm logo abaixo do título dela, não todos ao final.
    sequencia = [item["titulo"].split(" ")[0] for item in secoes if item.get("nivel", 1) >= 2]
    assert sequencia == [
        f"{numero_registros}.1", f"{numero_registros}.1.1", f"{numero_registros}.1.2",
        f"{numero_registros}.2", f"{numero_registros}.2.1",
        f"{numero_registros}.3", f"{numero_registros}.3.1",
        f"{numero_registros}.4", f"{numero_registros}.4.1",
    ]

    plano = next(item for item in secoes if "Plano" in item["titulo"])
    tabela = plano["tabelas"][0]
    assert tabela["cabecalhos"][1] == "Peça"
    assert [linha[1] for linha in tabela["linhas"]] == ["Coluna P2", "Coluna P1", "Viga V1", "Mão francesa MF-1", "-"]
    assert sum(tabela["larguras"]) == 9360
    assert "Registros técnicos" in " ".join(titulos)


def test_memorial_sem_componentes_mantem_lista_plana():
    projeto = _projeto_documentado()
    projeto["componentes"] = []
    projeto["registros_tecnicos"] = [
        normalizar_registro_tecnico(_registro("Um")),
        normalizar_registro_tecnico(_registro("Dois")),
    ]
    modelo = montar_modelo_relatorio(projeto, secoes_incluidas=["registros"])
    niveis = [item.get("nivel", 1) for item in modelo["secoes"]]
    assert niveis == [1, 2, 2]
    assert not any("sem peça vinculada" in item["titulo"] for item in modelo["secoes"])


def test_word_e_pdf_emitem_capitulos_por_peca():
    projeto = _projeto_mezanino()
    word = gerar_relatorio_industrial_word(projeto)
    documento = Document(BytesIO(word))
    titulos_h3 = [p.text for p in documento.paragraphs if p.style.name == "Heading 3"]
    # O gerador Word normaliza travessões para hífen; o que importa é o nível.
    assert any("Coluna P1" in texto and "Flambagem" in texto for texto in titulos_h3)
    texto_word = "\n".join(p.text for p in documento.paragraphs)
    assert "P1 · Coluna principal esquerda" in texto_word

    pdf = gerar_relatorio_industrial_pdf(projeto)
    texto_pdf = "\n".join(pagina.extract_text() or "" for pagina in PdfReader(BytesIO(pdf)).pages)
    assert "Coluna principal esquerda" in texto_pdf
    assert "Registros sem pe" in texto_pdf
