"""Vento em estruturas abertas: forças nos nós, cargas nas barras, plano de cargas, CSV e memorial."""

from __future__ import annotations

import csv
import io
import json
from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfReader

from core import plano_de_cargas as pc
from core import vento_aberto_registro as var
from core import vento_estrutura_aberta as va
from core.project_report import (
    gerar_relatorio_industrial_pdf,
    gerar_relatorio_industrial_word,
    montar_modelo_relatorio,
)
from core.project_store import novo_projeto_documento
from core.technical_records import avaliar_contrato_registro, identificar_peca_registro
from tests.test_vento_estrutura_aberta import VENTO, plataforma

PALAVRAS_PROIBIDAS = ("Não informado", "bloqueio", "Snapshot", "Apêndice")


@pytest.fixture(scope="module")
def dois_pisos() -> tuple[va.GeometriaAberta, va.ResultadoVentoAberto]:
    tanque = va.Equipamento("Tanque", 2, va.FORMA_CILINDRO, 2.0, 2.0, 3.0, peso_kN=80.0)
    g = plataforma(cotas_m=(3.0, 6.0), equipamentos=(tanque,))
    return g, va.calcular_vento_aberto(g, VENTO)


@pytest.mark.parametrize("direcao", ["X", "Y"])
def test_forcas_nos_nos_somam_o_total_da_direcao(dois_pisos, direcao):
    _, r = dois_pisos
    d = r.direcao(direcao)
    nos = va.cargas_nodais(d, incluir_base=True)
    equipamentos = sum(f for nivel in d.equipamentos_por_nivel for _, f, _ in nivel)
    assert sum(c.total_kN for c in nos) + equipamentos == pytest.approx(d.total_kN)
    # Sem a base, a soma é a dos níveis.
    sem_base = va.cargas_nodais(d)
    assert sum(c.total_kN for c in sem_base) + equipamentos == pytest.approx(
        sum(d.forcas_nos_niveis_kN)
    )
    assert len(sem_base) == len(d.planos) * len(d.niveis)


def test_parcela_de_cada_pilar(dois_pisos):
    _, r = dois_pisos
    c = va.cargas_nodais(r.x)[0]
    assert c.pilares == r.x.pilares_por_portico == 2  # pórtico em X: 1 vão em Y, 2 pilares
    assert c.por_pilar_kN == pytest.approx(c.total_kN / 2)
    assert c.portico == 1 and c.coordenada_m == 0.0


def test_guarda_corpo_so_nos_porticos_das_pontas(dois_pisos):
    _, r = dois_pisos
    por_portico = r.x.guarda_corpo_por_portico_kN
    assert all(v > 0 for v in por_portico[0]) and all(v > 0 for v in por_portico[-1])
    assert all(v == 0 for v in por_portico[1])


def test_cargas_distribuidas_conferem_com_o_coeficiente(dois_pisos):
    g, r = dois_pisos
    linhas = va.cargas_distribuidas(g, r.x)
    plano = r.x.planos[0]
    viga = next(c for c in linhas if c.portico == 1 and c.elemento == "Vigas do nível 1")
    q1 = r.x.niveis[0].q_N_m2 / 1e3
    assert viga.w_kN_m == pytest.approx(plano.eta * plano.ca * q1 * g.altura_viga_m)
    assert not any("iagona" in c.elemento for c in linhas)
    assert any(c.elemento.startswith("Guarda-corpo") for c in linhas)


def test_acoes_para_o_plano_levam_os_nos_e_o_equipamento(dois_pisos):
    _, r = dois_pisos
    acoes = var.acoes_para_o_plano(r, registro_id="abc")
    assert [a.codigo for a in acoes] == ["W0", "W90", "W180", "W270"]
    w0, _, w180, _ = acoes
    assert w0.registro_id == "abc" and w0.origem == var.ORIGEM
    total_por_pilar = sum(
        c.valor * (int(c.elemento.split("dos ")[1].split(" ")[0]) if "nós" in c.elemento else 1)
        for c in w0.cargas
    )
    assert total_por_pilar == pytest.approx(sum(r.x.forcas_nos_niveis_kN))
    assert any(c.observacao == "equipamento" for c in w0.cargas)
    assert w180.cargas[0].valor == pytest.approx(-w0.cargas[0].valor)
    plano = pc.PlanoDeCargas()
    for a in acoes:
        plano = pc.com_acao(plano, a)
    assert plano.codigos == ("W0", "W90", "W180", "W270")


def test_csv_dos_nos_e_das_barras(dois_pisos):
    g, r = dois_pisos
    nos = list(csv.reader(io.StringIO(var.csv_nos_nos(r).decode("utf-8-sig")), delimiter=";"))
    assert nos[0][0] == "Direção" and "Por pilar (kN)" in nos[0]
    assert any(linha[3] == "0" for linha in nos[1:])  # a faixa da base
    assert any(linha[1] == "Tanque" for linha in nos[1:])
    barras = var.csv_distribuidas(g, r).decode("utf-8-sig")
    assert barras.startswith("Direção;Pórtico;Elemento") and "Pilares (cada um)" in barras


def test_tres_tabelas_para_o_memorial(dois_pisos):
    _, r = dois_pisos
    tabelas = var.tabelas_para_memorial(r)
    assert len(tabelas) == 3
    for t in tabelas:
        assert t["cabecalhos"] and t["linhas"] and t["legenda"]
        assert all(len(linha) == len(t["cabecalhos"]) for linha in t["linhas"])
        assert sum(t["larguras"]) == var.LARGURA_UTIL_DXA


def test_registro_cumpre_o_contrato_e_vai_para_o_memorial(dois_pisos):
    _, r = dois_pisos
    registro = var.registro_vento_aberto(
        r,
        contexto={"obra": "Obra X", "tag": "PL-01", "data_do_calculo": "2026-10-08"},
        avisos_da_base=["Conflito de teste"],
    )
    assert avaliar_contrato_registro(registro)["valido"]
    assert registro["modulo_id"] == var.MODULO_ID and "PL-01" in registro["titulo"]
    json.dumps(registro)
    projeto = novo_projeto_documento("Plataforma", codigo="MC-PL-01")
    projeto["registros_tecnicos"] = [identificar_peca_registro(registro, peca="Plataforma P1")]
    modelo = montar_modelo_relatorio(projeto)
    capitulos = [s for s in modelo["secoes"] if "blocos" in s]
    assert len(capitulos) == 1
    blocos = capitulos[0]["blocos"]
    assert blocos[0]["tipo"] == "destaque" and blocos[0]["rotulo"] == "Resultado do cálculo."
    assert "kN em X" in blocos[0]["texto"]
    legendas = [b.get("legenda", "") for b in blocos if b["tipo"] == "tabela"]
    assert legendas[0] == "Dados de entrada."
    assert sum("nós" in legenda.lower() for legenda in legendas) == 1
    assert len(legendas) == 4  # entradas + níveis, pórticos e nós (nada de "Outros resultados")
    entradas = str(blocos[2])
    assert "V₀" in entradas or "V0" in entradas
    conferir = [b for b in blocos if b.get("tipo") == "bullets"]
    assert any("Conflito de teste" in " ".join(b["itens"]) for b in conferir)
    assert not any(str(b.get("texto", "")).startswith("Conclusão:") for b in blocos)
    word = Document(BytesIO(gerar_relatorio_industrial_word(projeto)))
    texto_word = "\n".join(
        [p.text for p in word.paragraphs]
        + [c.text for t in word.tables for linha in t.rows for c in linha.cells]
    )
    texto_pdf = "\n".join(
        p.extract_text() or ""
        for p in PdfReader(BytesIO(gerar_relatorio_industrial_pdf(projeto))).pages
    )
    for texto in (texto_word, texto_pdf):
        assert "Vento em estrutura aberta" in texto and "Resultado do cálculo" in texto
        for proibido in PALAVRAS_PROIBIDAS:
            assert proibido not in texto, proibido
        # Números de item não viram decimal: (item 8.5), não (8,5).
        assert "(8,5)" not in texto and "(item 8.5)" in texto
