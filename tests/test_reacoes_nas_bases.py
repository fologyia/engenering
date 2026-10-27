"""Reações nas bases para o projeto da fundação (core/reacoes_nas_bases.py e a página), com o
teste01 resolvido como pórtico espacial: cada pé, em cada caso, tem de ser a reação que o cálculo
dá naquele pé — venha da macro (-apoios.csv) ou da base do pilar (forças da viga)."""

from __future__ import annotations

import io
import json
import re
from dataclasses import replace
from datetime import date
from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from components.reacoes_nas_bases_help import AJUDA
from core import apoios_do_modelo as am
from core import diagnostico_do_modelo as dm
from core import direcoes_da_secao as ds
from core import esforcos_modelo as em
from core import leitura_do_modelo as ld
from core import plano_de_cargas as pc
from core import quadro_fundacoes as qf
from core import reacoes_nas_bases as rb
from core import vento_por_barras as vpb
from core.project_store import obter_projeto_ativo, salvar_projeto
from tests import portico_3d_teste01 as t
from tests.test_esforcos_modelo_pagina import TIPOS_COM_AJUDA, _ajuda, _rotulo
from tests.test_nucleo_esforcos_ponta_a_ponta import montar

RAIZ = Path(__file__).resolve().parent.parent
APP = str(RAIZ / "app.py")
PAGINA = "app_pages/reacoes_nas_bases.py"


def direcoes(dados: em.EsforcosDoModelo) -> ds.DirecoesDosPilares:
    """O triedro de cada pilar como o pórtico de teste escreve as forças: Cisalhamento 1 e
    Momento 1 no eixo y local, 2 no z, axial do Fim 1 ao Fim 2."""
    pilares = {}
    for p in qf.montar_quadro(dados).pilares:
        barra = next(b for b in t.BARRAS if b.viga == p.pilar)
        x, y, z = (
            vpb.para_a_planta(v, "Y") for v in t._eixos(np.array(barra.inicio), np.array(barra.fim))
        )
        sinal = -1.0 if p.base == qf.BASE_INICIO else 1.0
        pilares[p.pilar] = ds.DirecoesDaSecao(
            p.pilar, barra.corpo, x, y, z, sinal, "teste", "teste"
        )
    return ds.DirecoesDosPilares(pilares)


def _confere_com_o_portico(r: rb.ReacoesNasBases, casos=t.CASOS, tolerancia=1e-4) -> None:
    """Cada pé (A1…A4 = os pés do pórtico pela posição) é a reação do cálculo com o sinal
    trocado, nos eixos do modelo."""
    assert [p.rotulo for p in r.pes] == ["A1", "A2", "A3", "A4"]
    for pe, ponto in zip(r.pes, t.PES, strict=True):
        for caso in casos:
            acao = r.acao(pe.rotulo, caso)
            assert acao is not None, (pe.rotulo, caso)
            forca, momento = t.resolver(caso).reacoes[ponto]
            assert acao.forca_kN == pytest.approx(tuple(-forca / 1e3), abs=tolerancia)
            assert acao.momento_kNm == pytest.approx(tuple(-momento / 1e3), abs=tolerancia)


# ------------------------------------------------------------------ de onde vêm os números
def test_pela_macro_cada_pe_e_a_reacao_do_calculo():
    dados, apoios, geometria = montar()
    r = rb.montar(dados, apoios, plano=t.plano(), geometria=geometria)
    _confere_com_o_portico(r, tolerancia=1e-6)
    assert {a.fonte for a in r.acoes} == {rb.FONTE_APOIO}
    assert [p.pilar for p in r.pes] == [b.viga for b in t.BARRAS[:4]]
    assert all(p.comportamento == "engastado" and p.perfil == t.PILAR for p in r.pes)
    assert r.casos == t.CASOS and not r.faltam and not r.avisos
    assert len(r.conferencia) == len(t.CASOS)
    assert all(i.nivel == "ok" for i in r.conferencia)


def test_pela_base_dos_pilares_sem_a_macro():
    """Sem os -apoios.csv: a força na base de cada pilar, levada aos eixos do modelo pela direção
    da seção, é a mesma reação (forças e momentos) — o pilar é a única barra no pé."""
    dados, apoios, geometria = montar(com_apoios=False)
    r = rb.montar(dados, apoios, plano=t.plano(), geometria=geometria, direcoes=direcoes(dados))
    _confere_com_o_portico(r)
    assert {a.fonte for a in r.acoes} == {rb.FONTE_PILAR}
    assert r.pes[0].ponto == pytest.approx((0.0, -2.0, -1.0))
    assert all(i.nivel == "ok" for i in r.conferencia)


def test_caso_com_os_apoios_de_outra_rodada_sai_pela_base_do_pilar():
    dados, apoios, geometria = montar(apoios_de={"W0": "W90"})
    r = rb.montar(dados, apoios, plano=t.plano(), geometria=geometria, direcoes=direcoes(dados))
    _confere_com_o_portico(r)
    assert r.fontes("W0") == {rb.FONTE_PILAR}
    assert all(r.fontes(c) == {rb.FONTE_APOIO} for c in t.CASOS if c != "W0")
    assert any("W0" in a and "outra rodada" in a for a in r.avisos)


def test_sem_macro_e_sem_direcao_da_secao_fica_de_fora_com_o_motivo():
    dados, apoios, geometria = montar(com_apoios=False)
    r = rb.montar(dados, apoios, plano=t.plano(), geometria=geometria)
    assert not r.acoes and len(r.faltam) == len(t.CASOS)
    assert "sem a direção da seção do pilar" in r.faltam[0] and r.faltam[0].startswith("PP:")
    assert rb.linhas(r) == []


def _sem_posicao(apoios: am.ApoiosDoModelo) -> am.ApoiosDoModelo:
    """Os apoios como a macro v12 gravava: os nós sem a posição."""
    return replace(
        apoios,
        casos={
            c: replace(
                a, nos=tuple(am.ReacaoNoNo(n.no, None, n.forca_N, n.momento_Nm) for n in a.nos)
            )
            for c, a in apoios.casos.items()
        },
    )


def test_apoios_sem_posicao_ligados_aos_pilares_pelos_valores():
    """A macro v12 (nós sem posição) e nenhum esqueleto: cada apoio é ligado ao pilar cuja base
    leva, em todos os casos, a mesma compressão e a mesma cortante — e a ordem dos pés é a dos
    nós."""
    dados, apoios, _ = montar(com_esqueleto=False)
    r = rb.montar(dados, _sem_posicao(apoios), plano=t.plano())
    assert [p.pilar for p in r.pes] == [b.viga for b in t.BARRAS[:4]]
    assert all(p.ponto is None for p in r.pes)
    assert {a.fonte for a in r.acoes} == {rb.FONTE_APOIO}
    assert any("ligado ao seu pilar pelos valores" in a for a in r.avisos)
    _confere_com_o_portico(r, tolerancia=1e-6)


def test_o_teste01_real_apoios_v12_com_dois_casos_de_outra_rodada():
    """Como o banco do usuário (27/10): -apoios.csv sem posição, os do W180 e do W270 de outra
    rodada, e a direção da seção dos pilares desconhecida (vento só nos nós do topo). Os apoios
    se ligam aos pilares pelos valores, a direção da seção sai das reações da macro nos outros
    casos, e o W180 e o W270 saem pela base de cada pilar — iguais ao cálculo."""
    dados, apoios, _ = montar(com_esqueleto=False, apoios_de={"W180": "W90", "W270": "W0"})
    r = rb.montar(dados, _sem_posicao(apoios), plano=t.plano())
    assert r.fontes("W180") == r.fontes("W270") == {rb.FONTE_PILAR}
    assert r.fontes("W0") == {rb.FONTE_APOIO}
    assert any("saiu das reações da macro" in a for a in r.avisos)
    _confere_com_o_portico(r)
    assert all(i.nivel == "ok" for i in r.conferencia) and len(r.conferencia) == 6


def test_direcao_pela_macro_com_o_esqueleto_sem_a_direcao_da_pagina():
    dados, apoios, geometria = montar(apoios_de={"W0": "W90"})
    r = rb.montar(dados, apoios, plano=t.plano(), geometria=geometria)
    assert r.fontes("W0") == {rb.FONTE_PILAR}
    _confere_com_o_portico(r)


def test_sem_cortante_nos_casos_da_macro_nao_ha_direcao():
    """Só os casos de gravidade com os apoios: a cortante nos pés é quase nada e não orienta a
    seção — os ventos ficam de fora, com o motivo, em vez de sair com uma direção chutada."""
    dados, apoios, geometria = montar(apoios_de={c: "W90" for c in ("W0", "W90", "W180", "W270")})
    apoios = replace(apoios, casos={c: a for c, a in apoios.casos.items() if c in ("PP", "SC")})
    r = rb.montar(dados, apoios, plano=t.plano(), geometria=geometria)
    assert r.fontes("PP") == {rb.FONTE_APOIO} and not r.fontes("W0")
    assert any(f.startswith("W0:") and "sem a direção da seção" in f for f in r.faltam)


# ------------------------------------------------------------------ eixos e tabelas
def test_eixos_da_fundacao_e_do_solidworks():
    dados, apoios, geometria = montar()
    r = rb.montar(dados, apoios, plano=t.plano(), geometria=geometria)
    cab = rb.colunas(r)
    fundacao = {(l[0], l[1]): dict(zip(cab, l, strict=True)) for l in rb.linhas(r, t.plano())}
    modelo = {
        (l[0], l[1]): dict(zip(cab, l, strict=True))
        for l in rb.linhas(r, t.plano(), rb.EIXOS_MODELO)
    }
    # PP: o peso vai para baixo — Fz < 0 na fundação, Fy < 0 no SolidWorks; N positivo.
    pp = fundacao[("A1", "PP")]
    assert pp["Fz (kN)"] < 0 and pp["N (kN, compressão +)"] == pytest.approx(-pp["Fz (kN)"])
    assert modelo[("A1", "PP")]["Fy (kN)"] == pytest.approx(pp["Fz (kN)"])
    # W90 empurra para −Z do SolidWorks = +Y da fundação.
    soma_w90 = sum(v["Fy (kN)"] for k, v in fundacao.items() if k[1] == "W90")
    assert soma_w90 == pytest.approx(2.18, abs=1e-3)
    assert sum(v["Fz (kN)"] for k, v in modelo.items() if k[1] == "W90") == pytest.approx(
        -2.18, abs=1e-3
    )
    # O vento em X: pés de barlavento (x = 0) arrancam, os de sotavento comprimem.
    assert (
        fundacao[("A1", "W0")]["N (kN, compressão +)"]
        < 0
        < fundacao[("A3", "W0")]["N (kN, compressão +)"]
    )
    assert fundacao[("A1", "W0")]["Descrição do caso"] == t.plano().acao("W0").nome
    assert len(fundacao) == 4 * len(t.CASOS)
    # Locação: a planta da fundação tem o Y = −Z do SolidWorks.
    assert rb.linhas_da_locacao(r)[0][3:6] == [0.0, 1.0, -2.0]


def test_casos_com_a_natureza_e_os_coeficientes_para_combinar():
    dados, apoios, geometria = montar()
    r = rb.montar(dados, apoios, plano=t.plano(), geometria=geometria)
    casos = {
        l[0]: dict(zip(rb.COLUNAS_CASOS, l, strict=True)) for l in rb.linhas_dos_casos(r, t.plano())
    }
    assert casos["PP"]["Natureza"] == "Permanente" and casos["PP"]["γ desfavorável"] > 1
    assert casos["W0"]["ψ0"] == pytest.approx(0.6)
    assert casos["PP"]["ψ0"] == ""  # ψ não se aplica à ação permanente
    resumo = {l[0]: dict(zip(rb.COLUNAS_RESUMO, l, strict=True)) for l in rb.linhas_do_resumo(r)}
    assert resumo["A1"]["Caso"] == "SC"  # a maior compressão característica é a da sobrecarga
    # O maior arrancamento do pé A1 (x = 0, z = −1): o vento que empurra para +Z (W270) o levanta.
    assert resumo["A1"]["Caso "] == "W270" and resumo["A1"]["Maior arrancamento (kN)"] > 0


def test_excel_para_o_engenheiro_civil():
    dados, apoios, geometria = montar()
    r = rb.montar(dados, apoios, plano=t.plano(), geometria=geometria)
    conteudo = rb.xlsx(
        r, t.plano(), projeto="teste01", responsavel="Eng. Ana", hoje=date(2026, 10, 27)
    )
    livro = load_workbook(io.BytesIO(conteudo))
    assert livro.sheetnames == ["Leia-me", "Casos", "Locação", "Reações por pé", "Resumo por pé"]
    leia = [c.value for c in livro["Leia-me"]["A"] if c.value]
    assert leia[0] == "Reações nas bases para o projeto da fundação — teste01"
    assert "Data: 27/10/2026" in leia and "Responsável: Eng. Ana" in leia
    assert any("não devem ser combinadas nem majoradas" in x for x in leia)
    assert any(x.startswith("OK — W90") for x in leia)
    folha = livro["Reações por pé"]
    cabecalho = [c.value for c in folha[1]]
    assert cabecalho == rb.colunas(r) and folha.max_row == 1 + 4 * len(t.CASOS)
    assert folha.auto_filter.ref and folha.freeze_panes == "A2"
    assert rb.nome_do_arquivo("teste01 / Rev. 0") == "reacoes_nas_bases_teste01___Rev__0.xlsx"


# ------------------------------------------------------------------ pendências
def test_pendencias_que_mudam_as_reacoes():
    plano = t.plano()
    # Gravidade esquecida no W0: muda as reações → pendência e Excel PRELIMINAR.
    dados, apoios, geometria = montar(
        forcas={"W0": "PP+W0"}, reacoes={"W0": "PP+W0"}, apoios_de={"W0": "PP+W0"}
    )
    r = rb.montar(dados, apoios, plano=plano, geometria=geometria)
    d = dm.diagnosticar(dados, plano, apoios, geometria=geometria, ler_o_modelo=True)
    pendentes = rb.pendencias(d, r)
    assert [a.titulo for a in pendentes] == ["W0: ficou com a gravidade"]
    leia = [
        c.value
        for c in load_workbook(io.BytesIO(rb.xlsx(r, plano, pendentes=pendentes)))["Leia-me"]["A"]
    ]
    assert leia[0].endswith("PRELIMINAR (com pendências)")
    assert any(str(x).startswith("• W0: ficou com a gravidade") for x in leia)


def test_forcas_de_outro_caso_so_pesam_sem_a_macro():
    """O arquivo de forças do W0 com os números do W90: com os apoios da macro, as reações saem
    deles e estão certas; sem a macro, saem das forças e a pendência aparece."""
    plano = t.plano()
    dados, apoios, geometria = montar(forcas={"W0": "W90"})
    r = rb.montar(dados, apoios, plano=plano, geometria=geometria, direcoes=direcoes(dados))
    d = dm.diagnosticar(dados, plano, apoios, geometria=geometria)
    assert any(a.titulo == "W0: as forças são as do W90" for a in d.achados)
    assert rb.pendencias(d, r) == []
    _confere_com_o_portico(r)
    dados, apoios, geometria = montar(forcas={"W0": "W90"}, com_apoios=False)
    r = rb.montar(dados, apoios, plano=plano, geometria=geometria, direcoes=direcoes(dados))
    d = dm.diagnosticar(dados, plano, apoios, geometria=geometria)
    assert [a.titulo for a in rb.pendencias(d, r)] == ["W0: as forças são as do W90"]


def test_plano_com_a_grade_nao_muda_as_reacoes():

    grade = pc.nova_acao(
        "PE",
        cargas=[pc.CargaDoModelo("N1 · eixo A · 1–2", 0.585, "kN/m", "Z (vertical)", "", 4, 6.0)],
    )
    plano = pc.com_acao(t.plano(), grade)
    dados, apoios, geometria = montar()
    r = rb.montar(dados, apoios, plano=plano, geometria=geometria)
    d = dm.diagnosticar(dados, plano, apoios, geometria=geometria)
    assert any(a.titulo.startswith("Plano de cargas com cargas de exemplo") for a in d.achados)
    assert rb.pendencias(d, r) == []


# ------------------------------------------------------------------ página
def test_todo_texto_de_ajuda_e_usado():
    usadas = set(
        re.findall(r'AJUDA\["([A-Za-z_0-9]+)"\]', (RAIZ / PAGINA).read_text(encoding="utf-8"))
    )
    assert usadas == set(AJUDA)
    for chave, texto in AJUDA.items():
        assert texto.strip() == texto and 15 < len(texto) <= 900, chave


def abrir() -> AppTest:
    at = AppTest.from_file(APP, default_timeout=180)
    at.run()
    at.switch_page(PAGINA)
    at.run()
    assert not at.exception, [str(e.value) for e in at.exception]
    return at


def test_pagina_sem_resultados_aponta_para_esforcos_do_modelo(banco_com_projeto):
    at = abrir()
    assert any("ainda não tem os resultados do SolidWorks" in i.value for i in at.info)


def test_pagina_com_o_portico_e_o_excel(banco_com_projeto):
    dados, apoios, geometria = montar()
    projeto = obter_projeto_ativo()
    projeto["plano_de_cargas"] = pc.para_dicionario(t.plano())
    projeto["esforcos_do_modelo"] = em.para_dicionario(dados)
    projeto["apoios_do_modelo"] = am.para_dicionario(apoios)
    projeto["geometria_do_modelo"] = ld.para_dicionario(geometria)
    salvar_projeto(projeto, motivo="teste")
    at = abrir()
    sem_ajuda = {
        (tipo, _rotulo(e)) for tipo in TIPOS_COM_AJUDA for e in at.get(tipo) if not _ajuda(e)
    }
    sem_ajuda -= {("text_input", "Nome do novo projeto")}
    assert not sem_ajuda, sorted(sem_ajuda)
    metricas = {m.label: m.value for m in at.metric}
    assert metricas["Pés"] == "4" and metricas["Casos"] == "6"
    assert metricas["Pela macro"] == "24" and metricas["Pelos pilares"] == "0"
    assert any("Nada pendente" in s.value for s in at.success)
    tabela = next(d.value for d in at.dataframe if "De onde veio" in d.value.columns)
    assert len(tabela) == 24 and set(tabela["Pé"]) == {"A1", "A2", "A3", "A4"}
    botao = at.get("download_button")[0]
    assert "Excel" in _rotulo(botao)
    at.radio(key="rb_eixos").set_value(rb.EIXOS_MODELO).run()
    assert not at.exception
    assert any("Eixos do modelo do SolidWorks (Y para cima)" in c.value for c in at.caption)
    json.dumps(em.para_dicionario(dados))  # o projeto continua gravável
