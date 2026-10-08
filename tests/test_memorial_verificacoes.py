"""O memorial das análises com tabela de verificações: o que passou e, no fim, o que não passou.

Cobre o módulo ``core.memorial_verificacoes`` (classificação, resultado, tabelas, dados de entrada),
o modelo do memorial com várias análises no mesmo projeto (quadro-resumo, conclusão, títulos
repetidos, base repetida) e a saída em Word e PDF — inclusive o modo "para incluir em outro
documento".
"""

from __future__ import annotations

import math
import time
from io import BytesIO

import pytest
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader

from core import bolted_joint_check as chk
from core import column_buckling as cb
from core import memorial_verificacoes as mv
from core import section_catalog as catalogo
from core.memorial_blocos import PALETA_TONS, blocos_da_secao, contar_imagens
from core.project_report import (
    gerar_relatorio_industrial_pdf,
    gerar_relatorio_industrial_word,
    montar_modelo_relatorio,
)
from core.project_store import criar_item, novo_projeto_documento
from core.technical_records import criar_registro_tecnico, identificar_peca_registro

PERFIL = "W 200 x 35,9 (H)"


# ----------------------------------------------------------------------------- amostras reais
def registro_flambagem(comprimento_mm: float = 3000.0, forca_kN: float = 150.0, **extra) -> dict:
    secao = cb.secao_de_perfil(catalogo.obter_perfil(PERFIL))
    entrada = cb.EntradaColuna(
        secao=secao,
        fy_MPa=250.0,
        Lx_mm=comprimento_mm,
        Ly_mm=comprimento_mm,
        N_Sd_kN=forca_kN,
        **extra,
    )
    return cb.registro_coluna(entrada, cb.verificar_coluna(entrada))


def registro_ligacao(**mudancas) -> dict:
    base = {"designacao": "M22", "V": 32.0, "excentricidade_mm": 105.0}
    base.update(mudancas)
    entrada = chk.EntradaLigacao(**base)
    return chk.registro_ligacao(
        entrada, chk.verificar_ligacao(entrada), chk.menor_parafuso_que_atende(entrada)
    )


def projeto_com_analises() -> dict:
    projeto = novo_projeto_documento("Plataforma PM-01", codigo="MC-PM01-001")
    projeto.update(
        {
            "objetivo": "Verificar colunas e ligações.",
            "responsavel": "Eng. A",
            "verificador": "Eng. B",
            "aprovador": "Eng. C",
        }
    )
    p1 = criar_item(tag="P1", descricao="Coluna de canto", material="A36")
    p2 = criar_item(tag="P2", descricao="Coluna central", material="A36")
    l1 = criar_item(tag="L1", descricao="Ligação viga-coluna", material="A36")
    projeto["componentes"] = [p1, p2, l1]
    projeto["registros_tecnicos"] = [
        identificar_peca_registro(
            registro_flambagem(3000.0, 150.0, Mx_kNm=20.0),
            peca="Coluna P1",
            componentes_ids=[p1["id"]],
        ),
        identificar_peca_registro(
            registro_flambagem(6500.0, 700.0, Mx_kNm=35.0, My_kNm=12.0),
            peca="Coluna P2",
            componentes_ids=[p2["id"]],
        ),
        identificar_peca_registro(
            registro_ligacao(), peca="Ligação L1", componentes_ids=[l1["id"]]
        ),
        identificar_peca_registro(
            registro_ligacao(designacao="M16", V=95.0, excentricidade_mm=150.0),
            peca="Ligação L1 (reforço)",
            componentes_ids=[l1["id"]],
        ),
        criar_registro_tecnico(
            modulo="Análise estática",
            modulo_id="analise_estatica",
            titulo="Ponto crítico da chapa de base",
            status="Atende",
            resumo="von Mises no ponto crítico",
            metodo="Estado plano de tensões",
            entradas={"sigma_x_MPa": 120.0},
            resultados={"fator_seguranca": 2.1},
            conclusao="Atende ao critério informado.",
        ),
    ]
    return projeto


def texto_word(conteudo: bytes) -> str:
    documento = Document(BytesIO(conteudo))
    partes = [p.text for p in documento.paragraphs]
    partes += [c.text for t in documento.tables for linha in t.rows for c in linha.cells]
    return "\n".join(partes)


def texto_pdf(conteudo: bytes) -> str:
    return "\n".join(p.extract_text() or "" for p in PdfReader(BytesIO(conteudo)).pages)


def registro_minimo(linhas, **extra) -> dict:
    return {
        "id": extra.pop("id", "r1"),
        "modulo": "Teste",
        "modulo_id": "flambagem_colunas",
        "titulo": extra.pop("titulo", "Cálculo de teste"),
        "status": "Atende",
        "resultados": {"verificações": linhas},
        "entradas": {},
        **extra,
    }


def linha_bruta(nome, status, sol=None, res=None, aprov=None, formula="", ref="Ref"):
    return {
        "verificação": nome,
        "solicitante": sol if sol is not None else "—",
        "resistente": res if res is not None else "—",
        "unidade": "kN",
        "aproveitamento_pct": aprov if aprov is not None else "—",
        "status": status,
        "fórmula": formula or "—",
        "referência": ref,
    }


# --------------------------------------------------------------------------------- unidades
class TestStatusEDecimais:
    @pytest.mark.parametrize(
        ("bruto", "esperado"),
        [
            ("OK", "OK"),
            ("ok", "OK"),
            ("NÃO OK", "NÃO OK"),
            ("nao ok", "NÃO OK"),
            ("Não ok", "NÃO OK"),
            ("ALERTA", "ALERTA"),
            ("Atenção", "ALERTA"),
            ("INFO", "INFO"),
            ("N/A", "N/A"),
            ("", "N/A"),
            (None, "N/A"),
            ("coisa estranha", "N/A"),
        ],
    )
    def test_normaliza_status(self, bruto, esperado):
        assert mv.normalizar_status(bruto) == esperado

    @pytest.mark.parametrize(
        ("texto", "esperado"),
        [
            (
                "N_c,Rd = 0.270·1.000·1963·250/1.10 = 120.71 kN",
                "N_c,Rd = 0,270·1,000·1963·250/1,10 = 120,71 kN",
            ),
            (
                "π²·E·Ix/(Kx·Lx)² = π²·200000·3.437e+07/3000²",
                "π²·E·Ix/(Kx·Lx)² = π²·200000·3,437e+07/3000²",
            ),
            ("NBR 5.3.2 e Tab. A.3, item 6.3.12", "NBR 5.3.2 e Tab. A.3, item 6.3.12"),
            ("1,5·W·f_y = 1,10", "1,5·W·f_y = 1,10"),
            ("γ_a1 = 1.10 (Tabela 3)", "γ_a1 = 1,10 (Tabela 3)"),
            ("sem números", "sem números"),
            # item do critério Anglo e de norma: não é decimal
            ("espelho (Anglo 10.2): h = 180.5 mm", "espelho (Anglo 10.2): h = 180,5 mm"),
            ("conforme item 6.3 e Anglo 3.1", "conforme item 6.3 e Anglo 3.1"),
            ("itens 11.2", "itens 11.2"),
        ],
    )
    def test_decimal_ptbr_so_mexe_em_numero_solto(self, texto, esperado):
        assert mv.decimal_ptbr(texto) == esperado


class TestExtracaoEResumo:
    def test_registro_sem_tabela_nao_vira_capitulo_de_verificacoes(self):
        assert mv.extrair_linhas({"resultados": {"x": 1}}) is None
        assert mv.extrair_linhas({"resultados": {"verificações": []}}) is None
        assert mv.extrair_linhas({"resultados": {"verificações": "texto"}}) is None
        assert mv.extrair_linhas({"resultados": "nada"}) is None
        assert mv.extrair_linhas({}) is None

    def test_itens_malformados_sao_ignorados_e_o_resto_aproveitado(self):
        registro = registro_minimo(
            [
                "isto não é um mapa",
                {"status": "OK"},  # sem nome
                {"verificação": "   ", "status": "OK"},
                linha_bruta("Boa", "OK", 10, 20, 50.0),
            ]
        )
        linhas = mv.extrair_linhas(registro)
        assert [linha.nome for linha in linhas] == ["Boa"]
        assert linhas[0].aproveitamento == 50.0

    def test_chaves_sem_acento_tambem_servem(self):
        registro = registro_minimo(
            [
                {
                    "verificacao": "Sem acento",
                    "status": "ok",
                    "aproveitamento_pct": "12,5",
                    "formula": "x",
                }
            ]
        )
        linha = mv.extrair_linhas(registro)[0]
        assert (linha.nome, linha.status, linha.aproveitamento, linha.formula) == (
            "Sem acento",
            "OK",
            12.5,
            "x",
        )

    def test_infinito_e_traco(self):
        registro = registro_minimo(
            [linha_bruta("Zero", "NÃO OK", 5, 0, "∞"), linha_bruta("Qualitativa", "OK")]
        )
        zero, qualitativa = mv.extrair_linhas(registro)
        assert math.isinf(zero.aproveitamento) and qualitativa.aproveitamento is None

    def test_resultado_segue_o_pior_status(self):
        def resultado(*status):
            linhas = [linha_bruta(f"v{i}", s, 1, 2, 50.0) for i, s in enumerate(status)]
            return mv.resumir(mv.extrair_linhas(registro_minimo(linhas)))

        assert resultado("OK", "OK", "INFO").resultado == "ATENDE"
        assert resultado("OK", "ALERTA").resultado == "ATENÇÃO"
        assert resultado("OK", "ALERTA", "NÃO OK").resultado == "NÃO ATENDE"
        assert resultado("INFO", "N/A").resultado == "NÃO AVALIADO"
        assert resultado("NÃO OK").tom == "erro" and resultado("OK").tom == "ok"
        assert resultado("ALERTA").tom == "atencao" and resultado("N/A").tom == "neutro"

    def test_governante_e_a_maior_utilizacao_com_criterio(self):
        linhas = mv.extrair_linhas(
            registro_minimo(
                [
                    linha_bruta("a", "OK", 1, 2, 50.0),
                    linha_bruta("b", "NÃO OK", 3, 2, 150.0),
                    linha_bruta("info", "INFO", 1, 1, 999.0),
                ]
            )
        )
        governante = mv.resumir(linhas).governante
        assert governante.nome == "b" and governante.aproveitamento == 150.0

    def test_texto_do_resultado_conta_e_pluraliza(self):
        formatar = str
        um = mv.resumir(mv.extrair_linhas(registro_minimo([linha_bruta("a", "OK", 1, 2, 50.0)])))
        assert mv.texto_do_resultado({}, um, formatar).startswith(
            "1 verificação com critério: 1 passou."
        )
        varios = mv.resumir(
            mv.extrair_linhas(
                registro_minimo(
                    [
                        linha_bruta("a", "OK", 1, 2, 50.0),
                        linha_bruta("b", "NÃO OK", 3, 2, 150.0),
                        linha_bruta("c", "NÃO OK", 3, 2, 150.0),
                        linha_bruta("d", "ALERTA", 1, 2, 50.0),
                        linha_bruta("e", "N/A"),
                    ]
                )
            )
        )
        texto = mv.texto_do_resultado({}, varios, formatar)
        assert "4 verificações com critério: 1 passou, 2 não passaram, 1 em atenção." in texto
        assert "Maior aproveitamento: 150,0% em «b»." in texto and "1 não avaliada." in texto

    def test_parafuso_que_atende_e_a_ausencia_dele(self):
        resumo = mv.resumir(
            mv.extrair_linhas(registro_minimo([linha_bruta("a", "OK", 1, 2, 50.0)]))
        )
        com = mv.texto_do_resultado(
            {"resultados": {"menor_parafuso_que_atende": "M24"}}, resumo, str
        )
        sem = mv.texto_do_resultado(
            {"resultados": {"menor_parafuso_que_atende": "nenhum com esta geometria"}}, resumo, str
        )
        assert "Menor parafuso que atende esta geometria: M24." in com
        assert "Nenhum parafuso da tabela atende com esta geometria." in sem
        assert "nenhum com esta geometria" not in sem


class TestDadosDeEntrada:
    def test_rotulos_ordem_e_omissao_do_que_nao_foi_preenchido(self):
        registro = {
            "modulo_id": "flambagem_colunas",
            "entradas": {
                "kx": 1.0,
                "norma": "NBR8800_2008",
                "secao": "W 200 x 35,9",
                "excentricidade_x_mm": 0.0,  # se_preenchido: some
                "momento_x_kNm": 0.0,  # sempre: fica, mesmo zerado
                "secao_compacta_confirmada": False,
                "area_mm2": 4570.0,
                "mao_francesa": "não incluída",
                "material_id": "uuid-que-nao-deve-aparecer",
            },
        }
        pares = mv.pares_de_entrada(registro, str, lambda chave: f"auto:{chave}")
        rotulos = [rotulo for rotulo, _ in pares]
        assert rotulos == [
            "Norma",
            "Seção",
            "Área A [mm²]",
            "Coeficiente Kx",
            "Momento Mx,Sd [kN·m]",
        ]
        assert dict(pares)["Norma"] == "ABNT NBR 8800:2008"

    def test_dado_novo_que_o_mapa_nao_conhece_nao_some(self):
        registro = {
            "modulo_id": "flambagem_colunas",
            "entradas": {"secao": "W", "dado_do_futuro_mm": 12.5, "vazio_do_futuro": None},
        }
        pares = dict(mv.pares_de_entrada(registro, str, lambda chave: f"auto:{chave}"))
        assert pares["auto:dado_do_futuro_mm"] == "12.5" and "Seção" in pares
        assert not any("vazio" in rotulo for rotulo in pares)

    def test_booleanos_e_codigos(self):
        registro = {
            "modulo_id": "projeto_parafusos",
            "entradas": {
                "rosca_no_plano_corte": True,
                "ligacao_por_atrito": False,
                "superficie": "galvanizada_sem_tratamento",
            },
        }
        pares = dict(mv.pares_de_entrada(registro, str, str))
        assert pares["Rosca no plano de corte"] == "Sim"
        assert pares["Ligação por atrito"] == "Não"
        assert pares["Superfície de contato"] == "galvanizada sem tratamento"

    def test_modulo_desconhecido_mostra_tudo_o_que_nao_e_nulo(self):
        registro = {"modulo_id": "outro", "entradas": {"a_mm": 0, "b": None, "c_id": "x", "d": "t"}}
        pares = dict(mv.pares_de_entrada(registro, str, lambda chave: chave))
        assert pares == {"a_mm": "0", "d": "t"}

    def test_tabela_de_entradas_tem_quatro_colunas_e_completa_a_ultima_linha(self):
        tabela = mv.tabela_de_entradas([("a", "1"), ("b", "2"), ("c", "3")], legenda="x")
        assert tabela["cabecalhos"] == ["Dado", "Valor", "Dado", "Valor"]
        assert tabela["linhas"] == [["a", "1", "b", "2"], ["c", "3", "", ""]]
        assert sum(tabela["larguras"]) == 9360
        assert mv.tabela_de_entradas([], legenda="x")["linhas"] == []


class TestBaseEAlertas:
    def test_alerta_que_e_eco_de_uma_linha_nao_se_repete(self):
        registro = registro_minimo(
            [
                linha_bruta("Flexão", "NÃO OK", 2, 1, 200.0, "M > Mrd"),
                linha_bruta("Ok", "OK", 1, 2, 50.0),
            ],
            alertas=["Flexão: M > Mrd", "Conferir na NBR antes de emitir: Tab. A.3 e 5.2.5."],
        )
        linhas = mv.extrair_linhas(registro)
        assert mv.alertas_livres(registro, linhas) == [
            "Conferir na NBR antes de emitir: Tab. A.3 e 5.2.5."
        ]

    def test_equacoes_gerais_excluem_as_que_ja_estao_nas_linhas(self):
        registro = registro_minimo(
            [linha_bruta("A", "OK", 1, 2, 50.0, "F = 1.0·x")],
            equacoes=["N_Sd = 1.4·N_G", "F = 1.0·x"],
        )
        linhas = mv.extrair_linhas(registro)
        assert mv.equacoes_gerais(registro, linhas) == ["N_Sd = 1,4·N_G"]

    def test_assinatura_igual_para_bases_iguais_e_diferente_quando_algo_muda(self):
        a = registro_minimo([linha_bruta("A", "OK", 1, 2, 50.0)], metodo="m", premissas=["p"])
        b = registro_minimo([linha_bruta("B", "NÃO OK", 3, 2, 150.0)], metodo="m", premissas=["p"])
        c = registro_minimo([linha_bruta("A", "OK", 1, 2, 50.0)], metodo="m", premissas=["outra"])
        la, lb, lc = (mv.extrair_linhas(r) for r in (a, b, c))
        assert mv.assinatura_da_base(a, la) == mv.assinatura_da_base(b, lb)
        assert mv.assinatura_da_base(a, la) != mv.assinatura_da_base(c, lc)


# ------------------------------------------------------------------------------- o capítulo
def capitulo(registro, **extra):
    linhas = mv.extrair_linhas(registro)
    return mv.capitulo_de_verificacoes(
        registro,
        linhas,
        titulo="8.1.1 Cálculo",
        nivel=3,
        peca="Peça X",
        formatar=str,
        rotular=lambda chave: chave,
        **extra,
    )


def rotulos(secao):
    return [
        (bloco["tipo"], bloco.get("texto") or bloco.get("rotulo") or bloco.get("legenda"))
        for bloco in secao["blocos"]
    ]


class TestCapitulo:
    @pytest.fixture
    def com_reprovacao(self):
        return registro_minimo(
            [
                linha_bruta("Passa", "OK", 1, 2, 50.0, "a = 1.0", "Ref A"),
                linha_bruta("Reprova", "NÃO OK", 3, 2, 150.0, "motivo do erro", "Ref B"),
                linha_bruta("Pede conferência", "ALERTA", 1, 1, 100.0),
                linha_bruta("Sem dados", "N/A", formula="faltou o dado"),
                linha_bruta("Apoio", "INFO", None, 7.0),
            ],
            alertas=["Reprova: motivo do erro", "Conferir o item 5.3.2 antes de emitir."],
        )

    def test_abre_com_o_resultado_e_fecha_com_o_que_nao_passou(self, com_reprovacao):
        secao = capitulo(com_reprovacao)
        blocos = secao["blocos"]
        assert blocos[0]["tipo"] == "destaque" and blocos[0]["rotulo"] == "Resultado: NÃO ATENDE."
        assert blocos[0]["tom"] == "erro"
        textos = [b.get("texto") for b in blocos if b["tipo"] == "subtitulo"]
        # A ordem dos trechos é a do pedido: o que passou primeiro, e o que não passou — com o que
        # pede atenção e o que conferir — no fim.
        assert textos == [
            "O que passou (1)",
            "Não passou (1)",
            "Atenção — passou com ressalva ou precisa de conferência (1)",
            "Não avaliadas (1)",
            "Conferir antes de emitir",
        ]
        # Depois do "Não passou" só vêm tabelas e listas do próprio bloco final.
        indice_final = next(i for i, b in enumerate(blocos) if b.get("texto") == "Não passou (1)")
        assert {b["tipo"] for b in blocos[indice_final:]} <= {"subtitulo", "tabela", "bullets"}
        passou = next(i for i, b in enumerate(blocos) if b.get("texto") == "O que passou (1)")
        assert all(
            b.get("tom") != "erro" for b in blocos[passou:indice_final] if b["tipo"] == "tabela"
        )

    def test_a_tabela_dos_que_passaram_so_tem_os_que_passaram(self, com_reprovacao):
        tabelas = [b for b in capitulo(com_reprovacao)["blocos"] if b["tipo"] == "tabela"]
        passou = next(t for t in tabelas if t.get("tom") == "ok")
        reprovou = next(t for t in tabelas if t.get("tom") == "erro")
        assert [linha[0] for linha in passou["linhas"]] == ["Passa"]
        assert [linha[0] for linha in reprovou["linhas"]] == ["Reprova"]
        assert sum(passou["larguras"]) == 9360 and len(passou["cabecalhos"]) == len(
            passou["linhas"][0]
        )
        # O cálculo vai com a referência numa célula só, e o decimal vira vírgula.
        assert passou["linhas"][0][5] == ("a = 1,0", "Ref.: Ref A")

    def test_sem_reprovacao_declara_que_nada_reprovou(self):
        secao = capitulo(registro_minimo([linha_bruta("Passa", "OK", 1, 2, 50.0)]))
        ultimo = secao["blocos"][-1]
        assert ultimo["tipo"] == "destaque" and ultimo["rotulo"] == "Não passou: nenhuma."
        assert ultimo["tom"] == "ok"
        assert secao["blocos"][0]["rotulo"] == "Resultado: ATENDE."

    def test_so_informativas_nao_afirma_que_passou_nem_que_reprovou(self):
        secao = capitulo(registro_minimo([linha_bruta("Apoio", "INFO", None, 7.0)]))
        assert secao["blocos"][0]["rotulo"] == "Resultado: NÃO AVALIADO."
        assert secao["blocos"][-1]["tom"] == "neutro"
        assert "Nenhuma verificação com critério passou." in str(secao["blocos"])

    def test_imagem_fica_antes_do_que_passou(self):
        imagem = {"png": b"x", "legenda": "Curva", "largura_pol": 6.0}
        secao = capitulo(
            registro_minimo([linha_bruta("Passa", "OK", 1, 2, 50.0)]), imagens=[imagem]
        )
        tipos = [b["tipo"] for b in secao["blocos"]]
        assert tipos.index("imagem") < next(
            i for i, b in enumerate(secao["blocos"]) if b.get("texto") == "O que passou (1)"
        )
        assert contar_imagens([secao]) == 1

    def test_base_repetida_vira_referencia_ao_item_anterior(self):
        registro = registro_minimo(
            [linha_bruta("Passa", "OK", 1, 2, 50.0)],
            metodo="Método X",
            premissas=["P1"],
            criterios=["C1"],
            referencias=["R1"],
        )
        primeira = capitulo(registro)
        segunda = capitulo(registro, base_repetida_de="8.1.1")
        assert any(
            b["tipo"] == "bullets" and "Premissa: P1" in b["itens"] for b in primeira["blocos"]
        )
        assert not any(b["tipo"] == "bullets" for b in segunda["blocos"])
        assert "os mesmos do item 8.1.1" in str(segunda["blocos"])

    def test_identificacao_traz_modulo_peca_e_situacao(self):
        secao = capitulo(registro_minimo([linha_bruta("a", "OK", 1, 2, 50.0)]))
        assert "Módulo: Teste. Peça: Peça X. Situação registrada: Atende." in str(secao["blocos"])

    def test_blocos_explicitos_sao_usados_como_estao(self):
        secao = capitulo(registro_minimo([linha_bruta("a", "OK", 1, 2, 50.0)]))
        assert blocos_da_secao(secao) == secao["blocos"]

    def test_secao_antiga_mantem_a_ordem_historica(self):
        secao = {
            "paragrafos": ["p1", "p2"],
            "nota": "n",
            "bullets": ["b"],
            "formulas": ["f"],
            "imagens": [{"png": b"x"}],
            "tabelas": [{"cabecalhos": ["a"], "linhas": [["1"]], "larguras": [9360]}],
            "paragrafos_finais": ["fim"],
        }
        tipos = [b["tipo"] for b in blocos_da_secao(secao)]
        assert tipos == [
            "paragrafo",
            "paragrafo",
            "nota",
            "bullets",
            "formula",
            "imagem",
            "tabela",
            "paragrafo",
        ]
        assert blocos_da_secao(secao)[1]["manter_com_proximo"] is True  # antes de uma tabela
        assert blocos_da_secao(secao)[-1]["texto"] == "fim"


# --------------------------------------------------------------------- o memorial do projeto
@pytest.fixture(scope="module")
def modelo():
    return montar_modelo_relatorio(projeto_com_analises())


class TestModeloComVariasAnalises:
    def test_cada_analise_de_verificacao_e_um_capitulo_com_resultado(self, modelo):
        capitulos = [s for s in modelo["secoes"] if "blocos" in s]
        assert len(capitulos) == 4  # duas flambagens e duas ligações; a estática é genérica
        resultados = [s["blocos"][0]["rotulo"] for s in capitulos]
        assert resultados == [
            "Resultado: ATENDE.",
            "Resultado: NÃO ATENDE.",
            "Resultado: NÃO ATENDE.",
            "Resultado: NÃO ATENDE.",
        ]

    def test_analise_generica_continua_no_formato_de_sempre(self, modelo):
        generica = next(s for s in modelo["secoes"] if "Ponto crítico" in s["titulo"])
        assert "blocos" not in generica and "paragrafos_finais" in generica
        assert generica["paragrafos_finais"] == ["Conclusão: Atende ao critério informado."]

    def test_quadro_resumo_conta_o_que_passou_e_o_que_nao_passou(self, modelo):
        quadro = next(s for s in modelo["secoes"] if "Quadro-resumo" in s["titulo"])
        tabela = quadro["tabelas"][0]
        assert tabela["cabecalhos"] == [
            "Ordem",
            "Peça",
            "Cálculo",
            "Passou",
            "Não passou",
            "Aprov. máx.",
            "Situação",
        ]
        assert sum(tabela["larguras"]) == 9360
        por_peca = {linha[1]: linha for linha in tabela["linhas"]}
        assert por_peca["Coluna P1"][4] == "0" and por_peca["Coluna P1"][6] == "Atende"
        assert int(por_peca["Coluna P2"][4]) >= 1 and por_peca["Coluna P2"][6] == "Não atende"
        assert por_peca["Ligação L1 (reforço)"][3].isdigit()
        generica = next(linha for linha in tabela["linhas"] if linha[1] == "-")
        assert generica[3:6] == ["-", "-", "-"]
        # O módulo continua visível, sob o título do cálculo.
        assert por_peca["Coluna P1"][2][1] == "Flambagem de colunas"

    def test_conclusao_lista_o_que_nao_passou_em_todo_o_projeto(self, modelo):
        conclusao = modelo["secoes"][-1]
        tabela = conclusao["tabelas"][0]
        assert tabela["cabecalhos"] == ["Peça", "Cálculo", "Verificação", "Aprov."]
        assert tabela["tom"] == "erro"
        pecas = {linha[0] for linha in tabela["linhas"]}
        assert pecas == {"Coluna P2", "Ligação L1", "Ligação L1 (reforço)"}
        assert "Coluna P1" not in pecas
        # A análise que passou não aparece, e a conclusão geral segue em aberto.
        assert conclusao["paragrafos_finais"][0].startswith("Conclusão geral:")

    def test_resumo_executivo_aponta_os_calculos_que_nao_passaram(self, modelo):
        linhas = {linha[0]: linha for linha in modelo["resumo_executivo"]["linhas"]}
        assert linhas["Não passou"][1] == "3 cálculo(s)"
        assert "Coluna P2" in linhas["Não passou"][2] and "Coluna P1" not in linhas["Não passou"][2]

    def test_tudo_passando_nao_cria_tabela_nem_linha_de_reprovacao(self):
        projeto = projeto_com_analises()
        projeto["registros_tecnicos"] = projeto["registros_tecnicos"][:1]
        modelo = montar_modelo_relatorio(projeto)
        assert "tabelas" not in modelo["secoes"][-1] or not modelo["secoes"][-1]["tabelas"]
        assert "Não passou" not in {linha[0] for linha in modelo["resumo_executivo"]["linhas"]}

    def test_a_segunda_analise_do_mesmo_modulo_cita_a_base_da_primeira(self, modelo):
        capitulos = [s for s in modelo["secoes"] if "blocos" in s]
        segunda_flambagem = capitulos[1]
        primeira = capitulos[0]["titulo"].split(" ")[0]
        assert f"os mesmos do item {primeira}" in str(segunda_flambagem["blocos"])
        assert "os mesmos do item" not in str(capitulos[0]["blocos"])

    def test_titulos_repetidos_ganham_numero_do_calculo(self):
        projeto = projeto_com_analises()
        repetida = registro_flambagem(3000.0, 150.0, Mx_kNm=20.0)
        repetida["id"] = "outro-id"
        projeto["registros_tecnicos"] = [
            identificar_peca_registro(
                registro_flambagem(3000.0, 150.0, Mx_kNm=20.0), peca="Coluna P1"
            ),
            identificar_peca_registro(repetida, peca="Coluna P1"),
        ]
        modelo = montar_modelo_relatorio(projeto)
        titulos = [s["titulo"] for s in modelo["secoes"] if "blocos" in s]
        assert titulos[0].endswith("(cálculo 1 de 2)") and titulos[1].endswith("(cálculo 2 de 2)")

    def test_registros_com_o_mesmo_id_nao_se_confundem(self):
        projeto = projeto_com_analises()
        base = registro_flambagem(3000.0, 150.0)
        projeto["registros_tecnicos"] = [
            identificar_peca_registro(base, peca="Coluna A"),
            identificar_peca_registro(base, peca="Coluna B"),
        ]
        modelo = montar_modelo_relatorio(projeto)
        titulos = [s["titulo"] for s in modelo["secoes"] if "blocos" in s]
        assert titulos[0].split(" — ")[0].endswith("Coluna A")
        assert titulos[1].split(" — ")[0].endswith("Coluna B")

    def test_modo_para_incluir_em_outro_documento_numera_do_um_e_sem_aprovacoes(self):
        modelo = montar_modelo_relatorio(
            projeto_com_analises(),
            secoes_incluidas=["plano_calculo", "registros", "conclusao"],
            somente_capitulos=True,
        )
        titulos = [s["titulo"] for s in modelo["secoes"] if s.get("nivel", 1) == 1]
        assert titulos[0].startswith("1. ") and titulos[1].startswith("2. ")
        assert modelo["metadata"]["somente_capitulos"] is True
        completo = montar_modelo_relatorio(projeto_com_analises())
        assert completo["metadata"]["somente_capitulos"] is False
        assert [s["titulo"] for s in completo["secoes"]][0].startswith("3. ")


# ------------------------------------------------------------------------------ Word e PDF
@pytest.fixture(scope="module")
def projeto_completo():
    return projeto_com_analises()


class TestSaidaWord:
    def test_texto_do_que_passou_e_do_que_nao_passou(self, projeto_completo):
        texto = texto_word(gerar_relatorio_industrial_word(projeto_completo))
        for esperado in (
            "Resultado: ATENDE.",
            "Resultado: NÃO ATENDE.",
            "O que passou (",
            "Não passou (",
            "Não passou: nenhuma.",
            "O que não passou, por cálculo.",
            "Conferir antes de emitir",
        ):
            assert esperado in texto, esperado
        for proibido in (
            "Não informado",
            "bloqueio",
            "Snapshot",
            "Apêndice",
            "Central de validação",
        ):
            assert proibido not in texto, proibido

    def test_destaque_e_cabecalhos_usam_a_cor_do_tom(self, projeto_completo):
        documento = Document(BytesIO(gerar_relatorio_industrial_word(projeto_completo)))
        fundos_de_paragrafo = {
            p._p.find(".//" + qn("w:shd")).get(qn("w:fill"))
            for p in documento.paragraphs
            if p.text.startswith("Resultado:")
        }
        assert fundos_de_paragrafo == {PALETA_TONS["ok"]["fundo"], PALETA_TONS["erro"]["fundo"]}
        fundos_de_cabecalho = {
            t.rows[0].cells[0]._tc.find(".//" + qn("w:shd")).get(qn("w:fill"))
            for t in documento.tables
            if t.rows[0].cells[0].text == "Verificação"
        }
        assert {PALETA_TONS["ok"]["fundo"], PALETA_TONS["erro"]["fundo"]} <= fundos_de_cabecalho

    def test_o_fim_de_cada_analise_com_tabela_e_o_que_nao_passou(self, projeto_completo):
        documento = Document(BytesIO(gerar_relatorio_industrial_word(projeto_completo)))
        corpo = list(documento.element.body.iterchildren())
        textos = []
        for elemento in corpo:
            if elemento.tag == qn("w:p"):
                textos.append("".join(t.text or "" for t in elemento.iter(qn("w:t"))).strip())
        indices = {
            titulo: i
            for i, texto in enumerate(textos)
            for titulo in ("8.2.1", "8.3.1", "8.3.2", "8.4.1")
            if texto.startswith(titulo)
        }
        for inicio, fim in (("8.2.1", "8.3.1"), ("8.3.1", "8.3.2"), ("8.3.2", "8.4.1")):
            trecho = textos[indices[inicio] : indices[fim]]
            # Dentro do capítulo, "Não passou" vem depois de "O que passou".
            passou = next(i for i, t in enumerate(trecho) if t.startswith("O que passou"))
            nao = next(i for i, t in enumerate(trecho) if t.startswith("Não passou"))
            assert passou < nao

    def test_somente_capitulos_nao_leva_capa_resumo_controle_nem_aprovacoes(self, projeto_completo):
        completo = texto_word(gerar_relatorio_industrial_word(projeto_completo))
        so_capitulos = texto_word(
            gerar_relatorio_industrial_word(
                projeto_completo,
                secoes_incluidas=["plano_calculo", "registros", "conclusao"],
                somente_capitulos=True,
            )
        )
        for presente in (
            "Resumo executivo",
            "Controle do documento",
            "Aprovações",
            "Data de aprovação",
        ):
            assert presente in completo
            assert presente not in so_capitulos, presente
        assert "MEMÓRIA DE CÁLCULO · Plataforma PM-01 · MC-PM01-001" in so_capitulos
        assert "1. Quadro-resumo dos cálculos" in so_capitulos
        assert "O que passou (" in so_capitulos and "Não passou (" in so_capitulos


class TestSaidaPdf:
    def test_texto_do_que_passou_e_do_que_nao_passou(self, projeto_completo):
        pdf = gerar_relatorio_industrial_pdf(projeto_completo)
        assert pdf.startswith(b"%PDF")
        texto = texto_pdf(pdf)
        for esperado in (
            "Resultado: ATENDE.",
            "Resultado: NÃO ATENDE.",
            "O que passou (",
            "Não passou (",
            "Não passou: nenhuma.",
            "O que não passou, por cálculo.",
        ):
            assert esperado in texto, esperado
        for proibido in ("Não informado", "bloqueio", "Snapshot", "Apêndice"):
            assert proibido not in texto, proibido

    def test_somente_capitulos_no_pdf(self, projeto_completo):
        texto = texto_pdf(
            gerar_relatorio_industrial_pdf(
                projeto_completo,
                secoes_incluidas=["plano_calculo", "registros", "conclusao"],
                somente_capitulos=True,
            )
        )
        for ausente in ("Resumo executivo", "Controle do documento", "Aprovações"):
            assert ausente not in texto, ausente
        assert "MEMÓRIA DE CÁLCULO" in texto and "1. Quadro-resumo dos cálculos" in texto

    def test_celula_com_detalhe_vira_segunda_linha(self, projeto_completo):
        texto = texto_pdf(gerar_relatorio_industrial_pdf(projeto_completo))
        assert "Ref.: NBR 5.3.2 / AISC E3" in texto


class TestRobustez:
    def test_tabela_malformada_cai_no_capitulo_generico_sem_derrubar_o_memorial(self):
        projeto = projeto_com_analises()
        projeto["registros_tecnicos"] = [
            criar_registro_tecnico(
                modulo="Flambagem de colunas",
                modulo_id="flambagem_colunas",
                titulo="Registro quebrado",
                status="Atende",
                resumo="",
                entradas={"a": 1},
                resultados={"verificações": ["lixo", 3, None]},
                conclusao="Atende.",
            )
        ]
        modelo = montar_modelo_relatorio(projeto)
        assert not any("blocos" in s for s in modelo["secoes"])
        assert gerar_relatorio_industrial_word(projeto).startswith(b"PK")
        assert gerar_relatorio_industrial_pdf(projeto).startswith(b"%PDF")

    def test_texto_longo_e_caracteres_especiais_nao_quebram_word_nem_pdf(self):
        longo = "Verificação " + "muito longa " * 80 + "<b>&</b> \"aspas\" 'simples' ≤ ≥ γ_a1 μ"
        registro = registro_minimo(
            [
                linha_bruta(longo, "NÃO OK", 3, 2, 150.0, longo, longo),
                linha_bruta("Ok", "OK", 1, 2, 50.0),
            ],
            titulo=longo[:150],
            alertas=[longo],
        )
        projeto = projeto_com_analises()
        projeto["registros_tecnicos"] = [registro]
        assert gerar_relatorio_industrial_word(projeto).startswith(b"PK")
        assert gerar_relatorio_industrial_pdf(projeto).startswith(b"%PDF")

    def test_muitas_analises_no_mesmo_projeto(self):
        projeto = projeto_com_analises()
        modelo_base = projeto["registros_tecnicos"]
        registros = []
        for indice in range(30):
            registro = dict(modelo_base[indice % 4])
            registro["id"] = f"calc-{indice}"
            registros.append(identificar_peca_registro(registro, peca=f"Peça {indice}"))
        projeto["registros_tecnicos"] = registros
        inicio = time.perf_counter()
        modelo = montar_modelo_relatorio(projeto)
        capitulos = [s for s in modelo["secoes"] if "blocos" in s]
        assert len(capitulos) == 30
        titulos = [s["titulo"] for s in capitulos]
        assert len(set(titulos)) == 30
        assert gerar_relatorio_industrial_word(projeto).startswith(b"PK")
        assert gerar_relatorio_industrial_pdf(projeto).startswith(b"%PDF")
        assert time.perf_counter() - inicio < 120  # folga enorme: só pega laço infinito

    def test_registro_com_status_desconhecido_vira_nao_avaliado_e_nao_derruba(self):
        registro = registro_minimo([linha_bruta("Estranha", "talvez", 1, 2, 50.0)])
        secao = capitulo(registro)
        assert secao["blocos"][0]["rotulo"] == "Resultado: NÃO AVALIADO."
        assert "Estranha" in str(secao["blocos"])


def test_resultados_que_o_capitulo_nao_conhece_aparecem_em_outros_resultados():
    registro = registro_minimo([linha_bruta("a", "OK", 1, 2, 50.0)])
    registro["resultados"].update(
        {
            "pre_carga_nominal_kN": 55.5,
            "tem_folga": True,
            "status_geral": "OK",  # já está no destaque
            "chi": 0.7,  # já está na tabela de apoio
            "lista_qualquer": [1, 2],  # listas viram tabela própria, não entram aqui
            "id_interno_id": "uuid",
        }
    )
    secao = capitulo(registro)
    tabela = next(
        b for b in secao["blocos"] if b.get("legenda") == "Outros resultados registrados."
    )
    assert tabela["cabecalhos"] == ["Resultado", "Valor", "Resultado", "Valor"]
    assert tabela["linhas"] == [["pre_carga_nominal_kN", "55.5", "tem_folga", "Sim"]]


def test_nenhuma_linha_da_tabela_do_modulo_some_do_memorial():
    """Cada verificação do registro aparece exatamente uma vez no capítulo, na tabela do seu status."""
    modelo = montar_modelo_relatorio(projeto_com_analises())
    capitulos = [s for s in modelo["secoes"] if "blocos" in s]
    registros = [r for r in modelo["registros"] if mv.extrair_linhas(r) is not None]
    assert len(capitulos) == len(registros) == 4
    for secao, registro in zip(capitulos, registros, strict=True):
        resumo = mv.resumir(mv.extrair_linhas(registro))
        tabelas = [b for b in secao["blocos"] if b["tipo"] == "tabela"]
        nomes_por_tom = {
            tom: [linha[0] for t in tabelas if t.get("tom") == tom for linha in t["linhas"]]
            for tom in ("ok", "erro", "atencao")
        }
        assert nomes_por_tom["ok"] == [linha.nome for linha in resumo.passaram]
        assert nomes_por_tom["erro"] == [linha.nome for linha in resumo.reprovadas]
        assert nomes_por_tom["atencao"] == [linha.nome for linha in resumo.atencao]
        apoio = [
            linha[0] for t in tabelas if t["cabecalhos"][0] == "Grandeza" for linha in t["linhas"]
        ]
        assert apoio == [linha.nome for linha in resumo.informativas]
        if resumo.nao_avaliadas:
            blocos = secao["blocos"]
            titulo = f"Não avaliadas ({len(resumo.nao_avaliadas)})"
            indice = next(i for i, b in enumerate(blocos) if b.get("texto") == titulo)
            assert len(blocos[indice + 1]["itens"]) == len(resumo.nao_avaliadas)
        total = len(resumo.passaram) + len(resumo.reprovadas) + len(resumo.atencao)
        total += len(resumo.informativas) + len(resumo.nao_avaliadas)
        assert total == len(mv.extrair_linhas(registro))


def test_o_snapshot_distingue_o_documento_completo_do_so_com_capitulos():
    projeto = projeto_com_analises()
    secoes = ["plano_calculo", "registros", "conclusao"]
    completo = montar_modelo_relatorio(projeto, secoes_incluidas=secoes)
    capitulos = montar_modelo_relatorio(projeto, secoes_incluidas=secoes, somente_capitulos=True)
    assert completo["snapshot_hash"] != capitulos["snapshot_hash"]
    assert (
        capitulos["snapshot_hash"]
        == montar_modelo_relatorio(projeto, secoes_incluidas=secoes, somente_capitulos=True)[
            "snapshot_hash"
        ]
    )
