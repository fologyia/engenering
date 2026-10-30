"""O Word do fim da página Flambagem de colunas: tudo o que a página mostra, editável."""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from core import column_buckling as cb
from core import desenho_da_coluna as dc
from core import section_catalog as sc
from core.coluna_word import gerar_word_coluna, linhas_do_resumo

PALAVRAS_PROIBIDAS = ("Não informado", "bloqueio", "Snapshot", "Apêndice")


def texto(conteudo: bytes) -> str:
    documento = Document(BytesIO(conteudo))
    partes = [p.text for p in documento.paragraphs]
    partes += [c.text for t in documento.tables for linha in t.rows for c in linha.cells]
    return "\n".join(partes)


@pytest.fixture(scope="module")
def coluna() -> tuple[cb.EntradaColuna, cb.ResultadoColuna]:
    secao = cb.secao_de_perfil(sc.listar_perfis()["W 250 x 32,7"])
    entrada = cb.EntradaColuna(
        secao=secao, fy_MPa=345.0, norma="NBR8800_2008", Lx_mm=4000.0, Ly_mm=2000.0,
        N_Sd_kN=300.0, Mx_kNm=40.0, My_kNm=3.0, razao_m1_m2_x=0.0,
    )  # fmt: skip
    return entrada, cb.verificar_coluna(entrada)


def test_word_traz_todas_as_secoes_da_pagina(coluna):
    entrada, resultado = coluna
    conteudo = gerar_word_coluna(
        entrada, resultado, equacoes_de_carga=["N_Sd = 1,4·100 + 1,5·106,7 = 300 kN"],
        projeto={"nome": "Plataforma X", "codigo": "PX-01"},
    )  # fmt: skip
    cheio = texto(conteudo)
    for trecho in (
        "Flambagem de colunas — W 250 x 32,7",
        "Plataforma X",
        "PX-01",
        "O que esta página não faz",
        "Dados de entrada",
        "Propriedades da seção",
        "N_Sd = 1,4·100 + 1,5·106,7 = 300 kN",
        "1. Resumo",
        "2. Curva de flambagem",
        "3. Tabela de verificações",
        "4. Paredes da seção",
        "5. A mesma coluna nas três normas",
        "6. Regras adotadas nesta verificação",
        "Valores da NBR 8800:2008 conferidos na norma",
        "7. Método, premissas, critérios e referências",
    ):
        assert trecho in cheio, trecho
    for item in cb.FORA_DO_ESCOPO:
        assert item in cheio
    for v in resultado.verificacoes:
        assert v.nome in cheio
    for rotulo in cb.NORMAS_ROTULOS.values():
        assert rotulo in cheio
    for proibida in PALAVRAS_PROIBIDAS:
        assert proibida not in cheio, proibida
    assert "γ_a1 = 1,10" in cheio and "N_c,Rd = 835,24 kN" in cheio  # vírgula decimal
    assert "NBR 8800 4.7" in cheio  # item da norma não vira decimal
    documento = Document(BytesIO(conteudo))
    assert len(documento.inline_shapes) == 2  # o desenho da seção e a curva de flambagem
    assert "W 250 x 32,7: perfil I laminado" in cheio  # a legenda do desenho da seção


def test_word_leva_o_desenho_da_secao_girado(coluna):
    entrada, resultado = coluna
    girado = texto(gerar_word_coluna(entrada, resultado, giro_da_secao=90))
    assert "desenho girado 90° no sentido anti-horário" in girado


def test_word_leva_a_coluna_em_pe_com_a_mao_francesa(coluna):
    entrada, resultado = coluna
    em_pe = dc.ColunaEmPe(
        entrada.secao, 4000.0, 4000.0, 2000.0, 1.0, 1.0, "Biapoiada (pino-pino)",
        "Biapoiada (pino-pino)", dc.MaoFrancesaNoDesenho(20.0, 45.0, 3000.0, "x", 120.0),
    )  # fmt: skip
    conteudo = gerar_word_coluna(entrada, resultado, coluna_em_pe=em_pe)
    assert len(Document(BytesIO(conteudo)).inline_shapes) == 3  # seção, coluna em pé e curva
    assert "mão-francesa a 3000 mm da base" in texto(conteudo)


def test_resumo_igual_ao_da_pagina(coluna):
    entrada, resultado = coluna
    linhas = dict(map(tuple, linhas_do_resumo(entrada, resultado)))
    comp = resultado.compressao
    assert linhas["Status geral"] == resultado.status_geral
    assert linhas["N_c,Rd"] == f"{comp['Nc_Rd']:.2f} kN".replace(".", ",")
    assert linhas["χ"] == f"{comp['chi']:.3f}".replace(".", ",").rstrip("0")
    assert "B₁ em y" in linhas and "M_x,Rd" in linhas and "Índice de interação" in linhas
    assert linhas["Verificação que governa"] == resultado.governante


def test_word_sem_flexao_sem_projeto_e_noutra_norma():
    secao = cb.secao_de_perfil(sc.listar_perfis()["W 250 x 32,7"])
    entrada = cb.EntradaColuna(
        secao=secao, fy_MPa=250.0, norma="NBR8800_2024", Lx_mm=3000.0, Ly_mm=3000.0,
        N_Sd_kN=100.0,
    )  # fmt: skip
    conteudo = gerar_word_coluna(entrada, cb.verificar_coluna(entrada))
    cheio = texto(conteudo)
    assert "1. Resumo" in cheio and "B₁ em" not in cheio
    assert "Valores da NBR 8800:2008 conferidos" not in cheio
    assert "(norma escolhida)" in cheio and cb.NORMAS_ROTULOS["NBR8800_2024"] in cheio
    primeira = Document(BytesIO(conteudo)).tables[0]
    assert primeira.rows[0].cells[0].text == "Dado"  # sem projeto, sem a tabela do projeto


def test_caso_que_a_norma_nao_cobre_aparece_sem_a_palavra_proibida():
    # N_Sd acima de N_e: a norma não cobre o B₁ e o programa não estima
    secao = cb.secao_de_perfil(sc.listar_perfis()["W 250 x 32,7"])
    entrada = cb.EntradaColuna(
        secao=secao, fy_MPa=250.0, Lx_mm=9000.0, Ly_mm=9000.0, N_Sd_kN=2000.0, My_kNm=5.0
    )
    resultado = cb.verificar_coluna(entrada)
    cheio = texto(gerar_word_coluna(entrada, resultado))
    if resultado.bloqueios:
        assert "Casos que a norma não cobre" in cheio
    for proibida in PALAVRAS_PROIBIDAS:
        assert proibida not in cheio, proibida
