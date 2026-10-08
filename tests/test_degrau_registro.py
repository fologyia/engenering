"""Registro, memorial, CSV e PDF do degrau de escada em grade.

O registro leva as 38 verificações (com o item da norma em cada uma) e as tabelas para o memorial
da Central de relatórios; o PDF da página reúne tudo o que a página mostra. Os testes conferem o
contrato do registro, o capítulo do memorial (Word e PDF), os arquivos exportados e que nenhuma
linha some pelo caminho.
"""

from __future__ import annotations

import csv
import io
from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfReader

from core import degrau_escada as de
from core import degrau_registro as reg
from core import pdf_fonts
from core.degrau_relatorio import gerar_pdf_degrau
from core.project_report import (
    gerar_relatorio_industrial_pdf,
    gerar_relatorio_industrial_word,
    montar_modelo_relatorio,
)
from core.project_store import novo_projeto_documento
from core.technical_records import avaliar_contrato_registro, identificar_peca_registro

PALAVRAS_PROIBIDAS = ("Não informado", "bloqueio", "Snapshot", "Apêndice")


def resultado(**alteracoes) -> de.ResultadoEscada:
    return de.calcular_escada(de.EntradaDegrau(**alteracoes))


def registro(**alteracoes) -> dict:
    return reg.registro_degrau(
        resultado(**alteracoes),
        contexto={"obra": "Obra X", "tag": "ESC-01", "data_do_calculo": "2026-10-07"},
        responsavel="Eng. Y",
    )


def projeto_com(*registros: dict) -> dict:
    projeto = novo_projeto_documento("Escada ESC-01", codigo="MC-ESC-01")
    projeto.update(
        {
            "objetivo": "Verificar o degrau da escada.",
            "responsavel": "Eng. A",
            "verificador": "Eng. B",
            "aprovador": "Eng. C",
        }
    )
    projeto["registros_tecnicos"] = [
        identificar_peca_registro(r, peca=f"Degrau D{i}") for i, r in enumerate(registros, start=1)
    ]
    return projeto


def texto_word(conteudo: bytes) -> str:
    documento = Document(BytesIO(conteudo))
    partes = [p.text for p in documento.paragraphs]
    partes += [c.text for t in documento.tables for linha in t.rows for c in linha.cells]
    return "\n".join(partes)


def texto_pdf(conteudo: bytes) -> str:
    return "\n".join(p.extract_text() or "" for p in PdfReader(BytesIO(conteudo)).pages)


# ------------------------------------------------------------------ o registro
def test_registro_cumpre_o_contrato_do_programa():
    r = registro()
    assert avaliar_contrato_registro(r)["valido"]
    assert r["modulo_id"] == "degrau_escada" and r["modulo"] == "Degrau de escada em grade"
    assert r["responsavel"] == "Eng. Y"
    assert "ESC-01" in r["titulo"] and "DS-A4-35/3" in r["titulo"]


def test_registro_traz_as_38_verificacoes_com_norma_e_status():
    linhas = registro()["resultados"]["verificações"]
    assert len(linhas) == 38
    assert all(linha["referência"].strip() for linha in linhas)
    assert {linha["status"] for linha in linhas} <= {"OK", "NÃO OK", "ALERTA", "N/A", "INFO"}
    assert [linha["verificação"].split(".")[0] for linha in linhas] == [
        str(i) for i in range(1, 39)
    ]


@pytest.mark.parametrize(
    ("alteracoes", "status_geral", "status_registro"),
    [
        ({}, "ALERTA", "Atenção"),
        ({"parafuso": '1/2"'}, "NÃO OK", "Não atende"),
        ({"gc_superior_mm": 1300.0}, "NÃO OK", "Não atende"),  # a NR limita a 1200 mm
    ],
)
def test_status_do_registro_segue_o_status_geral(alteracoes, status_geral, status_registro):
    r = registro(**alteracoes)
    assert r["resultados"]["status_geral"] == status_geral
    assert r["status"] == status_registro


def test_resultados_escalares_e_utilizacao_maxima():
    resultados = registro()["resultados"]
    esperado = resultado()
    assert resultados["modelo_adotado"] == "DS-A4-35/3"
    assert resultados["número_de_espelhos"] == 20
    assert resultados["espelho_h_mm"] == pytest.approx(180.0)
    assert resultados["piso_b_mm"] == pytest.approx(270.0)
    assert resultados["profundidade_C_mm"] == pytest.approx(300.0)
    assert resultados["furação_F_mm"] == 135
    assert resultados["degraus_em_grade"] == 18
    assert resultados["projeção_horizontal_total_mm"] == pytest.approx(5760.0)
    assert resultados["utilizacao_maxima"] == pytest.approx(
        esperado.aproveitamento_maximo, abs=1e-4
    )
    assert resultados["peso_total_kg"] == pytest.approx(esperado.peso_total_kg, abs=0.01)


def test_entradas_do_registro_levam_o_que_foi_informado_e_omitem_o_vazio():
    entradas = registro(n_imposto=21, altura_max_lance_imposta_mm=2000.0)["entradas"]
    assert entradas["n_espelhos_imposto"] == 21
    assert entradas["altura_max_lance_imposta_mm"] == 2000.0
    assert entradas["obra"] == "Obra X" and entradas["tag"] == "ESC-01"
    assert "piso_b_imposto_mm" not in entradas and "modelo_manual" not in entradas
    assert entradas["enquadramento"] == "NR-12 – acesso a máquina/equipamento"
    manual = registro(selecao=de.SELECAO_MANUAL, modelo_manual="DS-F2-40/5")["entradas"]
    assert manual["modelo_manual"] == "DS-F2-40/5" and manual["selecao_do_modelo"] == "Manual"


def test_alertas_levam_o_que_nao_passou_e_os_avisos_fixos():
    alertas = " ".join(registro(parafuso='1/2"')["alertas"])
    assert "Parafuso A307" in alertas
    assert "carga-base" in alertas and "chapa xadrez" in alertas
    assert "Não faz: Longarina" in alertas
    assert "inoxidável" in alertas  # aviso fixo: vale para qualquer material
    inox = " ".join(registro(material="AISI 304", acabamento="passivado")["alertas"])
    assert "inoxidável" in inox


def test_tabelas_do_memorial_tem_larguras_que_somam_a_largura_util():
    tabelas = registro()["resultados"]["tabelas_memorial"]
    assert len(tabelas) == 4
    for tabela in tabelas:
        assert sum(tabela["larguras"]) == reg.LARGURA_UTIL_DXA
        assert len(tabela["larguras"]) == len(tabela["cabecalhos"])
        assert all(len(linha) == len(tabela["cabecalhos"]) for linha in tabela["linhas"])
    geometria, dimensionamento, catalogo, requisicao = tabelas
    assert any(linha[0] == "Piso b" for linha in geometria["linhas"])
    assert any(linha[0] == "M_Rd por barra" for linha in dimensionamento["linhas"])
    assert "DS-A4-35/3" in {linha[0] for linha in catalogo["linhas"]}  # o adotado sempre aparece
    assert requisicao["linhas"][0][0] == resultado().especificacao


def test_catalogo_do_memorial_lista_os_mais_leves_que_atendem_e_sempre_o_adotado():
    catalogo = registro()["resultados"]["tabelas_memorial"][2]
    pesos = [float(linha[1].replace(",", ".")) for linha in catalogo["linhas"]]
    assert pesos == sorted(pesos) and 1 <= len(pesos) <= 12
    assert [linha[5] for linha in catalogo["linhas"]].count("adotado") == 1
    assert catalogo["cabecalhos"][-1] == "Adotado"


def test_adotado_fora_dos_12_mais_leves_ainda_aparece_na_tabela():
    # Malha B: a preferência exclui famílias mais leves, e o adotado fica depois dos 12 primeiros.
    r = resultado(malha_preferida="B", ligacao_preferida="2", L_mm=1200.0)
    linhas = reg.tabelas_para_memorial(r)[2]["linhas"]
    assert len(linhas) <= 12
    assert [linha[0] for linha in linhas if linha[5] == "adotado"] == [r.adotado.modelo.nome]


def test_sem_modelo_que_atenda_a_tabela_do_catalogo_mostra_so_o_adotado_que_nao_atende():
    catalogo = registro(P_kN=12.0)["resultados"]["tabelas_memorial"][2]
    assert len(catalogo["linhas"]) == 1
    nome, _peso, _u, atende, _pref, marca = catalogo["linhas"][0]
    assert atende == "não" and marca == "adotado" and nome.startswith("DS-")


def test_sem_caracteres_que_nem_toda_fonte_tem():
    texto = str(registro()["resultados"]["tabelas_memorial"]) + str(registro()["equacoes"])
    assert "⌊" not in texto and "⌋" not in texto


# ------------------------------------------------------------------ a tabela de 7 colunas e o CSV
def test_tabela_de_verificacoes_tem_valor_limite_e_aproveitamento():
    linhas = reg.tabela_de_verificacoes(resultado())
    assert len(linhas) == 38 and list(linhas[0]) == list(reg.COLUNAS_VERIFICACAO)
    por_nome = {linha["Verificação"].split(". ", 1)[0]: linha for linha in linhas}
    flexao = por_nome["22"]
    assert flexao["Valor"] == "0,754 kN·m" and flexao["Limite"] == "0,775 kN·m"
    assert flexao["Aproveitamento (%)"] == pytest.approx(97.3, abs=0.1)
    espelho = por_nome["1"]
    assert espelho["Valor"] == "180 mm" and espelho["Limite"] == "250 mm"
    assert espelho["Aproveitamento (%)"] is None  # limite: só OK/NÃO OK
    faixa = por_nome["2"]
    assert faixa["Valor"] == "180 mm" and faixa["Limite"] == "—"
    largura = por_nome["14"]
    assert largura["Valor"] == "800 mm" and largura["Limite"] == "600 mm"  # mínimo: valor ≥ limite


def test_csv_das_verificacoes_tem_38_linhas_unidade_e_bom():
    conteudo = reg.csv_das_verificacoes(resultado())
    assert conteudo.startswith(b"\xef\xbb\xbf")
    linhas = list(csv.reader(io.StringIO(conteudo.decode("utf-8-sig"))))
    assert linhas[0][:5] == ["Verificação", "Norma/item", "Valor", "Limite", "Unidade"]
    assert len(linhas) == 39 and all(len(linha) == 8 for linha in linhas)
    assert linhas[1][1] == "NR-12 Anexo III 11 d)" and linhas[1][4] == "mm"


def test_csv_do_catalogo_tem_64_modelos_e_marca_o_adotado():
    conteudo = reg.csv_do_catalogo(resultado())
    linhas = list(csv.DictReader(io.StringIO(conteudo.decode("utf-8-sig"))))
    assert len(linhas) == 64
    assert [x["Modelo"] for x in linhas if x["Adotado"] == "ADOTADO"] == ["DS-A4-35/3"]
    assert {x["Atende"] for x in linhas} == {"SIM", "NÃO"}


def test_linhas_do_catalogo_tem_as_colunas_da_especificacao():
    linhas = reg.linhas_do_catalogo(resultado())
    assert len(linhas) == 64
    esperado = {"Modelo", "λ", "λp", "λr", "M_Rd (kN·m)", "Peso do degrau (kg)", "u máx", "Atende"}
    assert esperado <= set(linhas[0])


# ------------------------------------------------------------------ o memorial
@pytest.fixture()
def modelo():
    return montar_modelo_relatorio(
        projeto_com(registro(), registro(parafuso='1/2"', superficie=de.SUPERFICIE_LISA))
    )


def capitulos(modelo: dict) -> list[dict]:
    return [s for s in modelo["secoes"] if "blocos" in s]


def test_cada_degrau_registrado_vira_um_capitulo_com_resultado(modelo):
    caps = capitulos(modelo)
    assert len(caps) == 2
    primeiro = caps[0]["blocos"][0]
    assert primeiro["tipo"] == "destaque" and primeiro["rotulo"] == "Resultado: ATENÇÃO."
    assert "DS-A4-35/3" in primeiro["texto"] and "35 OK · 0 NÃO OK · 2 ALERTA" in primeiro["texto"]
    segundo = caps[1]["blocos"][0]
    assert segundo["rotulo"] == "Resultado: NÃO ATENDE."


def test_capitulo_tem_entradas_legiveis_o_que_passou_e_fecha_no_que_nao_passou(modelo):
    blocos = capitulos(modelo)[1]["blocos"]
    rotulos = [b.get("texto") or b.get("rotulo") for b in blocos]
    assert any(str(r).startswith("O que passou (") for r in rotulos)
    textos = [str(b.get("texto") or b.get("rotulo")) for b in blocos]
    posicao_passou = next(i for i, x in enumerate(textos) if x.startswith("O que passou ("))
    posicao_nao_passou = next(i for i, x in enumerate(textos) if x.startswith("Não passou ("))
    assert posicao_nao_passou > posicao_passou
    assert all("Não passou" not in x for x in textos[:posicao_passou])
    todas = str(blocos)
    assert "Desnível total H [mm]" in todas and "Enquadramento legal" in todas
    assert "NR-12 Anexo III 11 d)" in todas  # o item da norma em cada verificação
    assert "Conferir antes de emitir" in todas


def test_capitulo_mostra_as_quatro_tabelas_do_degrau(modelo):
    legendas = [b["legenda"] for b in capitulos(modelo)[0]["blocos"] if b.get("tipo") == "tabela"]
    todas = " ".join(legendas)
    assert "Geometria da escada" in todas
    assert "Dimensionamento do modelo adotado (DS-A4-35/3)" in todas
    assert "Modelos do catálogo que atendem" in todas
    assert "Texto para a requisição de compra" in todas


def test_valores_do_degrau_nao_viram_outros_resultados_repetidos(modelo):
    blocos = capitulos(modelo)[0]["blocos"]
    outros = [
        b for b in blocos if b.get("tipo") == "tabela" and "Outros resultados" in b["legenda"]
    ]
    assert outros == []  # tudo o que o módulo registra já tem tabela própria


def test_word_e_pdf_do_memorial_trazem_o_degrau_sem_palavras_proibidas():
    projeto = projeto_com(registro())
    word = texto_word(gerar_relatorio_industrial_word(projeto))
    pdf = texto_pdf(gerar_relatorio_industrial_pdf(projeto))
    for texto in (word, pdf):
        assert "Degrau de escada em grade" in texto and "DS-A4-35/3" in texto
        assert "O que passou" in texto and "Não passou" in texto
        for proibido in PALAVRAS_PROIBIDAS:
            assert proibido not in texto, proibido
    assert "Anexo III" in word and "Anglo 10.2" in word


# ------------------------------------------------------------------ o PDF da página
@pytest.fixture(scope="module")
def pdf() -> bytes:
    return gerar_pdf_degrau(
        resultado(),
        {"obra": "Obra X", "tag": "ESC-01", "responsavel": "Eng. Y", "data": "07/10/2026"},
    )


def test_pdf_e_um_documento_com_o_cabecalho_e_varias_paginas(pdf):
    assert pdf.startswith(b"%PDF")
    leitor = PdfReader(BytesIO(pdf))
    assert len(leitor.pages) >= 5
    assert "Degrau de escada em grade" in (leitor.metadata.title or "")


def test_pdf_traz_identificacao_resumo_verificacoes_e_requisicao(pdf):
    texto = texto_pdf(pdf)
    for trecho in (
        "Obra X",
        "ESC-01",
        "Eng. Y",
        "07/10/2026",
        "Quadro-resumo",
        "35 OK · 0 NÃO OK · 2 ALERTA · 0 N/A",
        "Texto para a requisição de compra",
        "Conflitos entre normas",
        "Esta página não faz",
    ):
        assert trecho in texto, trecho
    assert resultado().especificacao.split(" – ")[0] in texto


def test_pdf_traz_as_38_verificacoes_e_os_64_modelos(pdf):
    texto = texto_pdf(pdf)
    for numero in range(1, 39):
        assert f"{numero}. " in texto, numero
    assert "NR-12 Anexo III 11 d)" in texto and "NBR 8800:2008 5.4.2" in texto
    nomes = {m.nome for m in de.catalogo()}
    assert all(nome in texto for nome in nomes)


def test_pdf_com_helvetica_nao_quebra_e_troca_os_simbolos(monkeypatch):
    helvetica = pdf_fonts.FontePdf("Helvetica", "Helvetica-Bold", False)
    monkeypatch.setattr("core.degrau_relatorio.fonte_pdf", lambda: helvetica)
    conteudo = gerar_pdf_degrau(resultado(), {})
    assert conteudo.startswith(b"%PDF") and len(conteudo) > 10_000


def test_pdf_sem_identificacao_usa_tracos():
    conteudo = gerar_pdf_degrau(resultado(), None)
    assert "Obra" in texto_pdf(conteudo)


def test_pdf_de_um_caso_que_nao_atende_mostra_nao_ok():
    conteudo = gerar_pdf_degrau(resultado(P_kN=12.0), {})
    texto = texto_pdf(conteudo)
    assert "NÃO OK" in texto and "NENHUM modelo" in texto
