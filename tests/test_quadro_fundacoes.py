"""Quadro de cargas para as fundações a partir das forças das vigas do SolidWorks."""

from __future__ import annotations

import io
import json

import pytest
from openpyxl import load_workbook

from core import esforcos_modelo as em
from core import plano_de_cargas as pc
from core import quadro_fundacoes as qf
from core.project_report import montar_modelo_relatorio
from core.project_store import novo_projeto_documento
from core.technical_records import avaliar_contrato_registro
from tests import dados_solidworks as sw
from tests.test_esforcos_modelo import PILAR, importar, plano_do_portico

PILAR_2 = "Viga-3(Aparar/Estender12[2])"


def portico(*pilares: str, base: str = em.BASE_INICIO) -> em.EsforcosDoModelo:
    dados = importar("PP", "SC", "W0", "W180")
    return em.com_membros(
        dados, {p: em.ConfiguracaoDoMembro("W 200 x 35,9 (H)", "Pilar", base=base) for p in pilares}
    )


def _manual(primeiro: float, ultimo: float) -> dict[str, tuple[em.PontoDeEsforco, ...]]:
    return {
        "PP": (
            em.PontoDeEsforco("1", "1", primeiro, 0, 0, 0, 0, 0),
            em.PontoDeEsforco("2", "2", ultimo, 0, 0, 0, 0, 0),
        )
    }


@pytest.mark.parametrize(
    ("primeiro", "ultimo", "base"),
    [(-80.0, -20.0, em.BASE_INICIO), (-20.0, -80.0, em.BASE_FIM), (-50.0, -50.0, None)],
)
def test_base_automatica_e_a_ponta_mais_comprimida(primeiro, ultimo, base):
    achada, origem = qf.ponta_da_base(_manual(primeiro, ultimo), em.BASE_AUTOMATICA)
    assert achada == base
    assert ("escolha a base" in origem) is (base is None)


def test_base_escolhida_vale_mesmo_com_compressao_igual():
    assert qf.ponta_da_base(_manual(-50, -50), em.BASE_FIM) == (em.BASE_FIM, "escolhida")


def test_quadro_do_portico_bate_com_as_reacoes():
    quadro = qf.montar_quadro(portico(PILAR, PILAR_2))
    assert [p.pilar for p in quadro.pilares] == [PILAR, PILAR_2]
    sc_pilar = next(c for c in quadro.pilares[0].cargas if c.caso == "SC")
    resultado = sw.resolver("SC")
    reacao = next(r for r in resultado.reacoes_nodais if r["no"] == 1)
    assert sc_pilar.n == pytest.approx(float(reacao["ry_N"]) / 1e3)  # 30 kN de compressão
    assert abs(sc_pilar.m2) == pytest.approx(abs(float(reacao["mz_Nmm"])) / 1e6, rel=1e-4)
    w0 = next(c for c in quadro.pilares[0].cargas if c.caso == "W0")
    assert w0.n < 0  # o vento para +X levanta o pilar de barlavento
    textos = " ".join(i.texto for i in quadro.conferencia)
    assert "SC: a soma das bases (60,00 kN) bate com a reação vertical do modelo" in textos
    assert all(i.nivel == pc.NIVEL_OK for i in quadro.conferencia)


def test_conferencia_avisa_quando_falta_pilar_ou_base():
    so_um = qf.montar_quadro(portico(PILAR))
    assert any("há apoios que não são base de pilar marcado" in i.texto for i in so_um.conferencia)
    sem_base = qf.montar_quadro(portico(PILAR, base=em.BASE_AUTOMATICA))
    assert sem_base.pendentes and any("Escolha a base" in i.texto for i in sem_base.conferencia)
    nenhum = qf.montar_quadro(importar("PP"))
    assert "Nenhuma barra marcada como Pilar" in nenhum.conferencia[0].texto


def test_matriz_csv_excel_e_registro():
    quadro = qf.montar_quadro(portico(PILAR, PILAR_2))
    plano = plano_do_portico()
    cabecalho, linhas = qf.matriz_de_compressao(quadro)
    assert cabecalho == ["Pilar", "PP", "SC", "W0", "W180"]
    assert linhas[0][2] == pytest.approx(30.0)
    assert qf.csv_do_quadro(quadro, plano).decode("utf-8-sig").splitlines()[0].split(";") == list(
        qf.COLUNAS
    )
    livro = load_workbook(io.BytesIO(qf.xlsx_do_quadro(quadro, plano, titulo="Teste")))
    assert livro.sheetnames == ["Leia-me", "Quadro", "Compressão por pilar"]
    assert "5.9" in " ".join(str(c.value) for c in livro["Leia-me"]["A"] if c.value)
    assert len(list(livro["Quadro"].values)) == 1 + 2 * 4
    registro = qf.registro_do_quadro(quadro, plano)
    assert avaliar_contrato_registro(registro)["valido"]
    json.dumps(registro)
    assert "sem combinar nem majorar" in registro["resultados"]["destaque_memorial"]
    projeto = novo_projeto_documento("Plataforma", codigo="MC-01")
    projeto["registros_tecnicos"] = [registro]
    capitulo = next(s for s in montar_modelo_relatorio(projeto)["secoes"] if "blocos" in s)
    assert capitulo["blocos"][0]["rotulo"] == "Resultado do cálculo."
    legendas = " ".join(b.get("legenda", "") for b in capitulo["blocos"])
    assert "Compressão (N, kN) na base de cada pilar" in legendas


def test_base_vai_e_volta_pelo_projeto():
    dados = portico(PILAR, base=em.BASE_FIM)
    volta = em.de_dicionario(json.loads(json.dumps(em.para_dicionario(dados))))
    assert volta.membros[PILAR].base == em.BASE_FIM
    estranho = em.para_dicionario(dados)
    estranho["membros"][PILAR]["base"] = "qualquer"
    assert em.de_dicionario(estranho).membros[PILAR].base == em.BASE_AUTOMATICA
