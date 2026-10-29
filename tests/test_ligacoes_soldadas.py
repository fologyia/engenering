"""Barra soldada na chapa de nó e chapa soldada no apoio (NBR 8800 2025 e 2008).

Os quatro exercícios de solda que o usuário mandou (29/10/2026) refeitos pela norma: onde a apostila
usa a "resistência do metal-base" pela face de fusão (0,60·d_w·L·f_y/γ_a1), a Tabela 9 manda o
metal-base do filete atender à 6.5 — na espessura da peça —, e os comprimentos mudam.
"""

from __future__ import annotations

import io
import json
import math
from dataclasses import replace

import pytest
from docx import Document
from PIL import Image

from core import cantoneiras as ct
from core import ligacoes as lg
from core import ligacoes_soldadas as ls
from core import nbr8800 as nb
from core import section_catalog as sc
from core import soldas as sd
from core.record_charts import imagens_do_registro

Q4 = sd.resistencia_filete_por_mm(4.0, 485.0, 1.35)
Q5 = sd.resistencia_filete_por_mm(5.0, 485.0, 1.35)
Q6 = sd.resistencia_filete_por_mm(6.0, 485.0, 1.35)
PALAVRAS_PROIBIDAS = ("Não informado", "bloqueio", "Snapshot", "Apêndice", "Projeto de revisão")


def linha(r: lg.ResultadoLigacao, inicio: str):
    return next(v for v in r.verificacoes if v.nome.startswith(inicio))


# ------------------------------------------------------------------ exercício 4
def test_exercicio_4_cantoneira_balanceada_na_chapa_de_no():
    r = lg.calcular(ls.EntradaBarraSoldada())
    cant = ct.obter_cantoneira('L 2" × 1/4"')
    c_, x = cant.b_mm, cant.x_barra_mm
    assert r.status == "OK"
    # o metal da solda governa (0,61 kN/mm < metal-base da cantoneira 0,87 e da chapa 1,09)
    assert 100.0 / Q4 == pytest.approx(164.04, abs=0.01)
    f1, f2 = 100 * (c_ - x) / c_, 100 * x / c_
    assert linha(r, "Filete do dorso").solicitante == pytest.approx(f1)
    assert linha(r, "Filete da borda").solicitante == pytest.approx(f2)
    assert (r.grandezas["l_dorso_mm"], r.grandezas["l_borda_mm"]) == (120.0, 50.0)
    assert sd.adotar(f1 / Q4) == 120.0 and sd.adotar(f2 / Q4) == 50.0
    # a barra: C_t = 1 − x̄/ℓ_c com ℓ_c = média dos filetes
    ct_ = 1 - x / 85.0
    assert linha(r, "Barra: tração (ruptura").resistente == pytest.approx(
        ct_ * cant.area_mm2 * 400 / 1.35 / 1000
    )
    # a chapa de nó: Whitmore e a largura mínima pelo escoamento (55 mm da apostila)
    bw = c_ + 170 * math.tan(math.radians(30))
    assert r.grandezas["largura_whitmore_mm"] == pytest.approx(bw)
    assert r.grandezas["largura_minima_no_mm"] == pytest.approx(55.0)
    assert linha(r, "Chapa de nó: tração na seção de Whitmore (escoamento)").resistente == (
        pytest.approx(bw * 8 * 250 / 1.1 / 1000)
    )
    agv, ant = 8 * 170, 8 * c_
    assert linha(r, "Chapa de nó: colapso").resistente == pytest.approx(
        min(0.6 * 400 * agv + 400 * ant, 0.6 * 250 * agv + 400 * ant) / 1.35 / 1000
    )


def test_exercicio_4_chapa_de_no_soldada_no_pilar_a_60_graus():
    r = lg.calcular(ls.EntradaBarraSoldada())
    assert r.grandezas["l_apoio_mm"] == 50.0  # a apostila adota 70 mm
    solda = linha(r, "Chapa de nó no apoio — solda: metal da solda")
    assert solda.solicitante == pytest.approx(100 / (2 * 50) * 10)
    assert solda.resistente == pytest.approx(Q6 * sd.fator_direcao(60) * 10)
    # governa a chapa de nó à tração junto à solda: 86,6 kN em 8 × 50 mm
    tracao = linha(r, "Chapa de nó no apoio — chapa junto à solda: tração (escoamento)")
    assert tracao.solicitante == pytest.approx(100 * math.sin(math.radians(60)))
    assert tracao.resistente == pytest.approx(8 * 50 * 250 / 1.1 / 1000)
    assert linha(r, "Chapa de nó no apoio — solda: perna mínima").solicitante == 5.0
    # 45 mm não bastaria: o dimensionamento acha o menor
    menor = lg.calcular(replace(ls.EntradaBarraSoldada(), modo=ls.MODO_VERIFICAR, l_dorso_mm=120, l_borda_mm=50, l_apoio_mm=45))  # fmt: skip
    assert menor.status == "NÃO OK"


# ------------------------------------------------------------------ exercício 1
def test_exercicio_1_talas_sobrepostas_com_cobrimento_e_minimo_de_45_kN():
    e = ls.EntradaBarraSoldada(
        peca=ls.PECA_CHAPA, largura_mm=50.0, t_peca_mm=9.525, duas_pecas=True, t_no_mm=12.7,
        perna_mm=8.0, n_t_sd_kN=40.0, com_apoio=False,
    )  # fmt: skip
    r = lg.calcular(e)
    assert r.status == "OK" and r.grandezas["n_ligacao_kN"] == 45.0  # 6.1.5.2
    assert linha(r, "Força mínima da ligação").solicitante == 45.0
    # a chapa do meio recebe os filetes das duas talas nas mesmas linhas: governa ela
    mb = min(0.6 * 250 * 12.7 / 1.1, 0.6 * 400 * 12.7 / 1.35) / 1000 / 2
    assert linha(r, "Filete do dorso").resistente == pytest.approx(mb * 50)
    assert linha(r, "Filete do dorso").solicitante == pytest.approx(45 / 2 / 2)
    # 40 mm da apostila não passam: cobrimento ≥ 5·t = 47,6 mm e filete ≥ largura (50 mm)
    assert (r.grandezas["l_dorso_mm"], r.grandezas["l_borda_mm"]) == (50.0, 50.0)
    assert linha(r, "Solda: cobrimento").solicitante == pytest.approx(47.625)
    curta = lg.calcular(replace(e, modo=ls.MODO_VERIFICAR, l_dorso_mm=40, l_borda_mm=40))
    assert linha(curta, "Solda: cobrimento").status == "NÃO OK"
    assert linha(curta, "Solda: cada filete").status == "NÃO OK"
    assert linha(curta, "Barra: ligação permitida").status == "NÃO OK"  # ℓ_w < b (5.2.5 d)


# ------------------------------------------------------------------ exercícios 2 e 3
def test_exercicio_2_filete_intermitente_dos_dois_lados():
    r = lg.calcular(ls.EntradaChapaNoApoio())
    assert r.status == "OK"
    assert (r.grandezas["n_segmentos"], r.grandezas["segmento_mm"]) == (1.0, 40.0)
    solda = linha(r, "Solda: metal da solda")
    assert solda.solicitante == pytest.approx(45 / (2 * 40) * 10)  # 45 kN (6.1.5.2)
    assert solda.resistente == pytest.approx(Q5 * 1.5 * 10)  # θ = 90°: × 1,5
    assert linha(r, "Solda: perna mínima").solicitante == 5.0  # chapa de 1/2"
    assert linha(r, "Solda intermitente: permitida").status == "OK"
    assert linha(r, "Solda: comprimento mínimo do segmento").solicitante == 40.0


def test_exercicio_3_penetracao_total_e_compatibilidade():
    e = ls.EntradaChapaNoApoio(
        modo=ls.MODO_VERIFICAR, tipo_solda=ls.TIPO_TOTAL, intermitente=False, n_sd_kN=100.0
    )
    r = lg.calcular(e)
    assert linha(r, "Penetração total: tração normal").resistente == pytest.approx(
        288.636, abs=1e-3
    )
    assert linha(r, "Penetração total: metal da solda compatível").status == "OK"
    e60 = lg.calcular(replace(e, eletrodo="E60"))
    assert linha(e60, "Penetração total: metal da solda compatível").status == "NÃO OK"
    comprimida = lg.calcular(replace(e, eletrodo="E60", n_sd_kN=-100.0))
    assert linha(comprimida, "Penetração total: metal da solda compatível").status == "OK"
    cortante = lg.calcular(replace(e, n_sd_kN=0.0, v_sd_kN=100.0))
    assert linha(cortante, "Penetração total: cisalhamento").resistente == pytest.approx(
        0.6 * 288.636, abs=1e-3
    )


# ------------------------------------------------------------------ chapa no apoio: mais casos
def test_penetracao_parcial_garganta_e_resistencia():
    e = ls.EntradaChapaNoApoio(
        modo=ls.MODO_VERIFICAR, tipo_solda=ls.TIPO_PARCIAL, intermitente=False,
        profundidade_mm=6.0, n_sd_kN=100.0,
    )  # fmt: skip
    r = lg.calcular(e)
    solda = linha(r, "Solda: metal da solda (penetração parcial)")
    assert solda.resistente == pytest.approx(0.6 * 485 * 6 * 2 / 1.25 / 1000 * 10)  # γ_w1
    assert linha(r, "Solda: garganta mínima").solicitante == 5.0
    menos3 = lg.calcular(replace(e, descontar_3mm=True))
    assert linha(menos3, "Solda: garganta mínima").status == "NÃO OK"  # 3 mm < 5 mm
    with pytest.raises(lg.LigacaoInvalida):
        lg.calcular(replace(e, profundidade_mm=7.0))  # 2 × 7 > 12,7: é penetração total
    um_lado = lg.calcular(replace(e, lados=1, profundidade_mm=8.0))
    assert linha(um_lado, "Penetração parcial de um lado só").status == "ALERTA"


def test_dimensionar_acha_o_menor_comprimento_com_momento():
    e = ls.EntradaChapaNoApoio(
        intermitente=False, perna_mm=6.0, comprimento_mm=400.0, n_sd_kN=0.0, v_sd_kN=80.0,
        m_sd_kNm=4.0,
    )  # fmt: skip
    r = lg.calcular(e)
    adotado = r.grandezas["comprimento_efetivo_mm"]
    assert r.status == "OK" and adotado % 5 == 0
    antes = lg.calcular(replace(e, modo=ls.MODO_VERIFICAR, comprimento_mm=adotado - 5))
    assert antes.status == "NÃO OK"
    # com momento, sem o aumento pela direção (a carga fica fora do plano dos filetes)
    assert "6.2.5.1" in linha(r, "Solda: metal da solda").referencia


def test_apoio_na_mesa_de_um_perfil():
    perfil = "W 250 x 32,7"
    assert perfil in sc.listar_perfis()
    p = sc.listar_perfis()[perfil]
    e = ls.EntradaChapaNoApoio(
        modo=ls.MODO_VERIFICAR, intermitente=False, t_chapa_mm=8.0, perna_mm=5.0,
        comprimento_mm=150.0, n_sd_kN=120.0, apoio=ls.APOIO_PERFIL, perfil_apoio=perfil,
        aco_apoio="ASTM A572 Gr 50",
    )  # fmt: skip
    r = lg.calcular(e)
    # chapa ao longo do eixo: atua em t + 2·perna = 18 mm < 0,15·b_f (5.7.2.1)
    assert 8.0 + 2 * 5.0 < 0.15 * p.largura_mm
    assert linha(r, "Apoio: flexão local da mesa").status == "N/A"
    tf, tw = p.espessura_mesa_mm, p.espessura_alma_mm
    assert linha(r, "Apoio: escoamento local da alma").resistente == pytest.approx(
        1.10 * (5 * tf + 150) * 345 * tw / 1.10 / 1000
    )
    transversal = lg.calcular(replace(e, orientacao=ls.ORIENTACAO_TRANSVERSAL))
    assert linha(transversal, "Apoio: flexão local da mesa").resistente == pytest.approx(
        6.25 * tf**2 * 345 / 1.10 / 1000
    )
    comprimida = lg.calcular(replace(e, n_sd_kN=-120.0))
    assert linha(comprimida, "Apoio: enrugamento da alma").resistente > 0


# ------------------------------------------------------------------ barra: mais casos
def test_filete_na_ponta_divide_a_forca_pela_6_2_5_2_c():
    r = lg.calcular(ls.EntradaBarraSoldada(solda_na_ponta=True, com_apoio=False))
    cant = ct.obter_cantoneira('L 2" × 1/4"')
    f3 = Q4 * cant.b_mm
    f1 = 100 * (cant.b_mm - cant.x_barra_mm) / cant.b_mm - f3 / 2
    assert linha(r, "Filete do dorso").solicitante == pytest.approx(f1)
    grupo = linha(r, "Filetes da barra: grupo")
    r_long = Q4 * (r.grandezas["l_dorso_mm"] + r.grandezas["l_borda_mm"])
    assert grupo.resistente == pytest.approx(sd.combinacao_longitudinal_transversal(r_long, f3))
    assert r.grandezas["l_dorso_mm"] < 120.0  # o filete da ponta encurta os longitudinais


def test_compressao_na_chapa_de_no_e_criterio_anglo():
    e = ls.EntradaBarraSoldada(n_c_sd_kN=60.0, anglo=True, com_apoio=False)
    r = lg.calcular(e)
    cant = ct.obter_cantoneira('L 2" × 1/4"')
    assert r.grandezas["n_ligacao_kN"] == pytest.approx(
        max(100.0, 0.75 * cant.area_mm2 * 250 / 1.1 / 1000)
    )
    compressao = linha(r, "Chapa de nó: compressão")
    esbeltez = 1.2 * 100 * math.sqrt(12) / 8
    chi = nb.fator_chi(esbeltez / math.pi * math.sqrt(250 / 200_000))
    bw = r.grandezas["largura_whitmore_mm"]
    assert compressao.resistente == pytest.approx(chi * bw * 8 * 250 / 1.1 / 1000)
    assert compressao.solicitante == 60.0


def test_duas_cantoneiras_dividem_a_forca():
    r = lg.calcular(ls.EntradaBarraSoldada(duas_pecas=True, com_apoio=False))
    assert linha(r, "Barra: tração (escoamento").solicitante == pytest.approx(50.0)
    assert "por peça" in linha(r, "Filete do dorso").formula


def test_norma_2008_e_combinacao_excepcional_na_barra():
    r = lg.calcular(ls.EntradaBarraSoldada(norma=sd.NBR_2008, combinacao=sd.COMBINACAO_EXCEPCIONAL))
    refs = " ".join(v.referencia for v in r.verificacoes)
    assert "NBR 8800:2008" in refs and "a confirmar" not in refs and "2025" not in refs
    assert linha(r, "Barra: tração (escoamento").resistente == pytest.approx(
        ct.obter_cantoneira('L 2" × 1/4"').area_mm2 * 250 / 1.0 / 1000
    )


@pytest.mark.parametrize(
    "mudanca",
    [
        {"perna_mm": 0.0},
        {"n_t_sd_kN": 0.0, "n_c_sd_kN": 0.0},
        {"theta_graus": 0.0},
        {"cantoneira": "L inexistente"},
        {"aco_no": "aço de ninguém"},
    ],
)
def test_entrada_impossivel_da_barra_vira_erro_claro(mudanca):
    with pytest.raises(lg.LigacaoInvalida):
        lg.calcular(replace(ls.EntradaBarraSoldada(), **mudanca))


# ------------------------------------------------------------------ figura, registro e memorial
@pytest.mark.parametrize(
    "entrada",
    [
        ls.EntradaBarraSoldada(),
        ls.EntradaBarraSoldada(peca=ls.PECA_CHAPA, com_apoio=False, solda_na_ponta=True),
        ls.EntradaChapaNoApoio(),
        ls.EntradaChapaNoApoio(
            tipo_solda=ls.TIPO_PARCIAL, intermitente=False, modo=ls.MODO_VERIFICAR
        ),
        ls.EntradaChapaNoApoio(tipo_solda=ls.TIPO_TOTAL, intermitente=False, m_sd_kNm=2.0),
    ],
)
def test_figura_e_registro_redesenham_a_solda_adotada(entrada):
    r = lg.calcular(entrada)
    figura = lg.figura(r)
    imagem = Image.open(io.BytesIO(figura.png))
    assert imagem.format == "PNG" and imagem.size == (figura.largura_px, figura.altura_px)
    registro = json.loads(json.dumps(lg.registro_ligacao(r, contexto={"tag": "LG-02"})))
    assert registro["entradas"]["norma"] == sd.NORMAS_ROTULOS[entrada.norma]
    assert len(registro["resultados"]["verificações"]) == len(r.verificacoes)
    imagens, aviso = imagens_do_registro(registro, {})
    assert aviso == "" and len(imagens) == 1 and imagens[0]["png"] == figura.png


def test_memorial_da_ligacao_em_word_e_pdf():
    r = lg.calcular(ls.EntradaBarraSoldada())
    conteudo = lg.memorial_da_ligacao(r, formato="word")
    documento = Document(io.BytesIO(conteudo))
    texto = "\n".join(
        [p.text for p in documento.paragraphs]
        + [c.text for t in documento.tables for ln in t.rows for c in ln.cells]
    )
    assert "Barra soldada na chapa de nó" in texto and "Whitmore" in texto
    assert "ABNT NBR 8800:2025" in texto
    for proibida in PALAVRAS_PROIBIDAS:
        assert proibida not in texto, proibida
    assert lg.memorial_da_ligacao(r, formato="pdf").startswith(b"%PDF")
