"""Tração nos parafusos, interação com o corte, alavanca, furos alargados e alongados e chapa de
ligação (NBR 8800:2025 e 2008), com os casos de aceite A a E do pedido de 28/10/2026 —
tolerância de 0,5 %. O caso A já funcionava e não pode regredir."""

from __future__ import annotations

import json
import math
from pathlib import Path

import pytest

from core import bolted_connection as bc
from core import bolted_joint_check as chk

RAIZ = Path(__file__).resolve().parent.parent
NBR_2025 = "NBR8800_2024"  # o código interno da NBR 8800:2025
TOL = 0.005


def perto(valor: float | None, esperado: float, tol: float = TOL) -> bool:
    return valor is not None and math.isclose(valor, esperado, rel_tol=tol)


def linha(r: chk.ResultadoLigacao, trecho: str):
    return next(v for v in r.verificacoes if trecho in v.nome)


def tem_linha(r: chk.ResultadoLigacao, trecho: str) -> bool:
    return any(trecho in v.nome for v in r.verificacoes)


def caso_a(**mudancas) -> chk.EntradaLigacao:
    """M16 A325, t = 6,35 mm, A36, e = 35, s = 60, furo 18, grade 1 × 2, N_k = 5, V_k = 15,
    γ_f = 1,4, galvanizado (μ 0,20), C_e = 1, F_Tb = 91 kN."""
    base = dict(
        norma=NBR_2025,
        designacao="M16",
        grau="A325",
        t=6.35,
        aco="ASTM A36",
        e=35.0,
        s=60.0,
        g=60.0,
        n_lin=1,
        n_col=2,
        N=5.0,
        V=15.0,
        gama_f=1.4,
        superficie="galvanizada_sem_tratamento",
        C_e=1.0,
    )
    base.update(mudancas)
    return chk.EntradaLigacao(**base)


# ------------------------------------------------------------------ caso A (não pode regredir)
def test_caso_a_corte_contato_e_deslizamento():
    r = chk.verificar_ligacao(caso_a())
    assert perto(r.resistencias["F_v_Rd_kN"], 55.6)
    assert perto(r.resistencias["F_c_Rd_critico_kN"], 58.7)  # ℓ_f = 35 − 18/2 = 26 mm, fator 1,2
    assert perto(r.forcas["R_max_ELU_kN"], 11.07)
    assert perto(r.forcas["R_max_ELS_kN"], 7.91)
    assert perto(r.resistencias["F_f_Rk_kN"], 14.56)
    assert r.resistencias["F_Tb_kN"] == 91  # M16 = 91 kN; os 85 kN são do 5/8"
    assert bc.protensao_minima('5/8"', "A325") == 85


# ------------------------------------------------------------------ caso B: tração e interação
def test_caso_b_tracao_e_interacao():
    assert perto(bc.area_bruta(16), 201.06)
    assert perto(bc.area_efetiva(16), 150.80)
    F_t_Rd = bc.resist_tracao_parafuso(16, "A325", NBR_2025)
    assert perto(F_t_Rd, 92.7)
    F_v_Rd = bc.resist_corte(16, "A325", NBR_2025)
    assert perto(bc.interacao_tracao_corte(1.0, F_t_Rd, 15.0, F_v_Rd), 0.073, 0.01)
    assert perto(bc.limite_tracao_tabela(16, "A325", NBR_2025, True, 15.0), 95.1)
    assert bc.coeficiente_tabela_interacao("A325", rosca_no_plano=False) == 1.5
    assert bc.coeficiente_tabela_interacao("A307", rosca_no_plano=False) == 1.9


def test_caso_b_na_verificacao_por_parafuso():
    """2 parafusos, de cálculo: N_t = 2 kN do grupo → 1 kN por parafuso; V = 30 → 15 kN."""
    r = chk.verificar_ligacao(
        caso_a(N=0.0, V=30.0, N_t=2.0, valores_sao_de_calculo=True, interacao_por_tabela=True)
    )
    tracao = linha(r, "Tração no parafuso")
    assert perto(tracao.solicitante, 1.0) and perto(tracao.resistente, 92.7)
    interacao = linha(r, "Tração + corte (interação")
    assert perto(interacao.solicitante, 0.073, 0.01) and interacao.resistente == 1.0
    tabela = linha(r, "Tabela 12")  # a NBR 8800:2025 renumerou (era a Tabela 11 na 2008)
    assert perto(tabela.solicitante, 1.0) and perto(tabela.resistente, 95.1)
    assert perto(r.forcas["F_t_Sd_kN"], 1.0) and perto(r.forcas["N_t_Sd_kN"], 2.0)
    r2008 = chk.verificar_ligacao(
        caso_a(
            norma="NBR8800_2008",
            V=30.0,
            N_t=2.0,
            valores_sao_de_calculo=True,
            interacao_por_tabela=True,
        )
    )
    assert tem_linha(r2008, "Tabela 11")


def test_momento_fora_do_plano_traciona_os_de_cima():
    coords = bc.grade_retangular(2, 2, 70, 100)  # y = 0 e 100
    ft = bc.tracao_nos_parafusos(coords, N_t=0.0, M_t=2.0)  # 2 kN·m
    assert perto(max(ft), 2000 * 50 / (4 * 50**2)) and min(ft) == 0.0  # 10 kN em cima, 0 embaixo


def test_so_tracao_nao_inventa_linha_de_corte():
    r = chk.verificar_ligacao(caso_a(N=0.0, V=0.0, N_t=10.0))
    assert tem_linha(r, "Tração no parafuso") and not tem_linha(r, "Parafuso crítico")
    assert not tem_linha(r, "interação")  # sem corte não há interação


def test_tracao_nas_normas_americanas():
    linhas = chk.comparar_normas(caso_a(N_t=10.0))
    aisc = next(item for item in linhas if item["norma"] == "AISC360_LRFD")
    assert perto(aisc["F_t_Rd_kN"], 0.75 * 620 * bc.area_bruta(16) / 1000)
    r = chk.verificar_ligacao(caso_a(norma="AISC360_LRFD", N_t=10.0))
    assert linha(r, "Tração + corte").referencia == "AISC J3.7"


def test_tracao_negativa_e_erro():
    with pytest.raises(ValueError, match="N_t é a tração"):
        chk.verificar_ligacao(caso_a(N_t=-5.0))


# ------------------------------------------------------------------ caso C: alavanca
def test_caso_c_placa_rigida():
    kw = dict(a=35, b=35, p=60, d_b=16, d_f=18, t=6.35, F_t0_Sd=2.5, F_t_Rd=92.7)
    assert perto(bc.alavanca_2025(f_u=400, **kw).t_rigida_mm, 3.52)
    assert perto(bc.alavanca_2025(f_u=250, **kw).t_rigida_mm, 4.45)


def test_alavanca_placa_flexivel_beta_alfa():
    r = bc.alavanca_2025(a=60, b=35, p=60, d_b=16, d_f=18, t=3.0, f_u=400, F_t0_Sd=40, F_t_Rd=92.7)
    assert not r.rigida and r.a_calculo_mm == 1.25 * 35  # a limitado a 1,25·b
    b_l, a_l = 35 - 8, 1.25 * 35 + 8
    beta = a_l / b_l * (92.7 / 40 - 1)
    assert perto(r.beta, beta, 1e-9) and r.alfa == 1.0  # β ≥ 1 → α = 1
    assert perto(r.t_requerida_mm, r.t_rigida_mm / math.sqrt(1 + (1 - 18 / 60)))
    r2 = bc.alavanca_2025(a=35, b=35, p=60, d_b=16, d_f=18, t=3.0, f_u=400, F_t0_Sd=80, F_t_Rd=92.7)
    delta = 1 - 18 / 60
    beta2 = (43 / 27) * (92.7 / 80 - 1)
    assert 0 < beta2 < 1 and perto(r2.alfa, min(1, beta2 / (delta * (1 - beta2))), 1e-9)


def test_alavanca_beta_negativo_pede_outro_parafuso():
    r = bc.alavanca_2025(a=35, b=35, p=60, d_b=16, d_f=18, t=3, f_u=400, F_t0_Sd=100, F_t_Rd=92.7)
    assert r.beta < 0 and "altere o parafuso" in r.aviso


def test_largura_tributaria():
    assert bc.largura_tributaria_alavanca(35, e_ext=40, e_int=100) == 40 + 50
    assert bc.largura_tributaria_alavanca(20, e_ext=40, e_int=100) == 35 + 35  # 1,75·b
    assert bc.largura_tributaria_alavanca(20, e_int=100, extremo=False) == 70  # 3,5·b


def test_alavanca_na_verificacao():
    base = caso_a(N=0.0, V=0.0, N_t=5.0, valores_sao_de_calculo=True)  # 2,5 kN por parafuso
    rigida = chk.verificar_ligacao(
        base.__class__(
            **{**base.__dict__, "alavanca": True, "alav_a": 35.0, "alav_b": 35.0, "alav_p": 60.0}
        )
    )
    v = linha(rigida, "placa rígida")
    assert perto(v.solicitante, 3.52) and v.status == "OK"
    fina = dict(base.__dict__, alavanca=True, alav_a=35.0, alav_b=35.0, alav_p=60.0, alav_t=2.0)
    flexivel = chk.verificar_ligacao(chk.EntradaLigacao(**fina))
    assert linha(flexivel, "placa rígida").status == "INFO"
    assert tem_linha(flexivel, "placa flexível")
    assert (
        linha(flexivel, "Tração no parafuso com alavanca").status == "ALERTA"
    )  # falta a força com alavanca
    com_forca = chk.verificar_ligacao(chk.EntradaLigacao(**{**fina, "F_t_com_alavanca": 3.0}))
    assert perto(linha(com_forca, "Tração no parafuso").solicitante, 3.0)
    with pytest.raises(ValueError, match="não pode ser menor"):
        chk.verificar_ligacao(chk.EntradaLigacao(**{**fina, "F_t_com_alavanca": 2.0}))


def test_alavanca_2008_reduz_a_tracao_resistente():
    base = caso_a(norma="NBR8800_2008", N=0.0, V=0.0, N_t=5.0, valores_sao_de_calculo=True)
    r = chk.verificar_ligacao(
        chk.EntradaLigacao(
            **dict(base.__dict__, alavanca=True, alav_a=30.0, alav_b=35.0, alav_p=60.0)
        )
    )
    F_t_Rd = bc.resist_tracao_parafuso(16, "A325", "NBR8800_2008")
    assert perto(linha(r, "Tração no parafuso").resistente, 0.67 * F_t_Rd)
    assert perto(
        linha(r, "espessura da chapa fletida").solicitante,
        math.sqrt(4 * 2500 * 35 * 1.1 / (60 * 250)),
    )
    assert linha(r, "a ≥ b").status == "NÃO OK"  # a = 30 < b = 35


# ------------------------------------------------------------------ caso D: deslizamento
def test_caso_d_deslizamento():
    assert perto(
        bc.resist_deslizamento(91, 0.20, NBR_2025, nivel="elu", F_t=1.0, gama_e=1.20), 16.97
    )
    assert perto(bc.resist_deslizamento(91, 0.20, NBR_2025, nivel="servico", F_t=0.7), 14.42)


def test_caso_d_na_verificacao():
    """Furo alargado: deslizamento no ELU, com F_t,Sd = 1 kN por parafuso; furo-padrão: serviço
    com F_t,Sk = 0,7 kN (70 % de F_t,Sd, a simplificação da norma: γ_f = 1/0,7)."""
    alargado = chk.verificar_ligacao(caso_a(tipo_furo=bc.FURO_ALARGADO, N_t=1.4, gama_f=1 / 0.7))
    assert perto(linha(alargado, "Deslizamento (ELU").resistente, 16.97)
    padrao = chk.verificar_ligacao(caso_a(N_t=1.4, gama_f=1 / 0.7))
    assert perto(linha(padrao, "Deslizamento (ELS").resistente, 14.42)
    muito = chk.verificar_ligacao(caso_a(tipo_furo=bc.FURO_MUITO_PARALELO))
    assert "/1.40·" in linha(muito, "Deslizamento (ELU").formula  # γ_e = 1,40


def test_deslizamento_2008_usa_o_fator_de_furo():
    r = chk.verificar_ligacao(caso_a(norma="NBR8800_2008", tipo_furo=bc.FURO_ALARGADO, C_e=0.85))
    v = linha(r, "Deslizamento (ELU")
    # 2008: 1,13·μ·C_h·F_Tb·n_s/γ_e com C_h = 0,85 (alargado) e γ_e = 1,20; μ = 0,20 galvanizada.
    assert perto(v.resistente, 1.13 * 0.20 * 0.85 * 91 / 1.20)
    assert linha(r, "C_e na NBR 8800:2008").status == "INFO"
    transversal = chk.verificar_ligacao(
        caso_a(norma="NBR8800_2008", tipo_furo=bc.FURO_MUITO_TRANSVERSAL)
    )
    assert tem_linha(transversal, "Deslizamento (ELS")  # 2008: alongado transversal é serviço


# ------------------------------------------------------------------ caso E: chapa de ligação
def test_caso_e_chapa_de_ligacao():
    c = bc.chapa_de_ligacao(70, 6.35, 1, 18, 250, 400, NBR_2025)
    assert perto(c.A_g, 444.5) and perto(c.A_n, 317.5) and perto(0.85 * c.A_g, 377.8)
    assert c.A_e == c.A_n  # o menor: o erro comum é usar 0,85·A_g mesmo com A_n menor
    assert perto(c.escoamento_kN, 101.0) and perto(c.ruptura_kN, 94.1)


def test_caso_e_na_verificacao_com_espessura_minima_em_mm():
    r = chk.verificar_ligacao(
        caso_a(
            N=60.0,
            V=0.0,
            valores_sao_de_calculo=True,
            chapa_largura=70.0,
            chapa_t=6.35,
            chapa_furos=1,
        )
    )
    esc, rup = linha(r, "escoamento da seção bruta"), linha(r, "ruptura da seção líquida")
    assert perto(esc.resistente, 101.0) and perto(rup.resistente, 94.1)
    assert "governa A_n" in rup.formula and "C_t não se aplica" in rup.formula
    t_min = 60000 * 1.1 / (70 * 250)  # mm, nunca MPa
    assert f"t mínimo = {t_min:.2f} mm" in esc.formula


# ------------------------------------------------------------------ furos
def test_tabela_14_da_2025_bate_com_a_foto():
    dados = json.loads((RAIZ / "data" / "normas_parafusos.json").read_text(encoding="utf-8"))
    for d, linha_ in dados["furos"]["NBR8800_2025"]["mm"].items():
        nome = f"M{d}"
        if nome not in bc.PARAFUSOS:
            continue
        for familia, valor in linha_.items():
            tipo = familia if familia in ("padrao", "alargado") else f"{familia}_paralelo"
            par = valor if isinstance(valor, list) else [valor, valor]
            assert bc.dimensoes_do_furo(nome, tipo, NBR_2025) == pytest.approx(tuple(par)), (
                nome,
                tipo,
            )
    assert bc.dimensoes_do_furo('5/8"', bc.FURO_ALARGADO, NBR_2025)[0] == pytest.approx(
        13 / 16 * 25.4
    )
    assert bc.dimensoes_do_furo('1"', bc.FURO_MUITO_PARALELO, NBR_2025) == pytest.approx(
        (28.575, 63.5)
    )
    # O furo-padrão da tabela do programa é o da 2025.
    assert all(
        bc.dimensoes_do_furo(n, bc.FURO_PADRAO, NBR_2025)[0]
        == pytest.approx(bc.PARAFUSOS[n][1], abs=0.06)
        for n in bc.PARAFUSOS
    )


def test_tabela_12_da_2008():
    assert bc.dimensoes_do_furo("M16", bc.FURO_PADRAO, "NBR8800_2008") == (17.5, 17.5)
    assert bc.dimensoes_do_furo("M16", bc.FURO_ALARGADO, "NBR8800_2008") == (21.0, 21.0)
    assert bc.dimensoes_do_furo("M16", bc.FURO_POUCO_PARALELO, "NBR8800_2008") == (17.5, 22.0)
    assert bc.dimensoes_do_furo("M16", bc.FURO_MUITO_PARALELO, "NBR8800_2008") == (17.5, 40.0)
    assert bc.dimensoes_do_furo("M27", bc.FURO_PADRAO, "NBR8800_2008") == (28.5, 28.5)
    assert bc.dimensoes_do_furo("M30", bc.FURO_POUCO_PARALELO, "NBR8800_2008") == (31.5, 39.5)


def test_furo_alongado_muda_l_f_e_a_secao_liquida():
    paralelo = chk.verificar_ligacao(caso_a(tipo_furo=bc.FURO_POUCO_PARALELO))
    transversal = chk.verificar_ligacao(caso_a(tipo_furo=bc.FURO_POUCO_TRANSVERSAL))
    # Paralelo: ℓ_f = 35 − 22/2 = 24 mm; transversal: 35 − 18/2 = 26 mm (como o padrão).
    assert perto(paralelo.resistencias["F_c_Rd_critico_kN"], 1.2 * 24 * 6.35 * 400 / 1.35 / 1000)
    assert perto(transversal.resistencias["F_c_Rd_critico_kN"], 58.7)
    muito = chk.verificar_ligacao(caso_a(tipo_furo=bc.FURO_MUITO_TRANSVERSAL))
    assert perto(muito.resistencias["F_c_Rd_critico_kN"], 1.0 * 26 * 6.35 * 400 / 1.35 / 1000)


def test_furo_alargado_so_por_atrito():
    r = chk.verificar_ligacao(caso_a(tipo_furo=bc.FURO_ALARGADO, ligacao_por_atrito=False))
    assert linha(r, "Tipo de furo").status == "NÃO OK"
    ok = chk.verificar_ligacao(
        caso_a(tipo_furo=bc.FURO_POUCO_TRANSVERSAL, ligacao_por_atrito=False)
    )
    assert linha(ok, "Tipo de furo").status == "OK"


def test_borda_minima_com_incremento_na_2008_e_a_confirmar_na_2025():
    r = chk.verificar_ligacao(caso_a(norma="NBR8800_2008", tipo_furo=bc.FURO_ALARGADO))
    v = linha(r, "Borda mínima com furo alargado")
    assert perto(v.solicitante, 22 + 0.12 * 16) and v.resistente == 35.0
    r2025 = chk.verificar_ligacao(caso_a(tipo_furo=bc.FURO_ALARGADO))
    assert linha(r2025, "Borda mínima com furo alargado").status == "ALERTA"
    assert bc.incremento_de_borda_2008(16, bc.FURO_MUITO_PARALELO, True) == 0.75 * 16
    assert bc.incremento_de_borda_2008(16, bc.FURO_MUITO_PARALELO, False) == 0.0


# ------------------------------------------------------------------ entradas
def test_fu_menor_ou_igual_a_fy_e_erro_com_sugestao():
    with pytest.raises(ValueError, match="A36.*250.*400"):
        chk.verificar_ligacao(caso_a(aco=chk.ACO_INFORMADO, f_y=250.0, f_u=250.0))
    r = chk.verificar_ligacao(caso_a(aco=chk.ACO_INFORMADO, f_y=250.0, f_u=400.0))
    assert perto(r.resistencias["F_c_Rd_critico_kN"], 58.7)


def test_espacamento_maximo_compara_o_passo_real():
    r = chk.verificar_ligacao(caso_a(s=200.0))
    v = linha(r, "Espaçamento máximo")
    assert v.solicitante == 200.0 and v.resistente == min(24 * 6.35, 300) and v.status == "NÃO OK"


def test_todo_atende_compara_solicitacao_com_resistencia():
    r = chk.verificar_ligacao(
        caso_a(
            N_t=2.0,
            tipo_furo=bc.FURO_ALARGADO,
            chapa_largura=70.0,
            alavanca=True,
            alav_a=35.0,
            alav_b=35.0,
            alav_p=60.0,
        )
    )
    for v in r.verificacoes:
        if v.aproveitamento is not None:
            assert v.aproveitamento == pytest.approx(v.solicitante / v.resistente), v.nome


def test_registro_traz_a_tracao_e_nao_a_nota_de_versao():
    entrada = caso_a(N_t=2.0, tipo_furo=bc.FURO_ALARGADO)
    resultado = chk.verificar_ligacao(entrada)
    registro = chk.registro_ligacao(entrada, resultado, "M16")
    assert registro["entradas"]["tipo_de_furo"] == "Alargado"
    assert perto(registro["resultados"]["tracao_resistente_F_t_Rd_kN"], 92.7)
    assert chk.NOTA_DA_VERSAO not in json.dumps(registro, ensure_ascii=False)
    assert "ABNT NBR 8800:2025" in registro["titulo"] or "NBR 8800:2025" in " ".join(
        registro["referencias"]
    )


# ------------------------------------------------------------------ memorial desta ligação
def test_memorial_da_ligacao_em_word_e_pdf_sem_projeto():
    from io import BytesIO

    from docx import Document
    from pypdf import PdfReader

    entrada = caso_a(N_t=2.0, interacao_por_tabela=True, chapa_largura=70.0)
    resultado = chk.verificar_ligacao(entrada)
    word = chk.memorial_da_ligacao(entrada, resultado, "M16")
    doc = Document(BytesIO(word))
    textos = " ".join(
        [p.text for p in doc.paragraphs]
        + [c.text for tabela in doc.tables for linha_ in tabela.rows for c in linha_.cells]
    )
    for trecho in ("Tração no parafuso", "Tração + corte", "Chapa de ligação", "NBR 8800:2025"):
        assert trecho in textos, trecho
    assert chk.NOTA_DA_VERSAO not in textos  # a nota de versão é só da página
    pdf = chk.memorial_da_ligacao(entrada, resultado, "M16", formato="pdf")
    texto_pdf = " ".join(pagina.extract_text() or "" for pagina in PdfReader(BytesIO(pdf)).pages)
    assert "Tração no parafuso" in texto_pdf
