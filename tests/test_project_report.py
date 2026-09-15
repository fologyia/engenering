from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfReader

from core.materials_registry import criar_material_projeto
from core.pdf_fonts import fonte_pdf
from core.project_report import (
    A_PREENCHER,
    MODULOS_FORA_DO_MEMORIAL,
    SECOES_RELATORIO,
    avaliar_integridade_registro,
    gerar_relatorio_industrial_pdf,
    gerar_relatorio_industrial_word,
    montar_modelo_relatorio,
    sintetizar_calculos,
)
from core.project_store import criar_item
from core.technical_records import criar_registro_tecnico
from tests.test_project_validation import _projeto_documentado


def _texto_word(conteudo: bytes) -> str:
    documento = Document(BytesIO(conteudo))
    paragrafos = [paragrafo.text for paragrafo in documento.paragraphs]
    celulas = [
        celula.text
        for tabela in documento.tables
        for linha in tabela.rows
        for celula in linha.cells
    ]
    return "\n".join(paragrafos + celulas)


def _texto_pdf(conteudo: bytes) -> str:
    return "\n".join(pagina.extract_text() or "" for pagina in PdfReader(BytesIO(conteudo)).pages)


def test_modelo_relatorio_filtra_secoes_e_registros():
    projeto = _projeto_documentado()
    modelo = montar_modelo_relatorio(
        projeto,
        secoes_incluidas=["escopo", "plano_calculo", "conclusao"],
        registros_ids=[],
    )
    titulos = [item["titulo"] for item in modelo["secoes"]]
    assert len(titulos) == 3
    assert any("Quadro-resumo dos cálculos" in titulo for titulo in titulos)
    assert modelo["registros"] == []
    assert modelo["sintese"]["texto"] == "Nenhum cálculo anexado"


def test_relatorio_word_e_pdf_sao_validos_e_contem_identificacao():
    projeto = _projeto_documentado()
    word = gerar_relatorio_industrial_word(projeto)
    pdf = gerar_relatorio_industrial_pdf(projeto)

    assert word.startswith(b"PK")
    texto_word = _texto_word(word)
    assert "Memorial de cálculo do projeto industrial" in texto_word
    assert "Memória de cálculo" in texto_word
    assert "Barra B12" in texto_word

    assert pdf.startswith(b"%PDF")
    leitor = PdfReader(BytesIO(pdf))
    assert len(leitor.pages) >= 2
    texto_pdf = _texto_pdf(pdf)
    assert "PR-204" in texto_pdf
    assert "Resumo executivo" in texto_pdf
    assert "Barra B12" in texto_pdf


def test_memorial_nao_traz_validacao_checklist_normas_cargas_nem_apendices():
    # O memorial é um molde de cálculo: o que é gestão do programa (validação,
    # checklist, bloqueios, matriz normativa, casos de carga) fica nas páginas
    # de gestão, e os apêndices que duplicavam os números saíram.
    projeto = _projeto_documentado()
    for chave in ("validacao", "checklist", "normas", "carregamentos"):
        assert chave not in SECOES_RELATORIO
    modelo = montar_modelo_relatorio(projeto)
    assert "validacao" not in modelo
    assert "status" not in modelo

    for texto in (
        _texto_word(gerar_relatorio_industrial_word(projeto)),
        _texto_pdf(gerar_relatorio_industrial_pdf(projeto)),
    ):
        for proibido in (
            "Central de validação",
            "Pendências e checklist",
            "Matriz normativa",
            "Casos, combinações",
            "ABNT NBR 8800",
            "Revisão independente",
            "bloqueio",
            "PRONTO PARA",
            "Índice documental",
            "Integração com outras partes",
            "Apêndice",
            "Snapshot",
        ):
            assert proibido not in texto, proibido


def test_registros_de_mohr_e_casos_de_carga_ficam_fora_do_memorial():
    projeto = _projeto_documentado()
    mohr = criar_registro_tecnico(
        modulo="Círculo de Mohr",
        modulo_id="circulo_mohr",
        titulo="Transformação do estado plano",
        status="Calculado",
        resumo="",
        entradas={"sigma_x_MPa": 120.0, "sigma_y_MPa": 30.0, "tau_xy_MPa": 25.0},
        resultados={"sigma_1_MPa": 126.5, "sigma_2_MPa": 23.5},
    )
    cargas = criar_registro_tecnico(
        modulo="Casos de carga",
        modulo_id="casos_carga",
        titulo="Envelope dos carregamentos",
        status="Calculado",
        resumo="",
        entradas={},
        resultados={},
    )
    projeto["registros_tecnicos"] += [mohr, cargas]
    assert {"circulo_mohr", "casos_carga"} == set(MODULOS_FORA_DO_MEMORIAL)

    modelo = montar_modelo_relatorio(projeto)
    assert [item["titulo"] for item in modelo["registros"]] == ["Barra B12"]
    # Pedir o registro pelo id também não o traz de volta.
    forcado = montar_modelo_relatorio(projeto, registros_ids=[mohr["id"], cargas["id"]])
    assert forcado["registros"] == []

    for texto in (
        _texto_word(gerar_relatorio_industrial_word(projeto)),
        _texto_pdf(gerar_relatorio_industrial_pdf(projeto)),
    ):
        assert "Transformação do estado plano" not in texto
        assert "Envelope dos carregamentos" not in texto
        assert "Mohr" not in texto


def test_campos_vazios_saem_como_marca_para_preencher():
    projeto = _projeto_documentado()
    projeto.update({"objetivo": "", "descricao": "", "cliente": ""})
    projeto["base_projeto"]["limitacoes"] = ""
    projeto["registros_tecnicos"][0]["metodo"] = ""
    projeto["registros_tecnicos"][0]["conclusao"] = ""

    word = _texto_word(gerar_relatorio_industrial_word(projeto))
    pdf = _texto_pdf(gerar_relatorio_industrial_pdf(projeto))
    for texto in (word, pdf):
        assert f"Objetivo: {A_PREENCHER}" in texto
        assert f"Método: {A_PREENCHER}" in texto
        assert f"Conclusão: {A_PREENCHER}" in texto
        assert f"Conclusão geral: {A_PREENCHER}" in texto
        assert f"Recomendações: {A_PREENCHER}" in texto
        assert "Não informado" not in texto
        assert "ainda não documentad" not in texto
    # A assinatura fica em branco de verdade, e não "Não informado".
    assert "____/____/________" in pdf


def test_listas_de_resultados_viram_tabelas_proprias():
    projeto = _projeto_documentado()
    projeto["registros_tecnicos"][0]["resultados"] = {
        "momento_maximo_kNm": 67.5,
        "reacoes": [
            {"x (m)": 0.0, "Apoio": "pino", "Fy (kN)": 45.0},
            {"x (m)": 6.0, "Apoio": "rolete", "Fy (kN)": 45.0},
        ],
        "material_id": "uuid-que-nao-deve-aparecer",
    }
    modelo = montar_modelo_relatorio(projeto, secoes_incluidas=["registros"])
    capitulo = next(item for item in modelo["secoes"] if item.get("nivel", 1) >= 2)
    tabelas = {tabela["legenda"]: tabela for tabela in capitulo["tabelas"]}
    assert "Reacoes." in tabelas
    reacoes = tabelas["Reacoes."]
    assert reacoes["cabecalhos"] == ["X (m)", "Apoio", "Fy (kN)"]
    assert reacoes["linhas"] == [["0", "pino", "45"], ["6", "rolete", "45"]]
    assert sum(reacoes["larguras"]) == 9360
    resultados = tabelas["Resultados registrados."]
    assert [linha[0] for linha in resultados["linhas"]] == ["Momento maximo [kN·m]"]
    assert "uuid-que-nao-deve-aparecer" not in str(capitulo)


def test_secoes_sem_conteudo_sao_omitidas():
    # Sem viga nem análise de sensibilidade registradas, as seções delas não
    # entram — mesmo selecionadas — em vez de sair só com tabelas vazias.
    projeto = _projeto_documentado()
    modelo = montar_modelo_relatorio(projeto)
    titulos = " ".join(item["titulo"] for item in modelo["secoes"])
    assert "Vigas e eixos" not in titulos
    assert "Sensibilidade" not in titulos
    assert modelo["secoes"][-1]["titulo"].endswith("Conclusão e recomendações")
    assert (
        modelo["numero_aprovacoes"]
        == len([item for item in modelo["secoes"] if item.get("nivel", 1) == 1]) + 3
    )


def test_sintese_conta_as_situacoes_declaradas_pelos_modulos():
    registros = [
        {"status": "Atende"},
        {"status": "Atende"},
        {"status": "Atenção"},
        {"status": "Não atende"},
        {"status": "Calculado"},
    ]
    sintese = sintetizar_calculos(registros)
    assert sintese["atendem"] == 2
    assert sintese["nao_atendem"] == 1
    assert sintese["outros"] == 1
    assert (
        sintese["texto"]
        == "5 cálculos — 2 atendem · 1 com atenção · 1 não atende · 1 sem verificação de critério"
    )
    assert sintetizar_calculos([{"status": "Atende"}])["texto"] == "1 cálculo — 1 atende"


def test_conclusao_lista_cada_calculo_e_deixa_a_geral_em_aberto():
    projeto = _projeto_documentado()
    modelo = montar_modelo_relatorio(projeto, secoes_incluidas=["conclusao"])
    conclusao = modelo["secoes"][0]
    assert conclusao["paragrafos"] == ["Síntese dos cálculos anexados: 1 cálculo — 1 atende."]
    assert conclusao["bullets"] == ["Barra B12 (Atende): Atende aos critérios informados."]
    assert conclusao["paragrafos_finais"] == [
        f"Conclusão geral: {A_PREENCHER}",
        f"Recomendações: {A_PREENCHER}",
    ]


@pytest.mark.skipif(
    not fonte_pdf().unicode, reason="sem fonte TrueType no sistema, o PDF cai em Helvetica"
)
def test_pdf_imprime_simbolos_gregos_das_equacoes():
    projeto = _projeto_documentado()
    projeto["registros_tecnicos"][0]["equacoes"] = ["σvm = √(σx² − σxσy + σy² + 3τxy²)"]
    texto = _texto_pdf(gerar_relatorio_industrial_pdf(projeto))
    assert "σvm = √(σx² − σxσy + σy² + 3τxy²)" in texto
    assert "sigma_vm" not in texto


def test_memorial_modular_preserva_ordem_materiais_sensibilidade_e_hash():
    projeto = _projeto_documentado()
    material = criar_material_projeto(
        nome="Material rastreado",
        familia="Aço",
        condicao="Normalizado",
        forma_produto="Chapa 10 mm",
        lote="L-1",
        propriedades={
            "Sut_MPa": 430,
            "Sy_MPa": 280,
            "temperatura_min_C": -20,
            "temperatura_max_C": 100,
        },
        origem_tipo="Certificado do lote / MTR",
        fonte="Usina",
        documento="MTR-1",
        data_verificacao="2026-08-14",
        responsavel_verificacao="Eng. A",
        aplicabilidade="Produto do projeto",
    )
    projeto["materiais_projeto"] = [material]
    sensibilidade = criar_item(
        modulo="Análise de sensibilidade",
        titulo="Robustez da utilização",
        status="Atende",
        resumo="OAT e Monte Carlo",
        metodo="OAT ±10% e 1000 amostras",
        equacoes=["U = demanda/capacidade"],
        entradas={"demanda": 80, "capacidade": 100},
        resultados={
            "saida_nominal": 0.8,
            "p05": 0.7,
            "p95": 0.92,
            "probabilidade_nao_atendimento_pct": 1.2,
            "ranking_sensibilidade": [
                {
                    "variavel": "Demanda",
                    "impacto_percentual": 20,
                    "elasticidade": 1,
                    "direcao_critica": "aumentar",
                }
            ],
            "correlacoes_spearman": [{"variavel": "Demanda", "correlacao": 0.8}],
        },
        premissas=["Entradas independentes"],
        referencias=["MC-204"],
        conclusao="Risco inferior a 5%.",
    )
    projeto["registros_tecnicos"].append(sensibilidade)
    ordem = [sensibilidade["id"], projeto["registros_tecnicos"][0]["id"]]
    modelo = montar_modelo_relatorio(projeto, registros_ids=ordem)
    assert [item["id"] for item in modelo["registros"]] == ordem
    assert any(item["titulo"].endswith("Materiais") for item in modelo["secoes"])
    assert any("Sensibilidade, incertezas" in item["titulo"] for item in modelo["secoes"])
    assert len(modelo["snapshot_hash"]) == 64
    assert avaliar_integridade_registro(sensibilidade)["percentual"] == 100

    alterado = montar_modelo_relatorio(projeto, registros_ids=list(reversed(ordem)))
    assert alterado["snapshot_hash"] != modelo["snapshot_hash"]


def test_numeros_grandes_nao_saem_em_notacao_cientifica():
    from core.project_report import _valor

    assert _valor(200000.0) == "200.000"
    assert _valor(123456.7) == "123.457"
    assert _valor(1963.5) == "1963,5"
    assert _valor(0.66051) == "0,66051"
    assert _valor(1e-7) == "1e-07"
