"""Lista de material e peso da estrutura: para o orçamento e para conferir o peso próprio do modelo.

Cada item é um perfil (quantidade × comprimento × kg/m do catálogo), uma chapa (comprimento ×
largura × espessura × 7 850 kg/m³), uma grade ou piso (área × kg/m²) ou outro item (kg por
unidade). A lista dá a massa, o peso, a área de pintura, o resumo por perfil (com as barras
comerciais) e por espessura de chapa, e confere o peso com a reação vertical do caso PP do modelo.

A lista de corte do SolidWorks (CSV ou Excel) entra pela :func:`ler_lista_de_corte`, que acha as
colunas pelo nome (quantidade, descrição, comprimento) e casa a descrição com o catálogo.

É ferramenta de gestão: não vai para o memorial. Sem Streamlit.
"""

from __future__ import annotations

import csv
import io
import math
import re
import unicodedata
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, field, replace
from pathlib import Path
from typing import Any

from core import cantoneiras as ct
from core import esforcos_modelo as em
from core import plano_de_cargas as pc
from core import section_catalog as sc

TIPO_PERFIL = "Perfil"
TIPO_CHAPA = "Chapa"
TIPO_GRADE = "Grade / piso"
TIPO_OUTRO = "Outro"
TIPOS = (TIPO_PERFIL, TIPO_CHAPA, TIPO_GRADE, TIPO_OUTRO)
DENSIDADE_ACO = 7850.0  # kg/m³
G = 9.80665
ACRESCIMO_PADRAO = 5.0  # % de ligações, parafusos e soldas
BARRA_PADRAO_M = 12.0
TOLERANCIA_PP = 0.05  # diferença aceitável entre a lista e o caso PP do modelo
CASO_PP = "PP"


# ---------------------------------------------------------------------------------------------
# Itens
# ---------------------------------------------------------------------------------------------
@dataclass(frozen=True)
class ItemDaLista:
    marca: str = ""
    tipo: str = TIPO_PERFIL
    descricao: str = ""  # perfil do catálogo ou descrição livre
    quantidade: int = 1
    comprimento_m: float = 0.0
    largura_mm: float = 0.0  # chapa e grade
    espessura_mm: float = 0.0  # chapa
    #: kg/m (perfil fora do catálogo), kg/m² (grade) ou kg por unidade (outro). Vazio no perfil do
    #: catálogo e na chapa, que o programa calcula.
    massa_unitaria: float | None = None
    observacao: str = ""
    #: Massa de uma peça pela geometria do modelo (volume × 7 850 kg/m³), da macro do SolidWorks.
    #: Completa o item que o catálogo não conhece e é o peso que a gravidade do modelo enxerga.
    massa_geometria_kg: float | None = None


@dataclass(frozen=True)
class ListaDeMaterial:
    itens: tuple[ItemDaLista, ...] = ()
    acrescimo_pct: float = ACRESCIMO_PADRAO
    comprimento_barra_m: float = BARRA_PADRAO_M


@dataclass(frozen=True)
class PerfilConhecido:
    nome: str
    massa_kg_m: float
    perimetro_mm: float  # superfície a pintar por metro, em mm


@dataclass(frozen=True)
class LinhaCalculada:
    item: ItemDaLista
    perfil: PerfilConhecido | None
    massa_por_peca_kg: float | None
    massa_kg: float | None
    area_pintura_m2: float
    pendencia: str = ""


@dataclass(frozen=True)
class ResumoDoPerfil:
    perfil: str
    pecas: int
    comprimento_m: float
    massa_kg_m: float
    massa_kg: float
    barras: int


@dataclass(frozen=True)
class ResumoDaChapa:
    espessura_mm: float
    pecas: int
    area_m2: float
    massa_kg: float


@dataclass(frozen=True)
class ResumoDaLista:
    linhas: tuple[LinhaCalculada, ...]
    massa_itens_kg: float
    massa_acrescimo_kg: float
    area_pintura_m2: float
    por_tipo: Mapping[str, float]
    perfis: tuple[ResumoDoPerfil, ...]
    chapas: tuple[ResumoDaChapa, ...]
    acrescimo_pct: float = ACRESCIMO_PADRAO
    comprimento_barra_m: float = BARRA_PADRAO_M
    pendentes: tuple[LinhaCalculada, ...] = field(default=())

    @property
    def massa_geometria_kg(self) -> float | None:
        """Massa de todos os itens pela geometria do modelo, se todos com massa a tiverem."""
        com_massa = [x for x in self.linhas if x.massa_kg is not None]
        if not com_massa or any(x.item.massa_geometria_kg is None for x in com_massa):
            return None
        return sum((x.item.massa_geometria_kg or 0.0) * x.item.quantidade for x in com_massa)

    @property
    def massa_total_kg(self) -> float:
        return self.massa_itens_kg + self.massa_acrescimo_kg

    @property
    def peso_total_kN(self) -> float:
        return self.massa_total_kg * G / 1e3

    @property
    def peso_itens_kN(self) -> float:
        return self.massa_itens_kg * G / 1e3


# ---------------------------------------------------------------------------------------------
# Perfis: catálogo, cantoneiras e tubos ou barras descritos pelas medidas
# ---------------------------------------------------------------------------------------------
def _perimetro_do_catalogo(p: Any) -> float:
    """Superfície exposta por metro (mm) da seção do catálogo, pela família."""
    d, b, tw = p.altura_mm, p.largura_mm, p.espessura_alma_mm
    familia = p.familia
    if familia.startswith(("W", "HP", "I ", "U ")):
        return 2 * d + 4 * b - 2 * tw  # I e U: as faces externas e internas das mesas e da alma
    if familia.startswith("T ("):  # perfil tê (os tubos também começam com T)
        return 2 * d + 2 * b
    if familia.startswith("C enrijecido"):
        m = re.search(r"×\s*([\d.]+)\s*×\s*[\d.]+\s*$", p.nome)
        enrijecedor = float(m.group(1)) if m else 0.0
        return 2 * (d + 2 * b + 2 * enrijecedor)
    if familia in ("Tubo circular", "Barra circular"):
        return math.pi * d
    return 2 * (d + b)  # tubo retangular e barra retangular


def _sem_acento(texto: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn"
    ).lower()


def _decimal(texto: str) -> float:
    return float(texto.replace(",", "."))


def _por_medidas(texto: str) -> PerfilConhecido | None:
    """Tubos, barras e cantoneiras em mm descritos pelas medidas (como no SolidWorks)."""
    t = _sem_acento(texto).replace("×", "x")
    numero = r"(\d+(?:[.,]\d+)?)"
    m = re.search(
        rf"(?:tubo|tube|metalon)\s*(?:quadrado|retangular|square|rectangular|rect\.?)?\s*"
        rf"{numero}\s*x\s*{numero}\s*x\s*{numero}",
        t,
    )
    if m:
        h, b, e = (_decimal(m.group(i)) for i in (1, 2, 3))
        area = 2 * e * (h + b - 2 * e)
        return PerfilConhecido(texto.strip(), area * DENSIDADE_ACO * 1e-6, 2 * (h + b))
    m = re.search(
        rf"(?:tubo|tube|pipe)\s*(?:redondo|circular|round)?\s*[øo]?\s*{numero}\s*x\s*{numero}", t
    )
    if m:
        diametro, e = _decimal(m.group(1)), _decimal(m.group(2))
        area = math.pi * e * (diametro - e)
        return PerfilConhecido(texto.strip(), area * DENSIDADE_ACO * 1e-6, math.pi * diametro)
    m = re.search(
        rf"(?:barra|ferro|round bar|rod)\s*(?:redonda|redondo|circular)?\s*[øo]?\s*{numero}\b(?!\s*x)",
        t,
    )
    if m:
        diametro = _decimal(m.group(1))
        area = math.pi * diametro**2 / 4
        return PerfilConhecido(texto.strip(), area * DENSIDADE_ACO * 1e-6, math.pi * diametro)
    m = re.search(rf"(?:cantoneira|angle|l)\s*{numero}\s*x\s*{numero}\s*x\s*{numero}", t)
    if m and _decimal(m.group(1)) > 10:  # em mm (L 2 x 2 x 1/4 em polegadas fica para o catálogo)
        a, b, e = (_decimal(m.group(i)) for i in (1, 2, 3))
        area = e * (a + b - e)
        return PerfilConhecido(texto.strip(), area * DENSIDADE_ACO * 1e-6, 2 * (a + b))
    return None


def _sem_fracoes(texto: str) -> str:
    """``L 2 1/2 x 2 1/2 x 1/4`` → ``L 2.5 x 2.5 x 0.25``, como o SolidWorks escreve."""

    def mista(m: re.Match[str]) -> str:
        return f"{int(m.group(1)) + int(m.group(2)) / int(m.group(3)):g}"

    def simples(m: re.Match[str]) -> str:
        return f"{int(m.group(1)) / int(m.group(2)):g}"

    texto = re.sub(r"(\d+)\s+(\d+)/(\d+)", mista, texto)
    return re.sub(r"(\d+)/(\d+)", simples, texto).replace('"', "")


def perfil_conhecido(descricao: str) -> PerfilConhecido | None:
    """O perfil que a descrição indica: nome do catálogo, nome do SolidWorks ou medidas."""
    texto = descricao.strip()
    if not texto:
        return None
    perfis = sc.listar_perfis()
    nome = (
        texto
        if texto in perfis or texto in ct.CATALOGO_CANTONEIRAS
        else em.perfil_do_nome(_sem_fracoes(texto))
    )
    if nome in perfis:
        p = perfis[nome]
        return PerfilConhecido(nome, p.massa_kg_m, _perimetro_do_catalogo(p))
    if nome in ct.CATALOGO_CANTONEIRAS:
        c = ct.CATALOGO_CANTONEIRAS[nome]
        return PerfilConhecido(nome, c.massa_kg_m, 4 * c.b_mm)
    return _por_medidas(texto)


# ---------------------------------------------------------------------------------------------
# Cálculo
# ---------------------------------------------------------------------------------------------
def _n(valor: float, casas: int = 1) -> str:
    return f"{valor:,.{casas}f}".replace(",", " ").replace(".", ",")


def calcular_linha(item: ItemDaLista) -> LinhaCalculada:
    """Massa e pintura do item; sem os dados do tipo, vale a massa da geometria do modelo."""
    linha = _pelos_dados_do_item(item)
    geometria = item.massa_geometria_kg
    if linha.massa_kg is None and geometria:
        qtd = max(int(item.quantidade), 0)
        return LinhaCalculada(
            item, linha.perfil, geometria, geometria * qtd, 0.0, "massa pela geometria do modelo"
        )
    return linha


def _pelos_dados_do_item(item: ItemDaLista) -> LinhaCalculada:
    qtd = max(int(item.quantidade), 0)
    if item.tipo == TIPO_PERFIL:
        perfil = perfil_conhecido(item.descricao)
        massa_m = (
            item.massa_unitaria if item.massa_unitaria else (perfil.massa_kg_m if perfil else None)
        )
        if item.comprimento_m <= 0:
            return LinhaCalculada(item, perfil, None, None, 0.0, "informe o comprimento")
        if massa_m is None:
            return LinhaCalculada(
                item, None, None, None, 0.0, "perfil fora do catálogo: informe a massa (kg/m)"
            )
        por_peca = massa_m * item.comprimento_m
        pintura = (perfil.perimetro_mm / 1e3 if perfil else 0.0) * item.comprimento_m * qtd
        pendencia = "" if perfil else "sem a área de pintura (perfil fora do catálogo)"
        return LinhaCalculada(item, perfil, por_peca, por_peca * qtd, pintura, pendencia)
    if item.tipo == TIPO_CHAPA:
        if min(item.comprimento_m, item.largura_mm, item.espessura_mm) <= 0:
            return LinhaCalculada(
                item, None, None, None, 0.0, "informe comprimento, largura e espessura"
            )
        area = item.comprimento_m * item.largura_mm / 1e3
        por_peca = area * item.espessura_mm / 1e3 * DENSIDADE_ACO
        return LinhaCalculada(item, None, por_peca, por_peca * qtd, 2 * area * qtd)
    if item.tipo == TIPO_GRADE:
        if min(item.comprimento_m, item.largura_mm) <= 0 or not item.massa_unitaria:
            return LinhaCalculada(
                item, None, None, None, 0.0, "informe comprimento, largura e a massa (kg/m²)"
            )
        por_peca = item.comprimento_m * item.largura_mm / 1e3 * item.massa_unitaria
        return LinhaCalculada(item, None, por_peca, por_peca * qtd, 0.0)
    if not item.massa_unitaria:
        return LinhaCalculada(item, None, None, None, 0.0, "informe a massa por unidade (kg)")
    return LinhaCalculada(item, None, item.massa_unitaria, item.massa_unitaria * qtd, 0.0)


def resumir(lista: ListaDeMaterial) -> ResumoDaLista:
    linhas = tuple(calcular_linha(i) for i in lista.itens)
    massa = sum(linha.massa_kg or 0.0 for linha in linhas)
    por_tipo = {t: sum(x.massa_kg or 0.0 for x in linhas if x.item.tipo == t) for t in TIPOS}
    perfis: dict[str, list[LinhaCalculada]] = {}
    for linha in linhas:
        if linha.item.tipo == TIPO_PERFIL and linha.massa_kg is not None:
            nome = linha.perfil.nome if linha.perfil else linha.item.descricao.strip() or "—"
            perfis.setdefault(nome, []).append(linha)
    barra = lista.comprimento_barra_m if lista.comprimento_barra_m > 0 else BARRA_PADRAO_M
    resumo_perfis = []
    for nome, grupo in perfis.items():
        comprimento = sum(x.item.comprimento_m * x.item.quantidade for x in grupo)
        massa_grupo = sum(x.massa_kg or 0.0 for x in grupo)
        resumo_perfis.append(
            ResumoDoPerfil(
                nome,
                sum(x.item.quantidade for x in grupo),
                comprimento,
                massa_grupo / comprimento if comprimento else 0.0,
                massa_grupo,
                math.ceil(comprimento / barra - 1e-9),
            )
        )
    chapas: dict[float, list[LinhaCalculada]] = {}
    for linha in linhas:
        if linha.item.tipo == TIPO_CHAPA and linha.massa_kg is not None:
            chapas.setdefault(round(linha.item.espessura_mm, 2), []).append(linha)
    resumo_chapas = [
        ResumoDaChapa(
            espessura,
            sum(x.item.quantidade for x in grupo),
            sum(x.item.comprimento_m * x.item.largura_mm / 1e3 * x.item.quantidade for x in grupo),
            sum(x.massa_kg or 0.0 for x in grupo),
        )
        for espessura, grupo in sorted(chapas.items())
    ]
    return ResumoDaLista(
        linhas=linhas,
        massa_itens_kg=massa,
        massa_acrescimo_kg=massa * max(lista.acrescimo_pct, 0.0) / 100,
        area_pintura_m2=sum(x.area_pintura_m2 for x in linhas),
        por_tipo=por_tipo,
        perfis=tuple(sorted(resumo_perfis, key=lambda r: -r.massa_kg)),
        chapas=tuple(resumo_chapas),
        acrescimo_pct=lista.acrescimo_pct,
        comprimento_barra_m=barra,
        pendentes=tuple(x for x in linhas if x.massa_kg is None),
    )


# ---------------------------------------------------------------------------------------------
# Conferência com o peso próprio do modelo
# ---------------------------------------------------------------------------------------------
def conferir_com_o_modelo(
    resumo: ResumoDaLista, dados: em.EsforcosDoModelo, *, caso: str = CASO_PP
) -> list[pc.ItemDeConferencia]:
    """O peso dos itens (sem o acréscimo) contra a reação vertical do caso PP do modelo."""
    importado = dados.casos.get(caso)
    if importado is None or importado.reacoes is None:
        return [
            pc.ItemDeConferencia(
                pc.NIVEL_ATENCAO,
                f"Para conferir com o modelo, importe em Esforços do modelo o caso {caso} com o "
                "arquivo de reações (Listar forças resultantes).",
            )
        ]
    if resumo.massa_itens_kg <= 0:
        return [
            pc.ItemDeConferencia(pc.NIVEL_ATENCAO, "A lista ainda não tem massa para comparar.")
        ]
    _, _, vertical = em.para_o_plano(importado.reacoes.modelo, dados.eixo_vertical)
    geometria = resumo.massa_geometria_kg
    massa = geometria if geometria else resumo.massa_itens_kg
    lista_kN = massa * G / 1e3
    diferenca = (vertical - lista_kN) / lista_kN
    origem = (
        "pela geometria do modelo (aço, 7 850 kg/m³)"
        if geometria
        else "pela lista, sem o acréscimo"
    )
    texto = (
        f"Caso {caso} do modelo: reação vertical {_n(vertical, 2)} kN; peso {origem}: "
        f"{_n(lista_kN, 2)} kN ({_n(massa, 0)} kg) — diferença de {_n(100 * diferenca, 1)} %."
    )
    if abs(diferenca) <= TOLERANCIA_PP:
        return [
            pc.ItemDeConferencia(pc.NIVEL_OK, texto + " O peso próprio do modelo bate com a lista.")
        ]
    if diferenca < 0:
        motivo = (
            " O modelo pesa menos que a lista: falta barra no modelo, o peso próprio não está em "
            "todas as peças ou a lista tem itens que o modelo não representa (placas, grades)."
        )
    else:
        motivo = (
            " O modelo pesa mais que a lista: a lista está incompleta ou o caso PP tem outras "
            "cargas além do peso próprio do aço."
        )
    razao = vertical / lista_kN if lista_kN > 0 else 0.0
    if 9.5 <= razao <= 10.5 or 950 <= razao <= 1050:
        motivo += (
            f" A razão é quase {round(razao, -1 if razao < 100 else -3):.0f} vezes — típico de "
            "unidade trocada: confira no estudo a gravidade (9,81 m/s²) e a densidade do material "
            "(7 850 kg/m³)."
        )
    return [pc.ItemDeConferencia(pc.NIVEL_ATENCAO, texto + motivo)]


# ---------------------------------------------------------------------------------------------
# Projeto
# ---------------------------------------------------------------------------------------------
def _float(valor: Any, padrao: float = 0.0) -> float:
    try:
        numero = float(valor)
    except (TypeError, ValueError):
        return padrao
    return numero if math.isfinite(numero) else padrao


def item_de_dicionario(d: Mapping[str, Any]) -> ItemDaLista:
    tipo = str(d.get("tipo") or TIPO_PERFIL)
    massa = _float(d.get("massa_unitaria"), 0.0)
    geometria = _float(d.get("massa_geometria_kg"), 0.0)
    return ItemDaLista(
        marca=str(d.get("marca") or ""),
        tipo=tipo if tipo in TIPOS else TIPO_OUTRO,
        descricao=str(d.get("descricao") or ""),
        quantidade=max(int(_float(d.get("quantidade"), 1.0)), 0),
        comprimento_m=max(_float(d.get("comprimento_m")), 0.0),
        largura_mm=max(_float(d.get("largura_mm")), 0.0),
        espessura_mm=max(_float(d.get("espessura_mm")), 0.0),
        massa_unitaria=massa if massa > 0 else None,
        observacao=str(d.get("observacao") or ""),
        massa_geometria_kg=geometria if geometria > 0 else None,
    )


def para_dicionario(lista: ListaDeMaterial) -> dict[str, Any]:
    return {
        "itens": [dict(i.__dict__) for i in lista.itens],
        "acrescimo_pct": lista.acrescimo_pct,
        "comprimento_barra_m": lista.comprimento_barra_m,
    }


def de_dicionario(dados: Mapping[str, Any] | None) -> ListaDeMaterial:
    if not isinstance(dados, Mapping):
        return ListaDeMaterial()
    itens = tuple(item_de_dicionario(i) for i in dados.get("itens") or () if isinstance(i, Mapping))
    barra = _float(dados.get("comprimento_barra_m"), BARRA_PADRAO_M)
    return ListaDeMaterial(
        itens=itens,
        acrescimo_pct=max(_float(dados.get("acrescimo_pct"), ACRESCIMO_PADRAO), 0.0),
        comprimento_barra_m=barra if barra > 0 else BARRA_PADRAO_M,
    )


def lista_do_projeto(projeto: Mapping[str, Any] | None) -> ListaDeMaterial:
    if not isinstance(projeto, Mapping):
        return ListaDeMaterial()
    return de_dicionario(projeto.get("lista_de_material"))


def com_itens(lista: ListaDeMaterial, itens: Iterable[ItemDaLista]) -> ListaDeMaterial:
    return replace(lista, itens=tuple(itens))


# ---------------------------------------------------------------------------------------------
# Placas de base dos pilares (da página Esforços do modelo)
# ---------------------------------------------------------------------------------------------
MARCA_PLACA_DE_BASE = "PB"


def itens_das_placas_de_base(dados: em.EsforcosDoModelo) -> list[ItemDaLista]:
    """Uma linha de chapa com a placa padrão para cada pilar marcado no modelo."""
    from core import placa_base_pilares as pb

    pilares = [n for n, m in dados.membros.items() if m.tipo == "Pilar"]
    if not pilares:
        return []
    placa = pb.placa_do_modelo(dados)
    return [
        ItemDaLista(
            marca=MARCA_PLACA_DE_BASE,
            tipo=TIPO_CHAPA,
            descricao=f"Placa de base {placa.aco}",
            quantidade=len(pilares),
            comprimento_m=placa.comprimento_mm / 1e3,
            largura_mm=placa.largura_mm,
            espessura_mm=placa.espessura_mm,
            observacao="da verificação das placas de base (Esforços do modelo)",
        )
    ]


def com_placas_de_base(lista: ListaDeMaterial, dados: em.EsforcosDoModelo) -> ListaDeMaterial:
    """Troca as linhas de placa de base pela placa atual (não duplica ao repetir)."""
    outras = [i for i in lista.itens if i.marca != MARCA_PLACA_DE_BASE]
    return com_itens(lista, [*outras, *itens_das_placas_de_base(dados)])


# ---------------------------------------------------------------------------------------------
# Lista de corte do SolidWorks (CSV ou Excel)
# ---------------------------------------------------------------------------------------------
_COLUNA_QTD = re.compile(r"^(qtd|qtde|quant|qty|quantity|quantidade)")
#: Em ordem de preferência: "DESCRIÇÃO" vale mais que "PERFIL", que vale mais que "MATERIAL".
_COLUNAS_DESCRICAO = (
    re.compile(r"(descri|description)"),
    re.compile(r"(perfil|profile)"),
    re.compile(r"(nome|material)"),
)
_COLUNA_COMPRIMENTO = re.compile(r"(compr|length|tamanho)")
_COLUNA_MARCA = re.compile(r"^(n[o°º]?\.?\s*do\s*item|item|marca|pos|posicao|mark)")
_COLUNA_MASSA_GEOMETRIA = re.compile(r"^massa do aco")  # coluna da macro do SolidWorks
_CHAPA = re.compile(r"^\s*(chapa|placa|pl\b|ch\b|plate)", re.IGNORECASE)


class _PontoEVirgula(csv.excel):
    delimiter = ";"


def _celulas_do_arquivo(conteudo: bytes, nome_arquivo: str) -> list[list[str]]:
    if nome_arquivo.lower().endswith(".xls"):
        raise ValueError(
            f"{nome_arquivo} está no formato antigo do Excel (.xls), que o programa não lê: salve a "
            "tabela do SolidWorks como CSV ou abra o arquivo no Excel e salve como .xlsx."
        )
    if nome_arquivo.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            livro = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        except Exception as erro:  # arquivo corrompido ou que não é Excel
            raise ValueError(f"Não consegui abrir o Excel {nome_arquivo}: {erro}") from erro
        folha = livro.worksheets[0]
        return [
            ["" if c is None else str(c) for c in linha]
            for linha in folha.iter_rows(values_only=True)
        ]
    for codificacao in ("utf-8-sig", "cp1252"):
        try:
            texto = conteudo.decode(codificacao)
            break
        except UnicodeDecodeError:
            continue
    else:
        texto = conteudo.decode("latin-1")
    amostra = "\n".join(texto.splitlines()[:20])
    try:
        dialeto: Any = csv.Sniffer().sniff(amostra, delimiters=";,\t")
    except csv.Error:
        dialeto = _PontoEVirgula
    return [list(linha) for linha in csv.reader(io.StringIO(texto), dialeto)]


_MILHAR = re.compile(r"\d{1,3}(\.\d{3})+")


def numero_da_celula(texto: str, unidade: str = "mm") -> tuple[float | None, str]:
    """Número em pt-BR ou en (com unidade opcional) → (valor, unidade lida ou a padrão)."""
    t = str(texto).strip().lower()
    # Espaço só como separador de milhar ("1 234,5"); "2 1/2" não vira 21.
    m = re.search(r"(-?\d{1,3}(?:\s\d{3})+(?:[.,]\d+)?|-?[\d.,]*\d)\s*(mm|cm|m|in|pol|ft|\")?", t)
    if not m:
        return None, unidade
    if m.group(2) and m.group(2) != "mm" and t[m.end() : m.end() + 1].isalpha():
        return None, unidade  # "12 mx" não é medida
    corpo = m.group(1).replace(" ", "")
    if "," in corpo and "." in corpo:
        decimal = "," if corpo.rfind(",") > corpo.rfind(".") else "."
        milhar = "." if decimal == "," else ","
        corpo = corpo.replace(milhar, "").replace(decimal, ".")
    elif "," in corpo:
        corpo = corpo.replace(",", ".")
    elif unidade == "mm" and m.group(2) in (None, "mm") and _MILHAR.fullmatch(corpo):
        corpo = corpo.replace(".", "")  # 1.234 mm no padrão brasileiro: ponto de milhar
    try:
        return float(corpo), m.group(2) or unidade
    except ValueError:
        return None, unidade


def _em_metros(valor: float, unidade: str) -> float:
    por_metro = {
        "mm": 1000.0,
        "cm": 100.0,
        "m": 1.0,
        "in": 39.37008,
        "pol": 39.37008,
        '"': 39.37008,
    }
    return valor / por_metro.get(unidade, 1000.0) if unidade != "ft" else valor * 0.3048


MACRO_LISTA_DE_CORTE = (
    Path(__file__).resolve().parent.parent / "macros_solidworks" / "exportar_lista_de_corte.bas"
)


def macro_da_lista_de_corte() -> bytes:
    """A macro do SolidWorks que exporta a lista de corte da peça para CSV (texto ASCII, CRLF).

    O CSV que ela grava tem as colunas ITEM, QTD., DESCRICAO (o nome do item sem o ``<n>``),
    COMPRIMENTO (com a unidade do documento), o nome do item e todas as propriedades da lista de
    corte — o formato que :func:`ler_lista_de_corte` lê.
    """
    texto = MACRO_LISTA_DE_CORTE.read_text(encoding="ascii")
    return texto.replace("\r\n", "\n").replace("\n", "\r\n").encode("ascii")


def _leitor_de_celulas(linha: Sequence[str]) -> Any:
    def celula(j: int | None) -> str:
        return str(linha[j]).strip() if j is not None and j < len(linha) else ""

    return celula


def ler_lista_de_corte(
    conteudo: bytes, nome_arquivo: str, *, unidade: str = "mm"
) -> tuple[list[ItemDaLista], list[str]]:
    """Itens da lista de corte e os avisos (colunas não achadas, linhas puladas, perfis estranhos)."""
    linhas = _celulas_do_arquivo(conteudo, nome_arquivo)
    cabecalho_i = None
    for i, linha in enumerate(linhas[:30]):
        nomes = [_sem_acento(c).strip() for c in linha]
        if any(_COLUNA_QTD.match(n) for n in nomes) and any(
            padrao.search(n) for padrao in _COLUNAS_DESCRICAO for n in nomes
        ):
            cabecalho_i = i
            break
    if cabecalho_i is None:
        return [], [
            "Não achei o cabeçalho: a lista precisa das colunas de quantidade (QTD.) e de "
            "descrição (DESCRIÇÃO), e de preferência a de comprimento."
        ]
    nomes = [_sem_acento(c).strip() for c in linhas[cabecalho_i]]

    def coluna(padrao: re.Pattern[str], *, inicio: bool = False) -> int | None:
        for j, n in enumerate(nomes):
            if (padrao.match(n) if inicio else padrao.search(n)) is not None:
                return j
        return None

    c_qtd = coluna(_COLUNA_QTD, inicio=True)
    c_desc = next((j for padrao in _COLUNAS_DESCRICAO if (j := coluna(padrao)) is not None), None)
    c_comp = coluna(_COLUNA_COMPRIMENTO)
    c_marca = coluna(_COLUNA_MARCA, inicio=True)
    avisos: list[str] = []
    if c_comp is None:
        avisos.append("Sem coluna de comprimento: preencha os comprimentos na tabela.")
    unidade_coluna = unidade
    if c_comp is not None:
        m = re.search(r"\((mm|cm|m)\)", nomes[c_comp])
        unidade_coluna = m.group(1) if m else unidade
    c_geometria = coluna(_COLUNA_MASSA_GEOMETRIA)
    itens: list[ItemDaLista] = []
    estranhos: list[str] = []
    pela_geometria: list[str] = []
    vazios = 0
    for linha in linhas[cabecalho_i + 1 :]:
        celula = _leitor_de_celulas(linha)
        descricao = celula(c_desc)
        qtd, _ = numero_da_celula(celula(c_qtd))
        if not descricao or qtd is None or qtd <= 0:
            vazios += bool(descricao) and qtd == 0
            continue
        comprimento_m = 0.0
        if c_comp is not None:
            valor, lida = numero_da_celula(celula(c_comp), unidade_coluna)
            comprimento_m = _em_metros(valor, lida) if valor else 0.0
        geometria, _ = numero_da_celula(celula(c_geometria))
        massa_geometria = geometria if geometria and geometria > 0 else None
        base = ItemDaLista(
            marca=celula(c_marca),
            descricao=descricao,
            quantidade=int(round(qtd)),
            comprimento_m=comprimento_m,
            massa_geometria_kg=massa_geometria,
        )
        if _CHAPA.match(descricao):
            espessura, _ = numero_da_celula(_CHAPA.sub("", descricao))
            itens.append(
                replace(
                    base,
                    tipo=TIPO_CHAPA,
                    espessura_mm=espessura or 0.0,
                    observacao="confira a largura da chapa",
                )
            )
            continue
        perfil = perfil_conhecido(descricao)
        if perfil is not None:
            itens.append(replace(base, descricao=perfil.nome))
        elif massa_geometria is not None:
            pela_geometria.append(descricao)
            itens.append(
                replace(
                    base,
                    tipo=TIPO_PERFIL if comprimento_m > 0 else TIPO_OUTRO,
                    observacao="fora do catálogo: massa pela geometria do modelo, como aço",
                )
            )
        else:
            estranhos.append(descricao)
            itens.append(replace(base, observacao="fora do catálogo: informe a massa (kg/m)"))
    if estranhos:
        amostra = ", ".join(sorted(set(estranhos))[:5])
        avisos.append(
            f"{len(set(estranhos))} descrição(ões) fora do catálogo ({amostra}): informe a massa "
            "(kg/m) na tabela."
        )
    if pela_geometria:
        amostra = ", ".join(sorted(set(pela_geometria))[:5])
        avisos.append(
            f"{len(pela_geometria)} item(ns) fora do catálogo ({amostra}) entram com a massa da "
            "geometria do modelo, como aço (7 850 kg/m³): confira o material de cada um."
        )
    if vazios:
        avisos.append(
            f"{vazios} item(ns) da lista de corte sem nenhum corpo (quantidade 0) foram ignorados — "
            "sobras de alterações no modelo."
        )
    if not itens:
        avisos.append("Nenhuma linha com quantidade e descrição foi lida.")
    return itens, avisos


# ---------------------------------------------------------------------------------------------
# Saídas
# ---------------------------------------------------------------------------------------------
COLUNAS = (
    "Marca",
    "Tipo",
    "Descrição",
    "Qtd",
    "Comprimento (m)",
    "Largura (mm)",
    "Espessura (mm)",
    "Massa por peça (kg)",
    "Massa (kg)",
    "Pintura (m²)",
    "Observação",
)
COLUNAS_PERFIS = (
    "Perfil",
    "Peças",
    "Comprimento total (m)",
    "kg/m",
    "Massa (kg)",
    "Barras comerciais",
)
COLUNAS_CHAPAS = ("Espessura (mm)", "Peças", "Área (m²)", "Massa (kg)")


def linhas_da_tabela(resumo: ResumoDaLista) -> list[list[object]]:
    return [
        [
            x.item.marca,
            x.item.tipo,
            x.perfil.nome if x.perfil else x.item.descricao,
            x.item.quantidade,
            x.item.comprimento_m,
            x.item.largura_mm or "",
            x.item.espessura_mm or "",
            "" if x.massa_por_peca_kg is None else x.massa_por_peca_kg,
            "" if x.massa_kg is None else x.massa_kg,
            x.area_pintura_m2 or "",
            x.pendencia or x.item.observacao,
        ]
        for x in resumo.linhas
    ]


def linhas_dos_perfis(resumo: ResumoDaLista) -> list[list[object]]:
    return [
        [p.perfil, p.pecas, p.comprimento_m, p.massa_kg_m, p.massa_kg, p.barras]
        for p in resumo.perfis
    ]


def linhas_das_chapas(resumo: ResumoDaLista) -> list[list[object]]:
    return [[c.espessura_mm, c.pecas, c.area_m2, c.massa_kg] for c in resumo.chapas]


def totais(resumo: ResumoDaLista) -> list[tuple[str, str]]:
    return [
        ("Massa dos itens", f"{_n(resumo.massa_itens_kg, 0)} kg"),
        (
            f"Ligações, parafusos e soldas ({_n(resumo.acrescimo_pct, 1)} %)",
            f"{_n(resumo.massa_acrescimo_kg, 0)} kg",
        ),
        ("Massa total", f"{_n(resumo.massa_total_kg, 0)} kg"),
        ("Peso total", f"{_n(resumo.peso_total_kN, 2)} kN"),
        ("Área de pintura", f"{_n(resumo.area_pintura_m2, 1)} m²"),
    ] + (
        [("Massa pela geometria do modelo (aço)", f"{_n(resumo.massa_geometria_kg, 0)} kg")]
        if resumo.massa_geometria_kg
        else []
    )


def csv_da_lista(resumo: ResumoDaLista) -> bytes:
    return pc._csv(COLUNAS, linhas_da_tabela(resumo))


def xlsx_da_lista(resumo: ResumoDaLista, *, titulo: str = "") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    livro = Workbook()
    folha_resumo = livro.active
    assert folha_resumo is not None
    folha_resumo.title = "Resumo"
    folha_resumo.append([f"Lista de material — {titulo}" if titulo else "Lista de material"])
    folha_resumo["A1"].font = Font(bold=True)
    folha_resumo.append([])
    for rotulo, valor in totais(resumo):
        folha_resumo.append([rotulo, valor])
    folha_resumo.append([])
    folha_resumo.append(
        [
            f"Barras comerciais de {_n(resumo.comprimento_barra_m, 1)} m: estimativa pelo "
            "comprimento total, sem otimização de corte nem perdas."
        ]
    )
    if resumo.pendentes:
        folha_resumo.append(
            [f"{len(resumo.pendentes)} item(ns) sem massa — completar antes do orçamento."]
        )
    folha_resumo.column_dimensions["A"].width = 48
    folha_resumo.column_dimensions["B"].width = 20
    fonte = Font(bold=True, color="FFFFFF")
    fundo = PatternFill("solid", fgColor="1F4E8C")
    for nome, cabecalho, linhas in (
        ("Itens", COLUNAS, linhas_da_tabela(resumo)),
        ("Por perfil", COLUNAS_PERFIS, linhas_dos_perfis(resumo)),
        ("Chapas", COLUNAS_CHAPAS, linhas_das_chapas(resumo)),
    ):
        folha = livro.create_sheet(nome)
        folha.append(list(cabecalho))
        for linha in linhas:
            folha.append(list(linha))
        for celula in folha[1]:
            celula.font, celula.fill = fonte, fundo
            celula.alignment = Alignment(wrap_text=True, vertical="center")
        folha.freeze_panes = "A2"
        for indice, coluna in enumerate(folha.iter_cols(min_row=1, max_row=folha.max_row), 1):
            largura = max((len(str(c.value)) for c in coluna if c.value is not None), default=8)
            folha.column_dimensions[get_column_letter(indice)].width = min(max(10, largura + 2), 50)
            for celula in coluna[1:]:
                if isinstance(celula.value, float):
                    celula.number_format = "0.00"
    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()


def _celula_vazia(valor: Any) -> bool:
    if valor is None:
        return True
    if isinstance(valor, float):
        return math.isnan(valor)
    return type(valor).__name__ in ("NAType", "NaTType")


COLUNA_GEOMETRIA = "Pela geometria (kg/peça)"
COLUNAS_EDITAVEIS = (
    "Marca",
    "Tipo",
    "Descrição",
    "Qtd",
    "Comprimento (m)",
    "Largura (mm)",
    "Espessura (mm)",
    "Massa unitária",
    "Observação",
    COLUNA_GEOMETRIA,
)


def tabela_editavel(lista: ListaDeMaterial) -> list[dict[str, Any]]:
    """Os itens com os nomes de coluna da tabela editável da página."""
    return [
        {
            "Marca": i.marca,
            "Tipo": i.tipo,
            "Descrição": i.descricao,
            "Qtd": i.quantidade,
            "Comprimento (m)": i.comprimento_m,
            "Largura (mm)": i.largura_mm or None,
            "Espessura (mm)": i.espessura_mm or None,
            "Massa unitária": i.massa_unitaria,
            "Observação": i.observacao,
            COLUNA_GEOMETRIA: i.massa_geometria_kg,
        }
        for i in lista.itens
    ]


def itens_da_tabela(linhas: Sequence[Mapping[str, Any]]) -> list[ItemDaLista]:
    """As linhas da tabela editável da página (nomes de coluna da página) → itens.

    Célula vazia do pandas (NaN, ``pd.NA``) vira ``None``: ``bool(pd.NA)`` levanta erro e
    ``str(nan)`` daria a descrição "nan".
    """
    itens = []
    for bruta in linhas:
        linha = {chave: None if _celula_vazia(valor) else valor for chave, valor in bruta.items()}
        descricao = str(linha.get("Descrição") or "").strip()
        if not descricao and not linha.get("Marca"):
            continue
        itens.append(
            item_de_dicionario(
                {
                    "marca": linha.get("Marca"),
                    "tipo": linha.get("Tipo"),
                    "descricao": descricao,
                    "quantidade": linha.get("Qtd"),
                    "comprimento_m": linha.get("Comprimento (m)"),
                    "largura_mm": linha.get("Largura (mm)"),
                    "espessura_mm": linha.get("Espessura (mm)"),
                    "massa_unitaria": linha.get("Massa unitária"),
                    "observacao": linha.get("Observação"),
                    "massa_geometria_kg": linha.get(COLUNA_GEOMETRIA),
                }
            )
        )
    return itens
