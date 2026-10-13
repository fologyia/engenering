"""Exportação do plano de cargas para o modelo: SolidWorks Simulation, Robot e outros programas.

O plano guarda os valores **característicos** em kN e metros, com o eixo Z vertical. Este módulo
converte para o sistema de unidades e a orientação de eixos do programa de destino e monta:

* a tabela das **cargas por caso** — uma linha por carga, com o valor, as componentes F_x, F_y,
  F_z já com sinal, o sentido escrito por extenso e o comando equivalente no SolidWorks e no Robot;
* a **matriz das combinações** (uma linha por combinação, uma coluna por caso, o fator γ·ψ em cada
  célula) — o mesmo arranjo do gerenciador de casos de carga do SolidWorks e da tabela de
  combinações do Robot;
* a **lista das combinações** (combinação, caso, fator), o formato que SAP2000, ETABS, STAAD e
  planilhas importam;
* a planilha Excel com tudo isso e uma aba "Leia-me" com o passo a passo.

Convenções (as mesmas da página Plano de cargas)
-----------------------------------------------
* Eixos do programa: X e Y no plano horizontal, Z vertical para cima. O vento W0 sopra para +X,
  W90 para +Y, W180 para −X e W270 para −Y.
* Direção "Z (vertical)": valor positivo é **para baixo** (gravidade), logo F_z = −valor.
* Direções X e Y: o sinal está no valor (o "+" ou "−" escrito na direção é só rótulo).
* No SolidWorks a vertical costuma ser o eixo **Y**. Com "Y para cima", um vetor (F_x, F_y, F_z)
  do programa vira (F_x, F_z, −F_y): o X não muda, o Z do programa vira o Y e o Y do plano vira o
  −Z — na vista Superior do SolidWorks o desenho fica igual à planta do programa.

Sem Streamlit.
"""

from __future__ import annotations

import io
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import datetime

from core import load_combinations as comb
from core import plano_de_cargas as pc

# ---------------------------------------------------------------------------------------------
# Unidades e eixos do programa de destino
# ---------------------------------------------------------------------------------------------


@dataclass(frozen=True)
class SistemaDeUnidades:
    nome: str
    forca: str
    linha: str
    area: str
    momento: str
    comprimento: str
    fator_forca: float  # multiplica kN
    fator_linha: float  # multiplica kN/m
    fator_area: float  # multiplica kN/m²
    fator_momento: float  # multiplica kN·m


UNIDADES_KN_M = SistemaDeUnidades(
    "kN e m — Robot, SAP2000, Ftool", "kN", "kN/m", "kN/m²", "kN·m", "m", 1.0, 1.0, 1.0, 1.0
)
UNIDADES_N_M = SistemaDeUnidades(
    "N e m — SolidWorks (SI)", "N", "N/m", "N/m² (Pa)", "N·m", "m", 1e3, 1e3, 1e3, 1e3
)
UNIDADES_N_MM = SistemaDeUnidades(
    "N e mm — SolidWorks (MMGS)", "N", "N/mm", "N/mm² (MPa)", "N·mm", "mm", 1e3, 1.0, 1e-3, 1e6
)
SISTEMAS: dict[str, SistemaDeUnidades] = {
    s.nome: s for s in (UNIDADES_N_MM, UNIDADES_N_M, UNIDADES_KN_M)
}

EIXO_Z_PARA_CIMA = "Z para cima — Robot, SAP2000, padrão do programa"
EIXO_Y_PARA_CIMA = "Y para cima — SolidWorks"
EIXOS: tuple[str, ...] = (EIXO_Y_PARA_CIMA, EIXO_Z_PARA_CIMA)

GRAVIDADE_M_S2 = 9.81

# Tipos de carga e o comando equivalente em cada programa (rótulos da versão em português).
TIPO_NODAL = "Força concentrada"
TIPO_LINHA = "Carga distribuída em barra"
TIPO_AREA = "Carga por área"
TIPO_MOMENTO = "Momento concentrado"
TIPO_TEMPERATURA = "Temperatura"
TIPO_GRAVIDADE = "Gravidade"
TIPO_POR_UNIDADE = {
    "kN": TIPO_NODAL,
    "kN/m": TIPO_LINHA,
    "kN/m²": TIPO_AREA,
    "kN·m": TIPO_MOMENTO,
    "°C": TIPO_TEMPERATURA,
}
EQUIVALENTES: dict[str, tuple[str, str]] = {
    TIPO_NODAL: (
        "Cargas externas › Força, na junta da viga ou no vértice (direção selecionada)",
        "Carga nodal — força",
    ),
    TIPO_LINHA: (
        "Cargas externas › Força na viga, opção distribuída (por unidade de comprimento)",
        "Carga uniforme em barras",
    ),
    TIPO_AREA: (
        "Cargas externas › Pressão normal à face (ou força total na face)",
        "Carga uniforme em painel / planar",
    ),
    TIPO_MOMENTO: (
        "Cargas externas › Torque ou Carga remota (momento)",
        "Carga nodal — momento",
    ),
    TIPO_TEMPERATURA: (
        "Cargas externas › Temperatura (efeito térmico no estudo estático)",
        "Carga térmica em barras",
    ),
    TIPO_GRAVIDADE: (
        "Cargas externas › Gravidade, 9,81 m/s² para baixo",
        "Peso próprio automático do caso",
    ),
}


def _menos(valor: float) -> str:
    return "−" if valor < 0 else "+"


def eixo_da_direcao(direcao: str) -> str | None:
    """ "X", "Y" ou "Z" a partir do texto da direção ("+X", "−Y", "Z (vertical)"); ``None`` = sem."""
    texto = direcao.strip().lstrip("+−-").strip().upper()
    return texto[:1] if texto[:1] in ("X", "Y", "Z") else None


def componentes_no_programa(carga: pc.CargaDoModelo) -> tuple[float, float, float] | None:
    """(F_x, F_y, F_z) no sistema do programa (Z para cima); ``None`` para temperatura e sem eixo."""
    if carga.unidade == "°C":
        return None
    eixo = eixo_da_direcao(carga.direcao)
    if eixo == "X":
        return (carga.valor, 0.0, 0.0)
    if eixo == "Y":
        return (0.0, carga.valor, 0.0)
    if eixo == "Z":
        return (0.0, 0.0, -carga.valor)  # valor positivo = para baixo
    return None


def para_os_eixos(vetor: tuple[float, float, float], eixos: str) -> tuple[float, float, float]:
    """O vetor do programa (Z para cima) nos eixos do destino."""
    fx, fy, fz = vetor
    if eixos == EIXO_Y_PARA_CIMA:
        fx, fy, fz = fx, fz, -fy
    return (fx + 0.0, fy + 0.0, fz + 0.0)  # + 0.0 tira o "−0,0" da planilha


def eixo_vertical(eixos: str) -> str:
    return "Y" if eixos == EIXO_Y_PARA_CIMA else "Z"


def sentido(vetor: tuple[float, float, float] | None, eixos: str) -> str:
    """O sentido por extenso: "+X", "−Y (para baixo)"…"""
    if vetor is None:
        return "—"
    nomes = ("X", "Y", "Z")
    partes = []
    for nome, valor in zip(nomes, vetor, strict=True):
        if abs(valor) > 1e-12:
            texto = f"{_menos(valor)}{nome}"
            if nome == eixo_vertical(eixos):
                texto += " (para baixo)" if valor < 0 else " (para cima)"
            partes.append(texto)
    return ", ".join(partes) or "—"


def converter(valor: float, unidade: str, sistema: SistemaDeUnidades) -> tuple[float, str]:
    """O valor do plano (kN, kN/m, kN/m², kN·m, °C) na unidade do destino."""
    tabela = {
        "kN": (sistema.fator_forca, sistema.forca),
        "kN/m": (sistema.fator_linha, sistema.linha),
        "kN/m²": (sistema.fator_area, sistema.area),
        "kN·m": (sistema.fator_momento, sistema.momento),
    }
    if unidade in tabela:
        fator, nova = tabela[unidade]
        return valor * fator, nova
    return valor, unidade


# ---------------------------------------------------------------------------------------------
# Cargas por caso
# ---------------------------------------------------------------------------------------------
@dataclass(frozen=True)
class LinhaDeCarga:
    caso: str
    nome_do_caso: str
    tipo: str
    elemento: str
    valor: float
    unidade: str
    quantidade: int
    fx: float | None
    fy: float | None
    fz: float | None
    sentido: str
    solidworks: str
    robot: str
    observacao: str


COLUNAS_CARGAS = (
    "Caso",
    "Nome do caso",
    "Tipo da carga",
    "Onde aplicar",
    "Valor",
    "Unidade",
    "Quantidade",
    "Fx",
    "Fy",
    "Fz",
    "Sentido",
    "No SolidWorks",
    "No Robot",
    "Observação",
)


def linhas_de_carga(
    plano: pc.PlanoDeCargas, sistema: SistemaDeUnidades, eixos: str
) -> list[LinhaDeCarga]:
    """Uma linha por carga de cada ação, já nas unidades e nos eixos do destino.

    O peso próprio (PP) sem cargas informadas vira a linha "Gravidade": no modelo, o peso das
    peças sai do material, e o caso PP é só a gravidade.
    """
    linhas: list[LinhaDeCarga] = []
    vetor: tuple[float, float, float] | None
    for acao in plano.acoes:
        if acao.codigo == "PP" and not acao.cargas:
            fator = plano.fator_peso_proprio
            gravidade = GRAVIDADE_M_S2 * fator
            vetor = para_os_eixos((0.0, 0.0, -gravidade), eixos)
            sw, robot = EQUIVALENTES[TIPO_GRAVIDADE]
            observacao = "Confira a densidade do aço no material do modelo (7850 kg/m³)."
            if fator > 1:
                observacao = (
                    f"9,81 × {fator:.2f}".replace(".", ",")
                    + ": o acréscimo é o peso das ligações, chapas e parafusos que não estão no "
                    "modelo. " + observacao
                )
            linhas.append(
                LinhaDeCarga(
                    acao.codigo,
                    acao.nome,
                    TIPO_GRAVIDADE,
                    "Toda a estrutura (peso das peças do modelo)",
                    gravidade,
                    "m/s²",
                    1,
                    *vetor,
                    sentido(vetor, eixos),
                    sw.replace("9,81 m/s²", f"{gravidade:.2f} m/s²".replace(".", ",")),
                    robot,
                    observacao,
                )
            )
            continue
        for carga in acao.cargas:
            tipo = TIPO_POR_UNIDADE.get(carga.unidade, TIPO_NODAL)
            valor, unidade = converter(carga.valor, carga.unidade, sistema)
            no_programa = componentes_no_programa(carga)
            vetor = None
            if no_programa is not None:
                fator = converter(1.0, carga.unidade, sistema)[0]
                fx, fy, fz = no_programa
                vetor = para_os_eixos((fator * fx, fator * fy, fator * fz), eixos)
            sw, robot = EQUIVALENTES[tipo]
            observacao = carga.observacao
            if tipo == TIPO_MOMENTO:
                observacao = "; ".join(filter(None, ("vetor do momento", observacao)))
            linhas.append(
                LinhaDeCarga(
                    acao.codigo,
                    acao.nome,
                    tipo,
                    carga.elemento,
                    valor,
                    unidade,
                    carga.quantidade,
                    None if vetor is None else vetor[0],
                    None if vetor is None else vetor[1],
                    None if vetor is None else vetor[2],
                    sentido(vetor, eixos) if tipo != TIPO_TEMPERATURA else "ΔT",
                    sw,
                    robot,
                    observacao,
                )
            )
    return linhas


def _linha_csv(linha: LinhaDeCarga) -> list[object]:
    return [
        linha.caso,
        linha.nome_do_caso,
        linha.tipo,
        linha.elemento,
        linha.valor,
        linha.unidade,
        linha.quantidade,
        "" if linha.fx is None else linha.fx,
        "" if linha.fy is None else linha.fy,
        "" if linha.fz is None else linha.fz,
        linha.sentido,
        linha.solidworks,
        linha.robot,
        linha.observacao,
    ]


def csv_de_cargas(plano: pc.PlanoDeCargas, sistema: SistemaDeUnidades, eixos: str) -> bytes:
    return pc._csv(
        COLUNAS_CARGAS, [_linha_csv(linha) for linha in linhas_de_carga(plano, sistema, eixos)]
    )


# ---------------------------------------------------------------------------------------------
# Combinações
# ---------------------------------------------------------------------------------------------
SIGLAS: dict[str, str] = {
    comb.ELU_NORMAL: "ELU",
    comb.ELU_ESPECIAL: "ELUE",
    comb.ELU_EXCEPCIONAL: "ELUX",
    comb.ELS_RARA: "ELSR",
    comb.ELS_FREQUENTE: "ELSF",
    comb.ELS_QUASE_PERMANENTE: "ELSQ",
}


def nome_da_combinacao(c: pc.LinhaDeCombinacao, total: int) -> str:
    """Nome curto, sem espaço nem acento, que serve de nome do caso no modelo: ``C07-ELU``."""
    largura = max(2, len(str(total)))
    return f"C{c.numero:0{largura}d}-{SIGLAS.get(c.estado_limite, 'COMB')}"


def tipo_do_estado(estado: str) -> str:
    return "ELU" if estado in comb.COMBINACOES_ULTIMAS else "ELS"


def matriz_de_combinacoes(
    plano: pc.PlanoDeCargas, estados: Sequence[str]
) -> tuple[list[str], list[list[object]]]:
    """Cabeçalho e linhas: combinação, estado-limite, o fator de cada caso e a expressão."""
    lista = pc.combinacoes(plano, estados)
    codigos = list(plano.codigos)
    linhas: list[list[object]] = [
        [
            nome_da_combinacao(c, len(lista)),
            c.estado_limite,
            *[float(c.fatores.get(k, 0.0)) for k in codigos],
            c.expressao,
        ]
        for c in lista
    ]
    return ["Combinação", "Estado-limite", *codigos, "Expressão"], linhas


def lista_de_combinacoes(
    plano: pc.PlanoDeCargas, estados: Sequence[str]
) -> tuple[list[str], list[list[object]]]:
    """Uma linha por caso de cada combinação (só os fatores diferentes de zero)."""
    lista = pc.combinacoes(plano, estados)
    linhas: list[list[object]] = []
    for c in lista:
        nome = nome_da_combinacao(c, len(lista))
        for codigo in plano.codigos:
            fator = float(c.fatores.get(codigo, 0.0))
            if abs(fator) > 1e-12:
                linhas.append(
                    [nome, c.estado_limite, tipo_do_estado(c.estado_limite), codigo, fator]
                )
    return ["Combinação", "Estado-limite", "Tipo", "Caso", "Fator"], linhas


def csv_matriz(plano: pc.PlanoDeCargas, estados: Sequence[str]) -> bytes:
    cabecalho, linhas = matriz_de_combinacoes(plano, estados)
    return pc._csv(cabecalho, linhas)


def csv_lista(plano: pc.PlanoDeCargas, estados: Sequence[str]) -> bytes:
    cabecalho, linhas = lista_de_combinacoes(plano, estados)
    return pc._csv(cabecalho, linhas)


# ---------------------------------------------------------------------------------------------
# Gerenciador de casos de carga do SolidWorks (vista de cenário de carga)
# ---------------------------------------------------------------------------------------------
#: O nome que o SolidWorks dá a cada tipo de carga (Força-1, Pressão-1…), na ordem em que é criada.
NOME_NO_SOLIDWORKS = {
    TIPO_GRAVIDADE: "Gravidade",
    TIPO_NODAL: "Força",
    TIPO_LINHA: "Força",
    TIPO_AREA: "Pressão",
    TIPO_MOMENTO: "Torque",
    TIPO_TEMPERATURA: "Temperatura",
}
SUPRIMIR = "Suprimir"


def nome_no_solidworks(codigo: str) -> str:
    """O código do caso como se digita no SolidWorks: o "−" do programa vira hífen (T− → T-)."""
    return codigo.replace("−", "-")


@dataclass(frozen=True)
class CargaDoEstudo:
    """Uma carga criada uma vez no estudo e ligada ou suprimida em cada caso primário."""

    numero: int
    nome: str  # como o SolidWorks vai chamar, se criada nesta ordem (Força-2)
    tipo: str
    onde: str
    referencia: str  # o sentido em que a carga é criada ("−Y (para baixo)", "+X", "normal à face")
    unidade: str
    valor_ao_criar: float  # o do primeiro caso que usa a carga
    comando: str
    observacao: str

    @property
    def titulo(self) -> str:
        return f"{self.nome}: {self.onde} — {self.referencia} ({self.unidade})"


@dataclass(frozen=True)
class CenarioDeCarga:
    cargas: tuple[CargaDoEstudo, ...]
    casos: tuple[tuple[str, str], ...]  # (código, nome) dos casos primários
    valores: dict[tuple[str, int], float]  # (caso, carga) → valor no sentido de referência
    combinacoes: tuple[tuple[str, str], ...]  # (nome, expressão para o SolidWorks)

    def celula(self, caso: str, numero: int) -> str:
        valor = self.valores.get((caso, numero))
        if valor is None:
            return SUPRIMIR
        return f"{valor:.6g}".replace(".", ",")


def _sentido_no_eixo(eixo: str, sinal: float, eixos: str) -> str:
    texto = f"{'−' if sinal < 0 else '+'}{eixo}"
    if eixo == eixo_vertical(eixos):
        texto += " (para baixo)" if sinal < 0 else " (para cima)"
    return texto


def cenario_de_carga(
    plano: pc.PlanoDeCargas, sistema: SistemaDeUnidades, eixos: str, estados: Sequence[str]
) -> CenarioDeCarga:
    """O plano no arranjo do gerenciador de casos de carga do SolidWorks.

    Cada carga do estudo é criada uma vez (a gravidade, cada força ou carga distribuída num lugar e
    num eixo, cada pressão, cada temperatura) e cada caso primário diz o valor dela — no sentido em
    que foi criada — ou "Suprimir". O W0 e o W180 na mesma barra são a mesma carga em X, com sinais
    opostos. As combinações vêm junto, para ver no SolidWorks; o programa importa os casos
    primários (característicos) e combina ele mesmo.
    """
    linhas = linhas_de_carga(plano, sistema, eixos)
    chaves: dict[tuple[object, ...], dict[str, object]] = {}
    valores_por_chave: dict[tuple[object, ...], dict[str, float]] = {}
    for linha in linhas:
        vetor = (linha.fx, linha.fy, linha.fz)
        if linha.tipo == TIPO_GRAVIDADE:
            eixo = eixo_vertical(eixos)
            indice = "XYZ".index(eixo)
            chave: tuple[object, ...] = (TIPO_GRAVIDADE,)
            componente = float(vetor[indice] or 0.0)
            unidade = "m/s²"
        elif linha.tipo in (TIPO_AREA, TIPO_TEMPERATURA) or None in vetor:
            eixo = "normal à face" if linha.tipo == TIPO_AREA else "ΔT"
            chave = (linha.tipo, linha.elemento)
            componente = linha.valor
            unidade = linha.unidade
        else:
            componentes = [float(v or 0.0) for v in vetor]
            indice = max(range(3), key=lambda i: abs(componentes[i]))
            obliqua = sum(abs(c) > 1e-9 for c in componentes) > 1
            eixo = "XYZ"[indice]
            chave = (linha.tipo, linha.elemento, eixo if not obliqua else linha.sentido)
            componente = componentes[indice] if not obliqua else linha.valor
            if obliqua:
                eixo = linha.sentido
            unidade = linha.unidade
        if chave not in chaves:
            chaves[chave] = {
                "tipo": linha.tipo,
                "onde": linha.elemento,
                "eixo": eixo,
                "unidade": unidade,
                "comando": linha.solidworks,
                "observacao": linha.observacao,
                "quantidade": linha.quantidade,
            }
            valores_por_chave[chave] = {}
        valores = valores_por_chave[chave]
        valores[linha.caso] = valores.get(linha.caso, 0.0) + componente
    contagem: dict[str, int] = {}
    cargas: list[CargaDoEstudo] = []
    valores_finais: dict[tuple[str, int], float] = {}
    for numero, (chave, dados) in enumerate(chaves.items(), 1):
        tipo = str(dados["tipo"])
        base = NOME_NO_SOLIDWORKS.get(tipo, "Carga")
        contagem[base] = contagem.get(base, 0) + 1
        valores = valores_por_chave[chave]
        primeiro = next((v for v in valores.values() if abs(v) > 1e-12), 1.0)
        sinal = -1.0 if primeiro < 0 else 1.0
        eixo = str(dados["eixo"])
        if eixo in ("X", "Y", "Z"):
            referencia = _sentido_no_eixo(eixo, sinal, eixos)
        else:
            referencia = eixo
            sinal = 1.0
        observacao = str(dados["observacao"])
        if int(str(dados["quantidade"])) > 1:
            observacao = "; ".join(
                filter(None, (f"{dados['quantidade']} pontos: valor por item", observacao))
            )
        cargas.append(
            CargaDoEstudo(
                numero=numero,
                nome=f"{base}-{contagem[base]}",
                tipo=tipo,
                onde=str(dados["onde"]),
                referencia=referencia,
                unidade=str(dados["unidade"]),
                valor_ao_criar=round(abs(primeiro) if eixo in ("X", "Y", "Z") else primeiro, 3),
                comando=str(dados["comando"]),
                observacao=observacao,
            )
        )
        for caso, valor in valores.items():
            valores_finais[(caso, numero)] = valor * sinal
    lista = pc.combinacoes(plano, estados)
    combinacoes = tuple(
        (
            nome_da_combinacao(c, len(lista)),
            " + ".join(
                f'{fator:g}*"{nome_no_solidworks(codigo)}"'
                for codigo in plano.codigos
                if abs(fator := float(c.fatores.get(codigo, 0.0))) > 1e-12
            ),
        )
        for c in lista
    )
    return CenarioDeCarga(
        cargas=tuple(cargas),
        casos=tuple((a.codigo, a.nome) for a in plano.acoes),
        valores=valores_finais,
        combinacoes=combinacoes,
    )


def tabela_do_cenario(c: CenarioDeCarga) -> tuple[list[str], list[list[str]]]:
    """Cabeçalho e linhas como na tela: o caso primário e, em cada carga, o valor ou "Suprimir"."""
    cabecalho = ["Caso de carga primário", *(carga.titulo for carga in c.cargas)]
    linhas = [
        [nome_no_solidworks(codigo), *(c.celula(codigo, carga.numero) for carga in c.cargas)]
        for codigo, _ in c.casos
    ]
    return cabecalho, linhas


COLUNAS_CARGAS_DO_ESTUDO = (
    "Nº",
    "Nome no SolidWorks",
    "Tipo",
    "Onde aplicar",
    "Sentido em que é criada",
    "Valor ao criar",
    "Unidade",
    "No SolidWorks",
    "Observação",
)


def linhas_das_cargas_do_estudo(c: CenarioDeCarga) -> list[list[object]]:
    return [
        [
            carga.numero,
            carga.nome,
            carga.tipo,
            carga.onde,
            carga.referencia,
            carga.valor_ao_criar,
            carga.unidade,
            carga.comando,
            carga.observacao,
        ]
        for carga in c.cargas
    ]


def csv_do_cenario(c: CenarioDeCarga) -> bytes:
    """A tabela do gerenciador de casos de carga e, embaixo, as combinações."""
    cabecalho, linhas = tabela_do_cenario(c)
    linhas_csv: list[list[object]] = [list(linha) for linha in linhas]
    if c.combinacoes:
        linhas_csv.append([])
        linhas_csv.append(["Combinações de casos de carga (para ver no SolidWorks)", "Expressão"])
        linhas_csv.extend([nome, expressao] for nome, expressao in c.combinacoes)
    return pc._csv(cabecalho, linhas_csv)


# ---------------------------------------------------------------------------------------------
# Leia-me e planilha Excel
# ---------------------------------------------------------------------------------------------
def texto_leia_me(
    sistema: SistemaDeUnidades,
    eixos: str,
    *,
    titulo: str = "",
    anglo: bool = False,
    fator_peso_proprio: float = 1.0,
) -> list[str]:
    gravidade = f"{GRAVIDADE_M_S2 * fator_peso_proprio:.2f} m/s²".replace(".", ",")
    if fator_peso_proprio > 1:
        gravidade += (
            f" = 9,81 × {fator_peso_proprio:.2f}".replace(".", ",")
            + ", com as ligações, chapas e parafusos que não estão no modelo"
        )
    vertical = eixo_vertical(eixos)
    linhas = [
        f"Plano de cargas — {titulo}" if titulo else "Plano de cargas",
        f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')} pelo Mecânica Toolkit.",
        "",
        "UNIDADES E EIXOS",
        f"Unidades: {sistema.nome} (força em {sistema.forca}, comprimento em {sistema.comprimento}).",
        f"Eixos: {eixos}. Vertical = {vertical}; Fx, Fy e Fz já estão nesses eixos e com sinal.",
        "Vento: W0 sopra para +X e W90 para +Y da planta; W180 e W270 são os opostos.",
    ]
    if eixos == EIXO_Y_PARA_CIMA:
        linhas.append(
            "No SolidWorks (Y para cima) o Y da planta é o −Z: na vista Superior o desenho fica "
            "igual à planta do programa."
        )
    linhas += [
        "Valores característicos (sem coeficientes): os coeficientes estão nas combinações.",
        "",
        "NO SOLIDWORKS SIMULATION",
        "Um estudo estático por caso da aba Cargas, com o código do caso no nome do estudo, e as "
        "cargas CARACTERÍSTICAS: nunca majore no SolidWorks — os coeficientes estão nas "
        "combinações, que o programa monta na página Esforços do modelo.",
        "1. Estudo PP: material (aço A36 ou A572 Gr 50: 200 GPa, 7850 kg/m³), juntas das vigas, "
        f"fixações nas bases, malha e SÓ a Gravidade ({gravidade}, para baixo). Rode.",
        "2. Botão direito na aba do estudo PP › Duplicar, com o nome do próximo caso (SC, PE, EQ, "
        "W0, W90, W180, W270, T+ como T+, T− como T-). No estudo novo, exclua a Gravidade e aplique "
        "só as cargas desse caso da aba Cargas (colunas Tipo da carga e No SolidWorks; Fx, Fy e Fz "
        "já com o sinal nos eixos do modelo). Não refaça a malha: os casos são combinados ponto a "
        "ponto, na mesma malha. Rode.",
        "3. Em cada estudo: Resultados › Listar forças da viga (todas as vigas) e Listar forças "
        "resultantes (força de reação, modelo inteiro), salvando os dois em CSV.",
        "4. Envie todos os CSV na página Esforços do modelo: ela confere as reações de cada caso "
        "contra este plano (o PP contra o peso da Lista de material) e aponta estudo majorado ou "
        "carga faltando.",
        "5. Mudou a geometria? Refaça todos os estudos. Um estudo com as combinações (aba "
        "Combinações) serve só para olhar tensões e deslocamentos no SolidWorks — não entra na "
        "importação.",
        "",
        "OU NUM ESTUDO SÓ: GERENCIADOR DE CASOS DE CARGA",
        "1. No estudo, crie cada carga uma vez, na ordem da aba Cargas do estudo (assim os nomes "
        "Gravidade-1, Força-1, Pressão-1… batem com a planilha), onde a aba diz e no sentido em que "
        "ela diz, com o valor ao criar.",
        "2. Abra o Gerenciador de casos de carga e crie um caso primário por linha da aba Casos de "
        "carga (SolidWorks), com o código do caso (PP, SC, W0…): em cada carga, o valor da "
        "planilha (o sinal é no sentido em que a carga foi criada) ou Suprimir. As fixações ficam "
        "ativas em todos os casos. Os valores são CARACTERÍSTICOS.",
        "3. Rode. Na Vista de resultados, para cada caso primário: Listar forças da viga e Listar "
        "forças resultantes, salvando cada CSV com o código do caso no nome (SC-forcas.csv, "
        "SC-reacoes.csv) — todos saem do mesmo estudo, e é pelo nome do arquivo que o programa "
        "reconhece o caso.",
        "4. As combinações do gerenciador (embaixo da aba) são só para olhar no SolidWorks: o "
        "programa importa os casos primários e faz as combinações da NBR 8800.",
        "",
        "NO ROBOT STRUCTURAL ANALYSIS",
        "1. Crie os casos de carga com os códigos da aba Cargas (natureza: permanente, "
        "sobrecarga, vento, temperatura).",
        "2. Lance as cargas pela tabela de cargas (nodais, uniformes em barras, planares), "
        "copiando da aba Cargas.",
        "3. Crie as combinações manuais com os fatores da aba Combinações (ELU = ULS, ELS = SLS).",
        "",
        "OUTROS PROGRAMAS (SAP2000, ETABS, STAAD, Ftool, planilhas)",
        "A aba Combinações (lista) traz uma linha por caso de cada combinação, com o fator: é o "
        "formato das tabelas de definição de combinações desses programas.",
    ]
    if anglo:
        linhas += [
            "",
            "CRITÉRIO ANGLO (item 5.9)",
            "Os quadros de cargas para as fundações usam as ações sem combinar e sem majorar: "
            "entregue os casos primários.",
        ]
    return linhas


def xlsx_do_plano(
    plano: pc.PlanoDeCargas,
    *,
    sistema: SistemaDeUnidades,
    eixos: str,
    estados: Sequence[str],
    titulo: str = "",
    anglo: bool = False,
) -> bytes:
    """Planilha com as abas Leia-me, Ações, Cargas, Combinações e Combinações (lista)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    livro = Workbook()
    cabecalho_fonte = Font(bold=True, color="FFFFFF")
    cabecalho_fundo = PatternFill("solid", fgColor="1F4E8C")

    def aba(
        nome: str, cabecalho: Sequence[str], linhas: Sequence[Sequence[object]], *, fixar: str
    ) -> None:
        folha = livro.create_sheet(nome)
        folha.append(list(cabecalho))
        for linha in linhas:
            folha.append(list(linha))
        for celula in folha[1]:
            celula.font = cabecalho_fonte
            celula.fill = cabecalho_fundo
            celula.alignment = Alignment(vertical="center", wrap_text=True)
        folha.freeze_panes = fixar
        folha.auto_filter.ref = folha.dimensions
        for indice, coluna in enumerate(folha.iter_cols(min_row=1, max_row=folha.max_row), 1):
            largura = max((len(str(c.value)) for c in coluna if c.value is not None), default=8)
            folha.column_dimensions[get_column_letter(indice)].width = min(max(10, largura + 2), 60)
            for celula in coluna[1:]:
                if isinstance(celula.value, float):
                    celula.number_format = "0.000"

    leia = livro.active
    assert leia is not None
    leia.title = "Leia-me"
    for texto in texto_leia_me(
        sistema, eixos, titulo=titulo, anglo=anglo, fator_peso_proprio=plano.fator_peso_proprio
    ):
        leia.append([texto])
        if texto.isupper() or texto.startswith("Plano de cargas"):
            leia.cell(row=leia.max_row, column=1).font = Font(bold=True)
    leia.column_dimensions["A"].width = 120

    acoes = []
    for a in plano.acoes:
        cat = comb.categoria_nbr(a.categoria)
        acoes.append(
            [
                a.codigo,
                pc.simbolo(a),
                a.nome,
                cat.tipo,
                a.categoria,
                a.grupo,
                cat.gamma,
                "" if cat.gamma_favoravel is None else cat.gamma_favoravel,
                cat.psi0,
                cat.psi1,
                cat.psi2,
                a.origem,
                a.resumo,
            ]
        )
    aba(
        "Ações",
        (
            "Código",
            "Símbolo",
            "Nome",
            "Tipo",
            "Categoria NBR 8800",
            "Grupo exclusivo",
            "γ",
            "γ favorável",
            "ψ0",
            "ψ1",
            "ψ2",
            "Origem",
            "Resumo",
        ),
        acoes,
        fixar="B2",
    )
    aba(
        "Cargas",
        COLUNAS_CARGAS,
        [_linha_csv(linha) for linha in linhas_de_carga(plano, sistema, eixos)],
        fixar="E2",
    )
    cenario = cenario_de_carga(plano, sistema, eixos, estados)
    if cenario.cargas:
        cabecalho_cenario, linhas_cenario = tabela_do_cenario(cenario)
        linhas_com_combinacoes: list[list[object]] = [list(linha) for linha in linhas_cenario]
        if cenario.combinacoes:
            linhas_com_combinacoes.append([])
            linhas_com_combinacoes.append(
                ["Combinações de casos de carga (só para ver no SolidWorks)", "Expressão"]
            )
            linhas_com_combinacoes.extend([n, e] for n, e in cenario.combinacoes)
        aba("Casos de carga (SolidWorks)", cabecalho_cenario, linhas_com_combinacoes, fixar="B2")
        aba(
            "Cargas do estudo",
            COLUNAS_CARGAS_DO_ESTUDO,
            linhas_das_cargas_do_estudo(cenario),
            fixar="C2",
        )
    cabecalho, linhas = matriz_de_combinacoes(plano, estados)
    aba(
        "Combinações",
        cabecalho,
        [[v if v != 0.0 else None for v in linha] for linha in linhas],
        fixar="C2",
    )
    cabecalho, linhas = lista_de_combinacoes(plano, estados)
    aba("Combinações (lista)", cabecalho, linhas, fixar="B2")
    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()
