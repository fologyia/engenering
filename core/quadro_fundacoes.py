"""Quadro de cargas para as fundações a partir dos esforços do modelo (SolidWorks).

Para cada barra marcada como **Pilar** na tabela das barras, os esforços na ponta de **base** em
cada caso de carga importado — **sem combinar nem majorar** (critério Anglo AA-BR-DPST-DR-0001,
item 5.9: as ações dos quadros de cargas para as fundações não são combinadas nem majoradas; quem
projeta a fundação faz as combinações dela).

* A base é a ponta inicial ou final da lista de pontos do pilar no arquivo do SolidWorks. No modo
  automático é a ponta mais comprimida (o peso próprio faz a compressão crescer para baixo); quando
  as duas pontas têm a mesma compressão, o programa pede a escolha.
* Convenção: **N** = força vertical na fundação, compressão positiva; **V₁, V₂, M₁, M₂ e T** = os
  esforços internos na seção da base, nos eixos 1 e 2 da seção do SolidWorks (o eixo forte está na
  tabela das barras).
* Conferência: a soma de N nas bases de todos os pilares tem de bater com a reação vertical total do
  modelo (arquivo "Listar forças resultantes"), quando todos os apoios são bases de pilares.

Sem Streamlit.
"""

from __future__ import annotations

import io
import math
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from typing import Any

from core import criterio_anglo as ca
from core import esforcos_modelo as em
from core import plano_de_cargas as pc
from core.technical_records import criar_registro_tecnico

BASE_AUTOMATICA = em.BASE_AUTOMATICA
BASE_INICIO = em.BASE_INICIO
BASE_FIM = em.BASE_FIM
BASES = em.BASES
TOLERANCIA_BASE = 0.005  # diferença relativa mínima de compressão entre as pontas


@dataclass(frozen=True)
class CargaNaBase:
    pilar: str
    caso: str
    n: float  # compressão positiva
    v1: float
    v2: float
    m1: float
    m2: float
    t: float


@dataclass(frozen=True)
class PilarDoQuadro:
    pilar: str
    perfil: str
    eixo_forte: str
    base: str | None  # BASE_INICIO, BASE_FIM ou None (não deu para saber)
    origem_da_base: str  # "escolhida", "automática" ou o motivo de não saber
    cargas: tuple[CargaNaBase, ...]


@dataclass(frozen=True)
class QuadroDeFundacoes:
    pilares: tuple[PilarDoQuadro, ...]
    casos: tuple[str, ...]
    conferencia: tuple[pc.ItemDeConferencia, ...]

    @property
    def pendentes(self) -> tuple[PilarDoQuadro, ...]:
        return tuple(p for p in self.pilares if p.base is None)


def _n(valor: float, casas: int = 2) -> str:
    return f"{valor:.{casas}f}".replace(".", ",")


def ponta_da_base(
    pontos_por_caso: Mapping[str, Sequence[em.PontoDeEsforco]], escolha: str
) -> tuple[str | None, str]:
    """A ponta da base e de onde ela veio (escolhida, automática ou o motivo de não saber)."""
    if escolha in (BASE_INICIO, BASE_FIM):
        return escolha, "escolhida"
    melhor: tuple[float, str] | None = None
    for pontos in pontos_por_caso.values():
        if len(pontos) < 2:
            continue
        inicio, fim = pontos[0].n, pontos[-1].n
        referencia = max(abs(inicio), abs(fim))
        if referencia <= 1e-9:
            continue
        diferenca = abs(inicio - fim) / referencia
        if diferenca >= TOLERANCIA_BASE and (melhor is None or diferenca > melhor[0]):
            melhor = (diferenca, BASE_INICIO if inicio < fim else BASE_FIM)
    if melhor is None:
        return None, "as duas pontas têm a mesma compressão: escolha a base na tabela das barras"
    return melhor[1], "automática (ponta mais comprimida)"


def montar_quadro(dados: em.EsforcosDoModelo) -> QuadroDeFundacoes:
    pilares: list[PilarDoQuadro] = []
    casos = tuple(dados.casos)
    for nome in dados.nomes_dos_membros:
        cfg = dados.membros.get(nome, em.ConfiguracaoDoMembro())
        if cfg.tipo != "Pilar":
            continue
        pontos_por_caso = {k: c.membros[nome] for k, c in dados.casos.items() if nome in c.membros}
        base, origem = ponta_da_base(pontos_por_caso, cfg.base)
        cargas: list[CargaNaBase] = []
        if base is not None:
            for caso, pontos in pontos_por_caso.items():
                p = pontos[0] if base == BASE_INICIO else pontos[-1]
                cargas.append(CargaNaBase(nome, caso, -p.n + 0.0, p.v1, p.v2, p.m1, p.m2, p.t))
        pilares.append(PilarDoQuadro(nome, cfg.perfil, cfg.eixo_forte, base, origem, tuple(cargas)))
    return QuadroDeFundacoes(tuple(pilares), casos, tuple(_conferir(dados, pilares)))


def _conferir(
    dados: em.EsforcosDoModelo, pilares: Sequence[PilarDoQuadro]
) -> list[pc.ItemDeConferencia]:
    itens: list[pc.ItemDeConferencia] = []
    if not pilares:
        return [
            pc.ItemDeConferencia(
                pc.NIVEL_ATENCAO,
                "Nenhuma barra marcada como Pilar na tabela das barras: marque os pilares para "
                "montar o quadro.",
            )
        ]
    sem_base = [p.pilar for p in pilares if p.base is None]
    if sem_base:
        itens.append(
            pc.ItemDeConferencia(
                pc.NIVEL_ATENCAO,
                f"Escolha a base de {len(sem_base)} pilar(es) na tabela das barras: "
                f"{', '.join(sem_base[:4])}{'…' if len(sem_base) > 4 else ''}.",
            )
        )
    for caso, importado in dados.casos.items():
        if importado.reacoes is None:
            continue
        vertical = em.para_o_plano(importado.reacoes.modelo, dados.eixo_vertical)[2]
        soma = sum(c.n for p in pilares for c in p.cargas if c.caso == caso)
        if sem_base:
            continue
        diferenca = abs(soma - vertical)
        if diferenca <= max(0.02 * abs(vertical), 0.1):
            itens.append(
                pc.ItemDeConferencia(
                    pc.NIVEL_OK,
                    f"{caso}: a soma das bases ({_n(soma)} kN) bate com a reação vertical do modelo "
                    f"({_n(vertical)} kN).",
                )
            )
        else:
            itens.append(
                pc.ItemDeConferencia(
                    pc.NIVEL_ATENCAO,
                    f"{caso}: a soma das bases dos pilares dá {_n(soma)} kN e a reação vertical do "
                    f"modelo, {_n(vertical)} kN — há apoios que não são base de pilar marcado, ou "
                    "alguma base está na ponta errada.",
                )
            )
    return itens


COLUNAS = (
    "Pilar",
    "Perfil",
    "Caso",
    "Ação",
    "N (kN, compressão +)",
    "V1 (kN)",
    "V2 (kN)",
    "M1 (kN·m)",
    "M2 (kN·m)",
    "T (kN·m)",
    "Eixo forte",
    "Base",
)


def linhas_do_quadro(quadro: QuadroDeFundacoes, plano: pc.PlanoDeCargas) -> list[list[object]]:
    linhas: list[list[object]] = []
    for p in quadro.pilares:
        for c in p.cargas:
            acao = plano.acao(c.caso)
            linhas.append(
                [
                    p.pilar,
                    p.perfil or "—",
                    c.caso,
                    acao.nome if acao else "—",
                    c.n,
                    c.v1,
                    c.v2,
                    c.m1,
                    c.m2,
                    c.t,
                    p.eixo_forte,
                    f"{p.base} ({p.origem_da_base})",
                ]
            )
    return linhas


def matriz_de_compressao(quadro: QuadroDeFundacoes) -> tuple[list[str], list[list[object]]]:
    """Uma linha por pilar e uma coluna por caso, com N (compressão positiva)."""
    cabecalho = ["Pilar", *quadro.casos]
    linhas: list[list[object]] = []
    for p in quadro.pilares:
        por_caso = {c.caso: c.n for c in p.cargas}
        linhas.append([p.pilar, *(por_caso.get(caso, "") for caso in quadro.casos)])
    return cabecalho, linhas


def csv_do_quadro(quadro: QuadroDeFundacoes, plano: pc.PlanoDeCargas) -> bytes:
    return pc._csv(COLUNAS, linhas_do_quadro(quadro, plano))


NOTA_ANGLO = (
    f"{ca.QUADRO_DE_CARGAS_FUNDACOES} Valores característicos de cada caso; as combinações são de "
    "quem projeta a fundação."
)
CONVENCAO = (
    "N = força vertical na fundação, compressão positiva. V1, V2, M1, M2 e T = esforços internos na "
    "seção da base, nos eixos 1 e 2 da seção do SolidWorks (o eixo forte está na coluna Eixo forte)."
)


def xlsx_do_quadro(
    quadro: QuadroDeFundacoes, plano: pc.PlanoDeCargas, *, titulo: str = ""
) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    livro = Workbook()
    leia = livro.active
    assert leia is not None
    leia.title = "Leia-me"
    for texto in (
        f"Quadro de cargas para as fundações — {titulo}" if titulo else "Quadro de cargas",
        "",
        NOTA_ANGLO,
        CONVENCAO,
        f"Casos: {', '.join(quadro.casos)}.",
    ):
        leia.append([texto])
    leia["A1"].font = Font(bold=True)
    leia.column_dimensions["A"].width = 130
    fonte = Font(bold=True, color="FFFFFF")
    fundo = PatternFill("solid", fgColor="1F4E8C")
    cabecalho_m, linhas_m = matriz_de_compressao(quadro)
    for nome, cabecalho, linhas in (
        ("Quadro", COLUNAS, linhas_do_quadro(quadro, plano)),
        ("Compressão por pilar", cabecalho_m, linhas_m),
    ):
        folha = livro.create_sheet(nome)
        folha.append(list(cabecalho))
        for linha in linhas:
            folha.append(list(linha))
        for celula in folha[1]:
            celula.font, celula.fill = fonte, fundo
            celula.alignment = Alignment(wrap_text=True, vertical="center")
        folha.freeze_panes = "B2"
        for indice, coluna in enumerate(folha.iter_cols(min_row=1, max_row=folha.max_row), 1):
            largura = max((len(str(c.value)) for c in coluna if c.value is not None), default=8)
            folha.column_dimensions[get_column_letter(indice)].width = min(max(10, largura + 2), 50)
            for celula in coluna[1:]:
                if isinstance(celula.value, float):
                    celula.number_format = "0.00"
    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()


def registro_do_quadro(
    quadro: QuadroDeFundacoes,
    plano: pc.PlanoDeCargas,
    *,
    contexto: Mapping[str, Any] | None = None,
    responsavel: str = "",
) -> dict[str, Any]:
    cargas = [c for p in quadro.pilares for c in p.cargas]
    maior = max(cargas, key=lambda c: c.n, default=None)
    destaque = (
        f"Cargas características nas bases de {len(quadro.pilares)} pilar(es) para os casos "
        f"{', '.join(quadro.casos)}, sem combinar nem majorar (critério Anglo, item 5.9)."
        + (
            f" Maior compressão: {_n(maior.n)} kN no {maior.pilar} (caso {maior.caso})."
            if maior is not None
            else ""
        )
    )
    tabela = {
        "legenda": "Quadro de cargas para as fundações — valores característicos por caso. "
        + CONVENCAO,
        "cabecalhos": ["Pilar", "Caso", "N (kN)", "V1 (kN)", "V2 (kN)", "M1 (kN·m)", "M2 (kN·m)"],
        "linhas": [
            [c.pilar, c.caso, _n(c.n), _n(c.v1), _n(c.v2), _n(c.m1), _n(c.m2)] for c in cargas
        ],
        "larguras": [2560, 900, 1180, 1180, 1180, 1180, 1180],
        "fonte": 7.0,
    }
    cabecalho_m, linhas_m = matriz_de_compressao(quadro)
    matriz = {
        "legenda": "Compressão (N, kN) na base de cada pilar, por caso.",
        "cabecalhos": cabecalho_m,
        "linhas": [
            [str(linha[0]), *(_n(v) if isinstance(v, float) else "—" for v in linha[1:])]
            for linha in linhas_m
        ],
    }
    entradas: dict[str, Any] = {
        "casos_do_modelo": ", ".join(quadro.casos),
        "pilares": len(quadro.pilares),
        "pilares_sem_base": len(quadro.pendentes) or None,
    }
    entradas.update(dict(contexto or {}))
    entradas = {k: v for k, v in entradas.items() if v is not None}
    alertas = [i.texto for i in quadro.conferencia if i.nivel != pc.NIVEL_OK]
    return criar_registro_tecnico(
        modulo="Esforços do modelo",
        modulo_id="esforcos_modelo",
        titulo=f"Quadro de cargas para as fundações — {len(quadro.pilares)} pilar(es)",
        status="Calculado",
        resumo=destaque,
        entradas=entradas,
        resultados={
            "destaque_memorial": destaque,
            "tabelas_memorial": [matriz, tabela],
        },
        metodo=(
            "Esforços na seção de base de cada pilar, tirados das forças das vigas de cada estudo "
            "do SolidWorks Simulation (um estudo por caso de carga). " + CONVENCAO
        ),
        premissas=[NOTA_ANGLO],
        alertas=alertas,
        referencias=[f"{ca.REFERENCIA}, item 5.9."],
        conclusao=destaque,
        responsavel=responsavel,
    )


def resumo_numerico(quadro: QuadroDeFundacoes) -> dict[str, float | int | None]:
    cargas = [c for p in quadro.pilares for c in p.cargas]
    return {
        "pilares": len(quadro.pilares),
        "sem_base": len(quadro.pendentes),
        "maior_compressao": max((c.n for c in cargas), default=None),
        "maior_tracao": min((c.n for c in cargas if c.n < 0), default=None),
        "maior_momento": max((math.hypot(c.m1, c.m2) for c in cargas), default=None),
    }
