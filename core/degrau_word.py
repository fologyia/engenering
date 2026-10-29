"""Documento Word (editável) do degrau de escada em grade, com tudo o que a página mostra.

O botão no fim da página "Degrau de escada em grade" baixa este arquivo para reaproveitar os
dados em outro documento: identificação, todos os dados de entrada, o quadro-resumo, as 38
verificações, a geometria, os lances e patamares (com o esquemático e a escada em 3D), o memorial
do modelo adotado, a longarina (cargas, trechos, reações, verificações e diagramas), os 64 modelos,
o texto da requisição, os conflitos entre normas, os avisos, o que a página não faz e as
referências. As tabelas são tabelas do Word (copiáveis); as figuras, imagens PNG.

Sem Streamlit; usa as mesmas tabelas e figuras do PDF (:mod:`core.degrau_relatorio`).
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from datetime import date, datetime
from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from core import degrau_escada as de
from core import degrau_registro as reg
from core import escada_lances as el
from core import longarina_escada as le
from core import memorial_verificacoes as mv
from core import portico_figuras as pf

_AZUL = RGBColor(0x16, 0x32, 0x4F)
_CINZA = RGBColor(0x5F, 0x6B, 0x76)
_FUNDO_CABECALHO = "24577A"
_FUNDO_DESTAQUE = "FFF3C4"
_COR_STATUS = {
    "OK": "D9F2E3",
    "NÃO OK": "F8D4D4",
    "ALERTA": "FCEBC4",
    "INFO": "D8E8F8",
    "N/A": "E6E9ED",
}
_LARGURA_UTIL_MM = 273.0  # A4 em paisagem com margens de 12 mm


def _n(valor: float, casas: int = 2) -> str:
    return de.numero_pt(valor, casas)


def _data_texto(valor: Any) -> str:
    if isinstance(valor, (date, datetime)):
        return valor.strftime("%d/%m/%Y")
    return str(valor or "").strip()


def _celula_texto(valor: Any) -> str:
    if valor is None:
        return "—"
    if isinstance(valor, bool):
        return "Sim" if valor else "Não"
    if isinstance(valor, float):
        return _n(valor, 4)
    return str(valor)


def _sombrear(celula: Any, cor: str) -> None:
    propriedades = celula._tc.get_or_add_tcPr()
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), cor)
    propriedades.append(sombra)


def _repetir_cabecalho(linha: Any) -> None:
    propriedades = linha._tr.get_or_add_trPr()
    marca = OxmlElement("w:tblHeader")
    marca.set(qn("w:val"), "true")
    propriedades.append(marca)


class _Documento:
    """Um documento Word em paisagem com títulos, parágrafos, tabelas e figuras."""

    def __init__(self) -> None:
        self.doc = Document()
        secao = self.doc.sections[0]
        secao.orientation = WD_ORIENT.LANDSCAPE
        secao.page_width, secao.page_height = Mm(297), Mm(210)
        for margem in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(secao, margem, Mm(12))
        normal = self.doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(9)

    def titulo(self, texto: str, nivel: int = 1) -> None:
        paragrafo = self.doc.add_heading(texto, level=nivel)
        for trecho in paragrafo.runs:
            trecho.font.color.rgb = _AZUL

    def paragrafo(self, texto: str, *, legenda: bool = False, negrito: bool = False) -> None:
        paragrafo = self.doc.add_paragraph()
        trecho = paragrafo.add_run(texto)
        trecho.bold = negrito
        if legenda:
            trecho.font.size = Pt(8)
            trecho.font.color.rgb = _CINZA

    def itens(self, textos: Sequence[str]) -> None:
        for texto in textos:
            self.doc.add_paragraph(texto, style="List Bullet")

    def tabela(
        self,
        cabecalhos: Sequence[str],
        linhas: Sequence[Sequence[Any]],
        larguras_mm: Sequence[float] | None = None,
        *,
        status_coluna: int | None = None,
        destacar_linha: int | None = None,
        fonte: float = 8.0,
    ) -> None:
        tabela = self.doc.add_table(rows=1, cols=len(cabecalhos))
        tabela.style = "Table Grid"
        if larguras_mm is None:
            larguras_mm = [_LARGURA_UTIL_MM / len(cabecalhos)] * len(cabecalhos)
        escala = min(1.0, _LARGURA_UTIL_MM / sum(larguras_mm))
        larguras = [Mm(w * escala) for w in larguras_mm]

        def escrever(celula: Any, texto: str, *, cabecalho: bool = False) -> None:
            celula.text = ""
            trecho = celula.paragraphs[0].add_run(texto)
            trecho.font.size = Pt(fonte)
            if cabecalho:
                trecho.bold = True
                trecho.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        topo = tabela.rows[0]
        _repetir_cabecalho(topo)
        for celula, texto, largura in zip(topo.cells, cabecalhos, larguras, strict=True):
            escrever(celula, str(texto), cabecalho=True)
            _sombrear(celula, _FUNDO_CABECALHO)
            celula.width = largura
        for indice, linha in enumerate(linhas):
            celulas = tabela.add_row().cells
            for coluna, (celula, valor, largura) in enumerate(
                zip(celulas, linha, larguras, strict=True)
            ):
                texto = _celula_texto(valor)
                escrever(celula, texto)
                celula.width = largura
                if status_coluna is not None and coluna == status_coluna:
                    cor = _COR_STATUS.get(texto)
                    if cor:
                        _sombrear(celula, cor)
                elif destacar_linha is not None and indice == destacar_linha:
                    _sombrear(celula, _FUNDO_DESTAQUE)
        self.doc.add_paragraph()

    def figura(self, png: bytes, legenda: str, largura_mm: float = 230) -> None:
        self.doc.add_picture(BytesIO(png), width=Mm(min(largura_mm, _LARGURA_UTIL_MM)))
        self.paragrafo(legenda, legenda=True)

    def bytes(self) -> bytes:
        memoria = BytesIO()
        self.doc.save(memoria)
        return memoria.getvalue()


#: Rótulos das entradas da longarina (o memorial do projeto não as cura).
_ROTULOS_LONGARINA = {
    "longarina_perfil": "Longarina: perfil",
    "longarina_aco": "Longarina: aço",
    "longarina_apoio_no_pe": "Longarina: apoio no pé",
    "longarina_apoio_no_topo": "Longarina: apoio no topo",
    "longarina_trecho_horizontal_no_pe_m": "Longarina: trecho horizontal no pé [m]",
    "longarina_trecho_horizontal_no_topo_m": "Longarina: trecho horizontal no topo [m]",
    "longarina_distancia_horizontal_m": "Longarina: distância horizontal informada [m]",
    "longarina_guarda_corpo_kN_m": "Longarina: peso do guarda-corpo [kN/m]",
    "longarina_piso_do_patamar_kN_m2": "Longarina: piso do patamar [kN/m²]",
    "longarina_comprimento_destravado_m": "Longarina: comprimento destravado L_b [m]",
}


def _rotulo_livre(chave: str) -> str:
    if chave in _ROTULOS_LONGARINA:
        return _ROTULOS_LONGARINA[chave]
    texto = chave.replace("_", " ").strip()
    return texto[:1].upper() + texto[1:]


def _valor_de_entrada(valor: Any) -> str:
    if isinstance(valor, bool):
        return "Sim" if valor else "Não"
    if isinstance(valor, float):
        return _n(valor, 3)
    if isinstance(valor, (list, tuple)):
        return "; ".join(_valor_de_entrada(v) for v in valor)
    return str(valor)


def dados_de_entrada(
    r: de.ResultadoEscada,
    identificacao: Mapping[str, Any],
    escada: el.EscadaComPatamares | None,
    longarina: le.ResultadoLongarina | None,
) -> list[tuple[str, str]]:
    """Os dados de entrada da página como (rótulo, valor), com os rótulos do memorial."""
    contexto: dict[str, Any] = {}
    for chave in ("obra", "tag"):
        if identificacao.get(chave):
            contexto[chave] = identificacao[chave]
    if identificacao.get("data"):
        contexto["data_do_calculo"] = _data_texto(identificacao["data"])
    registro = reg.registro_degrau(
        r,
        contexto=contexto,
        responsavel=str(identificacao.get("responsavel") or ""),
        escada=escada,
        longarina=longarina,
    )
    pares = mv.pares_de_entrada(
        {"modulo_id": reg.MODULO_ID, "entradas": registro["entradas"]},
        _valor_de_entrada,
        _rotulo_livre,
    )
    if identificacao.get("responsavel"):
        pares.insert(0, ("Responsável", str(identificacao["responsavel"])))
    return pares


def _quadro_resumo(r: de.ResultadoEscada) -> list[list[str]]:
    e, g, ln, a = r.entrada, r.geometria, r.lances, r.adotado
    return [
        ["Modelo adotado", f"{a.modelo.nome} — {r.selecao.mensagem}"],
        [
            "Espelho h × piso b",
            f"{_n(g.h_mm, 1)} × {_n(g.b_mm, 0)} mm (2h + b = {_n(g.blondel_mm, 1)} mm; "
            f"α = {_n(g.alfa_graus, 2)}°)",
        ],
        [
            "Profundidade C × comprimento L",
            f"{_n(g.C_mm, 0)} × {_n(e.L_mm, 0)} mm (F = {g.F_mm} mm; r = {_n(g.r_mm, 0)} mm)",
        ],
        ["Espelhos / degraus em grade", f"{g.n} / {ln.n_degraus_grade} — {ln.texto}"],
        ["Aproveitamento máximo", f"{_n(100 * r.aproveitamento_maximo, 1)} %"],
        [
            "Peso: degrau / total",
            f"{_n(r.peso_unitario_kg, 2)} kg / {_n(r.peso_total_kg, 1)} kg "
            f"({ln.n_degraus_grade} degraus)",
        ],
        ["Situação da geometria", f"Nível {g.faixa.nivel} de 4 — {g.faixa.mensagem}"],
        ["Status geral", f"{r.status} — {de.texto_da_contagem(r.contagem)}"],
        [
            "Deslocamento do degrau",
            f"{_n(a.delta_d_mm, 2)} mm (carga distribuída) / {_n(a.delta_iso_mm, 2)} mm "
            "(carga concentrada da ISO)",
        ],
    ]


def _resumo_da_longarina(lg: le.LongarinaDoLance) -> list[list[str]]:
    inclinado = next(t for t in lg.trechos if t.barra == le.TRECHO_INCLINADO)
    rb = lg.elu.barra(le.TRECHO_INCLINADO)
    m_rd = inclinado.resultado.mx_rd
    ry = max(abs(x.ry) for x in lg.reacoes(le.ELU))
    rx = max(abs(x.rx) for x in lg.reacoes(le.ELU))
    return [
        [
            "Distância horizontal / inclinada",
            f"{_n(lg.vao_horizontal_m, 2)} / {_n(lg.vao_inclinado_m, 2)} m ("
            + ("informada" if lg.vao_informado else "linha dos bocéis, n·b")
            + f"; α = {_n(lg.alfa_graus, 1)}°)",
        ],
        [
            "Deslocamento vertical máximo",
            f"{_n(max(t.deslocamento_vertical_mm for t in lg.trechos), 2)} mm; flecha "
            f"{_n(abs(inclinado.flecha_mm), 2)} mm ≤ L/350 = {_n(inclinado.limite_flecha_mm, 1)} mm",
        ],
        [
            "M_Sd / M_Rd",
            f"{_n(abs(rb.m_max.m), 2)} / {_n(m_rd, 2) if m_rd else '—'} kN·m; aproveitamento "
            f"{_n(100 * lg.aproveitamento_maximo, 0)} % · {lg.status}",
        ],
        [
            "Maior reação vertical ELU",
            f"{_n(ry, 2)} kN; "
            + (f"empuxo horizontal ELU {_n(rx, 2)} kN" if rx > 0.005 else "sem empuxo horizontal"),
        ],
    ]


def gerar_word_degrau(
    r: de.ResultadoEscada,
    identificacao: Mapping[str, Any] | None = None,
    *,
    escada: el.EscadaComPatamares | None = None,
    longarina: le.ResultadoLongarina | None = None,
    escada_3d: Mapping[str, Any] | None = None,
) -> bytes:
    """O .docx com todas as informações da página. ``identificacao``: obra, tag, responsavel e
    data; ``escada``, ``longarina`` e ``escada_3d`` como no PDF da página."""
    ident = dict(identificacao or {})
    e, ln, a = r.entrada, r.lances, r.adotado
    doc = _Documento()
    doc.titulo("Degrau de escada em grade de piso eletrofundida", 0)
    doc.paragrafo(
        "Catálogo Selmec “Degraus” (DS) · NR-12 Anexo III · NR-22 · Critério Anglo "
        "AA-BR-DPST-DR-0001 Rev.1 · NBR 8800:2008 · NBR 6120 · ISO 14122-3 — modelo adotado "
        f"{a.modelo.nome}",
        legenda=True,
    )
    doc.tabela(
        ["Dado", "Valor", "Dado", "Valor"],
        [
            ["Obra", ident.get("obra") or "-", "TAG", ident.get("tag") or "-"],
            [
                "Responsável",
                ident.get("responsavel") or "-",
                "Data",
                _data_texto(ident.get("data")) or "-",
            ],
            [
                "Enquadramento",
                de.ENQUADRAMENTOS[e.enquadramento],
                "Uso (Anglo 10.2)",
                de.USOS[e.uso],
            ],
        ],
        [40, 96, 40, 97],
    )

    doc.titulo("O que esta página não faz", 1)
    doc.itens(list(reg.NAO_FAZ))

    doc.titulo("Dados de entrada", 1)
    doc.tabela(
        ["Dado", "Valor"], dados_de_entrada(r, ident, escada, longarina), [140, 133], fonte=8.5
    )

    doc.titulo("Quadro-resumo", 1)
    doc.tabela(["Item", "Resultado"], _quadro_resumo(r), [60, 213])
    for aviso in r.avisos:
        doc.paragrafo(f"Aviso: {aviso}")

    doc.titulo(f"1. Verificações ({len(r.verificacoes)}) com o item da norma", 1)
    doc.paragrafo(
        "Valor: o que atua (resistência) ou o valor adotado (limites). Limite: a resistência ou o "
        "que a norma exige; nas faixas, a faixa vem na observação. Aproveitamento só nas "
        "verificações de resistência.",
        legenda=True,
    )
    doc.tabela(
        list(reg.COLUNAS_VERIFICACAO),
        [
            [
                linha["Verificação"],
                linha["Norma/item"],
                linha["Valor"],
                linha["Limite"],
                "—"
                if linha["Aproveitamento (%)"] is None
                else f"{linha['Aproveitamento (%)']:.1f}".replace(".", ","),
                linha["Status"],
                linha["Observação"],
            ]
            for linha in reg.tabela_de_verificacoes(r)
        ],
        [58, 40, 21, 21, 15, 14, 104],
        status_coluna=5,
        fonte=7.5,
    )

    doc.titulo("2. Geometria da escada", 1)
    doc.tabela(
        ["Grandeza", "Valor", "Un.", "Cálculo"], reg.linhas_da_geometria(r), [70, 40, 15, 148]
    )

    if escada is not None:
        doc.titulo("2.1 Lances, patamares e esquemático", 1)
        esquema = el.figura_da_escada(escada)
        doc.figura(esquema.png, esquema.legenda)
        if escada_3d is not None:
            import importlib

            from core import cena_escada as cesc

            # record_charts ainda está fora da lista do mypy: importado pelo nome.
            imagem_da_cena = importlib.import_module("core.record_charts").imagem_da_cena
            modo = str(escada_3d.get("modo") or cesc.MODO_DETALHADO).lower()
            tres_d = imagem_da_cena(
                cesc.cena_do_registro(escada_3d), escada_3d, f"Escada em 3D ({modo})."
            )
            doc.figura(tres_d.png, tres_d.legenda, 200)
        doc.tabela(list(el.CABECALHO_LANCES), el.linhas_dos_lances(escada))
        if escada.patamares:
            doc.tabela(list(el.CABECALHO_PATAMARES), el.linhas_dos_patamares(escada))
        for aviso in escada.avisos:
            doc.paragrafo(f"Aviso: {aviso}")

    doc.titulo(f"3. Memorial do modelo adotado ({a.modelo.nome})", 1)
    doc.tabela(
        ["Grandeza", "Valor", "Un.", "Cálculo"], reg.linhas_do_dimensionamento(r), [70, 40, 15, 148]
    )

    if longarina is not None:
        doc.titulo(
            f"3.1 Longarina {longarina.entrada.perfil} ({longarina.entrada.aco}) — "
            f"{longarina.status}",
            1,
        )
        doc.paragrafo(
            "Núcleo de pórtico plano do programa: ELU γg·G + γq·Q para os esforços, ELS G + Q "
            f"para a flecha; apoios: pé {longarina.entrada.apoio_pe}, topo "
            f"{longarina.entrada.apoio_topo}.",
            legenda=True,
        )
        for aviso in longarina.avisos:
            doc.paragrafo(f"Aviso: {aviso}")
        for lg in longarina.longarinas:
            doc.titulo(lg.titulo, 2)
            doc.tabela(["Item", "Valor"], _resumo_da_longarina(lg), [70, 203])
            doc.paragrafo("Cargas por longarina", negrito=True)
            doc.tabela(
                ["Carga", "Valor", "Un.", "Origem"], le.linhas_das_cargas(lg), [70, 35, 30, 138]
            )
            doc.paragrafo("Trechos", negrito=True)
            doc.tabela(list(le.CABECALHO_TRECHOS), le.linhas_dos_trechos(lg))
            doc.paragrafo("Reações", negrito=True)
            doc.tabela(list(le.CABECALHO_REACOES), le.linhas_das_reacoes(lg))
            doc.paragrafo("Verificações", negrito=True)
            doc.tabela(
                [
                    "Verificação",
                    "Norma/item",
                    "Solicitante",
                    "Resistente",
                    "Un.",
                    "Aprov. (%)",
                    "Status",
                    "Cálculo",
                ],
                [
                    [
                        v.nome,
                        v.referencia,
                        "—" if v.solicitante is None else _n(v.solicitante, 3),
                        "—" if v.resistente is None else _n(v.resistente, 3),
                        v.unidade,
                        "—" if v.aproveitamento is None else _n(100 * v.aproveitamento, 1),
                        v.status,
                        v.formula,
                    ]
                    for v in lg.verificacoes
                ],
                [60, 30, 18, 18, 12, 15, 14, 106],
                status_coluna=6,
                fonte=7.5,
            )
            doc.paragrafo("Diagramas", negrito=True)
            for desenho in pf.figuras_do_portico(
                lg.elu,
                prefixo=lg.titulo,
                caso=f"ELU ({_n(lg.gamma_g)}·G + {_n(lg.gamma_q)}·Q)",
                deformada=lg.els,
                caso_deformada=le.ELS,
            ):
                doc.figura(desenho.png, desenho.legenda, 200)

    doc.titulo(f"4. Os {len(r.tabela_modelos)} modelos do catálogo", 1)
    doc.paragrafo(
        "u = aproveitamento de cada verificação do degrau. Atende: u máx ≤ 1 e L dentro do máximo "
        f"recomendado. O modelo adotado ({a.modelo.nome}) está em destaque.",
        legenda=True,
    )
    catalogo = reg.linhas_do_catalogo(r)
    colunas = [c for c in catalogo[0] if c != "Adotado"]
    destaque = next((i for i, linha in enumerate(catalogo) if linha["Adotado"]), None)
    casas = {"λ": 1, "λp": 1, "λr": 1, "M_Rd (kN·m)": 4, "Peso do degrau (kg)": 3}
    doc.tabela(
        colunas,
        [
            [
                _n(linha[c], casas.get(c, 3)) if isinstance(linha[c], float) else linha[c]
                for c in colunas
            ]
            for linha in catalogo
        ],
        destacar_linha=destaque,
        fonte=6.0,
    )

    doc.titulo("5. Texto para a requisição de compra", 1)
    doc.paragrafo(r.especificacao)
    doc.paragrafo(
        f"Peso unitário {_n(r.peso_unitario_kg)} kg; peso total {_n(r.peso_total_kg, 1)} kg "
        f"({ln.n_degraus_grade} degraus). Estimativa.",
        legenda=True,
    )

    doc.titulo("6. Conflitos entre normas, avisos e limites do cálculo", 1)
    doc.paragrafo("Conflitos entre normas", negrito=True)
    doc.itens([f"{i}. {item}" for i, item in enumerate(reg.CONFLITOS_ENTRE_NORMAS, start=1)])
    doc.paragrafo("Avisos e limites do cálculo — conferir antes de emitir", negrito=True)
    doc.itens(list(reg.AVISOS_FIXOS))
    doc.paragrafo("Referências", negrito=True)
    doc.itens(list(reg.REFERENCIAS))
    return doc.bytes()
