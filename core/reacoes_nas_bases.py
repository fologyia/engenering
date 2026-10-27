"""Reações nas bases para o projeto da fundação: a força e o momento de cada pé, caso a caso.

Quem projeta a fundação (o engenheiro civil) faz as combinações dele; recebe de cada pé, em cada
caso de carga (PP, PE, SC, W0, W90…), as três forças e os três momentos **característicos** — sem
combinar nem majorar (critério Anglo AA-BR-DPST-DR-0001, item 5.9) —, a locação dos pés e a
natureza de cada caso (permanente, variável, vento) com os coeficientes do plano, para combinar.

De onde vêm os números, em ordem de preferência, caso a caso:

1. **A reação em cada apoio** (os ``-apoios.csv`` da macro exportar_casos_de_carga): a reação de
   cada nó apoiado, já nos eixos do modelo, com os momentos — o número exato do SolidWorks. Vale
   o caso cujos apoios somam a reação total do mesmo caso (senão o arquivo é de outra rodada).
2. **A base de cada pilar** ("Listar forças da viga"): a força na ponta da base, levada aos eixos do
   modelo pela direção da seção de cada pilar (:mod:`core.direcoes_da_secao`). Vale quando só o
   pilar chega no apoio.

O pé sem nenhuma das duas fica de fora daquele caso, com o motivo — nada de célula vazia que pareça
zero. Cada caso é conferido: a soma dos pés tem de ser a reação total do caso.

Convenção: **ação da estrutura na fundação** (a reação com o sinal trocado). Nos eixos da fundação
(o padrão): X e Y na planta, Z para cima — o X é o do SolidWorks e o Y é o −Z do SolidWorks (Y para
cima no modelo); a força vertical para baixo (Fz < 0) é compressão, ``N = −Fz``. Ou nos próprios
eixos do SolidWorks.

Sem Streamlit.
"""

from __future__ import annotations

import io
import math
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import date

from core import apoios_do_modelo as am
from core import criterio_anglo as ca
from core import diagnostico_do_modelo as dm
from core import direcoes_da_secao as ds
from core import esforcos_modelo as em
from core import leitura_do_modelo as ld
from core import load_combinations as comb
from core import plano_de_cargas as pc
from core import quadro_fundacoes as qf
from core import vento_por_barras as vpb

Ponto = tuple[float, float, float]

FONTE_APOIO = "reação no apoio (macro)"
FONTE_PILAR = "base do pilar (forças da viga)"
EIXOS_FUNDACAO = "fundação"
EIXOS_MODELO = "modelo"
EIXOS = (EIXOS_FUNDACAO, EIXOS_MODELO)
#: A soma dos pés e a reação total do caso "são a mesma" até esta diferença.
TOLERANCIA_RELATIVA = 0.02
TOLERANCIA_ABSOLUTA_KN = 0.02
#: A cortante na base que orienta a seção do pilar pela macro (e pelo menos 10 % da compressão).
CORTANTE_QUE_ORIENTA_KN = 0.05

NOTA_COMBINACOES = (
    f"{ca.QUADRO_DE_CARGAS_FUNDACOES} Valores característicos de cada caso de carga: as "
    "combinações são de quem projeta a fundação (a aba Casos traz a natureza de cada caso e os "
    "coeficientes da NBR 8800 usados na estrutura metálica)."
)
#: Achados de corrigir que não mudam as reações (só o plano ou a verificação das barras).
NAO_MUDAM_AS_REACOES = (
    "Plano de cargas com cargas de exemplo",
    "Eixo forte trocado",
    "Perfil diferente do modelo",
)


@dataclass(frozen=True)
class Pe:
    rotulo: str  # A1, A2… pela posição (X, depois a outra horizontal)
    ponto: Ponto | None  # m, eixos do modelo
    pilar: str = ""  # o nome no Simulation
    perfil: str = ""
    apoio: str = ""  # o rótulo no quadro por apoio da macro
    comportamento: str = ""  # engastado ou articulado (pelos momentos da macro)


@dataclass(frozen=True)
class AcaoNoPe:
    pe: str
    caso: str
    forca_kN: Ponto  # ação da estrutura na fundação, eixos do modelo
    momento_kNm: Ponto
    fonte: str


@dataclass(frozen=True)
class ReacoesNasBases:
    pes: tuple[Pe, ...]
    acoes: tuple[AcaoNoPe, ...]
    casos: tuple[str, ...]
    eixo_vertical: str = "Y"
    conferencia: tuple[pc.ItemDeConferencia, ...] = ()
    faltam: tuple[str, ...] = ()  # o que ficou de fora e por quê
    avisos: tuple[str, ...] = ()

    def acao(self, pe: str, caso: str) -> AcaoNoPe | None:
        return next((a for a in self.acoes if a.pe == pe and a.caso == caso), None)

    def fontes(self, caso: str) -> set[str]:
        return {a.fonte for a in self.acoes if a.caso == caso}


def _n(valor: float, casas: int = 2) -> str:
    if round(valor, casas) == 0:
        valor = 0.0
    return f"{valor:.{casas}f}".replace(".", ",")


def _lista(itens: Sequence[str]) -> str:
    itens = list(itens)
    return itens[0] if len(itens) == 1 else ", ".join(itens[:-1]) + " e " + itens[-1]


# ---------------------------------------------------------------------------------------------
# Montagem
# ---------------------------------------------------------------------------------------------
def _apoios_que_valem(apoios: am.ApoiosDoModelo, dados: em.EsforcosDoModelo) -> set[str]:
    """Os casos cujos apoios somam a reação total do mesmo caso (sem a reação, valem)."""
    valem = set()
    for codigo, c in apoios.casos.items():
        caso = dados.casos.get(codigo)
        soma = [sum(n.forca_N[k] for n in c.nos) / 1e3 for k in range(3)]
        if caso is not None and caso.reacoes is not None:
            modelo = caso.reacoes.modelo
            escala = max(max(abs(v) for v in modelo), 1e-9)
            if max(abs(a - b) for a, b in zip(soma, modelo, strict=True)) > max(
                TOLERANCIA_RELATIVA * escala, TOLERANCIA_ABSOLUTA_KN
            ):
                continue
        valem.add(codigo)
    return valem


def _corpos_dos_pilares(
    dados: em.EsforcosDoModelo,
    geometria: ld.GeometriaDoModelo | None,
    leituras: Sequence[ld.LeituraDaBarra] | None,
) -> dict[str, ld.CorpoDoModelo]:
    if geometria is None or not geometria.corpos:
        return {}
    if leituras is not None:
        return {lb.barra: lb.corpo for lb in leituras if lb.corpo is not None}
    return dict(ld.correspondencia(geometria, dados.nomes_dos_membros, dados).corpos)


def _pe_do_corpo(corpo: ld.CorpoDoModelo, eixo_vertical: str) -> Ponto:
    vertical = 1 if eixo_vertical == "Y" else 2
    pe = min((corpo.inicio, corpo.fim), key=lambda p: p[vertical])
    return (float(pe[0]), float(pe[1]), float(pe[2]))


def _ordem(ponto: Ponto | None, eixo_vertical: str) -> tuple[float, float]:
    if ponto is None:
        return (math.inf, math.inf)
    return (ponto[0], ponto[2] if eixo_vertical == "Y" else ponto[1])


def _casos_na_ordem(dados: em.EsforcosDoModelo, plano: pc.PlanoDeCargas | None) -> tuple[str, ...]:
    do_plano = [c for c in (plano.codigos if plano else ()) if c in dados.casos]
    return (*do_plano, *(c for c in dados.casos if c not in do_plano))


def _tolerancia(*valores: float) -> float:
    escala = max((abs(v) for v in valores), default=0.0)
    return max(0.03 * escala, TOLERANCIA_ABSOLUTA_KN)


def _no_apoio(
    quadro_a: am.QuadroPorApoio, apoio: str, caso: str, eixo: str
) -> tuple[float, float] | None:
    """(N compressão +, horizontal resultante) que o apoio leva no caso, kN."""
    carga = quadro_a.carga(apoio, caso)
    if carga is None:
        return None
    a, b = (0, 2) if eixo == "Y" else (0, 1)
    return (quadro_a.n_kN(carga), math.hypot(carga.forca_kN[a], carga.forca_kN[b]))


def _na_base(pilar: qf.PilarDoQuadro, caso: str) -> qf.CargaNaBase | None:
    return next((c for c in pilar.cargas if c.caso == caso), None)


def _ligar_pelos_valores(
    quadro_a: am.QuadroPorApoio,
    valem: set[str],
    pilares: Mapping[str, qf.PilarDoQuadro],
    eixo: str,
) -> dict[str, str]:
    """Pilar → apoio, sem a posição dos nós: o único apoio cuja reação, em todos os casos que
    valem, é a compressão e a cortante na base do pilar. Vazio se algum pilar não tem um só."""
    casos = [c for c in quadro_a.casos if c in valem]
    rotulos = [a.rotulo for a in quadro_a.apoios]
    if not casos or len(rotulos) != len(pilares):
        return {}

    def base(pilar: str, caso: str) -> tuple[float, float] | None:
        carga = _na_base(pilares[pilar], caso)
        return None if carga is None else (carga.n, math.hypot(carga.v1, carga.v2))

    tolerancia = {
        caso: _tolerancia(*(v for p in pilares if (b := base(p, caso)) is not None for v in b))
        for caso in casos
    }
    mapa: dict[str, str] = {}
    for pilar in pilares:
        candidatos = []
        for rotulo in rotulos:
            for caso in casos:
                b, a = base(pilar, caso), _no_apoio(quadro_a, rotulo, caso, eixo)
                if (
                    a is None
                    or b is None
                    or max(abs(b[0] - a[0]), abs(b[1] - a[1])) > tolerancia[caso]
                ):
                    break
            else:
                candidatos.append(rotulo)
        if len(candidatos) != 1:
            return {}
        mapa[pilar] = candidatos[0]
    return mapa if len(set(mapa.values())) == len(mapa) else {}


def _unitario(v: Sequence[float]) -> Ponto | None:
    n = math.sqrt(sum(c * c for c in v))
    return None if n <= 1e-12 else (v[0] / n, v[1] / n, v[2] / n)


def _vetorial(a: Sequence[float], b: Sequence[float]) -> Ponto:
    return (a[1] * b[2] - a[2] * b[1], a[2] * b[0] - a[0] * b[2], a[0] * b[1] - a[1] * b[0])


def _escalar(a: Sequence[float], b: Sequence[float]) -> float:
    return a[0] * b[0] + a[1] * b[1] + a[2] * b[2]


def _direcao_pela_macro(
    pilar: qf.PilarDoQuadro,
    corpo: ld.CorpoDoModelo | None,
    quadro_a: am.QuadroPorApoio,
    apoio: str,
    valem: set[str],
    eixo: str,
) -> ds.DirecoesDaSecao | None:
    """O triedro do Simulation de um pilar pelas reações da macro: a cortante na base (V₁, V₂, nos
    eixos da seção) gira até ser a reação horizontal no pé dele, nos casos em que os apoios valem
    (mínimos quadrados no ângulo da direção 1). Só vale se, girada, bate em todos esses casos —
    e o axial também. ``None`` sem cortante para orientar ou se não bate."""
    sinal = -1.0 if pilar.base == qf.BASE_INICIO else 1.0
    para_cima: Ponto | None = (0.0, 0.0, 1.0)
    if corpo is not None:
        vertical = 1 if eixo == "Y" else 2
        pe, topo = sorted((corpo.inicio, corpo.fim), key=lambda p: p[vertical])
        para_cima = _unitario(
            [
                t - p
                for t, p in zip(
                    vpb.para_a_planta(topo, eixo), vpb.para_a_planta(pe, eixo), strict=True
                )
            ]
        )
    if para_cima is None:
        return None
    # O axial vai do Fim 1 ao Fim 2: sobe quando a base é o início da lista.
    axial = (
        para_cima if pilar.base == qf.BASE_INICIO else (-para_cima[0], -para_cima[1], -para_cima[2])
    )
    e1 = _unitario(_vetorial(axial, (1.0, 0.0, 0.0))) or _unitario(
        _vetorial(axial, (0.0, 1.0, 0.0))
    )
    if e1 is None:  # pragma: no cover - o axial não é paralelo aos dois ao mesmo tempo
        return None
    e2 = _vetorial(axial, e1)
    pares: list[tuple[complex, complex, float, float]] = []
    for caso in quadro_a.casos:
        carga, no_apoio = _na_base(pilar, caso), quadro_a.carga(apoio, caso)
        if caso not in valem or carga is None or no_apoio is None:
            continue
        reacao = vpb.para_a_planta(tuple(-v for v in no_apoio.forca_kN), eixo)
        pares.append(
            (
                complex(_escalar(reacao, e1), _escalar(reacao, e2)),
                complex(sinal * carga.v1, sinal * carga.v2),
                _escalar(reacao, axial),
                sinal * -carga.n,
            )
        )
    # Só orienta o caso com cortante de verdade (o vento): a dos casos de gravidade é resto do
    # pórtico, pequena perto da compressão, e qualquer arredondamento a vira de lado.
    orientam = [
        (r, v) for r, v, _, n in pares if abs(v) >= max(CORTANTE_QUE_ORIENTA_KN, 0.1 * abs(n))
    ]
    soma = sum((r * v.conjugate() for r, v in orientam), complex(0.0))
    if not orientam or abs(soma) <= 1e-12:
        return None
    giro = soma / abs(soma)
    erro = max(max(abs(r - giro * v), abs(fa - esperado)) for r, v, fa, esperado in pares)
    if erro > _tolerancia(*(x for r, v, fa, e in pares for x in (abs(r), abs(v), fa, e))):
        return None
    d1 = tuple(giro.real * a + giro.imag * b for a, b in zip(e1, e2, strict=True))
    casos = [c for c in quadro_a.casos if c in valem]
    return ds.DirecoesDaSecao(
        pilar.pilar,
        corpo.nome if corpo is not None else "",
        axial,
        (d1[0], d1[1], d1[2]),
        _vetorial(axial, d1),
        sinal,
        "reações da macro",
        f"pelas reações da macro em {_lista(casos)}",
    )


def montar(
    dados: em.EsforcosDoModelo,
    apoios: am.ApoiosDoModelo | None = None,
    *,
    plano: pc.PlanoDeCargas | None = None,
    geometria: ld.GeometriaDoModelo | None = None,
    leituras: Sequence[ld.LeituraDaBarra] | None = None,
    direcoes: ds.DirecoesDosPilares | None = None,
) -> ReacoesNasBases:
    """Os pés, a ação em cada um caso a caso e as conferências."""
    apoios = apoios or am.ApoiosDoModelo()
    eixo = dados.eixo_vertical
    casos = _casos_na_ordem(dados, plano)
    avisos: list[str] = []
    # A macro: os apoios e os casos cujos arquivos valem.
    quadro_a = (
        am.montar_quadro(apoios, geometria=geometria, esforcos=dados, eixo_vertical=eixo)
        if apoios.casos
        else None
    )
    valem = _apoios_que_valem(apoios, dados)
    de_outra_rodada = sorted(set(apoios.casos) - valem)
    if de_outra_rodada:
        avisos.append(
            f"Os -apoios.csv de {_lista(de_outra_rodada)} não somam a reação total do caso (são "
            "de outra rodada): nesses casos o programa usa a base de cada pilar."
        )
    # Os pilares: a base de cada um e o pé no esqueleto.
    quadro_f = qf.montar_quadro(dados)
    corpos = _corpos_dos_pilares(dados, geometria, leituras)
    pilares = {p.pilar: p for p in quadro_f.pilares if p.base is not None}
    # Os pés: os apoios da macro (com o pilar que chega neles) ou os pilares.
    pes: list[Pe] = []
    apoio_do_pilar: dict[str, str] = {}
    if quadro_a is not None:
        for a in quadro_a.apoios:
            pilar = next(
                (p for p in pilares if p in corpos and corpos[p].nome == a.barra and a.barra), ""
            )
            if pilar:
                apoio_do_pilar[pilar] = a.rotulo
        if pilares and set(apoio_do_pilar) != set(pilares):
            # Sem a posição dos nós (macro v12) ou sem o esqueleto: pelos valores.
            pelos_valores = _ligar_pelos_valores(quadro_a, valem, pilares, eixo)
            if pelos_valores:
                apoio_do_pilar = pelos_valores
                avisos.append(
                    "Cada apoio da macro foi ligado ao seu pilar pelos valores: a compressão e a "
                    "cortante na base do pilar batem com a reação do apoio em "
                    f"{_lista([c for c in casos if c in valem])}."
                )
    ligados = bool(pilares) and set(apoio_do_pilar) == set(pilares)
    so_pela_macro = quadro_a is not None and all(c in valem for c in dados.casos)
    # A direção da seção de cada pilar: a da página (esqueleto) ou, sem ela, a que as reações da
    # macro dão (casos que valem) — para os casos com os apoios de outra rodada.
    direcoes_dos_pilares = dict(direcoes.pilares) if direcoes is not None else {}
    if quadro_a is not None and ligados:
        calibradas = {
            pilar: d
            for pilar, apoio in apoio_do_pilar.items()
            if not (pilar in direcoes_dos_pilares and direcoes_dos_pilares[pilar].sabida)
            and (
                d := _direcao_pela_macro(
                    pilares[pilar], corpos.get(pilar), quadro_a, apoio, valem, eixo
                )
            )
            is not None
        }
        if calibradas and any(c not in valem for c in dados.casos):
            avisos.append(
                f"A direção da seção de {len(calibradas)} pilar(es) saiu das reações da macro "
                "(o que a base de cada pilar leva bate com a reação no pé dele): os casos sem os "
                "apoios saem pela base dos pilares."
            )
        direcoes_dos_pilares.update(calibradas)
    if quadro_a is not None and (ligados or not pilares or so_pela_macro):
        if pilares and not ligados:
            avisos.append(
                "Os pilares não foram ligados aos apoios da macro (sem o esqueleto ou sem a posição "
                "dos nós): os pés são os apoios, sem o nome do pilar."
            )
        pilar_do_apoio = {v: k for k, v in apoio_do_pilar.items()}
        for a in quadro_a.apoios:
            pilar = pilar_do_apoio.get(a.rotulo, "")
            ponto = a.ponto
            if ponto is None and pilar in corpos:
                ponto = _pe_do_corpo(corpos[pilar], eixo)
            pes.append(
                Pe(
                    a.rotulo,
                    ponto,
                    pilar,
                    pilares[pilar].perfil if pilar else "",
                    a.rotulo,
                    a.comportamento,
                )
            )
    else:
        if quadro_a is not None:
            avisos.append(
                "Os apoios da macro não puderam ser ligados aos pilares (sem o esqueleto ou sem a "
                "posição dos nós) e não valem para todos os casos: os pés são as bases dos pilares."
            )
        ordenados = sorted(
            pilares,
            key=lambda p: (
                _ordem(_pe_do_corpo(corpos[p], eixo) if p in corpos else None, eixo),
                p,
            ),
        )
        for k, pilar in enumerate(ordenados, start=1):
            ponto = _pe_do_corpo(corpos[pilar], eixo) if pilar in corpos else None
            pes.append(Pe(f"A{k}", ponto, pilar, pilares[pilar].perfil))
    # A ação em cada pé, caso a caso.
    acoes: list[AcaoNoPe] = []
    faltam: list[str] = []
    for caso in casos:
        sem: list[str] = []
        motivos: set[str] = set()
        for pe in pes:
            carga_a = (
                quadro_a.carga(pe.apoio, caso)
                if quadro_a is not None and pe.apoio and caso in valem
                else None
            )
            if carga_a is not None:
                acoes.append(
                    AcaoNoPe(pe.rotulo, caso, carga_a.forca_kN, carga_a.momento_kNm, FONTE_APOIO)
                )
                continue
            p = pilares.get(pe.pilar)
            carga_p = next((c for c in p.cargas if c.caso == caso), None) if p else None
            d = direcoes_dos_pilares.get(pe.pilar) if pe.pilar else None
            na_fundacao = d.na_fundacao(carga_p) if d is not None and carga_p is not None else None
            if na_fundacao is not None:
                forca, momento = (vpb.do_modelo(v, eixo) for v in na_fundacao)
                acoes.append(AcaoNoPe(pe.rotulo, caso, forca, momento, FONTE_PILAR))
                continue
            sem.append(pe.rotulo)
            if not pe.pilar:
                motivos.add("sem o -apoios.csv desse caso (ou de outra rodada)")
            elif carga_p is None:
                motivos.add("sem a força na base do pilar nesse caso")
            else:
                motivos.add(
                    "sem a direção da seção do pilar (esqueleto da macro v6 e um caso com vento "
                    "no pilar) e sem o -apoios.csv do caso"
                )
        if sem:
            faltam.append(
                f"{caso}: {_lista(sem)} fora — {'; '.join(sorted(motivos))}. Rode a macro "
                "exportar_casos_de_carga (grava o -apoios.csv de cada caso) e envie em Esforços do "
                "modelo."
            )
    return ReacoesNasBases(
        tuple(pes),
        tuple(acoes),
        casos,
        eixo,
        tuple(_conferir(dados, pes, acoes, casos)),
        tuple(faltam),
        tuple(avisos),
    )


def _conferir(
    dados: em.EsforcosDoModelo,
    pes: Sequence[Pe],
    acoes: Sequence[AcaoNoPe],
    casos: Sequence[str],
) -> list[pc.ItemDeConferencia]:
    """A soma dos pés de cada caso contra a reação total do caso (com o sinal trocado)."""
    itens: list[pc.ItemDeConferencia] = []
    for caso in casos:
        do_caso = [a for a in acoes if a.caso == caso]
        reacoes = dados.casos[caso].reacoes if caso in dados.casos else None
        if reacoes is None or len(do_caso) < len(pes) or not do_caso:
            continue
        soma = tuple(sum(a.forca_kN[k] for a in do_caso) for k in range(3))
        esperado = tuple(-v for v in reacoes.modelo)
        diferenca = math.dist(soma, esperado)
        referencia = max(math.hypot(*esperado), 1e-9)
        if diferenca <= max(TOLERANCIA_RELATIVA * referencia, TOLERANCIA_ABSOLUTA_KN):
            itens.append(
                pc.ItemDeConferencia(
                    pc.NIVEL_OK,
                    f"{caso}: a soma dos {len(do_caso)} pés ({_n(math.hypot(*soma))} kN) é a reação "
                    "total do caso.",
                )
            )
        else:
            itens.append(
                pc.ItemDeConferencia(
                    pc.NIVEL_ATENCAO,
                    f"{caso}: a soma dos pés ({'; '.join(_n(v) for v in soma)} kN) não é a reação "
                    f"total do caso ({'; '.join(_n(v) for v in esperado)} kN, eixos do modelo) — "
                    "há apoio fora dos pés, diagonal chegando na base de um pilar, ou arquivos de "
                    "rodadas diferentes.",
                )
            )
    return itens


def pendencias(diagnostico: dm.Diagnostico, r: ReacoesNasBases) -> list[dm.Achado]:
    """Os problemas para corrigir que mudam as reações exportadas (do painel Antes de verificar).

    Os do arquivo de forças de um caso só contam se algum pé daquele caso saiu da base do pilar
    (com os apoios da macro, as forças da viga não entram)."""
    saida = []
    for a in diagnostico.achados:
        if a.nivel != dm.NIVEL_ERRO or a.titulo.startswith(NAO_MUDAM_AS_REACOES):
            continue
        if a.detalhe.get("sobre") == dm.SOBRE_AS_FORCAS:
            casos = a.detalhe.get("casos") or [a.detalhe.get("caso")]
            if not any(FONTE_PILAR in r.fontes(str(c)) for c in casos):
                continue
        saida.append(a)
    return saida


# ---------------------------------------------------------------------------------------------
# Tabelas
# ---------------------------------------------------------------------------------------------
def nos_eixos(vetor: Ponto, eixo_vertical: str, eixos: str) -> Ponto:
    """O vetor do modelo nos eixos pedidos (fundação: X, Y na planta e Z para cima)."""
    if eixos == EIXOS_MODELO:
        return (vetor[0] + 0.0, vetor[1] + 0.0, vetor[2] + 0.0)
    return em.para_o_plano(vetor, eixo_vertical)


def _vertical(eixo_vertical: str, eixos: str) -> int:
    return 1 if eixos == EIXOS_MODELO and eixo_vertical == "Y" else 2


def colunas(r: ReacoesNasBases, eixos: str = EIXOS_FUNDACAO) -> list[str]:
    nomes = ("x", "y", "z")
    return [
        "Pé",
        "Caso",
        *(f"F{e} (kN)" for e in nomes),
        *(f"M{e} (kN·m)" for e in nomes),
        "N (kN, compressão +)",
        "H (kN, horizontal)",
        "M horizontal (kN·m)",
        "Descrição do caso",
        "Pilar",
        "De onde veio",
    ]


def _arredondar(v: float, casas: int = 3) -> float:
    valor = round(v, casas)
    return 0.0 if valor == 0 else valor


def linhas(
    r: ReacoesNasBases, plano: pc.PlanoDeCargas | None = None, eixos: str = EIXOS_FUNDACAO
) -> list[list[object]]:
    """Uma linha por pé e caso, na ordem dos pés e dos casos."""
    vertical = _vertical(r.eixo_vertical, eixos)
    horizontais = [k for k in range(3) if k != vertical]
    saida: list[list[object]] = []
    for pe in r.pes:
        for caso in r.casos:
            a = r.acao(pe.rotulo, caso)
            if a is None:
                continue
            f = nos_eixos(a.forca_kN, r.eixo_vertical, eixos)
            m = nos_eixos(a.momento_kNm, r.eixo_vertical, eixos)
            acao = plano.acao(caso) if plano is not None else None
            saida.append(
                [
                    pe.rotulo,
                    caso,
                    *(_arredondar(v) for v in f),
                    *(_arredondar(v) for v in m),
                    _arredondar(-f[vertical]),
                    _arredondar(math.hypot(*(f[k] for k in horizontais))),
                    _arredondar(math.hypot(*(m[k] for k in horizontais))),
                    acao.nome if acao is not None else "—",
                    pe.pilar or "—",
                    a.fonte,
                ]
            )
    return saida


COLUNAS_LOCACAO = ["Pé", "Pilar", "Perfil", "X (m)", "Y (m)", "Z (m)", "Trabalha como"]


def linhas_da_locacao(r: ReacoesNasBases, eixos: str = EIXOS_FUNDACAO) -> list[list[object]]:
    saida: list[list[object]] = []
    for pe in r.pes:
        p = nos_eixos(pe.ponto, r.eixo_vertical, eixos) if pe.ponto is not None else None
        saida.append(
            [
                pe.rotulo,
                pe.pilar or "—",
                pe.perfil or "—",
                *((_arredondar(v) for v in p) if p is not None else ("—", "—", "—")),
                pe.comportamento or "—",
            ]
        )
    return saida


COLUNAS_CASOS = [
    "Caso",
    "Descrição",
    "Natureza",
    "Categoria (NBR 8800)",
    "γ desfavorável",
    "γ favorável",
    "ψ0",
    "ψ1",
    "ψ2",
]


def linhas_dos_casos(r: ReacoesNasBases, plano: pc.PlanoDeCargas | None) -> list[list[object]]:
    saida: list[list[object]] = []
    for caso in r.casos:
        acao = plano.acao(caso) if plano is not None else None
        if acao is None:
            saida.append([caso, "não está no plano de cargas", "—", "—", "", "", "", "", ""])
            continue
        try:
            categoria = comb.categoria_nbr(acao.categoria)
        except ValueError:
            saida.append([caso, acao.nome, "—", acao.categoria, "", "", "", "", ""])
            continue
        saida.append(
            [
                caso,
                acao.nome,
                categoria.tipo,
                categoria.rotulo,
                categoria.gamma,
                "" if categoria.gamma_favoravel is None else categoria.gamma_favoravel,
                # ψ só vale para as ações variáveis (a permanente entra sempre inteira).
                *(
                    ("", "", "")
                    if categoria.tipo.casefold() == "permanente"
                    else (categoria.psi0, categoria.psi1, categoria.psi2)
                ),
            ]
        )
    return saida


COLUNAS_RESUMO = [
    "Pé",
    "Pilar",
    "Maior compressão (kN)",
    "Caso",
    "Maior arrancamento (kN)",
    "Caso ",
    "Maior H (kN)",
    "Caso  ",
    "Maior M horizontal (kN·m)",
    "Caso   ",
]


def linhas_do_resumo(r: ReacoesNasBases) -> list[list[object]]:
    """O maior de cada pé entre os casos (característico, sem combinar: só para ter a ordem de
    grandeza)."""
    vertical = _vertical(r.eixo_vertical, EIXOS_FUNDACAO)
    saida: list[list[object]] = []
    for pe in r.pes:
        do_pe = [a for a in r.acoes if a.pe == pe.rotulo]
        if not do_pe:
            continue
        valores = []
        for a in do_pe:
            f = nos_eixos(a.forca_kN, r.eixo_vertical, EIXOS_FUNDACAO)
            m = nos_eixos(a.momento_kNm, r.eixo_vertical, EIXOS_FUNDACAO)
            valores.append((a.caso, -f[vertical], math.hypot(f[0], f[1]), math.hypot(m[0], m[1])))
        comp = max(valores, key=lambda v: v[1])
        arr = min(valores, key=lambda v: v[1])
        h = max(valores, key=lambda v: v[2])
        momento = max(valores, key=lambda v: v[3])
        saida.append(
            [
                pe.rotulo,
                pe.pilar or "—",
                _arredondar(comp[1]) if comp[1] > 0 else "—",
                comp[0] if comp[1] > 0 else "—",
                _arredondar(-arr[1]) if arr[1] < 0 else "—",
                arr[0] if arr[1] < 0 else "—",
                _arredondar(h[2]),
                h[0],
                _arredondar(momento[3]),
                momento[0],
            ]
        )
    return saida


def convencao(r: ReacoesNasBases, eixos: str) -> str:
    if eixos == EIXOS_MODELO:
        vertical = "Y" if r.eixo_vertical == "Y" else "Z"
        return (
            f"Eixos do modelo do SolidWorks ({vertical} para cima). Força e momento que a estrutura "
            f"faz na fundação (a reação com o sinal trocado): F{vertical} negativa é para baixo, "
            "compressão. N = compressão (positiva) e H = a força horizontal resultante."
        )
    troca = (
        " X é o X do SolidWorks, Y é o −Z do SolidWorks e Z é o Y do SolidWorks (no modelo, Y para "
        "cima) — a vista Superior do SolidWorks com o Y da planta para cima."
        if r.eixo_vertical == "Y"
        else " Os eixos são os do modelo (Z para cima)."
    )
    return (
        "Eixos da fundação: X e Y na planta, Z para cima." + troca + " Força e momento que a "
        "estrutura faz na fundação (a reação com o sinal trocado): Fz negativa é para baixo, "
        "compressão; N = −Fz (compressão positiva). H = a força horizontal resultante; "
        "M horizontal = a resultante de Mx e My (o Mz é a torção). Momentos pela regra da mão "
        "direita."
    )


# ---------------------------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------------------------
def xlsx(
    r: ReacoesNasBases,
    plano: pc.PlanoDeCargas | None = None,
    *,
    eixos: str = EIXOS_FUNDACAO,
    projeto: str = "",
    responsavel: str = "",
    pendentes: Sequence[dm.Achado] = (),
    hoje: date | None = None,
) -> bytes:
    """A planilha para quem projeta a fundação: Leia-me, Casos, Locação, Reações por pé, Resumo."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    livro = Workbook()
    fonte_cabecalho = Font(bold=True, color="FFFFFF")
    fundo = PatternFill("solid", fgColor="1F4E8C")
    aviso = PatternFill("solid", fgColor="FDE9E7")

    def tabela(nome: str, cabecalho: Sequence[str], dados_: Sequence[Sequence[object]]) -> None:
        folha = livro.create_sheet(nome)
        folha.append(list(cabecalho))
        for linha in dados_:
            folha.append(list(linha))
        for celula in folha[1]:
            celula.font, celula.fill = fonte_cabecalho, fundo
            celula.alignment = Alignment(wrap_text=True, vertical="center")
        folha.freeze_panes = "A2"
        if dados_:
            folha.auto_filter.ref = folha.dimensions
        for indice, coluna in enumerate(folha.iter_cols(min_row=1, max_row=folha.max_row), 1):
            largura = max((len(str(c.value)) for c in coluna if c.value is not None), default=8)
            folha.column_dimensions[get_column_letter(indice)].width = min(max(9, largura + 2), 45)
            for celula in coluna[1:]:
                if isinstance(celula.value, float):
                    celula.number_format = "0.000"

    leia = livro.active
    assert leia is not None
    leia.title = "Leia-me"
    titulo = "Reações nas bases para o projeto da fundação" + (f" — {projeto}" if projeto else "")
    if pendentes:
        titulo += " — PRELIMINAR (com pendências)"
    leia.append([titulo])
    leia["A1"].font = Font(bold=True, size=13)
    leia.append([f"Data: {(hoje or date.today()).strftime('%d/%m/%Y')}"])
    if responsavel:
        leia.append([f"Responsável: {responsavel}"])
    leia.append([])
    for texto in (
        NOTA_COMBINACOES,
        convencao(r, eixos),
        "Unidades: kN e kN·m; coordenadas em m.",
        f"Casos: {', '.join(r.casos)}. Pés: {len(r.pes)} (locação na aba Locação).",
        "De onde veio: "
        + f"“{FONTE_APOIO}” = a reação de cada nó apoiado, lida do SolidWorks pela macro; "
        + f"“{FONTE_PILAR}” = a força na base do pilar, levada aos eixos do modelo pela direção "
        + "da seção do pilar.",
    ):
        leia.append([texto])
    if r.conferencia:
        leia.append([])
        leia.append(["Conferência: a soma dos pés de cada caso contra a reação total do caso"])
        leia.cell(leia.max_row, 1).font = Font(bold=True)
        for item in r.conferencia:
            leia.append([("OK — " if item.nivel == pc.NIVEL_OK else "ATENÇÃO — ") + item.texto])
    if r.faltam or r.avisos:
        leia.append([])
        leia.append(["Avisos"])
        leia.cell(leia.max_row, 1).font = Font(bold=True)
        for texto in (*r.faltam, *r.avisos):
            leia.append([texto])
    if pendentes:
        leia.append([])
        leia.append(["PENDÊNCIAS — os valores podem mudar depois de corrigir:"])
        leia.cell(leia.max_row, 1).font = Font(bold=True, color="C00000")
        for a in pendentes:
            leia.append([f"• {a.titulo}. {a.problema}"])
            leia.cell(leia.max_row, 1).fill = aviso
    leia.column_dimensions["A"].width = 150
    for linha in leia.iter_rows():
        for celula in linha:
            celula.alignment = Alignment(wrap_text=True, vertical="top")

    tabela("Casos", COLUNAS_CASOS, linhas_dos_casos(r, plano))
    tabela("Locação", COLUNAS_LOCACAO, linhas_da_locacao(r, eixos))
    tabela("Reações por pé", colunas(r, eixos), linhas(r, plano, eixos))
    tabela("Resumo por pé", COLUNAS_RESUMO, linhas_do_resumo(r))
    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()


def nome_do_arquivo(projeto: str = "") -> str:
    base = "".join(c if c.isalnum() or c in "-_" else "_" for c in projeto).strip("_")
    return f"reacoes_nas_bases{('_' + base) if base else ''}.xlsx"
