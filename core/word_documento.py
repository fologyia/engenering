"""Documento Word (editável) em paisagem, com títulos, tabelas copiáveis e figuras.

Base dos botões "Gerar o Word" do fim das páginas (degrau de escada, flambagem de colunas): cada
página monta o seu conteúdo com esta classe. Sem Streamlit.
"""

from __future__ import annotations

from collections.abc import Sequence
from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

AZUL = RGBColor(0x16, 0x32, 0x4F)
CINZA = RGBColor(0x5F, 0x6B, 0x76)
FUNDO_CABECALHO = "24577A"
FUNDO_DESTAQUE = "FFF3C4"
COR_STATUS = {
    "OK": "D9F2E3",
    "NÃO OK": "F8D4D4",
    "ALERTA": "FCEBC4",
    "INFO": "D8E8F8",
    "N/A": "E6E9ED",
}
LARGURA_UTIL_MM = 273.0  # A4 em paisagem com margens de 12 mm


def numero_pt(x: float, casas: int = 2) -> str:
    """Número com vírgula decimal e sem zeros inúteis: 180 → "180", 166,667 → "166,67"."""
    texto = f"{x:.{casas}f}"
    if "." in texto:
        texto = texto.rstrip("0").rstrip(".")
    return texto.replace(".", ",") or "0"


def celula_texto(valor: Any) -> str:
    if valor is None:
        return "—"
    if isinstance(valor, bool):
        return "Sim" if valor else "Não"
    if isinstance(valor, float):
        return numero_pt(valor, 4)
    return str(valor)


def _sombrear(celula: Any, cor: str) -> None:
    propriedades = celula._tc.get_or_add_tcPr()
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), cor)
    propriedades.append(sombra)


def _repetir_cabecalho(linha: Any) -> None:
    propriedades = linha._tr.get_or_add_trPr()
    marca = OxmlElement("w:tblHeader")
    marca.set(qn("w:val"), "true")
    propriedades.append(marca)


class DocumentoWord:
    """Um documento Word em paisagem com títulos, parágrafos, tabelas e figuras."""

    def __init__(self) -> None:
        self.doc = Document()
        secao = self.doc.sections[0]
        secao.orientation = WD_ORIENT.LANDSCAPE
        secao.page_width, secao.page_height = Mm(297), Mm(210)
        for margem in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(secao, margem, Mm(12))
        normal = self.doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(9)

    def titulo(self, texto: str, nivel: int = 1) -> None:
        paragrafo = self.doc.add_heading(texto, level=nivel)
        for trecho in paragrafo.runs:
            trecho.font.color.rgb = AZUL

    def paragrafo(self, texto: str, *, legenda: bool = False, negrito: bool = False) -> None:
        paragrafo = self.doc.add_paragraph()
        trecho = paragrafo.add_run(texto)
        trecho.bold = negrito
        if legenda:
            trecho.font.size = Pt(8)
            trecho.font.color.rgb = CINZA

    def itens(self, textos: Sequence[str]) -> None:
        for texto in textos:
            self.doc.add_paragraph(texto, style="List Bullet")

    def tabela(
        self,
        cabecalhos: Sequence[str],
        linhas: Sequence[Sequence[Any]],
        larguras_mm: Sequence[float] | None = None,
        *,
        status_coluna: int | None = None,
        destacar_linha: int | None = None,
        fonte: float = 8.0,
    ) -> None:
        tabela = self.doc.add_table(rows=1, cols=len(cabecalhos))
        tabela.style = "Table Grid"
        if larguras_mm is None:
            larguras_mm = [LARGURA_UTIL_MM / len(cabecalhos)] * len(cabecalhos)
        escala = min(1.0, LARGURA_UTIL_MM / sum(larguras_mm))
        larguras = [Mm(w * escala) for w in larguras_mm]

        def escrever(celula: Any, texto: str, *, cabecalho: bool = False) -> None:
            celula.text = ""
            trecho = celula.paragraphs[0].add_run(texto)
            trecho.font.size = Pt(fonte)
            if cabecalho:
                trecho.bold = True
                trecho.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        topo = tabela.rows[0]
        _repetir_cabecalho(topo)
        for celula, texto, largura in zip(topo.cells, cabecalhos, larguras, strict=True):
            escrever(celula, str(texto), cabecalho=True)
            _sombrear(celula, FUNDO_CABECALHO)
            celula.width = largura
        for indice, linha in enumerate(linhas):
            celulas = tabela.add_row().cells
            for coluna, (celula, valor, largura) in enumerate(
                zip(celulas, linha, larguras, strict=True)
            ):
                texto = celula_texto(valor)
                escrever(celula, texto)
                celula.width = largura
                if status_coluna is not None and coluna == status_coluna:
                    cor = COR_STATUS.get(texto)
                    if cor:
                        _sombrear(celula, cor)
                elif destacar_linha is not None and indice == destacar_linha:
                    _sombrear(celula, FUNDO_DESTAQUE)
        self.doc.add_paragraph()

    def figura(self, png: bytes, legenda: str, largura_mm: float = 230) -> None:
        self.doc.add_picture(BytesIO(png), width=Mm(min(largura_mm, LARGURA_UTIL_MM)))
        self.paragrafo(legenda, legenda=True)

    def bytes(self) -> bytes:
        memoria = BytesIO()
        self.doc.save(memoria)
        return memoria.getvalue()
